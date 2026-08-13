from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from typing import Any, Iterable, Sequence
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session

from app.models.entities import (
    Asset,
    AssetHistory,
    ComponentReplacement,
    ReplacementRecord,
    WorkRecord,
)
from app.modules.it_activity.models import (
    ITHandoverRecord,
    ITPurchaseRecord,
    ITPurchaseRequest,
    ITPurchaseRequestHistory,
)
from app.services.monthly_snapshot_service import (
    assets_for_month,
    month_end,
    month_start,
    next_month,
    parse_month_key,
)
from app.services.asset_lifecycle_service import (
    canonical_device_type,
    inventory_summary,
    is_primary_device_type,
)


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
IST = ZoneInfo("Asia/Kolkata")
NAVY = "0F2744"
TEAL = "087E8B"
GREEN = "059669"
LIGHT_BLUE = "EAF2F8"
LIGHT_ORANGE = "FFF4E5"
BORDER = "D5E1EB"
TEXT = "172B3A"


@dataclass(frozen=True)
class ReportPeriod:
    kind: str
    key: str
    label: str
    start_date: date
    end_date: date
    months: tuple[date, ...]

    @property
    def month_keys(self) -> tuple[str, ...]:
        return tuple(item.strftime("%Y-%m") for item in self.months)


def _month_sequence(start: date, end: date) -> tuple[date, ...]:
    months: list[date] = []
    current = start.replace(day=1)
    final = end.replace(day=1)
    while current <= final:
        months.append(current)
        current = next_month(current)
    return tuple(months)


def monthly_period(month_key: str) -> ReportPeriod:
    start = parse_month_key(month_key)
    return ReportPeriod(
        kind="monthly",
        key=month_key,
        label=start.strftime("%B %Y"),
        start_date=start,
        end_date=month_end(start),
        months=(start,),
    )


def yearly_period(year: int, period_type: str) -> ReportPeriod:
    if year < 2000 or year > 2100:
        raise ValueError("Year must be between 2000 and 2100")
    normalized = (period_type or "calendar").strip().lower()
    if normalized == "calendar":
        start = date(year, 1, 1)
        end = date(year, 12, 31)
        label = f"Calendar Year {year}"
        key = f"CY-{year}"
    elif normalized == "financial":
        start = date(year, 4, 1)
        end = date(year + 1, 3, 31)
        label = f"Financial Year {year}-{str(year + 1)[-2:]}"
        key = f"FY-{year}-{str(year + 1)[-2:]}"
    else:
        raise ValueError("Period type must be calendar or financial")
    return ReportPeriod(
        kind=normalized,
        key=key,
        label=label,
        start_date=start,
        end_date=end,
        months=_month_sequence(start, end),
    )


def _safe(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (int, float, bool, date, datetime)):
        return value
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text[:32000]


def _local_datetime(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    aware = value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
    return aware.astimezone(IST).replace(tzinfo=None)


def _utc_bounds(period: ReportPeriod) -> tuple[datetime, datetime]:
    start_local = datetime.combine(period.start_date, datetime.min.time(), tzinfo=IST)
    end_local = datetime.combine(period.end_date + timedelta(days=1), datetime.min.time(), tzinfo=IST)
    return (
        start_local.astimezone(timezone.utc).replace(tzinfo=None),
        end_local.astimezone(timezone.utc).replace(tzinfo=None),
    )


def _blank_month(column):
    return or_(column.is_(None), column == "")


def _period_filter(reporting_month_column, *, period: ReportPeriod, datetime_column=None, date_column=None):
    month_match = reporting_month_column.in_(period.month_keys)
    fallback_parts = [_blank_month(reporting_month_column)]
    if date_column is not None:
        fallback_parts.append(date_column >= period.start_date)
        fallback_parts.append(date_column <= period.end_date)
    elif datetime_column is not None:
        utc_start, utc_end = _utc_bounds(period)
        fallback_parts.append(datetime_column >= utc_start)
        fallback_parts.append(datetime_column < utc_end)
    else:
        raise ValueError("A fallback date or datetime column is required")
    return or_(month_match, and_(*fallback_parts))


def _period_filter_with_optional_business_date(
    reporting_month_column,
    business_date_column,
    datetime_column,
    *,
    period: ReportPeriod,
):
    utc_start, utc_end = _utc_bounds(period)
    return or_(
        reporting_month_column.in_(period.month_keys),
        and_(
            _blank_month(reporting_month_column),
            or_(
                and_(
                    business_date_column.is_not(None),
                    business_date_column >= period.start_date,
                    business_date_column <= period.end_date,
                ),
                and_(
                    business_date_column.is_(None),
                    datetime_column >= utc_start,
                    datetime_column < utc_end,
                ),
            ),
        ),
    )


def _header(sheet, columns: Sequence[str]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column_index, title in enumerate(columns, 1):
        cell = sheet.cell(1, column_index, title)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color=BORDER))
    sheet.row_dimensions[1].height = 34


def _write_rows(sheet, columns: Sequence[str], rows: Iterable[dict[str, Any]]) -> int:
    _header(sheet, columns)
    count = 0
    for row_index, row in enumerate(rows, 2):
        count += 1
        for column_index, column in enumerate(columns, 1):
            cell = sheet.cell(row_index, column_index, _safe(row.get(column)))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FBFD")
            if isinstance(cell.value, datetime):
                cell.number_format = "DD-MM-YYYY HH:MM:SS"
            elif isinstance(cell.value, date):
                cell.number_format = "DD-MM-YYYY"
            elif isinstance(cell.value, float) and any(token in column.lower() for token in ("price", "amount", "cost", "value")):
                cell.number_format = '#,##0.00;[Red](#,##0.00);-'
        sheet.row_dimensions[row_index].height = 28
    last_row = max(sheet.max_row, 1)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"
    for column_index, column in enumerate(columns, 1):
        sample_lengths = [
            len(str(sheet.cell(row_index, column_index).value or ""))
            for row_index in range(1, min(last_row, 250) + 1)
        ]
        width = min(max([len(column), *sample_lengths], default=len(column)) + 2, 48)
        sheet.column_dimensions[get_column_letter(column_index)].width = max(width, 12)
    return count


def _status_label(value: str | None) -> str | None:
    if not value:
        return None
    return value.replace("_", " ").title()


def _asset_category(asset: Any) -> str:
    return canonical_device_type(getattr(asset, "device_type", None))


def _is_external_hdd(asset: Any) -> bool:
    return _asset_category(asset) == "External HDD"


def _asset_source_label(source: str) -> str:
    return {
        "live": "Live register",
        "snapshot": "Finalized system snapshot",
        "template": "Original historical workbook",
        "missing": "No source available",
    }.get(source, source.replace("_", " ").title())


def _asset_summary_row(month: date, assets: Sequence[Any], source: str) -> dict[str, Any]:
    if source == "missing":
        return {
            "Month": month.strftime("%B %Y"),
            "Month Key": month.strftime("%Y-%m"),
            "Source": _asset_source_label(source),
            "Source Available": "No",
            "Primary IT Assets": None,
            "Computers": None,
            "Laptops": None,
            "Smartphones": None,
            "Printers": None,
            "External HDDs": None,
            "Assigned / In Use": None,
            "Available": None,
            "Under Repair": None,
            "Replacement Pending": None,
            "All Asset Rows": None,
        }
    primary = [asset for asset in assets if is_primary_device_type(getattr(asset, "device_type", None))]
    primary_summary = inventory_summary(primary)
    all_summary = inventory_summary(assets)
    return {
        "Month": month.strftime("%B %Y"),
        "Month Key": month.strftime("%Y-%m"),
        "Source": _asset_source_label(source),
        "Source Available": "Yes",
        "Primary IT Assets": primary_summary["total"],
        "Computers": primary_summary["computers"],
        "Laptops": primary_summary["laptops"],
        "Smartphones": primary_summary["smartphones"],
        "Printers": primary_summary["printers"],
        "External HDDs": all_summary["external_hdds"],
        "Assigned / In Use": primary_summary["assigned"],
        "Available": primary_summary["available"],
        "Under Repair": primary_summary["repair"],
        "Replacement Pending": primary_summary["replacement_pending"],
        "All Asset Rows": len(assets),
    }


def _assets_for_closing_period(db: Session, period: ReportPeriod) -> tuple[list[Any], str, date | None]:
    current = month_start()
    candidate_months = [month for month in period.months if month <= current]
    for month in reversed(candidate_months):
        assets, source = assets_for_month(db, month)
        if source != "missing":
            return list(assets), source, month
    return [], "missing", None


def _asset_rows(assets: Sequence[Any], *, include: str) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for asset in assets:
        external_hdd = _is_external_hdd(asset)
        printer = _asset_category(asset) == "Printer"
        if include == "primary" and (external_hdd or printer):
            continue
        if include == "printers" and not printer:
            continue
        if include == "external_hdds" and not external_hdd:
            continue
        result.append({
            "Internal Asset ID": getattr(asset, "asset_code", None),
            "CPU / Asset Tag": getattr(asset, "cpu_asset_tag", None),
            "Device Type": _asset_category(asset),
            "Assigned User": getattr(asset, "used_by", None),
            "Current Holder": getattr(asset, "current_holder", None),
            "Workstation": getattr(asset, "workstation_no", None),
            "Department": getattr(asset, "department", None),
            "Location / Floor": getattr(asset, "location", None),
            "Work Mode": _status_label(getattr(asset, "work_mode", None)),
            "Status": _status_label(getattr(asset, "status", None)),
            "Brand": getattr(asset, "brand", None),
            "Model": getattr(asset, "model", None),
            "Serial Number": getattr(asset, "serial_number", None),
            "Connection": getattr(asset, "connection_type", None),
            "Capacity": getattr(asset, "capacity", None),
            "Ownership": getattr(asset, "ownership", None),
            "Client": getattr(asset, "client_name", None),
            "Project ID": getattr(asset, "project_id", None),
            "Processor": getattr(asset, "processor", None),
            "Memory": getattr(asset, "memory_gb", None),
            "SSD": getattr(asset, "ssd", None),
            "HDD": getattr(asset, "hdd", None),
            "Monitor Asset Tags": getattr(asset, "monitor_asset_tags", None),
            "Mouse Asset Tag": getattr(asset, "mouse_asset_tag", None),
            "Keyboard Asset Tag": getattr(asset, "keyboard_asset_tag", None),
            "System Name": getattr(asset, "system_name", None),
            "IP Address": getattr(asset, "ip_address", None),
            "MAC Address": getattr(asset, "mac_address", None),
            "Graphics Card": getattr(asset, "graphics_card", None),
            "Operating System": getattr(asset, "operating_system", None),
            "Antivirus": getattr(asset, "antivirus", None),
            "Network Type": getattr(asset, "network_type", None),
            "Price": getattr(asset, "price", None),
            "Original Asset Date": getattr(asset, "original_asset_date", None) or getattr(asset, "asset_date", None),
            "Last Asset Date": getattr(asset, "asset_date", None),
            "Remarks": getattr(asset, "remarks", None),
        })
    return sorted(
        result,
        key=lambda row: (
            str(row.get("CPU / Asset Tag") or "").casefold(),
            str(row.get("Internal Asset ID") or "").casefold(),
        ),
    )


def _query_period_records(db: Session, period: ReportPeriod) -> dict[str, list[Any]]:
    histories = list(db.scalars(
        select(AssetHistory)
        .where(_period_filter(AssetHistory.reporting_month, period=period, datetime_column=AssetHistory.created_at))
        .order_by(AssetHistory.created_at, AssetHistory.id)
    ).all())
    work_records = list(db.scalars(
        select(WorkRecord)
        .where(
            WorkRecord.module == "it",
            _period_filter_with_optional_business_date(
                WorkRecord.reporting_month,
                WorkRecord.start_date,
                WorkRecord.created_at,
                period=period,
            ),
        )
        .order_by(WorkRecord.created_at, WorkRecord.id)
    ).all())
    components = list(db.scalars(
        select(ComponentReplacement)
        .where(_period_filter_with_optional_business_date(
            ComponentReplacement.reporting_month,
            ComponentReplacement.replacement_date,
            ComponentReplacement.created_at,
            period=period,
        ))
        .order_by(ComponentReplacement.created_at, ComponentReplacement.id)
    ).all())
    replacements = list(db.scalars(
        select(ReplacementRecord)
        .where(_period_filter(ReplacementRecord.reporting_month, period=period, datetime_column=ReplacementRecord.created_at))
        .order_by(ReplacementRecord.created_at, ReplacementRecord.id)
    ).all())
    handovers = list(db.scalars(
        select(ITHandoverRecord)
        .where(_period_filter(ITHandoverRecord.reporting_month, period=period, date_column=ITHandoverRecord.activity_date))
        .order_by(ITHandoverRecord.activity_date, ITHandoverRecord.activity_time, ITHandoverRecord.id)
    ).all())
    purchases = list(db.scalars(
        select(ITPurchaseRecord)
        .where(_period_filter(ITPurchaseRecord.reporting_month, period=period, date_column=ITPurchaseRecord.purchase_date))
        .order_by(ITPurchaseRecord.purchase_date, ITPurchaseRecord.id)
    ).all())
    requests = list(db.scalars(
        select(ITPurchaseRequest)
        .where(_period_filter(ITPurchaseRequest.reporting_month, period=period, datetime_column=ITPurchaseRequest.requested_at))
        .order_by(ITPurchaseRequest.requested_at, ITPurchaseRequest.id)
    ).unique().all())
    utc_start, utc_end = _utc_bounds(period)
    request_histories = list(db.scalars(
        select(ITPurchaseRequestHistory)
        .where(ITPurchaseRequestHistory.created_at >= utc_start, ITPurchaseRequestHistory.created_at < utc_end)
        .order_by(ITPurchaseRequestHistory.created_at, ITPurchaseRequestHistory.id)
    ).all())
    return {
        "histories": histories,
        "work_records": work_records,
        "components": components,
        "replacements": replacements,
        "handovers": handovers,
        "purchases": purchases,
        "requests": requests,
        "request_histories": request_histories,
    }


def _build_summary_sheet(
    workbook: Workbook,
    *,
    period: ReportPeriod,
    closing_month: date | None,
    closing_source: str,
    closing_assets: Sequence[Any],
    records: dict[str, list[Any]],
    monthly_rows: Sequence[dict[str, Any]],
) -> None:
    ws = workbook.create_sheet("Report Summary")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F2")
    ws["A1"] = f"NAKSHATECH COMPLETE IT REPORT — {period.label.upper()}"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18

    metadata = [
        ("Report Period", period.label),
        ("From", period.start_date),
        ("To", period.end_date),
        ("Timezone", "Asia/Kolkata"),
        ("Generated At", datetime.now(IST).replace(tzinfo=None)),
        ("Closing Asset Month", closing_month.strftime("%B %Y") if closing_month else "No asset source available"),
        ("Closing Asset Source", _asset_source_label(closing_source)),
    ]
    for row_index, (label, value) in enumerate(metadata, 4):
        ws.cell(row_index, 1, label).font = Font(bold=True, color=TEXT)
        ws.cell(row_index, 2, value)
        if isinstance(value, datetime):
            ws.cell(row_index, 2).number_format = "DD-MM-YYYY HH:MM:SS"
        elif isinstance(value, date):
            ws.cell(row_index, 2).number_format = "DD-MM-YYYY"

    primary_assets = [asset for asset in closing_assets if is_primary_device_type(getattr(asset, "device_type", None))]
    primary_summary = inventory_summary(primary_assets)
    all_summary = inventory_summary(closing_assets)
    metrics = [
        ("Primary IT Assets at Closing", primary_summary["total"]),
        ("Computers at Closing", primary_summary["computers"]),
        ("Laptops at Closing", primary_summary["laptops"]),
        ("Smartphones at Closing", primary_summary["smartphones"]),
        ("Printers at Closing", primary_summary["printers"]),
        ("External HDDs at Closing", all_summary["external_hdds"]),
        ("Servers at Closing", primary_summary["servers"]),
        ("Network Devices at Closing", primary_summary["network_devices"]),
        ("Other Device Types at Closing", primary_summary["other"]),
        ("Assigned / In Use at Closing", primary_summary["assigned"]),
        ("Available at Closing", primary_summary["available"]),
        ("Under Repair at Closing", primary_summary["repair"]),
        ("Replacement Pending at Closing", primary_summary["replacement_pending"]),
        ("Damaged / At Risk at Closing", primary_summary["damaged"]),
        ("Retired / Finalized at Closing", primary_summary["terminal"]),
        ("Asset Change Operations", len(records["histories"])),
        ("Work Records", len(records["work_records"])),
        ("Component Changes", len(records["components"])),
        ("Full Asset Replacements", len(records["replacements"])),
        ("Handover / Return Operations", len(records["handovers"])),
        ("Purchases", len(records["purchases"])),
        ("Purchase Requests", len(records["requests"])),
        ("Purchase Approval Actions", len(records["request_histories"])),
        ("Purchase Value (INR)", round(sum(float(item.total_price or 0) for item in records["purchases"]), 2)),
        ("Months with Asset Source", sum(row["Source Available"] == "Yes" for row in monthly_rows)),
        ("Months Missing Asset Source", sum(row["Source Available"] == "No" for row in monthly_rows)),
    ]
    header_row = 13
    ws.cell(header_row, 1, "Metric")
    ws.cell(header_row, 2, "Count / Value")
    for cell in ws[header_row]:
        if cell.column <= 2:
            cell.fill = PatternFill("solid", fgColor=TEAL)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, (label, value) in enumerate(metrics, header_row + 1):
        ws.cell(row_index, 1, label)
        ws.cell(row_index, 2, value)
        if row_index % 2 == 0:
            ws.cell(row_index, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            ws.cell(row_index, 2).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        if "Value" in label:
            ws.cell(row_index, 2).number_format = '#,##0.00;[Red](#,##0.00);-'

    note_row = header_row + len(metrics) + 3
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 2, end_column=6)
    ws.cell(note_row, 1, (
        "Accuracy rule: asset totals represent the closing state for each month and are never summed as yearly assets. "
        "Operational rows such as edits, work records, handovers, returns and purchases are counted inside the selected period. "
        "A month marked as missing has not been estimated or replaced with current live data."
    ))
    ws.cell(note_row, 1).fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
    ws.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(note_row, 1).font = Font(color=TEXT, italic=True)

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 26
    for column in "CDEF":
        ws.column_dimensions[column].width = 18


def _history_rows(records: Sequence[AssetHistory], assets_by_id: dict[int, Asset]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in records:
        asset = assets_by_id.get(item.asset_id)
        rows.append({
            "Reporting Month": item.reporting_month or _local_datetime(item.created_at).strftime("%Y-%m"),
            "System Recorded At": _local_datetime(item.created_at),
            "Internal Asset ID": getattr(asset, "asset_code", None),
            "CPU / Asset Tag": getattr(asset, "cpu_asset_tag", None),
            "Device Type": getattr(asset, "device_type", None),
            "Assigned User": getattr(asset, "used_by", None),
            "Department": getattr(asset, "department", None),
            "Action": item.action,
            "Change Type": _status_label(item.change_type),
            "Batch ID": item.batch_code,
            "Fields Changed": item.field_count,
            "Old Value": item.old_value,
            "New Value": item.new_value,
            "Reason": item.reason,
            "Remarks": item.remarks,
            "Changed By": item.changed_by_name or item.changed_by,
            "Changed By Email": item.changed_by,
            "Changed By Role": _status_label(item.changed_by_role),
        })
    return rows


def _work_rows(records: Sequence[WorkRecord], assets_by_id: dict[int, Asset]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in records:
        asset = assets_by_id.get(item.asset_id) if item.asset_id else None
        rows.append({
            "Reporting Month": item.reporting_month or (item.start_date or _local_datetime(item.created_at).date()).strftime("%Y-%m"),
            "Work Code": item.work_code,
            "Internal Asset ID": getattr(asset, "asset_code", None),
            "CPU / Asset Tag": getattr(asset, "cpu_asset_tag", None),
            "Title": item.title,
            "Work Type": item.work_type,
            "Project": item.project,
            "Assigned To": item.assigned_to,
            "Technician": item.technician,
            "Priority": _status_label(item.priority),
            "Issue Description": item.issue_description,
            "Details": item.details,
            "Status": _status_label(item.status),
            "Root Cause": item.root_cause,
            "Resolution": item.resolution,
            "Replaced Component": item.replaced_component,
            "Replacement Asset Tag": item.replacement_asset_tag,
            "Cost": item.cost,
            "Approval Status": _status_label(item.approval_status),
            "Start Date": item.start_date,
            "Expected Completion": item.expected_completion_date,
            "Completed At": _local_datetime(item.completed_at),
            "System Recorded At": _local_datetime(item.created_at),
            "Last Updated At": _local_datetime(item.updated_at),
        })
    return rows


def _component_rows(records: Sequence[ComponentReplacement], assets_by_id: dict[int, Asset]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in records:
        asset = assets_by_id.get(item.asset_id)
        rows.append({
            "Reporting Month": item.reporting_month or (item.replacement_date or _local_datetime(item.created_at).date()).strftime("%Y-%m"),
            "Replacement Code": item.replacement_code,
            "Batch ID": item.batch_code,
            "Sequence": item.sequence_no,
            "Internal Asset ID": getattr(asset, "asset_code", None),
            "CPU / Asset Tag": item.cpu_asset_tag or getattr(asset, "cpu_asset_tag", None),
            "Workstation": item.workstation_no or getattr(asset, "workstation_no", None),
            "Component Type": item.component_type,
            "Change Type": _status_label(item.change_type),
            "Field": item.field_name,
            "Old Value": item.old_value,
            "New Value": item.new_value,
            "Old Condition": item.old_condition,
            "Reason": item.reason,
            "Technician": item.technician,
            "Replacement Date": item.replacement_date,
            "Performed By": item.performed_by,
            "Performed By Email": item.performed_by_email,
            "Performed By Role": _status_label(item.performed_by_role),
            "Approved By": item.approved_by,
            "Remarks": item.remarks,
            "System Recorded At": _local_datetime(item.created_at),
        })
    return rows


def _replacement_rows(records: Sequence[ReplacementRecord], assets_by_id: dict[int, Asset]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in records:
        old_asset = assets_by_id.get(item.old_asset_id)
        new_asset = assets_by_id.get(item.new_asset_id) if item.new_asset_id else None
        rows.append({
            "Reporting Month": item.reporting_month or _local_datetime(item.created_at).strftime("%Y-%m"),
            "Replacement Code": item.replacement_code,
            "Old Asset ID": getattr(old_asset, "asset_code", None),
            "Old CPU / Asset Tag": getattr(old_asset, "cpu_asset_tag", None),
            "New Asset ID": getattr(new_asset, "asset_code", None),
            "New CPU / Asset Tag": getattr(new_asset, "cpu_asset_tag", None),
            "Reason": item.reason,
            "Damage Category": _status_label(item.damage_category),
            "Inspection Finding": item.inspection_finding,
            "Approval Status": _status_label(item.approval_status),
            "Final Action": _status_label(item.final_action),
            "Requested By": item.requested_by,
            "Requested By Email": item.requested_by_email,
            "Requested By Role": _status_label(item.requested_by_role),
            "Approved By": item.approved_by,
            "Approved By Email": item.approved_by_email,
            "Approved By Role": _status_label(item.approved_by_role),
            "Requested At": _local_datetime(item.created_at),
            "Approved At": _local_datetime(item.approved_at),
        })
    return rows


def _handover_rows(records: Sequence[ITHandoverRecord]) -> list[dict[str, Any]]:
    return [{
        "Reporting Month": item.reporting_month or item.activity_date.strftime("%Y-%m"),
        "Activity Code": item.activity_code,
        "Device Category": _status_label(item.device_category),
        "Action": _status_label(item.action_type),
        "Original Action": item.action_raw,
        "Activity Date": item.activity_date,
        "Activity Time": item.activity_time.strftime("%I:%M:%S %p") if item.activity_time else None,
        "Internal Asset ID": item.asset_code_snapshot,
        "Internal Asset No.": item.internal_asset_no,
        "DC / Workstation No.": item.dc_number,
        "Employee Name": item.employee_name,
        "Department": item.department,
        "Work Mode": _status_label(item.work_mode),
        "Specification": item.specification,
        "Serial Number": item.serial_number,
        "Accessories Provided": item.accessories_provided,
        "Condition": item.condition,
        "Issued / Performed By": item.issued_by or item.performed_by,
        "Performed By Email": item.performed_by_email,
        "Performed By Role": _status_label(item.performed_by_role),
        "Remarks": item.remarks,
        "Asset Updated Status": _status_label(item.asset_updated_status),
        "Source File": item.source_file,
        "Source Sheet": item.source_sheet,
        "Source Row": item.source_row,
        "Imported": "Yes" if item.imported else "No",
        "System Recorded At": _local_datetime(item.created_at),
    } for item in records]


def _purchase_rows(records: Sequence[ITPurchaseRecord]) -> list[dict[str, Any]]:
    return [{
        "Reporting Month": item.reporting_month or item.purchase_date.strftime("%Y-%m"),
        "Purchase Code": item.purchase_code,
        "Purchase Request ID": item.purchase_request_id,
        "Linked Asset ID": item.linked_asset_code_snapshot,
        "Purchase Date": item.purchase_date,
        "PO Number": item.po_number,
        "Asset Number": item.asset_number,
        "Supplier": item.supplier_name,
        "Supplier Contact": item.supplier_contact,
        "Item Description": item.item_description,
        "Warranty Number": item.warranty_number,
        "Quantity": item.quantity,
        "Unit Price": item.unit_price,
        "Total Price": item.total_price,
        "Received Date": item.received_date,
        "Inspection Status": _status_label(item.inspection_status),
        "Approved By": item.approved_by,
        "Department": item.department,
        "Remarks": item.remarks,
        "Created By": item.created_by,
        "Created By Email": item.created_by_email,
        "Created By Role": _status_label(item.created_by_role),
        "Source File": item.source_file,
        "Source Sheet": item.source_sheet,
        "Source Row": item.source_row,
        "Imported": "Yes" if item.imported else "No",
        "System Recorded At": _local_datetime(item.created_at),
    } for item in records]


def _purchase_request_rows(records: Sequence[ITPurchaseRequest]) -> list[dict[str, Any]]:
    return [{
        "Reporting Month": item.reporting_month or _local_datetime(item.requested_at).strftime("%Y-%m"),
        "Request Code": item.request_code,
        "Requesting Department": item.requesting_department,
        "Requested Employee": item.requested_employee,
        "Item Type": _status_label(item.item_type),
        "Item Name": item.item_name,
        "Item Description": item.item_description,
        "Quantity": item.quantity,
        "Estimated Unit Price": item.estimated_unit_price,
        "Estimated Total Amount": item.estimated_total_amount,
        "Business Reason": item.business_reason,
        "Required By Date": item.required_by_date,
        "Priority": _status_label(item.priority),
        "IT Remarks": item.it_remarks,
        "Status": _status_label(item.status),
        "Branch": item.branch,
        "Requested By": item.requested_by_name,
        "Requested By Email": item.requested_by_email,
        "Requested By Role": _status_label(item.requested_by_role),
        "Requested At": _local_datetime(item.requested_at),
        "Approved Amount": item.approved_amount,
        "Management Remarks": item.management_remarks,
        "Decided By": item.decided_by_name,
        "Decided By Email": item.decided_by_email,
        "Decided By Role": _status_label(item.decided_by_role),
        "Decided At": _local_datetime(item.decided_at),
        "Purchase Completed At": _local_datetime(item.purchase_completed_at),
        "Last Updated At": _local_datetime(item.updated_at),
    } for item in records]


def _approval_rows(records: Sequence[ITPurchaseRequestHistory], requests_by_id: dict[int, ITPurchaseRequest]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in records:
        request = requests_by_id.get(item.request_id)
        rows.append({
            "Action Month": _local_datetime(item.created_at).strftime("%Y-%m"),
            "Request Code": getattr(request, "request_code", None),
            "Request Reporting Month": getattr(request, "reporting_month", None),
            "Action": _status_label(item.action),
            "From Status": _status_label(item.from_status),
            "To Status": _status_label(item.to_status),
            "Remarks": item.remarks,
            "Performed By": item.performed_by_name,
            "Performed By Email": item.performed_by_email,
            "Performed By Role": _status_label(item.performed_by_role),
            "System Recorded At": _local_datetime(item.created_at),
        })
    return rows


def _build_reconciliation_sheet(
    workbook: Workbook,
    *,
    monthly_rows: Sequence[dict[str, Any]],
    closing_assets: Sequence[Any],
    closing_month: date | None,
    closing_source: str,
    records: dict[str, list[Any]],
    row_counts: dict[str, int],
) -> None:
    ws = workbook.create_sheet("Reconciliation")
    ws.sheet_view.showGridLines = False
    ws.append(["Check", "Expected", "Workbook Rows", "Result", "Explanation"])
    _header(ws, ["Check", "Expected", "Workbook Rows", "Result", "Explanation"])

    primary = [asset for asset in closing_assets if not _is_external_hdd(asset)]
    printers = [asset for asset in closing_assets if _asset_category(asset) == "Printer"]
    hdds = [asset for asset in closing_assets if _is_external_hdd(asset)]
    checks = [
        ("Primary IT asset register", len([asset for asset in primary if _asset_category(asset) != "Printer"]), row_counts.get("Primary IT Assets", 0), "Closing state excluding Printer and External HDD"),
        ("Printer register", len(printers), row_counts.get("Printers", 0), "Closing-state Printer rows"),
        ("External HDD register", len(hdds), row_counts.get("External HDDs", 0), "Closing-state External HDD rows"),
        ("Asset change history", len(records["histories"]), row_counts.get("Asset Change History", 0), "Period-filtered audit operations"),
        ("Work records", len(records["work_records"]), row_counts.get("Work Records", 0), "IT module work records"),
        ("Component changes", len(records["components"]), row_counts.get("Component Changes", 0), "Period-filtered component operations"),
        ("Full replacements", len(records["replacements"]), row_counts.get("Full Replacements", 0), "Period-filtered complete-asset replacements"),
        ("Handover and return records", len(records["handovers"]), row_counts.get("Handovers Returns", 0), "Period-filtered laptop and desktop operations"),
        ("Purchase records", len(records["purchases"]), row_counts.get("Purchases", 0), "Period-filtered purchase records"),
        ("Purchase requests", len(records["requests"]), row_counts.get("Purchase Requests", 0), "Period-filtered requests"),
        ("Purchase approval history", len(records["request_histories"]), row_counts.get("Approval History", 0), "Approval actions recorded inside the selected period"),
    ]
    for row_index, (label, expected, actual, explanation) in enumerate(checks, 2):
        ws.cell(row_index, 1, label)
        ws.cell(row_index, 2, expected)
        ws.cell(row_index, 3, actual)
        ws.cell(row_index, 4, f'=IF(B{row_index}=C{row_index},"PASS","REVIEW")')
        ws.cell(row_index, 5, explanation)
        if row_index % 2 == 0:
            for column in range(1, 6):
                ws.cell(row_index, column).fill = PatternFill("solid", fgColor="F8FBFD")

    source_start = len(checks) + 4
    ws.cell(source_start, 1, "Monthly Asset Source Coverage")
    ws.cell(source_start, 1).fill = PatternFill("solid", fgColor=TEAL)
    ws.cell(source_start, 1).font = Font(color="FFFFFF", bold=True)
    ws.cell(source_start + 1, 1, "Closing Asset Month")
    ws.cell(source_start + 1, 2, closing_month.strftime("%B %Y") if closing_month else "Missing")
    ws.cell(source_start + 2, 1, "Closing Asset Source")
    ws.cell(source_start + 2, 2, _asset_source_label(closing_source))
    ws.cell(source_start + 3, 1, "Available Months")
    ws.cell(source_start + 3, 2, sum(row["Source Available"] == "Yes" for row in monthly_rows))
    ws.cell(source_start + 4, 1, "Missing Months")
    ws.cell(source_start + 4, 2, sum(row["Source Available"] == "No" for row in monthly_rows))
    ws.cell(source_start + 5, 1, "Missing Month Keys")
    ws.cell(source_start + 5, 2, ", ".join(row["Month Key"] for row in monthly_rows if row["Source Available"] == "No") or "None")

    for column, width in {"A": 34, "B": 18, "C": 18, "D": 14, "E": 62}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"


def build_complete_period_report(db: Session, period: ReportPeriod) -> BytesIO:
    monthly_rows: list[dict[str, Any]] = []
    for month in period.months:
        assets, source = assets_for_month(db, month)
        monthly_rows.append(_asset_summary_row(month, list(assets), source))

    closing_assets, closing_source, closing_month = _assets_for_closing_period(db, period)
    records = _query_period_records(db, period)
    live_assets = list(db.scalars(select(Asset).order_by(Asset.id)).all())
    assets_by_id = {asset.id: asset for asset in live_assets}
    requests_by_id = {item.id: item for item in db.scalars(select(ITPurchaseRequest)).unique().all()}

    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except AttributeError:
        pass

    _build_summary_sheet(
        workbook,
        period=period,
        closing_month=closing_month,
        closing_source=closing_source,
        closing_assets=closing_assets,
        records=records,
        monthly_rows=monthly_rows,
    )

    row_counts: dict[str, int] = {}
    monthly_columns = [
        "Month", "Month Key", "Source", "Source Available", "Primary IT Assets",
        "Computers", "Laptops", "Smartphones", "Printers", "External HDDs",
        "Assigned / In Use", "Available", "Under Repair", "Replacement Pending", "All Asset Rows",
    ]
    row_counts["Monthly Asset Totals"] = _write_rows(
        workbook.create_sheet("Monthly Asset Totals"), monthly_columns, monthly_rows
    )

    asset_columns = [
        "Internal Asset ID", "CPU / Asset Tag", "Device Type", "Assigned User", "Current Holder",
        "Workstation", "Department", "Location / Floor", "Work Mode", "Status", "Brand", "Model",
        "Serial Number", "Connection", "Capacity", "Ownership", "Client", "Project ID", "Processor",
        "Memory", "SSD", "HDD", "Monitor Asset Tags", "Mouse Asset Tag", "Keyboard Asset Tag",
        "System Name", "IP Address", "MAC Address", "Graphics Card", "Operating System", "Antivirus",
        "Network Type", "Price", "Original Asset Date", "Last Asset Date", "Remarks",
    ]
    row_counts["Primary IT Assets"] = _write_rows(
        workbook.create_sheet("Primary IT Assets"), asset_columns, _asset_rows(closing_assets, include="primary")
    )
    row_counts["Printers"] = _write_rows(
        workbook.create_sheet("Printers"), asset_columns, _asset_rows(closing_assets, include="printers")
    )
    row_counts["External HDDs"] = _write_rows(
        workbook.create_sheet("External HDDs"), asset_columns, _asset_rows(closing_assets, include="external_hdds")
    )

    history_columns = [
        "Reporting Month", "System Recorded At", "Internal Asset ID", "CPU / Asset Tag", "Device Type",
        "Assigned User", "Department", "Action", "Change Type", "Batch ID", "Fields Changed",
        "Old Value", "New Value", "Reason", "Remarks", "Changed By", "Changed By Email", "Changed By Role",
    ]
    row_counts["Asset Change History"] = _write_rows(
        workbook.create_sheet("Asset Change History"), history_columns, _history_rows(records["histories"], assets_by_id)
    )

    work_columns = [
        "Reporting Month", "Work Code", "Internal Asset ID", "CPU / Asset Tag", "Title", "Work Type",
        "Project", "Assigned To", "Technician", "Priority", "Issue Description", "Details", "Status",
        "Root Cause", "Resolution", "Replaced Component", "Replacement Asset Tag", "Cost", "Approval Status",
        "Start Date", "Expected Completion", "Completed At", "System Recorded At", "Last Updated At",
    ]
    row_counts["Work Records"] = _write_rows(
        workbook.create_sheet("Work Records"), work_columns, _work_rows(records["work_records"], assets_by_id)
    )

    component_columns = [
        "Reporting Month", "Replacement Code", "Batch ID", "Sequence", "Internal Asset ID", "CPU / Asset Tag",
        "Workstation", "Component Type", "Change Type", "Field", "Old Value", "New Value", "Old Condition",
        "Reason", "Technician", "Replacement Date", "Performed By", "Performed By Email", "Performed By Role",
        "Approved By", "Remarks", "System Recorded At",
    ]
    row_counts["Component Changes"] = _write_rows(
        workbook.create_sheet("Component Changes"), component_columns, _component_rows(records["components"], assets_by_id)
    )

    replacement_columns = [
        "Reporting Month", "Replacement Code", "Old Asset ID", "Old CPU / Asset Tag", "New Asset ID",
        "New CPU / Asset Tag", "Reason", "Damage Category", "Inspection Finding", "Approval Status", "Final Action",
        "Requested By", "Requested By Email", "Requested By Role", "Approved By", "Approved By Email",
        "Approved By Role", "Requested At", "Approved At",
    ]
    row_counts["Full Replacements"] = _write_rows(
        workbook.create_sheet("Full Replacements"), replacement_columns, _replacement_rows(records["replacements"], assets_by_id)
    )

    handover_columns = [
        "Reporting Month", "Activity Code", "Device Category", "Action", "Original Action", "Activity Date",
        "Activity Time", "Internal Asset ID", "Internal Asset No.", "DC / Workstation No.", "Employee Name",
        "Department", "Work Mode", "Specification", "Serial Number", "Accessories Provided", "Condition",
        "Issued / Performed By", "Performed By Email", "Performed By Role", "Remarks", "Asset Updated Status",
        "Source File", "Source Sheet", "Source Row", "Imported", "System Recorded At",
    ]
    row_counts["Handovers Returns"] = _write_rows(
        workbook.create_sheet("Handovers Returns"), handover_columns, _handover_rows(records["handovers"])
    )

    purchase_columns = [
        "Reporting Month", "Purchase Code", "Purchase Request ID", "Linked Asset ID", "Purchase Date", "PO Number",
        "Asset Number", "Supplier", "Supplier Contact", "Item Description", "Warranty Number", "Quantity",
        "Unit Price", "Total Price", "Received Date", "Inspection Status", "Approved By", "Department", "Remarks",
        "Created By", "Created By Email", "Created By Role", "Source File", "Source Sheet", "Source Row", "Imported",
        "System Recorded At",
    ]
    row_counts["Purchases"] = _write_rows(
        workbook.create_sheet("Purchases"), purchase_columns, _purchase_rows(records["purchases"])
    )

    request_columns = [
        "Reporting Month", "Request Code", "Requesting Department", "Requested Employee", "Item Type", "Item Name",
        "Item Description", "Quantity", "Estimated Unit Price", "Estimated Total Amount", "Business Reason",
        "Required By Date", "Priority", "IT Remarks", "Status", "Branch", "Requested By", "Requested By Email",
        "Requested By Role", "Requested At", "Approved Amount", "Management Remarks", "Decided By", "Decided By Email",
        "Decided By Role", "Decided At", "Purchase Completed At", "Last Updated At",
    ]
    row_counts["Purchase Requests"] = _write_rows(
        workbook.create_sheet("Purchase Requests"), request_columns, _purchase_request_rows(records["requests"])
    )

    approval_columns = [
        "Action Month", "Request Code", "Request Reporting Month", "Action", "From Status", "To Status", "Remarks",
        "Performed By", "Performed By Email", "Performed By Role", "System Recorded At",
    ]
    row_counts["Approval History"] = _write_rows(
        workbook.create_sheet("Approval History"), approval_columns, _approval_rows(records["request_histories"], requests_by_id)
    )

    _build_reconciliation_sheet(
        workbook,
        monthly_rows=monthly_rows,
        closing_assets=closing_assets,
        closing_month=closing_month,
        closing_source=closing_source,
        records=records,
        row_counts=row_counts,
    )

    for sheet in workbook.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_view.zoomScale = 85

    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream
