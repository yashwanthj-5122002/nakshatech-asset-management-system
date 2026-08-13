from __future__ import annotations

from calendar import monthrange
from functools import lru_cache
from datetime import date, datetime, timedelta, timezone
import json
from pathlib import Path
import re
from types import SimpleNamespace

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.services.asset_lifecycle_service import canonical_device_type, is_supported_device_type
from app.models.entities import Asset, MonthlyAssetSnapshot, MonthlySnapshotRun


SNAPSHOT_FIELDS = [
    "id", "asset_code", "source_sheet", "source_row", "used_by", "workstation_no", "department",
    "cpu_asset_tag", "monitor_asset_tags", "mouse_asset_tag", "keyboard_asset_tag", "system_name",
    "brand", "model", "serial_number", "connection_type", "capacity", "ownership",
    "client_name", "project_id", "current_holder", "device_type", "processor",
    "memory_gb", "ssd", "hdd", "ip_address", "mac_address",
    "graphics_card", "operating_system", "antivirus", "network_type", "performed_by", "approved_by",
    "price", "remarks", "asset_date", "original_asset_date", "location", "work_mode", "status", "created_at", "updated_at",
]
MONTH_FORMATS = ("%B %Y", "%b %Y", "%B-%Y", "%b-%Y")




def parse_template_sheet_month(sheet_name: str) -> date | None:
    cleaned = sheet_name.strip().title()
    for format_string in MONTH_FORMATS:
        try:
            return datetime.strptime(cleaned, format_string).date().replace(day=1)
        except ValueError:
            continue
    return None

def month_start(value: date | None = None) -> date:
    value = value or date.today()
    return value.replace(day=1)


def parse_month_key(month_key: str) -> date:
    try:
        parsed = datetime.strptime(month_key, "%Y-%m").date()
    except ValueError as exc:
        raise ValueError("Month must use YYYY-MM format") from exc
    return parsed.replace(day=1)


def month_end(start: date) -> date:
    return start.replace(day=monthrange(start.year, start.month)[1])


def previous_month(start: date) -> date:
    return (start - timedelta(days=1)).replace(day=1)


def next_month(start: date) -> date:
    end = month_end(start)
    return end + timedelta(days=1)


def _serialise_asset(asset: Asset) -> str:
    payload = {}
    for field in SNAPSHOT_FIELDS:
        value = getattr(asset, field, None)
        payload[field] = value.isoformat() if isinstance(value, (date, datetime)) else value
    return json.dumps(payload, ensure_ascii=False)


def _snapshot_namespace(payload: str) -> SimpleNamespace:
    values = json.loads(payload)
    for field in ("asset_date", "original_asset_date"):
        if values.get(field):
            values[field] = date.fromisoformat(values[field])
    for field in ("created_at", "updated_at"):
        if values.get(field):
            values[field] = datetime.fromisoformat(values[field])
    return SimpleNamespace(**values)


def get_snapshot_run(db: Session, start: date) -> MonthlySnapshotRun | None:
    return db.scalar(select(MonthlySnapshotRun).where(MonthlySnapshotRun.month_start == start))


def finalize_month_snapshot(
    db: Session,
    start: date,
    finalized_by: str,
    source: str = "manual",
    replace_existing: bool = False,
) -> MonthlySnapshotRun:
    start = start.replace(day=1)
    current = month_start()
    if start > current:
        raise ValueError("A future month cannot be finalized")

    existing = get_snapshot_run(db, start)
    if existing and not replace_existing:
        return existing
    if existing and replace_existing:
        db.delete(existing)
        db.flush()

    assets = list(db.scalars(select(Asset).order_by(Asset.id)).all())
    previous = get_snapshot_run(db, previous_month(start))
    opening_count = previous.closing_count if previous else len([asset for asset in assets if asset.source_sheet])
    run = MonthlySnapshotRun(
        month_start=start,
        status="finalized",
        source=source,
        opening_count=opening_count,
        closing_count=len(assets),
        finalized_by=finalized_by,
        finalized_at=datetime.now(timezone.utc).replace(tzinfo=None),
    )
    db.add(run)
    db.flush()
    for asset in assets:
        db.add(MonthlyAssetSnapshot(run_id=run.id, asset_code=asset.asset_code, payload=_serialise_asset(asset)))
    db.commit()
    db.refresh(run)
    return run


def ensure_previous_month_snapshot(db: Session) -> MonthlySnapshotRun | None:
    current = month_start()
    previous = previous_month(current)
    existing = get_snapshot_run(db, previous)
    if existing:
        return existing
    oldest_created = db.scalar(select(Asset.created_at).order_by(Asset.created_at).limit(1))
    if not oldest_created:
        return None
    # A fresh installation made during the current month must not invent a closing
    # snapshot for the month before the system existed. On the first startup after
    # a real month boundary, oldest_created will be earlier than current month.
    if oldest_created.date() >= current:
        return None
    return finalize_month_snapshot(db, previous, "System automatic month close", source="automatic")


def snapshot_assets(db: Session, start: date) -> list[SimpleNamespace]:
    run = get_snapshot_run(db, start)
    if not run:
        return []
    rows = db.scalars(
        select(MonthlyAssetSnapshot)
        .where(MonthlyAssetSnapshot.run_id == run.id)
        .order_by(MonthlyAssetSnapshot.id)
    ).all()
    return [_snapshot_namespace(row.payload) for row in rows]



def _cell_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return None if text in {"", "-"} else text


def _template_asset_status(used_by: str | None, remarks: str | None) -> str:
    note = (remarks or "").lower()
    for token, status in (
        ("replacement pending", "replacement_pending"),
        ("under repair", "repair"),
        ("not working", "repair"),
        ("returned", "returned"),
        ("missing", "missing"),
        ("damaged", "damaged"),
        ("retired", "retired"),
        ("disposed", "disposed"),
    ):
        if token in note:
            return status
    return "assigned" if used_by else "available"


@lru_cache(maxsize=48)
def template_assets(start: date) -> list[SimpleNamespace]:
    """Read one historical company workbook sheet into read-only asset-like rows.

    These rows are never written back to the database. Synthetic negative IDs make
    them safe to display in the historical Asset Register without colliding with live
    asset IDs.
    """
    template_path = Path(settings.seed_excel_path)
    if not template_path.exists():
        return []
    workbook = load_workbook(template_path, read_only=True, data_only=True, keep_links=False)
    sheet = next((item for item in workbook.worksheets if parse_template_sheet_month(item.title) == start), None)
    if sheet is None:
        workbook.close()
        return []

    rows: list[SimpleNamespace] = []
    recorded_at = datetime.combine(month_end(start), datetime.min.time())
    known_devices = {
        key: canonical_device_type(key)
        for key in (
            "computer", "desktop", "pc", "laptop", "notebook", "smartphone", "mobile",
            "printer", "server", "network device", "external hdd",
        )
        if is_supported_device_type(key)
    }
    for row_index, row in enumerate(sheet.iter_rows(min_row=2, max_row=min(sheet.max_row or 1000, 1000), max_col=25, values_only=True), start=2):
        values = list(row) + [None] * (25 - len(row))
        serial = _cell_text(values[0])
        cpu_tag = _cell_text(values[4])
        raw_device = (_cell_text(values[9]) or "").lower()
        device_type = known_devices.get(raw_device)
        is_serial_row = bool(serial and serial.replace(".", "", 1).isdigit())
        if not device_type and not (cpu_tag and is_serial_row):
            continue
        device_type = device_type or "Computer"
        used_by = _cell_text(values[1])
        if used_by and used_by.lower() in {"own", "na", "n/a"}:
            used_by = None
        workstation = _cell_text(values[2])
        remarks = _cell_text(values[23])
        work_hint = f"{workstation or ''} {remarks or ''}".lower()
        work_mode = "wfh" if "wfh" in work_hint or "work from home" in work_hint else ("field" if "field" in work_hint else "office")
        location = workstation if workstation and "floor" in workstation.lower() else "Head Office"
        asset_date = values[24]
        if isinstance(asset_date, datetime):
            asset_date = asset_date.date()
        if not isinstance(asset_date, date):
            asset_date = month_end(start)
        try:
            price = float(values[22]) if values[22] not in (None, "", "-") else None
        except (TypeError, ValueError):
            price = None
        synthetic_id = -(start.year * 1_000_000 + start.month * 10_000 + row_index)
        rows.append(SimpleNamespace(
            id=synthetic_id,
            asset_code=f"HIST-{start.strftime('%Y%m')}-{row_index:04d}",
            source_sheet=sheet.title,
            source_row=row_index,
            used_by=used_by,
            workstation_no=workstation,
            department=_cell_text(values[3]),
            cpu_asset_tag=cpu_tag,
            monitor_asset_tags=_cell_text(values[5]),
            mouse_asset_tag=_cell_text(values[6]),
            keyboard_asset_tag=_cell_text(values[7]),
            system_name=_cell_text(values[8]),
            brand=None,
            model=None,
            serial_number=None,
            connection_type=None,
            device_type=device_type,
            processor=_cell_text(values[10]),
            memory_gb=_cell_text(values[11]),
            ssd=_cell_text(values[12]),
            hdd=_cell_text(values[13]),
            ip_address=_cell_text(values[14]),
            mac_address=_cell_text(values[15]),
            graphics_card=_cell_text(values[16]),
            operating_system=_cell_text(values[17]),
            antivirus=_cell_text(values[18]),
            network_type=_cell_text(values[19]),
            performed_by=_cell_text(values[20]),
            approved_by=_cell_text(values[21]),
            price=price,
            remarks=remarks,
            asset_date=asset_date,
            original_asset_date=asset_date,
            location=location,
            work_mode=work_mode,
            status=_template_asset_status(used_by, remarks),
            created_at=recorded_at,
            updated_at=recorded_at,
        ))
    workbook.close()
    return rows


@lru_cache(maxsize=1)
def template_months() -> list[dict]:
    template_path = Path(settings.seed_excel_path)
    if not template_path.exists():
        return []
    workbook = load_workbook(template_path, read_only=True, data_only=True)
    result = []
    for sheet_name in workbook.sheetnames:
        parsed = parse_template_sheet_month(sheet_name)
        if not parsed:
            continue
        result.append({
            "key": parsed.strftime("%Y-%m"),
            "label": parsed.strftime("%B %Y"),
            "status": "historical",
            "source": "original_excel",
            "sheet_name": sheet_name,
            "is_current": parsed == month_start(),
        })
    workbook.close()
    return result


def available_months(db: Session) -> list[dict]:
    ensure_previous_month_snapshot(db)
    merged: dict[str, dict] = {item["key"]: item for item in template_months()}
    for run in db.scalars(select(MonthlySnapshotRun).order_by(MonthlySnapshotRun.month_start.desc())).all():
        key = run.month_start.strftime("%Y-%m")
        merged[key] = {
            "key": key,
            "label": run.month_start.strftime("%B %Y"),
            "status": "finalized",
            "source": "system_snapshot",
            "opening_count": run.opening_count,
            "closing_count": run.closing_count,
            "is_current": run.month_start == month_start(),
        }
    current = month_start()
    current_key = current.strftime("%Y-%m")
    if current_key not in merged or merged[current_key].get("source") == "original_excel":
        merged[current_key] = {
            "key": current_key,
            "label": current.strftime("%B %Y"),
            "status": "live",
            "source": "live_register",
            "is_current": True,
        }
    if current_key in merged:
        merged[current_key]["is_current"] = True
    return sorted(merged.values(), key=lambda item: item["key"], reverse=True)


def assets_for_month(db: Session, start: date):
    current = month_start()
    # The present register must stay live even when an administrator has created a
    # reference snapshot for the same month.
    if start == current:
        return list(db.scalars(select(Asset).order_by(Asset.id)).all()), "live"
    snapshots = snapshot_assets(db, start)
    if snapshots:
        return snapshots, "snapshot"
    historical = template_assets(start)
    if historical:
        return historical, "template"
    return [], "missing"
