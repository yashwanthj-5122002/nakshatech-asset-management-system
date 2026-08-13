from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import json
import re

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.entities import Asset, AssetHistory
from app.services.asset_lifecycle_service import canonical_device_type, is_supported_device_type


EXTERNAL_HDD_HEADERS = [
    "Asset ID",
    "Brand",
    "Capacity",
    "Serial No.",
    "Ownership",
    "Department",
    "Client Name",
    "Project ID",
    "Current Holder",
    "Status",
    "Remarks",
]

PRINTER_HEADERS = [
    "Asset ID",
    "Assigned User",
    "Brand",
    "Model",
    "Serial No.",
    "Connection",
    "Department",
    "Floor",
    "Status",
    "Remarks",
    "Last Updated",
]

DEPARTMENT_MAP = {
    "ldr": "LiDAR",
    "lidar": "LiDAR",
    "lider": "LiDAR",
    "mapping": "Mapping",
    "ortho": "Ortho / GIS",
    "orthophoto": "Ortho / GIS",
    "ortho/gis": "Ortho / GIS",
    "pg": "Photogrammetry",
    "photogramettry": "Photogrammetry",
    "photogrammetry": "Photogrammetry",
    "bim": "BIM",
    "civil": "Civil",
    "drone": "Drone",
    "drone assembly": "Drone Assembly",
    "it": "IT",
    "management": "Management",
    "hr-management": "HR / Management",
    "hr- management": "HR / Management",
    "finance": "Finance",
    "hr": "HR",
    "admin": "Admin",
    "bd": "Business Development",
    "gis cordinator": "GIS Coordination",
}


def _text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not text or text.lower() in {"none", "nan"}:
        return None
    return text


def _clean_optional(value) -> str | None:
    text = _text(value)
    if text in {"-", "--", "N/a", "N/A", "n/a"}:
        return None
    return text


def _normalise_department(value) -> str | None:
    text = _clean_optional(value)
    if not text:
        return None
    key = re.sub(r"\s+", " ", text).strip().lower()
    return DEPARTMENT_MAP.get(key, text.title())


def _to_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _device_type(value) -> str | None:
    text = _clean_optional(value)
    if not text or not is_supported_device_type(text):
        return None
    return canonical_device_type(text)


def _work_mode(workstation: str | None, remarks: str | None) -> str:
    combined = f"{workstation or ''} {remarks or ''}".lower()
    if "wfh" in combined or "work from home" in combined:
        return "wfh"
    if "field" in combined or "feld" in combined:
        return "field"
    return "office"


def _location(workstation: str | None, work_mode: str) -> str:
    if work_mode == "wfh":
        return "Work From Home"
    if work_mode == "field":
        return "Field"
    text = (workstation or "").strip()
    if "floor" in text.lower() or text.lower() == "office":
        return text.title()
    return "Head Office"


def _status(used_by: str | None, work_mode: str, remarks: str | None) -> str:
    note = (remarks or "").lower()
    if "replacement pending" in note or "replace pending" in note:
        return "replacement_pending"
    if "missing" in note:
        return "missing"
    if "destroy" in note or "beyond repair" in note:
        return "beyond_repair"
    if "repair" in note or "not working" in note or "slow working" in note or "problem" in note:
        return "repair"
    if "returned" in note or "return to" in note or "returnrd" in note:
        return "returned"
    if work_mode == "wfh":
        return "wfh"
    if work_mode == "field":
        return "field_deployment"
    if not used_by or used_by == "-":
        return "available"
    return "assigned"


def _next_asset_number(db: Session, prefix: str) -> int:
    existing = db.scalars(select(Asset.asset_code).where(Asset.asset_code.like(f"{prefix}-%"))).all()
    numbers = []
    for code in existing:
        match = re.search(r"(\d+)$", code)
        if match:
            numbers.append(int(match.group(1)))
    return max(numbers, default=0) + 1


def _code_prefix(device_type: str) -> str:
    return {"Computer": "NT-PC", "Laptop": "NT-LAP", "Smartphone": "NT-MOB", "Printer": "NT-PRN", "External HDD": "NT-HDD"}.get(device_type, "NT-IT")


def _latest_sheet_name(workbook) -> str:
    month_numbers = {
        "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
        "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
    }
    ranked: list[tuple[int, int, str]] = []
    for name in workbook.sheetnames:
        match = re.search(r"([A-Za-z]+)[ -](\d{4})", name)
        if not match:
            continue
        month = month_numbers.get(match.group(1)[:3].lower())
        if month:
            ranked.append((int(match.group(2)), month, name))
    return max(ranked)[2] if ranked else workbook.sheetnames[0]


def import_nakshatech_workbook(db: Session, source: bytes | str | Path) -> dict:
    workbook = load_workbook(BytesIO(source) if isinstance(source, bytes) else source, data_only=True)
    sheet_name = _latest_sheet_name(workbook)
    ws = workbook[sheet_name]

    counters = {
        "NT-PC": _next_asset_number(db, "NT-PC"),
        "NT-LAP": _next_asset_number(db, "NT-LAP"),
        "NT-MOB": _next_asset_number(db, "NT-MOB"),
        "NT-IT": _next_asset_number(db, "NT-IT"),
    }
    result = {"sheet": sheet_name, "created": 0, "updated": 0, "skipped": 0, "errors": []}

    for row_number in range(2, ws.max_row + 1):
        values = [ws.cell(row_number, column).value for column in range(1, 26)]
        device_type = _device_type(values[9])
        if not device_type:
            continue

        used_by = _clean_optional(values[1])
        workstation = _clean_optional(values[2])
        remarks = _clean_optional(values[23])
        cpu_tag = _clean_optional(values[4])
        system_name = _clean_optional(values[8])
        work_mode = _work_mode(workstation, remarks)

        existing = None
        if cpu_tag and cpu_tag.upper() != "OWN":
            existing = db.scalar(select(Asset).where(Asset.cpu_asset_tag == cpu_tag))
        if not existing and system_name:
            existing = db.scalar(select(Asset).where(Asset.system_name == system_name, Asset.device_type == device_type))

        payload = dict(
            source_sheet=sheet_name,
            source_row=row_number,
            used_by=used_by,
            workstation_no=workstation,
            department=_normalise_department(values[3]),
            cpu_asset_tag=cpu_tag,
            monitor_asset_tags=_clean_optional(values[5]),
            mouse_asset_tag=_clean_optional(values[6]),
            keyboard_asset_tag=_clean_optional(values[7]),
            system_name=system_name,
            device_type=device_type,
            processor=_clean_optional(values[10]),
            memory_gb=_clean_optional(values[11]),
            ssd=_clean_optional(values[12]),
            hdd=_clean_optional(values[13]),
            ip_address=_clean_optional(values[14]),
            mac_address=_clean_optional(values[15]),
            graphics_card=_clean_optional(values[16]),
            operating_system=_clean_optional(values[17]),
            antivirus=_clean_optional(values[18]),
            network_type=_clean_optional(values[19]),
            performed_by=_clean_optional(values[20]),
            approved_by=_clean_optional(values[21]),
            price=float(values[22]) if isinstance(values[22], (int, float)) else None,
            remarks=remarks,
            asset_date=_to_date(values[24]),
            location=_location(workstation, work_mode),
            work_mode=work_mode,
            status=_status(used_by, work_mode, remarks),
        )

        if existing:
            if existing.original_asset_date is None:
                existing.original_asset_date = existing.asset_date or payload.get("asset_date")
            for key, value in payload.items():
                setattr(existing, key, value)
            result["updated"] += 1
        else:
            prefix = _code_prefix(device_type)
            asset_code = f"{prefix}-{counters[prefix]:04d}"
            counters[prefix] += 1
            db.add(Asset(asset_code=asset_code, original_asset_date=payload.get("asset_date"), **payload))
            result["created"] += 1

    db.commit()
    return result


def import_assets_workbook(db: Session, content: bytes) -> dict:
    return import_nakshatech_workbook(db, content)


def _printer_status(value, used_by: str | None, department: str | None) -> str:
    text = (_clean_optional(value) or "active").strip().lower().replace("-", " ").replace("_", " ")
    if text in {"repair", "under repair", "service", "under service"}:
        return "repair"
    if text in {"replacement pending", "replace pending"}:
        return "replacement_pending"
    if text in {"retired", "inactive"}:
        return "retired"
    if text in {"disposed", "scrapped"}:
        return "disposed"
    if text in {"available", "stock", "in stock"}:
        return "available"
    if text in {"active", "assigned", "in use", "working"}:
        return "assigned" if used_by or department else "available"
    return "assigned" if used_by or department else "available"


def _header_map(ws) -> dict[str, int]:
    result: dict[str, int] = {}
    for column in range(1, ws.max_column + 1):
        value = _text(ws.cell(1, column).value)
        if value:
            result[value.casefold()] = column
    return result


def import_printer_assets_workbook(
    db: Session,
    content: bytes,
    *,
    performed_by: str | None = None,
) -> dict:
    workbook = load_workbook(BytesIO(content), data_only=True, keep_links=False)
    ws = workbook["Printer Assets"] if "Printer Assets" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    headers = _header_map(ws)
    missing = [header for header in PRINTER_HEADERS if header.casefold() not in headers]
    if missing:
        workbook.close()
        raise ValueError(f"Printer workbook is missing required columns: {', '.join(missing)}")

    result = {
        "sheet": ws.title,
        "created": 0,
        "updated": 0,
        "skipped": 0,
        "warnings": [],
        "errors": [],
    }
    next_number = _next_asset_number(db, "NT-PRN")

    def cell(row: int, header: str):
        return ws.cell(row, headers[header.casefold()]).value

    for row_number in range(2, ws.max_row + 1):
        asset_id = _clean_optional(cell(row_number, "Asset ID"))
        assigned_user = _clean_optional(cell(row_number, "Assigned User"))
        brand = _clean_optional(cell(row_number, "Brand"))
        model = _clean_optional(cell(row_number, "Model"))
        serial_number = _clean_optional(cell(row_number, "Serial No."))
        connection_type = _clean_optional(cell(row_number, "Connection"))
        department = _normalise_department(cell(row_number, "Department"))
        floor = _clean_optional(cell(row_number, "Floor"))
        remarks = _clean_optional(cell(row_number, "Remarks"))
        last_updated = _to_date(cell(row_number, "Last Updated"))

        if not any((asset_id, assigned_user, brand, model, serial_number, connection_type, department, floor, remarks, last_updated)):
            continue
        if not asset_id or not brand or not model:
            result["skipped"] += 1
            result["errors"].append({
                "row": row_number,
                "message": "Asset ID, Brand and Model are required.",
            })
            continue

        existing = db.scalar(select(Asset).where(func.lower(Asset.cpu_asset_tag) == asset_id.lower()))
        if existing and existing.device_type != "Printer":
            result["skipped"] += 1
            result["errors"].append({
                "row": row_number,
                "asset_id": asset_id,
                "message": f"Asset ID already belongs to {existing.device_type}.",
            })
            continue

        if serial_number:
            serial_owner = db.scalar(
                select(Asset).where(
                    func.lower(Asset.serial_number) == serial_number.lower(),
                    Asset.id != (existing.id if existing else -1),
                )
            )
            if serial_owner:
                result["skipped"] += 1
                result["errors"].append({
                    "row": row_number,
                    "asset_id": asset_id,
                    "message": f"Serial number already belongs to {serial_owner.cpu_asset_tag or serial_owner.asset_code}.",
                })
                continue
        else:
            result["warnings"].append({
                "row": row_number,
                "asset_id": asset_id,
                "message": "Serial number is blank; record imported with a warning.",
            })

        status = _printer_status(cell(row_number, "Status"), assigned_user, department)
        payload = {
            "source_sheet": ws.title,
            "source_row": row_number,
            "used_by": assigned_user,
            "workstation_no": None,
            "department": department,
            "cpu_asset_tag": asset_id,
            "system_name": None,
            "brand": brand,
            "model": model,
            "serial_number": serial_number,
            "connection_type": connection_type,
            "device_type": "Printer",
            "performed_by": performed_by or "Printer Excel Import",
            "remarks": remarks,
            "asset_date": last_updated or date.today(),
            "location": floor,
            "work_mode": "office",
            "status": status,
        }

        if existing:
            if existing.original_asset_date is None:
                existing.original_asset_date = existing.asset_date or payload["asset_date"]
            for key, value in payload.items():
                setattr(existing, key, value)
            result["updated"] += 1
        else:
            asset_code = f"NT-PRN-{next_number:04d}"
            next_number += 1
            db.add(Asset(asset_code=asset_code, original_asset_date=payload["asset_date"], **payload))
            result["created"] += 1

    db.commit()
    workbook.close()
    result["warning_count"] = len(result["warnings"])
    result["error_count"] = len(result["errors"])
    return result


def _external_hdd_ownership(value) -> str | None:
    text = (_clean_optional(value) or "").strip().casefold()
    if text in {"nakshatech", "naksha tech", "company", "company owned", "internal"}:
        return "NakshaTech"
    if text in {"client", "client owned", "customer", "customer owned"}:
        return "Client"
    return None


def _external_hdd_status(value) -> str | None:
    text = (_clean_optional(value) or "").strip().lower().replace("-", " ").replace("_", " ")
    mapping = {
        "in use": "in_use",
        "active": "in_use",
        "issued": "issued",
        "temporarily issued": "issued",
        "permanently issued": "permanently_issued",
        "permanent issued": "permanently_issued",
        "returned": "returned",
        "available": "available",
        "stock": "available",
        "under repair": "repair",
        "repair": "repair",
        "replacement pending": "replacement_pending",
        "retired": "retired",
        "disposed": "disposed",
        "missing": "missing",
    }
    return mapping.get(text)


def import_external_hdd_assets_workbook(
    db: Session,
    content: bytes,
    *,
    performed_by: str | None = None,
) -> dict:
    workbook = load_workbook(BytesIO(content), data_only=True, keep_links=False)
    sheet_name = "External HDD Asset Register"
    ws = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    headers = _header_map(ws)
    missing = [header for header in EXTERNAL_HDD_HEADERS if header.casefold() not in headers]
    if missing:
        workbook.close()
        raise ValueError(f"External HDD workbook is missing required columns: {', '.join(missing)}")

    result = {
        "sheet": ws.title,
        "created": 0,
        "updated": 0,
        "skipped": 0,
        "warnings": [],
        "errors": [],
    }
    next_number = _next_asset_number(db, "NT-HDD")

    def cell(row: int, header: str):
        return ws.cell(row, headers[header.casefold()]).value

    for row_number in range(2, ws.max_row + 1):
        asset_id = _clean_optional(cell(row_number, "Asset ID"))
        brand = _clean_optional(cell(row_number, "Brand"))
        capacity = _clean_optional(cell(row_number, "Capacity"))
        serial_number = _clean_optional(cell(row_number, "Serial No."))
        ownership = _external_hdd_ownership(cell(row_number, "Ownership"))
        department = _normalise_department(cell(row_number, "Department"))
        client_name = _clean_optional(cell(row_number, "Client Name"))
        project_id = _clean_optional(cell(row_number, "Project ID"))
        current_holder = _clean_optional(cell(row_number, "Current Holder"))
        status = _external_hdd_status(cell(row_number, "Status"))
        remarks = _clean_optional(cell(row_number, "Remarks"))

        if not any((asset_id, brand, capacity, serial_number, ownership, department, client_name, project_id, current_holder, status, remarks)):
            continue

        required = []
        if not asset_id:
            required.append("Asset ID")
        if not brand:
            required.append("Brand")
        if not capacity:
            required.append("Capacity")
        if not serial_number:
            required.append("Serial No.")
        if not ownership:
            required.append("Ownership")
        if not status:
            required.append("Status")
        if required:
            result["skipped"] += 1
            result["errors"].append({
                "row": row_number,
                "asset_id": asset_id,
                "message": f"Required or invalid values: {', '.join(required)}.",
            })
            continue
        if ownership == "Client" and not client_name:
            result["skipped"] += 1
            result["errors"].append({
                "row": row_number,
                "asset_id": asset_id,
                "message": "Client Name is required for a client-owned External HDD.",
            })
            continue
        if status in {"in_use", "issued", "permanently_issued"} and not current_holder:
            result["skipped"] += 1
            result["errors"].append({
                "row": row_number,
                "asset_id": asset_id,
                "message": "Current Holder is required for an in-use or issued External HDD.",
            })
            continue

        existing = db.scalar(select(Asset).where(func.lower(Asset.cpu_asset_tag) == asset_id.lower()))
        if existing and existing.device_type != "External HDD":
            result["skipped"] += 1
            result["errors"].append({
                "row": row_number,
                "asset_id": asset_id,
                "message": f"Asset ID already belongs to {existing.device_type}.",
            })
            continue

        serial_owner = db.scalar(
            select(Asset).where(
                func.lower(Asset.serial_number) == serial_number.lower(),
                Asset.id != (existing.id if existing else -1),
            )
        )
        if serial_owner:
            result["skipped"] += 1
            result["errors"].append({
                "row": row_number,
                "asset_id": asset_id,
                "message": f"Serial number already belongs to {serial_owner.cpu_asset_tag or serial_owner.asset_code}.",
            })
            continue

        payload = {
            "source_sheet": ws.title,
            "source_row": row_number,
            "used_by": None,
            "workstation_no": None,
            "department": department,
            "cpu_asset_tag": asset_id,
            "system_name": None,
            "brand": brand,
            "model": None,
            "serial_number": serial_number,
            "connection_type": None,
            "capacity": capacity,
            "ownership": ownership,
            "client_name": client_name,
            "project_id": project_id,
            "current_holder": current_holder,
            "device_type": "External HDD",
            "performed_by": performed_by or "External HDD Excel Import",
            "remarks": remarks,
            "asset_date": date.today(),
            "location": current_holder or department,
            "work_mode": "office",
            "status": status,
        }

        if existing:
            before = {key: getattr(existing, key, None) for key in payload}
            changed = {key: {"from": before[key], "to": value} for key, value in payload.items() if before[key] != value}
            if existing.original_asset_date is None:
                existing.original_asset_date = existing.asset_date or payload["asset_date"]
            for key, value in payload.items():
                setattr(existing, key, value)
            if changed:
                db.add(AssetHistory(
                    asset_id=existing.id,
                    action="External HDD updated from Excel",
                    change_type="external_hdd_import_update",
                    old_value=json.dumps({key: value["from"] for key, value in changed.items()}, default=str),
                    new_value=json.dumps({key: value["to"] for key, value in changed.items()}, default=str),
                    remarks=f"Updated from {ws.title} row {row_number}",
                    reason="External HDD Excel import",
                    changed_by_name=performed_by or "External HDD Excel Import",
                    field_count=len(changed),
                ))
            result["updated"] += 1
        else:
            asset_code = f"NT-HDD-{next_number:04d}"
            next_number += 1
            asset = Asset(asset_code=asset_code, original_asset_date=payload["asset_date"], **payload)
            db.add(asset)
            db.flush()
            db.add(AssetHistory(
                asset_id=asset.id,
                action="External HDD created from Excel",
                change_type="asset_created",
                new_value=json.dumps({"asset_code": asset_code, **payload}, default=str),
                remarks=f"Imported from {ws.title} row {row_number}",
                reason="External HDD Excel import",
                changed_by_name=performed_by or "External HDD Excel Import",
                field_count=len(payload),
            ))
            result["created"] += 1

    db.commit()
    workbook.close()
    result["warning_count"] = len(result["warnings"])
    result["error_count"] = len(result["errors"])
    return result
