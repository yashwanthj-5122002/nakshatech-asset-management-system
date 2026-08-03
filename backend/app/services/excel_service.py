from __future__ import annotations

from collections import Counter
from copy import copy
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import json
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.services.monthly_snapshot_service import assets_for_month, get_snapshot_run, month_end, month_start, parse_month_key, parse_template_sheet_month
from app.models.entities import Asset, AssetHistory, ComponentReplacement, MonthlySnapshotRun, ReplacementRecord, User, WorkRecord


NAVY = "082B57"
CYAN = "08A8C0"
LIGHT = "EAF4F8"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(color="FFFFFF", bold=True)

IST = ZoneInfo("Asia/Kolkata")
FIELD_LABELS = {
    "asset_code": "Internal Asset ID",
    "used_by": "Used By",
    "workstation_no": "Workstation Number",
    "department": "Department",
    "cpu_asset_tag": "CPU / Asset Tag",
    "monitor_asset_tags": "Monitor Asset Tag(s)",
    "mouse_asset_tag": "Mouse Asset Tag",
    "keyboard_asset_tag": "Keyboard Asset Tag",
    "system_name": "System Name",
    "device_type": "Device Type",
    "processor": "Processor",
    "memory_gb": "Memory",
    "ssd": "SSD",
    "hdd": "HDD",
    "ip_address": "IP Address",
    "mac_address": "MAC Address",
    "graphics_card": "Graphics Card",
    "operating_system": "Operating System",
    "antivirus": "Antivirus",
    "network_type": "Network Type",
    "approved_by": "Approved By",
    "price": "Price",
    "remarks": "Remarks",
    "asset_date": "Original Asset Date",
    "location": "Location",
    "work_mode": "Work Mode",
    "status": "Status",
}
LIVE_ASSET_HEADERS = [
    "SL", "USED BY", "WS No", "DEPARTMENT", "cpu", "MONITOR", "MOUSE", "KB",
    "SYSTEM NAME", "COMPUTER / LAPTOP/Smartphone", "PROCESSOR", "MEMORY IN GB",
    "SSD", "HDD", "IP ADDRESS", "MAC ADDRESS", "GC", "OS", "ANTIVIRUS", "DHCP",
    "PERFORMED BY", "APPROVED BY", "PRICE", "REMARKS", "date",
    "ORIGINAL ASSET DATE", "LAST UPDATED TIME", "UPDATED BY", "UPDATED BY ROLE",
    "LAST CHANGE TYPE", "FIELDS CHANGED", "EDIT REASON", "CHANGE BATCH ID",
]
LIVE_ASSET_WIDTHS = [
    7, 22, 14, 20, 14, 25, 14, 14, 22, 24, 22, 16, 14, 14, 18, 20, 22, 20,
    20, 14, 20, 20, 14, 34, 18, 18, 18, 24, 20, 26, 38, 40, 24,
]


def _local_datetime(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    aware = value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
    return aware.astimezone(IST)


def _json_dict(raw: str | None) -> dict:
    if not raw:
        return {}
    try:
        value = json.loads(raw)
    except (TypeError, ValueError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def _history_fields(history: AssetHistory | None) -> list[str]:
    if history is None:
        return []
    old_values = _json_dict(history.old_value)
    new_values = _json_dict(history.new_value)
    return sorted(set(old_values) | set(new_values))


def _latest_history_map(db: Session) -> dict[int, AssetHistory]:
    latest: dict[int, AssetHistory] = {}
    rows = db.scalars(select(AssetHistory).order_by(AssetHistory.created_at, AssetHistory.id)).all()
    for row in rows:
        latest[row.asset_id] = row
    return latest


def _user_maps(db: Session) -> tuple[dict[str, User], dict[str, User]]:
    by_email: dict[str, User] = {}
    by_name: dict[str, User] = {}
    for user in db.scalars(select(User)).all():
        by_email[user.email.lower()] = user
        by_name[user.full_name.strip().lower()] = user
    return by_email, by_name


def _history_actor(history: AssetHistory, users_by_email: dict[str, User]) -> tuple[str | None, str | None, str | None]:
    fallback = users_by_email.get((history.changed_by or "").lower())
    name = history.changed_by_name or (fallback.full_name if fallback else None)
    role = history.changed_by_role or (fallback.role if fallback else None)
    return name, history.changed_by, role


def _component_actor(record: ComponentReplacement, users_by_email: dict[str, User], users_by_name: dict[str, User]) -> tuple[str | None, str | None, str | None]:
    fallback = users_by_email.get((record.performed_by_email or "").lower())
    if fallback is None:
        fallback = users_by_name.get((record.performed_by or "").strip().lower())
    name = record.performed_by or (fallback.full_name if fallback else None)
    email = record.performed_by_email or (fallback.email if fallback else None)
    role = record.performed_by_role or (fallback.role if fallback else None)
    return name, email, role



def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        if cell.value is not None:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _copy_row_style(ws, source_row: int, target_row: int, max_column: int = 25) -> None:
    for column in range(1, max_column + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _excel_value(value):
    if value is None:
        return None
    return value


def _asset_row(asset: Asset, serial: int, latest_history: AssetHistory | None = None) -> list:
    remarks = asset.remarks or ""
    if asset.status not in {"assigned", "available", "wfh", "field_deployment"}:
        status_text = asset.status.replace("_", " ").title()
        if status_text.lower() not in remarks.lower():
            remarks = f"{remarks} | System Status: {status_text}".strip(" |")

    changed_at = _local_datetime(latest_history.created_at) if latest_history else None
    changed_fields = _history_fields(latest_history)
    latest_date = changed_at.date() if changed_at else asset.asset_date
    latest_time = changed_at.strftime("%I:%M:%S %p") if changed_at else None
    updated_by = None
    updated_role = None
    change_type = None
    reason = None
    batch_code = None
    if latest_history:
        updated_by = latest_history.changed_by_name or latest_history.changed_by
        updated_role = latest_history.changed_by_role
        change_type = (latest_history.change_type or latest_history.action).replace("_", " ").title()
        reason = latest_history.reason or latest_history.remarks
        batch_code = latest_history.batch_code

    return [
        serial,
        asset.used_by,
        asset.workstation_no,
        asset.department,
        asset.cpu_asset_tag,
        asset.monitor_asset_tags,
        asset.mouse_asset_tag,
        asset.keyboard_asset_tag,
        asset.system_name,
        asset.device_type,
        asset.processor,
        asset.memory_gb,
        asset.ssd,
        asset.hdd,
        asset.ip_address,
        asset.mac_address,
        asset.graphics_card,
        asset.operating_system,
        asset.antivirus,
        asset.network_type,
        asset.performed_by,
        asset.approved_by,
        asset.price or 0,
        remarks or None,
        latest_date,
        getattr(asset, "original_asset_date", None) or asset.asset_date,
        latest_time,
        updated_by or asset.performed_by,
        updated_role,
        change_type,
        ", ".join(FIELD_LABELS.get(field, field.replace("_", " ").title()) for field in changed_fields) or None,
        reason,
        batch_code,
    ]


def _latest_sheet_name(workbook) -> str:
    return workbook.sheetnames[0]


def build_asset_report(db: Session) -> BytesIO:
    """Build the live master workbook with exact audit date/time and user details.

    The original 25 company columns remain in the same order. The existing ``date``
    column now represents the latest successful system change date. The original
    asset date is preserved in a separate audit column so no source information is
    lost. Historical workbook sheets remain unchanged.
    """
    workbook = Workbook()
    ws = workbook.active
    current_sheet_name = datetime.now(IST).strftime("%B %Y")
    ws.title = current_sheet_name
    ws.append(LIVE_ASSET_HEADERS)
    yellow_fill = PatternFill("solid", fgColor="FFF200")
    for cell in ws[1]:
        cell.fill = yellow_fill
        cell.font = Font(color="000000", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    assets = list(db.scalars(select(Asset).order_by(Asset.id)).all())
    latest_history = _latest_history_map(db)
    office_assets = [asset for asset in assets if asset.work_mode == "office"]
    remote_assets = [asset for asset in assets if asset.work_mode in {"wfh", "field"}]
    serial = 1
    for asset in office_assets:
        ws.append(_asset_row(asset, serial, latest_history.get(asset.id)))
        serial += 1
    ws.append([])
    ws.append(["Work From Home Systems and laptops"])
    heading_row = ws.max_row
    ws.merge_cells(start_row=heading_row, start_column=1, end_row=heading_row, end_column=len(LIVE_ASSET_HEADERS))
    ws.cell(heading_row, 1).fill = HEADER_FILL
    ws.cell(heading_row, 1).font = HEADER_FONT
    serial = 1
    for asset in remote_assets:
        ws.append(_asset_row(asset, serial, latest_history.get(asset.id)))
        serial += 1

    ws.append([])
    counts = Counter(asset.device_type for asset in assets)
    ws.append([None, "Current Inventory Summary", "Computer", "Laptop", "Mobile"])
    ws.append([None, "NakshaTech", counts.get("Computer", 0), counts.get("Laptop", 0), counts.get("Smartphone", 0)])
    ws.append([None, "Total", counts.get("Computer", 0), counts.get("Laptop", 0), counts.get("Smartphone", 0)])
    ws.freeze_panes = "A2"
    last_column = get_column_letter(len(LIVE_ASSET_HEADERS))
    ws.auto_filter.ref = f"A1:{last_column}{max(len(office_assets) + 1, 2)}"
    ws.sheet_view.showGridLines = False
    for index, width in enumerate(LIVE_ASSET_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=25, max_col=26):
        for cell in row:
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "DD-MM-YYYY"

    template_path = Path(settings.seed_excel_path)
    if template_path.exists():
        source_book = load_workbook(template_path, read_only=True, data_only=False, keep_links=False)
        historical_widths = LIVE_ASSET_WIDTHS[:25]
        for source_sheet in source_book.worksheets:
            if source_sheet.title == current_sheet_name:
                continue
            target = workbook.create_sheet(source_sheet.title)
            last_nonempty = 0
            cached_rows: list[tuple] = []
            for row_index, values in enumerate(source_sheet.iter_rows(values_only=True), start=1):
                values = tuple(values[:25])
                cached_rows.append(values)
                if any(value not in (None, "") for value in values):
                    last_nonempty = row_index
            for values in cached_rows[:last_nonempty]:
                target.append(values)
            if target.max_row:
                for cell in target[1]:
                    if cell.value is not None:
                        cell.fill = yellow_fill
                        cell.font = Font(color="000000", bold=True)
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            target.freeze_panes = "A2"
            target.sheet_view.showGridLines = False
            for index, width in enumerate(historical_widths, start=1):
                target.column_dimensions[get_column_letter(index)].width = width
        source_book.close()

    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream

def build_dashboard_report(db: Session) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Dashboard Summary"
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:H2")
    summary["A1"] = "NAKSHATECH IT ASSET MANAGEMENT DASHBOARD"
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")

    assets = list(db.scalars(select(Asset)).all())
    works = list(db.scalars(select(WorkRecord).where(WorkRecord.module == "it")).all())
    status_counts = Counter(asset.status for asset in assets)
    device_counts = Counter(asset.device_type for asset in assets)
    department_counts = Counter(asset.department or "Unassigned" for asset in assets)

    kpis = [
        ("Total Assets", len(assets)),
        ("Computers", device_counts.get("Computer", 0)),
        ("Laptops", device_counts.get("Laptop", 0)),
        ("Smartphones", device_counts.get("Smartphone", 0)),
        ("Assigned / In Use", status_counts.get("assigned", 0)),
        ("Available", status_counts.get("available", 0)),
        ("Under Repair", status_counts.get("repair", 0)),
        ("Replacement Pending", status_counts.get("replacement_pending", 0)),
    ]
    summary.append([])
    summary.append(["KPI", "Value"])
    _style_header(summary, 4)
    for label, value in kpis:
        summary.append([label, value])

    asset_sheet = workbook.create_sheet("Asset Register")
    asset_sheet.append([
        "Asset ID", "Used By", "Workstation", "Department", "Device Type", "CPU Asset Tag",
        "Monitor Asset Tags", "Mouse Asset Tag", "Keyboard Asset Tag", "System Name", "Processor",
        "Memory", "SSD", "HDD", "IP Address", "MAC Address", "Graphics Card", "Operating System",
        "Antivirus", "Network Type", "Work Mode", "Location", "Status", "Remarks", "Asset Date",
    ])
    _style_header(asset_sheet)
    for asset in sorted(assets, key=lambda item: item.asset_code):
        asset_sheet.append([
            asset.asset_code, asset.used_by, asset.workstation_no, asset.department, asset.device_type,
            asset.cpu_asset_tag, asset.monitor_asset_tags, asset.mouse_asset_tag, asset.keyboard_asset_tag,
            asset.system_name, asset.processor, asset.memory_gb, asset.ssd, asset.hdd, asset.ip_address,
            asset.mac_address, asset.graphics_card, asset.operating_system, asset.antivirus,
            asset.network_type, asset.work_mode, asset.location, asset.status.replace("_", " ").title(),
            asset.remarks, asset.asset_date,
        ])
    asset_sheet.auto_filter.ref = f"A1:Y{asset_sheet.max_row}"

    status_sheet = workbook.create_sheet("Asset Status")
    status_sheet.append(["Status", "Count"])
    _style_header(status_sheet)
    for status, count in sorted(status_counts.items()):
        status_sheet.append([status.replace("_", " ").title(), count])
    pie = PieChart()
    pie.title = "Asset Status Distribution"
    pie.add_data(Reference(status_sheet, min_col=2, min_row=1, max_row=status_sheet.max_row), titles_from_data=True)
    pie.set_categories(Reference(status_sheet, min_col=1, min_row=2, max_row=status_sheet.max_row))
    status_sheet.add_chart(pie, "D2")

    department_sheet = workbook.create_sheet("Department Summary")
    department_sheet.append(["Department", "Asset Count"])
    _style_header(department_sheet)
    for department, count in department_counts.most_common():
        department_sheet.append([department, count])
    bar = BarChart()
    bar.type = "bar"
    bar.title = "Assets by Department"
    bar.add_data(Reference(department_sheet, min_col=2, min_row=1, max_row=department_sheet.max_row), titles_from_data=True)
    bar.set_categories(Reference(department_sheet, min_col=1, min_row=2, max_row=department_sheet.max_row))
    department_sheet.add_chart(bar, "D2")

    work_sheet = workbook.create_sheet("Work Records")
    work_sheet.append(["Work ID", "Asset", "Title", "Type", "Technician", "Priority", "Status", "Approval", "Created"])
    _style_header(work_sheet)
    asset_by_id = {asset.id: asset.asset_code for asset in assets}
    for work in works:
        work_sheet.append([
            work.work_code,
            asset_by_id.get(work.asset_id),
            work.title,
            work.work_type,
            work.technician,
            work.priority,
            work.status,
            work.approval_status,
            work.created_at,
        ])

    data_quality_sheet = workbook.create_sheet("Data Quality")
    data_quality_sheet.append(["Check", "Count", "Details"])
    _style_header(data_quality_sheet)
    ip_counts = Counter(asset.ip_address for asset in assets if asset.ip_address and asset.ip_address not in {"-", "Dynamic"})
    duplicate_ips = sorted(ip for ip, count in ip_counts.items() if count > 1)
    checks = [
        ("Missing employee assignment", sum(not asset.used_by for asset in assets), "Complete employee allocation"),
        ("Missing MAC address", sum(not asset.mac_address for asset in assets), "Capture device MAC address"),
        ("Missing IP address", sum(not asset.ip_address for asset in assets), "Capture IP or mark DHCP"),
        ("Missing asset date", sum(not asset.asset_date for asset in assets), "Verify purchase or assignment date"),
        ("Duplicate IP addresses", len(duplicate_ips), ", ".join(duplicate_ips) or "None"),
    ]
    for check, count, details in checks:
        data_quality_sheet.append([check, count, details])

    for ws in workbook.worksheets:
        ws.freeze_panes = "A2" if ws.title != "Dashboard Summary" else "A4"
        for column_index, column in enumerate(ws.columns, start=1):
            cells = list(column)
            if not cells:
                continue
            max_length = max(len(str(getattr(cell, "value", None) or "")) for cell in cells)
            ws.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 42)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def build_upload_template() -> BytesIO:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "IT Asset Upload"
    headers = [
        "USED BY", "WS No", "DEPARTMENT", "cpu", "MONITOR", "MOUSE", "KB", "SYSTEM NAME",
        "COMPUTER / LAPTOP / Smartphone", "PROCESSOR", "MEMORY IN GB", "SSD", "HDD", "IP ADDRESS",
        "MAC ADDRESS", "GC", "OS", "ANTIVIRUS", "DHCP", "PERFORMED BY", "APPROVED BY", "PRICE",
        "REMARKS", "date",
    ]
    ws.append(headers)
    ws.append([
        "Employee Name", "NW001", "IT", "2400", "4944, 4934", "7151", "7299", "NW001-IT",
        "Computer", "i7-12700F", "32GB", "512GB", "1TB", "192.168.1.100",
        "00:11:22:33:44:55", "RTX 3050", "Windows 11", "Yes", "STATIC", "IT-Admin", "", 0,
        "Example asset", datetime.now().date(),
    ])
    _style_header(ws)
    ws.freeze_panes = "A2"
    for column in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max_length + 2, 34)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def build_replacement_history_report(db: Session) -> BytesIO:
    workbook = Workbook()
    component_sheet = workbook.active
    component_sheet.title = "Component Replacements"
    component_sheet.sheet_view.showGridLines = False
    component_headers = [
        "Replacement ID", "CPU / Asset Tag", "Workstation No.", "System Asset ID",
        "Component / Field", "Old Value / Tag", "New Value / Tag", "Reason",
        "Old Condition", "Replacement Date", "Technician", "Performed By",
        "Approved By", "Work Record", "Remarks", "Created At", "Change Type", "Batch ID", "Sequence",
    ]
    component_sheet.append(component_headers)
    _style_header(component_sheet)
    component_records = list(
        db.scalars(select(ComponentReplacement).order_by(ComponentReplacement.created_at.desc())).all()
    )
    for record in component_records:
        component_sheet.append([
            record.replacement_code, record.cpu_asset_tag, record.workstation_no,
            record.asset.asset_code if record.asset else None, record.component_type,
            record.old_value or "Not Previously Recorded", record.new_value, record.reason,
            record.old_condition, record.replacement_date, record.technician, record.performed_by,
            record.approved_by, record.work_record.work_code if record.work_record else None,
            record.remarks, record.created_at, (record.change_type or "replacement").replace("_", " ").title(),
            record.batch_code, record.sequence_no or 1,
        ])
    component_sheet.auto_filter.ref = f"A1:S{max(component_sheet.max_row, 1)}"
    component_sheet.freeze_panes = "A2"

    asset_sheet = workbook.create_sheet("Complete Asset Replacements")
    asset_sheet.sheet_view.showGridLines = False
    asset_sheet.append([
        "Replacement ID", "Old CPU / Asset Tag", "Old Workstation", "Old System Asset ID",
        "New CPU / Asset Tag", "New Workstation", "New System Asset ID", "Employee",
        "Department", "Reason", "Damage Category", "Inspection Finding", "Approval Status",
        "Final Action", "Requested By", "Approved By", "Requested At", "Approved At",
    ])
    _style_header(asset_sheet)
    full_records = list(db.scalars(select(ReplacementRecord).order_by(ReplacementRecord.created_at.desc())).all())
    for record in full_records:
        asset_sheet.append([
            record.replacement_code, record.old_asset.cpu_asset_tag, record.old_asset.workstation_no,
            record.old_asset.asset_code, record.new_asset.cpu_asset_tag if record.new_asset else None,
            record.new_asset.workstation_no if record.new_asset else None,
            record.new_asset.asset_code if record.new_asset else None,
            (record.new_asset.used_by if record.new_asset else record.old_asset.used_by),
            (record.new_asset.department if record.new_asset else record.old_asset.department),
            record.reason, record.damage_category, record.inspection_finding, record.approval_status,
            record.final_action, record.requested_by, record.approved_by, record.created_at, record.approved_at,
        ])
    asset_sheet.auto_filter.ref = f"A1:R{max(asset_sheet.max_row, 1)}"
    asset_sheet.freeze_panes = "A2"

    summary = workbook.create_sheet("Replacement Summary", 0)
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:F2")
    summary["A1"] = "NAKSHATECH REPLACEMENT HISTORY"
    summary["A1"].fill = HEADER_FILL
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary.append([])
    summary.append(["Metric", "Count"])
    _style_header(summary, 4)
    change_counts = Counter((record.change_type or "replacement") for record in component_records)
    summary.append(["Component / Configuration Changes", len(component_records)])
    summary.append(["Upgrades", change_counts.get("upgrade", 0)])
    summary.append(["Replacements", change_counts.get("replacement", 0)])
    summary.append(["Downgrades", change_counts.get("downgrade", 0)])
    summary.append(["Upgrade + Replacement", change_counts.get("upgrade_replacement", 0)])
    summary.append(["Complete Asset Replacements", len(full_records)])
    summary.append(["Approved Complete Replacements", sum(r.approval_status == "approved" for r in full_records)])
    summary.append(["Pending Complete Replacements", sum(r.approval_status == "pending" for r in full_records)])

    for sheet in workbook.worksheets:
        for column in range(1, sheet.max_column + 1):
            letter = get_column_letter(column)
            max_length = max((len(str(sheet.cell(row, column).value or "")) for row in range(1, min(sheet.max_row, 200) + 1)), default=10)
            sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 38)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream



def _write_asset_register_sheet(
    workbook,
    ws,
    assets: list,
    sheet_name: str,
    latest_history: dict[int, AssetHistory] | None = None,
) -> None:
    ws.title = sheet_name
    latest_history = latest_history or {}
    max_columns = len(LIVE_ASSET_HEADERS)
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.max_row >= 2 and merged_range.min_row <= max(260, ws.max_row):
            ws.unmerge_cells(str(merged_range))
    for row in ws.iter_rows(min_row=1, max_row=max(260, ws.max_row), min_col=1, max_col=max_columns):
        for cell in row:
            cell.value = None

    for column, header in enumerate(LIVE_ASSET_HEADERS, start=1):
        cell = ws.cell(1, column, header)
        cell.fill = PatternFill("solid", fgColor="FFF200")
        cell.font = Font(color="000000", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    office_assets = [asset for asset in assets if getattr(asset, "work_mode", "office") == "office"]
    remote_assets = [asset for asset in assets if getattr(asset, "work_mode", "office") in {"wfh", "field"}]
    row_number = 2
    serial = 1
    for asset in office_assets:
        for column, value in enumerate(_asset_row(asset, serial, latest_history.get(getattr(asset, "id", -1))), start=1):
            ws.cell(row_number, column).value = _excel_value(value)
        row_number += 1
        serial += 1

    row_number += 1
    heading_row = row_number
    ws.cell(heading_row, 1).value = "Work From Home Systems and laptops"
    ws.cell(heading_row, 1).font = Font(bold=True, color="FFFFFF")
    ws.cell(heading_row, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=heading_row, start_column=1, end_row=heading_row, end_column=max_columns)
    row_number += 1

    serial = 1
    for asset in remote_assets:
        for column, value in enumerate(_asset_row(asset, serial, latest_history.get(getattr(asset, "id", -1))), start=1):
            ws.cell(row_number, column).value = _excel_value(value)
        row_number += 1
        serial += 1

    summary_row = row_number + 2
    counts = Counter(getattr(asset, "device_type", "Other") for asset in assets)
    ws.cell(summary_row, 2).value = "Current Inventory Summary"
    ws.cell(summary_row, 2).font = Font(bold=True, color="FFFFFF")
    ws.cell(summary_row, 2).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(summary_row + 1, 3).value = "Computer"
    ws.cell(summary_row + 1, 4).value = "Laptop"
    ws.cell(summary_row + 1, 5).value = "Mobile"
    ws.cell(summary_row + 2, 2).value = "NakshaTech"
    ws.cell(summary_row + 2, 3).value = counts.get("Computer", 0)
    ws.cell(summary_row + 2, 4).value = counts.get("Laptop", 0)
    ws.cell(summary_row + 2, 5).value = counts.get("Smartphone", 0)
    ws.cell(summary_row + 3, 2).value = "Total"
    ws.cell(summary_row + 3, 3).value = counts.get("Computer", 0)
    ws.cell(summary_row + 3, 4).value = counts.get("Laptop", 0)
    ws.cell(summary_row + 3, 5).value = counts.get("Smartphone", 0)

    department_row = summary_row + 6
    ws.cell(department_row, 12).value = "Department"
    ws.cell(department_row, 13).value = "Asset Count"
    for column in (12, 13):
        ws.cell(department_row, column).fill = HEADER_FILL
        ws.cell(department_row, column).font = HEADER_FONT
    departments = Counter(getattr(asset, "department", None) or "Unassigned" for asset in assets)
    for offset, (department, count) in enumerate(sorted(departments.items()), start=1):
        ws.cell(department_row + offset, 12).value = department
        ws.cell(department_row + offset, 13).value = count
    ws.freeze_panes = "A2"
    last_column = get_column_letter(max_columns)
    ws.auto_filter.ref = f"A1:{last_column}{max(row_number - 1, 2)}"
    ws.sheet_view.showGridLines = False
    for index, width in enumerate(LIVE_ASSET_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=25, max_col=26):
        for cell in row:
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "DD-MM-YYYY"

def build_monthly_asset_report(db: Session, month_key: str) -> BytesIO:
    start = parse_month_key(month_key)
    template_path = Path(settings.seed_excel_path)
    assets, source = assets_for_month(db, start)
    desired_sheet = start.strftime("%B %Y")

    if source == "template":
        if not template_path.exists():
            raise ValueError("The selected historical month is not available")
        source_book = load_workbook(template_path, read_only=True, data_only=False, keep_links=False)
        source_sheet = next(
            (sheet for sheet in source_book.worksheets if parse_template_sheet_month(sheet.title) == start),
            None,
        )
        if source_sheet is None:
            source_book.close()
            raise ValueError("The selected historical month is not available")
        workbook = Workbook()
        target = workbook.active
        target.title = desired_sheet
        last_nonempty = 0
        cached_rows: list[tuple] = []
        for row_index, values in enumerate(source_sheet.iter_rows(values_only=True), start=1):
            values = tuple(values[:25])
            cached_rows.append(values)
            if any(value not in (None, "") for value in values):
                last_nonempty = row_index
        for values in cached_rows[:last_nonempty]:
            target.append(values)
        source_book.close()
        yellow_fill = PatternFill("solid", fgColor="FFF200")
        if target.max_row:
            for cell in target[1]:
                if cell.value is not None:
                    cell.fill = yellow_fill
                    cell.font = Font(color="000000", bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        target.freeze_panes = "A2"
        target.sheet_view.showGridLines = False
        widths = [7, 22, 14, 20, 14, 25, 14, 14, 22, 24, 22, 16, 14, 14, 18, 20, 22, 20, 20, 14, 20, 20, 14, 34, 14]
        for index, width in enumerate(widths, start=1):
            target.column_dimensions[get_column_letter(index)].width = width
    else:
        workbook = load_workbook(template_path) if template_path.exists() else Workbook()
        source_sheet = workbook[workbook.sheetnames[0]]
        if desired_sheet in workbook.sheetnames:
            ws = workbook[desired_sheet]
        else:
            ws = workbook.copy_worksheet(source_sheet)
        audit_history = _latest_history_map(db) if source == "live" else {}
        _write_asset_register_sheet(workbook, ws, assets, desired_sheet, audit_history)
        for sheet in list(workbook.worksheets):
            if sheet.title != desired_sheet:
                workbook.remove(sheet)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _month_bounds(month_key: str) -> tuple[date, date, datetime, datetime]:
    start = parse_month_key(month_key)
    end = month_end(start)
    # Database timestamps are naive UTC. Convert the selected India-calendar
    # month boundaries to UTC before querying so midnight edge cases are exact.
    start_local = datetime.combine(start, datetime.min.time(), tzinfo=IST)
    end_local = datetime.combine(end + timedelta(days=1), datetime.min.time(), tzinfo=IST)
    start_dt = start_local.astimezone(timezone.utc).replace(tzinfo=None)
    end_dt = end_local.astimezone(timezone.utc).replace(tzinfo=None)
    return start, end, start_dt, end_dt


def build_monthly_change_history_report(db: Session, month_key: str) -> BytesIO:
    """Export every system-recorded asset and component change made in a month.

    Month membership uses ``reporting_month`` for new records. Older records without
    that field fall back to the immutable server-recorded ``created_at`` timestamp.
    The workbook shows both the effective reporting month and actual system time.
    """
    start, end, start_dt, end_dt = _month_bounds(month_key)
    histories = list(db.scalars(
        select(AssetHistory)
        .where(or_(
            AssetHistory.reporting_month == month_key,
            and_(AssetHistory.reporting_month.is_(None), AssetHistory.created_at >= start_dt, AssetHistory.created_at < end_dt),
        ))
        .order_by(AssetHistory.created_at, AssetHistory.id)
    ).all())
    component_records = list(db.scalars(
        select(ComponentReplacement)
        .where(or_(
            ComponentReplacement.reporting_month == month_key,
            and_(ComponentReplacement.reporting_month.is_(None), ComponentReplacement.created_at >= start_dt, ComponentReplacement.created_at < end_dt),
        ))
        .order_by(ComponentReplacement.created_at, ComponentReplacement.id)
    ).all())
    complete_records = list(db.scalars(
        select(ReplacementRecord)
        .where(or_(
            ReplacementRecord.reporting_month == month_key,
            and_(ReplacementRecord.reporting_month.is_(None), ReplacementRecord.created_at >= start_dt, ReplacementRecord.created_at < end_dt),
        ))
        .order_by(ReplacementRecord.created_at, ReplacementRecord.id)
    ).all())
    assets_by_id = {asset.id: asset for asset in db.scalars(select(Asset)).all()}
    users_by_email, users_by_name = _user_maps(db)

    full_edits = [item for item in histories if item.change_type == "full_edit"]
    audit_histories = [item for item in histories if item.change_type != "asset_created"]
    component_counts = Counter((record.change_type or "replacement") for record in component_records)
    unique_asset_ids = {item.asset_id for item in audit_histories}
    unique_asset_ids.update(record.asset_id for record in component_records)
    unique_asset_ids.update(record.old_asset_id for record in complete_records)
    unique_asset_ids.update(record.new_asset_id for record in complete_records if record.new_asset_id)
    total_fields = sum(item.field_count or len(_history_fields(item)) for item in audit_histories)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Monthly Summary"
    summary.merge_cells("A1:G2")
    summary["A1"] = f"NAKSHATECH IT ASSET CHANGES - {start.strftime('%B %Y').upper()}"
    summary["A1"].fill = HEADER_FILL
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary.append([])
    summary.append(["Metric", "Count", "Explanation"])
    _style_header(summary, 4)
    metrics = [
        ("Assets Changed", len(unique_asset_ids), "Unique assets with a recorded edit, component operation or complete replacement"),
        ("Full Asset Edit Operations", len(full_edits), "Saves made through Asset Register → Edit Full Record"),
        ("Asset Audit Operations", len(audit_histories), "All recorded asset actions excluding initial creation"),
        ("Fields Changed", total_fields, "Old-to-new field rows captured in asset audit history"),
        ("Upgrade Items", component_counts.get("upgrade", 0), "Individual component/configuration upgrades"),
        ("Replacement Items", component_counts.get("replacement", 0), "Individual component/configuration replacements"),
        ("Downgrade Items", component_counts.get("downgrade", 0), "Individual component/configuration downgrades"),
        ("Upgrade + Replacement Items", component_counts.get("upgrade_replacement", 0), "Combined upgrade/replacement items"),
        ("Complete Asset Replacements", len(complete_records), "Whole system replacement requests created in the month"),
    ]
    for row in metrics:
        summary.append(row)

    operations = workbook.create_sheet("Asset Edit Operations")
    operations.append([
        "Batch ID", "Internal Asset ID", "CPU / Asset Tag", "Workstation", "Action",
        "Change Type", "Fields Changed", "Effective Reporting Month", "Changed By", "User Email", "User Role",
        "System Change Date", "System Change Time", "Reason", "Remarks",
    ])
    _style_header(operations)
    for history in histories:
        asset = assets_by_id.get(history.asset_id)
        actor_name, actor_email, actor_role = _history_actor(history, users_by_email)
        changed_at = _local_datetime(history.created_at)
        operations.append([
            history.batch_code,
            asset.asset_code if asset else history.asset_id,
            asset.cpu_asset_tag if asset else None,
            asset.workstation_no if asset else None,
            history.action,
            (history.change_type or "asset_activity").replace("_", " ").title(),
            history.field_count or len(_history_fields(history)),
            history.reporting_month or (changed_at.strftime("%Y-%m") if changed_at else None),
            actor_name,
            actor_email,
            (actor_role or "").replace("_", " ").title() or None,
            changed_at.date() if changed_at else None,
            changed_at.strftime("%I:%M:%S %p") if changed_at else None,
            history.reason,
            history.remarks,
        ])

    details = workbook.create_sheet("Detailed Field Changes")
    details.append([
        "Batch ID", "Internal Asset ID", "CPU / Asset Tag", "Workstation", "Action",
        "Field / Component", "Old Value", "New Value", "Changed By", "User Email",
        "User Role", "Change Date", "Change Time", "Reason", "Remarks",
    ])
    _style_header(details)
    for history in histories:
        asset = assets_by_id.get(history.asset_id)
        actor_name, actor_email, actor_role = _history_actor(history, users_by_email)
        changed_at = _local_datetime(history.created_at)
        old_values = _json_dict(history.old_value)
        new_values = _json_dict(history.new_value)
        fields = sorted(set(old_values) | set(new_values))
        if not fields:
            fields = [history.action]
            old_values = {history.action: history.old_value}
            new_values = {history.action: history.new_value}
        for field in fields:
            old_value = old_values.get(field)
            new_value = new_values.get(field)
            if old_value == new_value and field != history.action:
                continue
            details.append([
                history.batch_code,
                asset.asset_code if asset else history.asset_id,
                asset.cpu_asset_tag if asset else None,
                asset.workstation_no if asset else None,
                history.action,
                FIELD_LABELS.get(field, field.replace("_", " ").title()),
                old_value,
                new_value,
                actor_name,
                actor_email,
                (actor_role or "").replace("_", " ").title() or None,
                changed_at.date() if changed_at else None,
                changed_at.strftime("%I:%M:%S %p") if changed_at else None,
                history.reason,
                history.remarks,
            ])

    components = workbook.create_sheet("Component Changes")
    components.append([
        "Batch ID", "Change ID", "Work Record", "Internal Asset ID", "CPU / Asset Tag",
        "Workstation", "Change Type", "Component / Field", "Old Value / Tag", "New Value / Tag",
        "Reason", "Old Condition", "Changed By", "User Email", "User Role", "Change Date",
        "Change Time", "Operational Date", "Approved By", "Remarks",
    ])
    _style_header(components)
    for record in component_records:
        actor_name, actor_email, actor_role = _component_actor(record, users_by_email, users_by_name)
        changed_at = _local_datetime(record.created_at)
        components.append([
            record.batch_code,
            record.replacement_code,
            record.work_record.work_code if record.work_record else None,
            record.asset.asset_code if record.asset else record.asset_id,
            record.cpu_asset_tag,
            record.workstation_no,
            (record.change_type or "replacement").replace("_", " ").title(),
            record.component_type,
            record.old_value or "Not Previously Recorded",
            record.new_value,
            record.reason,
            record.old_condition,
            actor_name,
            actor_email,
            (actor_role or "").replace("_", " ").title() or None,
            changed_at.date() if changed_at else None,
            changed_at.strftime("%I:%M:%S %p") if changed_at else None,
            record.replacement_date,
            record.approved_by,
            record.remarks,
        ])

    complete = workbook.create_sheet("Complete Asset Replacements")
    complete.append([
        "Replacement ID", "Old CPU / Asset Tag", "Old Workstation", "New CPU / Asset Tag",
        "New Workstation", "Employee", "Department", "Reason", "Damage Category",
        "Approval Status", "Final Action", "Requested By", "Requested By Email",
        "Requested By Role", "Approved By", "Approved By Email", "Approved By Role",
        "Request Date", "Request Time", "Approved At",
    ])
    _style_header(complete)
    for record in complete_records:
        requested_at = _local_datetime(record.created_at)
        approved_at = _local_datetime(record.approved_at)
        complete.append([
            record.replacement_code,
            record.old_asset.cpu_asset_tag,
            record.old_asset.workstation_no,
            record.new_asset.cpu_asset_tag if record.new_asset else None,
            record.new_asset.workstation_no if record.new_asset else None,
            record.new_asset.used_by if record.new_asset else record.old_asset.used_by,
            record.new_asset.department if record.new_asset else record.old_asset.department,
            record.reason,
            record.damage_category,
            record.approval_status,
            record.final_action,
            record.requested_by,
            record.requested_by_email,
            (record.requested_by_role or "").replace("_", " ").title() or None,
            record.approved_by,
            record.approved_by_email,
            (record.approved_by_role or "").replace("_", " ").title() or None,
            requested_at.date() if requested_at else None,
            requested_at.strftime("%I:%M:%S %p") if requested_at else None,
            approved_at.strftime("%d-%m-%Y %I:%M:%S %p") if approved_at else None,
        ])

    user_summary = workbook.create_sheet("User Activity Summary")
    user_summary.append([
        "Changed By", "User Email", "User Role", "Asset Audit Operations",
        "Fields Changed", "Component Items", "First Activity", "Last Activity",
    ])
    _style_header(user_summary)
    activity: dict[tuple[str, str, str], dict] = {}
    for history in histories:
        name, email, role = _history_actor(history, users_by_email)
        key = (name or "Unknown", email or "", role or "")
        row = activity.setdefault(key, {"operations": 0, "fields": 0, "components": 0, "times": []})
        row["operations"] += 1
        row["fields"] += history.field_count or len(_history_fields(history))
        row["times"].append(history.created_at)
    for record in component_records:
        name, email, role = _component_actor(record, users_by_email, users_by_name)
        key = (name or "Unknown", email or "", role or "")
        row = activity.setdefault(key, {"operations": 0, "fields": 0, "components": 0, "times": []})
        row["components"] += 1
        row["times"].append(record.created_at)
    for (name, email, role), values in sorted(activity.items(), key=lambda item: item[0][0].lower()):
        times = [_local_datetime(value) for value in values["times"] if value]
        user_summary.append([
            name,
            email or None,
            role.replace("_", " ").title() if role else None,
            values["operations"],
            values["fields"],
            values["components"],
            min(times).strftime("%d-%m-%Y %I:%M:%S %p") if times else None,
            max(times).strftime("%d-%m-%Y %I:%M:%S %p") if times else None,
        ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2" if sheet.title != "Monthly Summary" else "A4"
        sheet.sheet_view.showGridLines = False
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, date) and not isinstance(cell.value, datetime):
                    cell.number_format = "DD-MM-YYYY"
        for column in range(1, sheet.max_column + 1):
            letter = get_column_letter(column)
            max_length = max(
                (len(str(sheet.cell(row, column).value or "")) for row in range(1, min(sheet.max_row, 500) + 1)),
                default=10,
            )
            sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 45)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream

def build_monthly_summary_report(db: Session, month_key: str) -> BytesIO:
    start, end, start_dt, end_dt = _month_bounds(month_key)
    run = get_snapshot_run(db, start)

    # Months that predate system snapshots are represented by the original workbook.
    # Do not incorrectly apply today's live database totals to an old historical month.
    if start != month_start() and run is None:
        template_path = Path(settings.seed_excel_path)
        if template_path.exists():
            historical_book = load_workbook(template_path, read_only=True, data_only=True, keep_links=False)
            source_sheet = next(
                (sheet for sheet in historical_book.worksheets if parse_template_sheet_month(sheet.title) == start),
                None,
            )
            if source_sheet is not None:
                device_counts = Counter()
                closing_count = 0
                for row in source_sheet.iter_rows(min_row=2, values_only=True):
                    cpu_tag = row[4] if len(row) > 4 else None
                    device = str(row[9] or "").strip() if len(row) > 9 else ""
                    if device.lower() in {"computer", "laptop", "smartphone", "mobile"} or (cpu_tag not in (None, "", "-") and str(row[0] or "").strip().isdigit()):
                        closing_count += 1
                        normalized = "Smartphone" if device.lower() in {"smartphone", "mobile"} else (device.title() if device else "Other")
                        device_counts[normalized] += 1
                historical_book.close()

                workbook = Workbook()
                ws = workbook.active
                ws.title = "Monthly Asset Summary"
                ws.merge_cells("A1:F2")
                ws["A1"] = f"NAKSHATECH HISTORICAL ASSET SUMMARY - {start.strftime('%B %Y').upper()}"
                ws["A1"].fill = HEADER_FILL
                ws["A1"].font = Font(color="FFFFFF", bold=True, size=16)
                ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
                ws.append([])
                ws.append(["Metric", "Count", "Explanation"])
                _style_header(ws, 4)
                ws.append(["Historical Closing Assets", closing_count, "Count derived from the original monthly workbook"])
                ws.append(["Computers", device_counts.get("Computer", 0), "Original monthly workbook"])
                ws.append(["Laptops", device_counts.get("Laptop", 0), "Original monthly workbook"])
                ws.append(["Smartphones", device_counts.get("Smartphone", 0), "Original monthly workbook"])
                ws.append(["System-recorded component changes", 0, "Detailed upgrade, replacement and downgrade history was not captured in the application before system adoption"])
                ws.column_dimensions["A"].width = 38
                ws.column_dimensions["B"].width = 14
                ws.column_dimensions["C"].width = 72
                ws.freeze_panes = "A4"
                stream = BytesIO()
                workbook.save(stream)
                stream.seek(0)
                return stream
    current_assets = list(db.scalars(select(Asset)).all())
    imported_baseline = sum(bool(asset.source_sheet) for asset in current_assets)
    manual_added_ids = set(db.scalars(select(AssetHistory.asset_id).where(
        AssetHistory.change_type == "asset_created",
        or_(
            AssetHistory.reporting_month == month_key,
            and_(AssetHistory.reporting_month.is_(None), AssetHistory.created_at >= start_dt, AssetHistory.created_at < end_dt),
        ),
    )).all())
    manual_added = [asset for asset in current_assets if asset.id in manual_added_ids]
    component_records = list(db.scalars(
        select(ComponentReplacement).where(or_(
            ComponentReplacement.reporting_month == month_key,
            and_(ComponentReplacement.reporting_month.is_(None), ComponentReplacement.created_at >= start_dt, ComponentReplacement.created_at < end_dt),
        ))
    ).all())
    complete_replacements = db.scalar(select(func.count(ReplacementRecord.id)).where(or_(
        ReplacementRecord.reporting_month == month_key,
        and_(ReplacementRecord.reporting_month.is_(None), ReplacementRecord.created_at >= start_dt, ReplacementRecord.created_at < end_dt),
    ))) or 0
    histories = list(db.scalars(select(AssetHistory).where(or_(
        AssetHistory.reporting_month == month_key,
        and_(AssetHistory.reporting_month.is_(None), AssetHistory.created_at >= start_dt, AssetHistory.created_at < end_dt),
    ))).all())
    change_counts = Counter((record.change_type or "replacement") for record in component_records)
    full_edit_histories = [history for history in histories if history.change_type == "full_edit"]
    edited_asset_count = len({history.asset_id for history in full_edit_histories})
    edited_field_count = sum(history.field_count or len(_history_fields(history)) for history in full_edit_histories)
    opening_count = run.opening_count if run else imported_baseline
    closing_count = run.closing_count if run else len(current_assets)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Monthly Asset Summary"
    ws.merge_cells("A1:F2")
    ws["A1"] = f"NAKSHATECH MONTHLY ASSET SUMMARY - {start.strftime('%B %Y').upper()}"
    ws["A1"].fill = HEADER_FILL
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.append([])
    ws.append(["Metric", "Count", "Explanation"])
    _style_header(ws, 4)
    rows = [
        ("Opening / Imported Baseline", opening_count, "Assets carried into the selected month"),
        ("New Assets Added", len(manual_added), "Manually registered assets during the selected month"),
        ("Assets Edited", edited_asset_count, "Unique assets updated through Full Edit during the selected month"),
        ("Full Asset Edit Operations", len(full_edit_histories), "Successful Full Edit save operations during the selected month"),
        ("Asset Fields Edited", edited_field_count, "Individual old-to-new fields captured by edit audit history"),
        ("Component Upgrades", change_counts.get("upgrade", 0), "Individual upgraded fields/components"),
        ("Component Replacements", change_counts.get("replacement", 0), "Individual replacement fields/components"),
        ("Component Downgrades", change_counts.get("downgrade", 0), "Individual fields/components changed to a lower or reassigned specification"),
        ("Upgrade + Replacement Items", change_counts.get("upgrade_replacement", 0), "Failed items replaced with higher specification"),
        ("Complete Asset Replacements", complete_replacements, "Whole computer/laptop replacements"),
        ("Returns", sum("returned" in history.action.lower() for history in histories), "Asset return actions"),
        ("Retired / Archived", sum("retired" in history.action.lower() or "archived" in history.action.lower() for history in histories), "Assets removed from active use"),
        ("Closing Assets", closing_count, "Final/live inventory count for the selected month"),
    ]
    for row in rows:
        ws.append(row)

    added = workbook.create_sheet("New Assets Added")
    added.append(["Asset ID", "CPU / Asset Tag", "Workstation", "Device", "Used By", "Department", "Created At"])
    _style_header(added)
    for asset in manual_added:
        added.append([asset.asset_code, asset.cpu_asset_tag, asset.workstation_no, asset.device_type, asset.used_by, asset.department, asset.created_at])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2" if sheet.title != "Monthly Asset Summary" else "A4"
        for column in range(1, sheet.max_column + 1):
            letter = get_column_letter(column)
            max_length = max((len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)), default=10)
            sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 45)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream
