from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from io import BytesIO
from math import ceil
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.services.monthly_snapshot_service import assets_for_month, month_start, parse_month_key
from app.services.asset_lifecycle_service import (
    canonical_device_type,
    inventory_summary,
    is_primary_device_type,
    lifecycle_status_distribution,
    status_matches,
)

VALID_SCOPES = {"all", "primary", "device", "status", "department"}
SORT_FIELDS = {
    "asset_code",
    "cpu_asset_tag",
    "workstation_no",
    "device_type",
    "department",
    "used_by",
    "status",
    "location",
    "updated_at",
    "capacity",
    "ownership",
    "client_name",
    "project_id",
    "current_holder",
}

def _text(value: Any) -> str:
    return str(value or "").strip()


def _same(left: Any, right: Any) -> bool:
    return _text(left).casefold() == _text(right).casefold()


def _status_matches(asset_status: Any, requested: str | None) -> bool:
    return status_matches(asset_status, requested)


def _scope_matches(asset: Any, scope: str, scope_value: str | None) -> bool:
    if scope == "all":
        return True
    if scope == "primary":
        return is_primary_device_type(getattr(asset, "device_type", None))
    if not scope_value:
        return False
    if scope == "device":
        return canonical_device_type(getattr(asset, "device_type", None)) == canonical_device_type(scope_value)
    if scope == "department":
        department = _text(getattr(asset, "department", None)) or "Unassigned"
        return _same(department, scope_value)
    if scope == "status":
        return _status_matches(getattr(asset, "status", None), scope_value)
    return False


def _search_matches(asset: Any, query: str | None) -> bool:
    normalized = _text(query).casefold()
    if not normalized:
        return True
    values = (
        getattr(asset, "asset_code", None),
        getattr(asset, "cpu_asset_tag", None),
        getattr(asset, "system_name", None),
        getattr(asset, "brand", None),
        getattr(asset, "model", None),
        getattr(asset, "serial_number", None),
        getattr(asset, "connection_type", None),
        getattr(asset, "capacity", None),
        getattr(asset, "ownership", None),
        getattr(asset, "client_name", None),
        getattr(asset, "project_id", None),
        getattr(asset, "current_holder", None),
        getattr(asset, "workstation_no", None),
        getattr(asset, "used_by", None),
        getattr(asset, "department", None),
        getattr(asset, "device_type", None),
        getattr(asset, "processor", None),
        getattr(asset, "memory_gb", None),
        getattr(asset, "ssd", None),
        getattr(asset, "hdd", None),
        getattr(asset, "ip_address", None),
        getattr(asset, "mac_address", None),
        getattr(asset, "location", None),
        getattr(asset, "remarks", None),
    )
    return normalized in " ".join(_text(value) for value in values).casefold()


def _summary(assets: Iterable[Any]) -> dict[str, int]:
    rows = list(assets)
    summary = inventory_summary(rows)
    return {
        **summary,
        "nakshatech_owned": sum(
            1 for asset in rows
            if canonical_device_type(getattr(asset, "device_type", None)) == "External HDD"
            and _same(getattr(asset, "ownership", None), "NakshaTech")
        ),
        "client_owned": sum(
            1 for asset in rows
            if canonical_device_type(getattr(asset, "device_type", None)) == "External HDD"
            and _same(getattr(asset, "ownership", None), "Client")
        ),
    }



def _distribution(assets: Iterable[Any], field: str, fallback: str) -> list[dict[str, Any]]:
    if field == "device_type":
        counts = Counter(canonical_device_type(getattr(asset, field, None)) for asset in assets)
    else:
        counts = Counter(_text(getattr(asset, field, None)) or fallback for asset in assets)
    return [
        {"name": name, "value": value, "key": name}
        for name, value in sorted(counts.items(), key=lambda item: (-item[1], item[0].casefold()))
    ]


def _status_distribution(assets: Iterable[Any]) -> list[dict[str, Any]]:
    return lifecycle_status_distribution(assets)

def _options(assets: Iterable[Any]) -> dict[str, list[str]]:
    rows = list(assets)

    def values(field: str, fallback: str | None = None) -> list[str]:
        result = {_text(getattr(asset, field, None)) or fallback or "" for asset in rows}
        return sorted((value for value in result if value), key=str.casefold)

    device_values = sorted({canonical_device_type(getattr(asset, "device_type", None)) for asset in rows}, key=str.casefold)
    return {
        "devices": device_values,
        "departments": values("department", "Unassigned"),
        "statuses": values("status"),
        "locations": values("location", "Unknown"),
        "work_modes": values("work_mode"),
        "ownerships": values("ownership"),
        "clients": values("client_name"),
        "projects": values("project_id"),
        "holders": values("current_holder"),
        "brands": values("brand"),
        "capacities": values("capacity"),
    }


def _sort_key(asset: Any, field: str) -> tuple[int, Any]:
    value = getattr(asset, field, None)
    if value is None or value == "":
        return (1, "")
    if isinstance(value, (date, datetime)):
        return (0, value.isoformat())
    return (0, str(value).casefold())


def _scope_label(scope: str, value: str | None) -> str:
    if scope == "all":
        return "All Tracked IT Records"
    if scope == "primary":
        return "Primary IT Assets"
    if scope == "device":
        return f"{value or 'Device'} Assets"
    if scope == "department":
        return f"{value or 'Department'} Department Assets"
    normalized = _text(value).lower()
    labels = {
        "assigned": "Assigned / In Use Assets",
        "assigned_in_use": "Assigned / In Use Assets",
        "available": "Available Assets",
        "repair": "Assets Under Repair",
        "replacement_pending": "Replacement Pending Assets",
    }
    return labels.get(normalized, f"{_text(value).replace('_', ' ').title()} Assets")


def query_asset_drilldown(
    db: Session,
    *,
    month_key: str,
    scope: str = "all",
    scope_value: str | None = None,
    search: str | None = None,
    department: str | None = None,
    device_type: str | None = None,
    status: str | None = None,
    location: str | None = None,
    work_mode: str | None = None,
    ownership: str | None = None,
    client_name: str | None = None,
    project_id: str | None = None,
    current_holder: str | None = None,
    brand: str | None = None,
    capacity: str | None = None,
    sort_by: str = "asset_code",
    sort_dir: str = "asc",
    page: int = 1,
    page_size: int = 20,
) -> dict[str, Any]:
    if scope not in VALID_SCOPES:
        raise ValueError("Scope must be all, primary, device, status or department")
    if scope not in {"all", "primary"} and not _text(scope_value):
        raise ValueError("A scope value is required for this drill-down")
    if sort_by not in SORT_FIELDS:
        raise ValueError("Unsupported asset sort field")
    if sort_dir not in {"asc", "desc"}:
        raise ValueError("Sort direction must be asc or desc")
    if page < 1:
        raise ValueError("Page must be at least 1")
    if page_size < 1 or page_size > 100:
        raise ValueError("Page size must be between 1 and 100")

    selected_start = parse_month_key(month_key)
    assets, source = assets_for_month(db, selected_start)
    if source == "missing":
        raise ValueError("No asset register is available for the selected month")

    scoped_assets = [asset for asset in assets if _scope_matches(asset, scope, scope_value)]
    filtered_assets = []
    for asset in scoped_assets:
        asset_department = _text(getattr(asset, "department", None)) or "Unassigned"
        asset_location = _text(getattr(asset, "location", None)) or "Unknown"
        if not _search_matches(asset, search):
            continue
        if department and not _same(asset_department, department):
            continue
        if device_type and canonical_device_type(getattr(asset, "device_type", None)) != canonical_device_type(device_type):
            continue
        if not _status_matches(getattr(asset, "status", None), status):
            continue
        if location and not _same(asset_location, location):
            continue
        if work_mode and not _same(getattr(asset, "work_mode", None), work_mode):
            continue
        if ownership and not _same(getattr(asset, "ownership", None), ownership):
            continue
        if client_name and not _same(getattr(asset, "client_name", None), client_name):
            continue
        if project_id and not _same(getattr(asset, "project_id", None), project_id):
            continue
        if current_holder and not _same(getattr(asset, "current_holder", None), current_holder):
            continue
        if brand and not _same(getattr(asset, "brand", None), brand):
            continue
        if capacity and not _same(getattr(asset, "capacity", None), capacity):
            continue
        filtered_assets.append(asset)

    filtered_assets.sort(key=lambda asset: _sort_key(asset, sort_by), reverse=sort_dir == "desc")
    filtered_total = len(filtered_assets)
    total_pages = max(1, ceil(filtered_total / page_size))
    safe_page = min(page, total_pages)
    start_index = (safe_page - 1) * page_size
    page_assets = filtered_assets[start_index:start_index + page_size]
    current = month_start()

    return {
        "month": {
            "key": selected_start.strftime("%Y-%m"),
            "label": selected_start.strftime("%B %Y"),
            "source": source,
            "is_live": selected_start == current,
        },
        "scope": {
            "type": scope,
            "value": scope_value,
            "label": _scope_label(scope, scope_value),
        },
        "scope_total": len(scoped_assets),
        "filtered_total": filtered_total,
        "page": safe_page,
        "page_size": page_size,
        "pages": total_pages,
        "summary": _summary(scoped_assets),
        "filtered_summary": _summary(filtered_assets),
        "visuals": {
            "device_distribution": _distribution(filtered_assets, "device_type", "Other"),
            "status_distribution": _status_distribution(filtered_assets),
            "department_distribution": _distribution(filtered_assets, "department", "Unassigned"),
        },
        "filter_options": _options(scoped_assets),
        "assets": page_assets,
        "all_filtered_assets": filtered_assets,
    }


def build_asset_drilldown_workbook(result: dict[str, Any]) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Drilldown Summary"
    navy = PatternFill("solid", fgColor="082B57")
    cyan = PatternFill("solid", fgColor="DDF6FB")
    white_bold = Font(color="FFFFFF", bold=True)

    summary.merge_cells("A1:F2")
    summary["A1"] = f"NAKSHATECH IT DASHBOARD DRILL-DOWN - {result['scope']['label'].upper()}"
    summary["A1"].fill = navy
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=15)
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary.append([])
    summary.append(["Reporting Month", result["month"]["label"]])
    summary.append(["Scope", result["scope"]["label"]])
    summary.append(["Scope Total", result["scope_total"]])
    summary.append(["Filtered Records", result["filtered_total"]])
    summary.append([])
    summary.append(["Metric", "Count"])
    for cell in summary[9]:
        cell.fill = navy
        cell.font = white_bold
    metric_labels = (
        ("Total", "total"),
        ("Computers", "computers"),
        ("Laptops", "laptops"),
        ("Smartphones", "smartphones"),
        ("Printers", "printers"),
        ("External HDDs", "external_hdds"),
        ("Servers", "servers"),
        ("Network Devices", "network_devices"),
        ("Other Device Types", "other"),
        ("NakshaTech Owned HDDs", "nakshatech_owned"),
        ("Client Owned HDDs", "client_owned"),
        ("Issued HDDs", "issued"),
        ("Permanently Issued HDDs", "permanently_issued"),
        ("Returned HDDs", "returned"),
        ("Assigned / In Use", "assigned"),
        ("Available", "available"),
        ("Under Repair", "repair"),
        ("Replacement Pending", "replacement_pending"),
        ("Damaged / At Risk", "damaged"),
        ("Retired / Finalized", "terminal"),
        ("Other / Legacy Lifecycle", "other_lifecycle"),
    )
    for label, key in metric_labels:
        summary.append([label, result["filtered_summary"][key]])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 24
    summary.sheet_view.showGridLines = False

    assets_sheet = workbook.create_sheet("Asset Details")
    headers = [
        "SL", "Internal Asset ID", "CPU / Asset Tag", "System Name", "Workstation", "Device Type",
        "Department", "Used By", "Status", "Work Mode", "Location", "Processor", "Memory", "SSD", "HDD",
        "Graphics Card", "Operating System", "IP Address", "MAC Address", "Monitor Tag(s)", "Mouse Tag",
        "Keyboard Tag", "Antivirus", "Network Type", "Brand", "Model", "Serial Number",
        "Connection Type", "Capacity", "Ownership", "Client Name", "Project ID", "Current Holder",
        "Approved By", "Price", "Asset Master Remarks", "Original Asset Date",
        "Last Updated",
    ]
    assets_sheet.append(headers)
    for cell in assets_sheet[1]:
        cell.fill = navy
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for index, asset in enumerate(result["all_filtered_assets"], start=1):
        assets_sheet.append([
            index,
            getattr(asset, "asset_code", None),
            getattr(asset, "cpu_asset_tag", None),
            getattr(asset, "system_name", None),
            getattr(asset, "workstation_no", None),
            getattr(asset, "device_type", None),
            getattr(asset, "department", None),
            getattr(asset, "used_by", None),
            _text(getattr(asset, "status", None)).replace("_", " ").title(),
            _text(getattr(asset, "work_mode", None)).replace("_", " ").title(),
            getattr(asset, "location", None),
            getattr(asset, "processor", None),
            getattr(asset, "memory_gb", None),
            getattr(asset, "ssd", None),
            getattr(asset, "hdd", None),
            getattr(asset, "graphics_card", None),
            getattr(asset, "operating_system", None),
            getattr(asset, "ip_address", None),
            getattr(asset, "mac_address", None),
            getattr(asset, "monitor_asset_tags", None),
            getattr(asset, "mouse_asset_tag", None),
            getattr(asset, "keyboard_asset_tag", None),
            getattr(asset, "antivirus", None),
            getattr(asset, "network_type", None),
            getattr(asset, "brand", None),
            getattr(asset, "model", None),
            getattr(asset, "serial_number", None),
            getattr(asset, "connection_type", None),
            getattr(asset, "capacity", None),
            getattr(asset, "ownership", None),
            getattr(asset, "client_name", None),
            getattr(asset, "project_id", None),
            getattr(asset, "current_holder", None),
            getattr(asset, "approved_by", None),
            getattr(asset, "price", None),
            getattr(asset, "remarks", None),
            getattr(asset, "original_asset_date", None) or getattr(asset, "asset_date", None),
            getattr(asset, "updated_at", None),
        ])
        if index % 2 == 0:
            for cell in assets_sheet[assets_sheet.max_row]:
                cell.fill = cyan

    assets_sheet.freeze_panes = "A2"
    assets_sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(assets_sheet.max_row, 2)}"
    assets_sheet.sheet_view.showGridLines = False
    widths = [7, 20, 18, 20, 16, 15, 22, 22, 20, 14, 18, 24, 14, 14, 14, 22, 20, 16, 20, 24, 16, 16, 16, 16, 18, 22, 22, 18, 14, 18, 24, 16, 24, 20, 14, 38, 18, 22]
    for index, width in enumerate(widths, start=1):
        assets_sheet.column_dimensions[get_column_letter(index)].width = width

    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream
