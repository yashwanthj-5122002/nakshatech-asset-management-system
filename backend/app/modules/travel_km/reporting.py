from __future__ import annotations

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.entities import User
from app.modules.travel_km.service import claim_payload, dashboard_payload, list_visible_claims

EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _fit(ws):
    for col in range(1, ws.max_column + 1):
        width = 12
        for cell in ws.iter_cols(min_col=col, max_col=col):
            for item in cell:
                width = max(width, min(55, len(str(item.value or "")) + 2))
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="123B5D")
        cell.alignment = Alignment(vertical="center")


def build_travel_km_workbook(db, *, viewer: User, effective_role: str) -> BytesIO:
    claims = list_visible_claims(db, viewer=viewer, effective_role=effective_role)
    wb = Workbook()
    ws = wb.active
    ws.title = "Claims"
    headers = [
        "Claim ID", "Employee", "Employee Email", "Department", "Project ID", "Project Name", "Travel Date",
        "Start KM", "End KM", "Odometer KM", "GPS Straight-Line KM", "Variance KM", "Variance %",
        "Start Latitude", "Start Longitude", "Start GPS Accuracy m", "Start Evidence Verification",
        "End Latitude", "End Longitude", "End GPS Accuracy m", "End Evidence Verification",
        "Rate per KM", "Calculated Allowance", "Admin Eligible KM", "HR Eligible KM", "Final Allowance",
        "Status", "Admin Remarks", "HR Remarks", "Finance Remarks", "Finance Decision At", "Submitted At"
    ]
    ws.append(headers)
    all_payloads = [claim_payload(db, c, viewer=viewer, effective_role=effective_role) for c in claims]
    for p in all_payloads:
        evidence = {item["phase"]: item for item in p["attachments"]}
        start_evidence = evidence.get("start", {})
        end_evidence = evidence.get("end", {})
        ws.append([
            p["claim_code"], p["employee_name"], p["employee_email"], p["department"], p["project_code"], p["project_name"], p["travel_date"],
            p["start_km"], p["end_km"], p["odometer_km"], p["gps_straight_line_km"], p["distance_variance_km"], p["distance_variance_percent"],
            p["start_latitude"], p["start_longitude"], p["start_accuracy_m"], start_evidence.get("verification_flag"),
            p["end_latitude"], p["end_longitude"], p["end_accuracy_m"], end_evidence.get("verification_flag"),
            p["rate_per_km"], p["calculated_allowance"], p["admin_eligible_km"], p["hr_eligible_km"], p["final_allowance"],
            p["status"], p["admin_comments"], p["hr_comments"], p["finance_comments"], p["finance_decision_at"], p["submitted_at"],
        ])
    _header(ws); _fit(ws)

    ev = wb.create_sheet("Workflow Events")
    ev.append(["Claim ID", "Action", "From", "To", "Actor", "Role", "Comments", "Date Time"])
    for p in all_payloads:
        for e in p["events"]:
            ev.append([p["claim_code"], e["action"], e["from_status"], e["to_status"], e["actor_name"], e["actor_role"], e["comments"], e["created_at"]])
    _header(ev); _fit(ev)

    dash = dashboard_payload(db, viewer=viewer, effective_role=effective_role)
    sm = wb.create_sheet("Summary")
    sm.append(["Metric", "Value"])
    for key in ("total_claims", "total_km", "approved_allowance", "salary_approved_amount", "pending_admin", "pending_hr", "pending_finance", "finance_approved_count", "rejected_count", "sent_back_count"):
        sm.append([key.replace("_", " ").title(), dash[key]])
    _header(sm); _fit(sm)

    for title, key in (("Project Summary", "project_summary"), ("Employee Summary", "employee_summary"), ("Monthly Summary", "monthly_summary")):
        sh = wb.create_sheet(title)
        sh.append(["Name", "Claims", "KM", "Allowance"])
        for row in dash[key]:
            sh.append([row["name"], row["claims"], row["km"], row["allowance"]])
        _header(sh); _fit(sh)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
