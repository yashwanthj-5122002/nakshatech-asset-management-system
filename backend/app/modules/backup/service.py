from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import subprocess
import tarfile
import tempfile
from typing import Any, Iterable
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import Date, DateTime, select
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.entities import (
    Asset,
    AssetHistory,
    ComponentReplacement,
    Drone,
    DroneLocation,
    MonthlySnapshotRun,
    ReplacementRecord,
    User,
    WorkRecord,
    utc_now,
)
from app.modules.backup.models import BackupRun
from app.modules.it_activity.models import ITHandoverRecord, ITPurchaseRecord
from app.modules.drone.models import (
    DroneAssetKit,
    DroneAssetMovement,
    DroneAuditLog,
    DroneHDDDelivery,
    DroneImportBatch,
    DroneImportException,
    DroneKitComponent,
    DroneOperation,
    DroneOperationItem,
    DroneProject,
    DroneSurveyAsset,
    DroneTelecomConnection,
    DroneUINRegistration,
    DroneWorkRecord,
)
from app.services.monthly_snapshot_service import assets_for_month, month_start


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
VALID_BACKUP_TYPES = {"daily", "monthly", "current_month", "yearly", "financial_year", "full"}
VALID_SCOPES = {"all", "it", "drone"}
MAX_EXCEL_TEXT = 32000


@dataclass(frozen=True)
class Period:
    backup_type: str
    start: date | None
    end: date | None
    key: str
    label: str


@dataclass(frozen=True)
class ExportSpec:
    sheet_name: str
    model: type
    scope: str
    date_field: str | None = None
    master: bool = False
    exclude: tuple[str, ...] = ()
    admin_only: bool = False
    special: str | None = None


EXPORT_SPECS: tuple[ExportSpec, ...] = (
    ExportSpec("IT Asset Register", Asset, "it", master=True, special="it_assets"),
    ExportSpec("IT Work Records", WorkRecord, "it", date_field="created_at"),
    ExportSpec("IT Component Changes", ComponentReplacement, "it", date_field="created_at"),
    ExportSpec("IT Replacements", ReplacementRecord, "it", date_field="created_at"),
    ExportSpec("IT Asset History", AssetHistory, "it", date_field="created_at"),
    ExportSpec("IT Handover Return", ITHandoverRecord, "it", date_field="activity_date"),
    ExportSpec("IT Purchase Records", ITPurchaseRecord, "it", date_field="purchase_date"),
    ExportSpec("IT Monthly Snapshots", MonthlySnapshotRun, "it", date_field="month_start"),
    ExportSpec("Drone Asset Register", DroneSurveyAsset, "drone", master=True, exclude=("original_raw_payload", "original_header_map")),
    ExportSpec("Drone Kits", DroneAssetKit, "drone", master=True, exclude=("original_raw_payload",)),
    ExportSpec("Drone Kit Components", DroneKitComponent, "drone", master=True),
    ExportSpec("Drone Projects", DroneProject, "drone", master=True),
    ExportSpec("Drone Operations", DroneOperation, "drone", date_field="operation_date"),
    ExportSpec("Drone Operation Items", DroneOperationItem, "drone", date_field="created_at", special="operation_items"),
    ExportSpec("Drone Movements", DroneAssetMovement, "drone", date_field="occurred_at"),
    ExportSpec("Drone Work Records", DroneWorkRecord, "drone", date_field="created_at"),
    ExportSpec("Drone UIN Register", DroneUINRegistration, "drone", master=True, exclude=("original_raw_payload",)),
    ExportSpec("Drone HDD Deliveries", DroneHDDDelivery, "drone", date_field="delivery_date", exclude=("original_raw_payload",)),
    ExportSpec("Drone Telecom", DroneTelecomConnection, "drone", master=True, exclude=("original_raw_payload",)),
    ExportSpec("Drone Audit", DroneAuditLog, "drone", date_field="created_at"),
    ExportSpec("Drone Import Batches", DroneImportBatch, "drone", date_field="created_at"),
    ExportSpec("Drone Import Exceptions", DroneImportException, "drone", date_field="created_at"),
    ExportSpec("Telemetry Drones", Drone, "drone", master=True),
    ExportSpec("Drone Locations", DroneLocation, "drone", date_field="recorded_at"),
)


def local_now() -> datetime:
    try:
        zone = ZoneInfo(settings.backup_timezone)
    except Exception:
        zone = timezone.utc
    return datetime.now(zone)


def resolve_period(backup_type: str, period: str | None = None, today: date | None = None) -> Period:
    backup_type = backup_type.strip().lower()
    if backup_type not in VALID_BACKUP_TYPES:
        raise ValueError(f"Backup type must be one of: {', '.join(sorted(VALID_BACKUP_TYPES))}")
    current = today or local_now().date()

    if backup_type == "full":
        return Period(backup_type, None, None, "all-time", "Complete Current Backup")

    if backup_type == "daily":
        selected = date.fromisoformat(period) if period else current
        return Period(backup_type, selected, selected, selected.isoformat(), selected.strftime("%d %B %Y"))

    if backup_type in {"monthly", "current_month"}:
        if backup_type == "current_month" or not period:
            selected = current.replace(day=1)
        else:
            selected = datetime.strptime(period, "%Y-%m").date().replace(day=1)
        next_month = (selected.replace(day=28) + timedelta(days=4)).replace(day=1)
        end = next_month - timedelta(days=1)
        key = selected.strftime("%Y-%m")
        label_prefix = "Current Month Live" if backup_type == "current_month" else "Monthly"
        return Period(backup_type, selected, end, key, f"{label_prefix} — {selected.strftime('%B %Y')}")

    if backup_type == "yearly":
        year = int(period) if period else current.year
        return Period(backup_type, date(year, 1, 1), date(year, 12, 31), str(year), f"Calendar Year {year}")

    # Financial year defaults to the FY containing today's date (April to March).
    if period:
        start_year = int(period.split("-")[0])
    else:
        start_year = current.year if current.month >= 4 else current.year - 1
    end_year = start_year + 1
    key = f"{start_year}-{str(end_year)[-2:]}"
    return Period(backup_type, date(start_year, 4, 1), date(end_year, 3, 31), key, f"Financial Year {key}")


def _model_columns(model: type, excluded: Iterable[str] = ()) -> list[str]:
    excluded_set = set(excluded)
    return [column.name for column in model.__table__.columns if column.name not in excluded_set]


def _serialize(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return value
    if isinstance(value, (dict, list, tuple, set)):
        text = json.dumps(value, ensure_ascii=False, default=str)
        return text[:MAX_EXCEL_TEXT]
    if isinstance(value, bytes):
        return f"<binary {len(value)} bytes>"
    text = str(value) if not isinstance(value, (int, float, bool)) else value
    if isinstance(text, str):
        if len(text) > MAX_EXCEL_TEXT:
            text = text[:MAX_EXCEL_TEXT] + "…"
        # Prevent imported/user-entered text from becoming an executable Excel formula.
        if text.startswith(("=", "+", "-", "@")):
            text = "'" + text
    return text


def _filter_query(model: type, field_name: str | None, period: Period):
    query = select(model)
    if not field_name or not period.start or not period.end:
        return query
    column = getattr(model, field_name)
    column_type = model.__table__.columns[field_name].type
    if isinstance(column_type, DateTime):
        start_dt = datetime.combine(period.start, time.min)
        end_dt = datetime.combine(period.end + timedelta(days=1), time.min)
        return query.where(column >= start_dt, column < end_dt)
    if isinstance(column_type, Date):
        return query.where(column >= period.start, column <= period.end)
    return query


def _rows_for_spec(db: Session, spec: ExportSpec, period: Period) -> list[Any]:
    if spec.special == "it_assets" and period.backup_type in {"monthly", "current_month"} and period.start:
        assets, _source = assets_for_month(db, period.start.replace(day=1))
        return list(assets)

    if spec.special == "operation_items" and period.start and period.end:
        operation_query = _filter_query(DroneOperation, "operation_date", period)
        operation_ids = list(db.scalars(operation_query.with_only_columns(DroneOperation.id)).all())
        if not operation_ids:
            return []
        return list(db.scalars(select(DroneOperationItem).where(DroneOperationItem.operation_id.in_(operation_ids)).order_by(DroneOperationItem.id)).all())

    query = _filter_query(spec.model, None if spec.master else spec.date_field, period)
    primary_key = spec.model.__table__.primary_key.columns.values()[0]
    query = query.order_by(primary_key)
    return list(db.scalars(query).all())


def _sheet_title(value: str) -> str:
    invalid = set('[]:*?/\\')
    title = ''.join('_' if char in invalid else char for char in value)
    return title[:31]


def _friendly_header(name: str) -> str:
    return name.replace("_", " ").strip().title().replace("Uin", "UIN").replace("Hdd", "HDD").replace("Id", "ID")


def _write_table_sheet(workbook: Workbook, sheet_name: str, columns: list[str], rows: list[Any]) -> int:
    sheet = workbook.create_sheet(_sheet_title(sheet_name))
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="082A52")
    header_font = Font(color="FFFFFF", bold=True)
    thin_blue = Side(style="thin", color="D8E4EF")

    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(1, index, _friendly_header(column))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_blue)
    sheet.row_dimensions[1].height = 30

    for row_index, row in enumerate(rows, start=2):
        for column_index, column in enumerate(columns, start=1):
            value = _serialize(getattr(row, column, None))
            cell = sheet.cell(row_index, column_index, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
        if row_index % 2 == 0:
            for column_index in range(1, len(columns) + 1):
                sheet.cell(row_index, column_index).fill = PatternFill("solid", fgColor="F8FBFE")

    if columns:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    for index, column in enumerate(columns, start=1):
        max_length = len(_friendly_header(column))
        for row_index in range(2, min(len(rows) + 2, 202)):
            value = sheet.cell(row_index, index).value
            if value is not None:
                max_length = max(max_length, min(len(str(value)), 60))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 12), 45)
    return len(rows)


def _write_summary_sheet(workbook: Workbook, period: Period, scope: str, generated_by: str, row_counts: dict[str, int], warnings: list[str]) -> None:
    sheet = workbook.active
    sheet.title = "Backup Summary"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 31
    sheet.column_dimensions["B"].width = 68

    sheet.merge_cells("A1:B1")
    sheet["A1"] = "NakshaTech Asset Management Backup"
    sheet["A1"].fill = PatternFill("solid", fgColor="082A52")
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 36

    values = [
        ("Backup Type", period.backup_type.replace("_", " ").title()),
        ("Reporting Period", period.label),
        ("Period Start", period.start),
        ("Period End", period.end),
        ("Scope", scope.upper()),
        ("Generated At", local_now().replace(tzinfo=None)),
        ("Generated By", generated_by),
        ("Timezone", settings.backup_timezone),
        ("Master Sheet Rule", "Master sheets show current state; a finalized IT monthly snapshot is used when available."),
        ("Activity Sheet Rule", "Activity/history sheets are filtered to the selected reporting period."),
        ("Disaster Recovery", "Use the database dump for full restoration. Excel files are readable operational copies."),
    ]
    for row_number, (label, value) in enumerate(values, start=3):
        sheet.cell(row_number, 1, label).font = Font(bold=True, color="082A52")
        sheet.cell(row_number, 1).fill = PatternFill("solid", fgColor="EDF4FA")
        sheet.cell(row_number, 2, value)
        sheet.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(value, datetime):
            sheet.cell(row_number, 2).number_format = "yyyy-mm-dd hh:mm:ss"
        elif isinstance(value, date):
            sheet.cell(row_number, 2).number_format = "yyyy-mm-dd"

    start_row = len(values) + 5
    sheet.cell(start_row, 1, "Sheet").font = Font(bold=True, color="FFFFFF")
    sheet.cell(start_row, 2, "Rows Exported").font = Font(bold=True, color="FFFFFF")
    sheet.cell(start_row, 1).fill = PatternFill("solid", fgColor="126AE8")
    sheet.cell(start_row, 2).fill = PatternFill("solid", fgColor="126AE8")
    for offset, (name, count) in enumerate(row_counts.items(), start=1):
        sheet.cell(start_row + offset, 1, name)
        sheet.cell(start_row + offset, 2, count)

    if warnings:
        warning_row = start_row + len(row_counts) + 3
        sheet.cell(warning_row, 1, "Warnings / Notes").font = Font(bold=True, color="E49A18")
        for offset, warning in enumerate(warnings, start=1):
            sheet.cell(warning_row + offset, 1, "•")
            sheet.cell(warning_row + offset, 2, warning)
            sheet.cell(warning_row + offset, 2).alignment = Alignment(wrap_text=True)



def build_backup_workbook(db: Session, backup_type: str, scope: str, period_value: str | None, generated_by: str, include_admin_data: bool = False) -> tuple[BytesIO, Period, dict[str, int], list[str]]:
    if scope not in VALID_SCOPES:
        raise ValueError("Scope must be all, it or drone")
    period = resolve_period(backup_type, period_value)
    workbook = Workbook()
    row_counts: dict[str, int] = {}
    warnings: list[str] = []

    for spec in EXPORT_SPECS:
        if spec.admin_only and not include_admin_data:
            continue
        if scope != "all" and spec.scope != scope:
            continue
        if scope == "all" and spec.scope not in {"all", "it", "drone"}:
            continue
        rows = _rows_for_spec(db, spec, period)
        columns = _model_columns(spec.model, spec.exclude)
        count = _write_table_sheet(workbook, spec.sheet_name, columns, rows)
        row_counts[spec.sheet_name] = count

    if period.backup_type == "monthly" and period.start and period.start != month_start():
        _assets, source = assets_for_month(db, period.start)
        if source == "missing":
            warnings.append("No frozen IT asset snapshot was available for this month; the IT Asset Register sheet may be empty.")
        elif source == "template":
            warnings.append("The IT Asset Register sheet was reconstructed from the original historical workbook for this month.")

    _write_summary_sheet(workbook, period, scope, generated_by, row_counts, warnings)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream, period, row_counts, warnings


def backup_root() -> Path:
    root = Path(settings.backup_root).expanduser()
    root.mkdir(parents=True, exist_ok=True)
    return root.resolve()


def _period_directory(root: Path, period: Period) -> Path:
    if period.backup_type == "daily" and period.start:
        return root / "Daily" / str(period.start.year) / f"{period.start.month:02d}" / f"{period.start.day:02d}"
    if period.backup_type in {"monthly", "current_month"} and period.start:
        folder = "Current_Month" if period.backup_type == "current_month" else "Monthly"
        return root / folder / str(period.start.year) / f"{period.start.month:02d}"
    if period.backup_type == "yearly" and period.start:
        return root / "Yearly" / str(period.start.year)
    if period.backup_type == "financial_year":
        return root / "Financial_Year" / period.key
    return root / "Full"


def _filename_prefix(period: Period, scope: str) -> str:
    scope_label = {"all": "Complete", "it": "IT", "drone": "Drone"}[scope]
    type_label = {
        "daily": "Daily",
        "monthly": "Monthly",
        "current_month": "Current_Month_Live",
        "yearly": "Yearly",
        "financial_year": "Financial_Year",
        "full": "Full_Backup",
    }[period.backup_type]
    return f"NakshaTech_{scope_label}_{type_label}_{period.key}"


def _atomic_write(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_bytes(data)
    os.replace(temporary, path)


def file_checksum(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _database_dump(output_path: Path) -> tuple[Path | None, str | None]:
    if not settings.backup_database_enabled:
        return None, "Database dump disabled by configuration."
    url = make_url(settings.database_url)
    if url.drivername.startswith("sqlite"):
        database_path = Path(url.database or "")
        if not database_path.exists():
            return None, f"SQLite database file not found: {database_path}"
        destination = output_path.with_suffix(".sqlite3")
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(database_path, destination)
        return destination, None
    if not url.drivername.startswith("postgresql"):
        return None, f"Database dump is not implemented for {url.drivername}."

    binary = shutil.which(settings.pg_dump_bin) or (settings.pg_dump_bin if Path(settings.pg_dump_bin).exists() else None)
    if not binary:
        return None, f"pg_dump was not found. Configure PG_DUMP_BIN on the deployment server."

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(output_path.suffix + ".tmp")
    database_url = url.set(drivername="postgresql").render_as_string(hide_password=False)
    environment = os.environ.copy()
    if url.password:
        environment["PGPASSWORD"] = str(url.password)
    command = [
        str(binary),
        "--format=custom",
        "--no-owner",
        "--no-privileges",
        "--file",
        str(temporary),
        database_url,
    ]
    try:
        completed = subprocess.run(command, env=environment, capture_output=True, text=True, timeout=settings.backup_command_timeout_seconds, check=False)
    except Exception as exc:
        temporary.unlink(missing_ok=True)
        return None, f"Database dump command failed to start: {exc}"
    if completed.returncode != 0:
        temporary.unlink(missing_ok=True)
        detail = (completed.stderr or completed.stdout or "Unknown pg_dump error").strip()
        return None, f"pg_dump failed: {detail[:1000]}"
    os.replace(temporary, output_path)
    return output_path, None


def _minio_dump(output_path: Path) -> tuple[Path | None, str | None]:
    if not settings.backup_minio_enabled:
        return None, "MinIO backup disabled by configuration."
    try:
        from minio import Minio
    except Exception as exc:
        return None, f"MinIO client unavailable: {exc}"

    client = Minio(
        settings.minio_endpoint,
        access_key=settings.minio_root_user,
        secret_key=settings.minio_root_password,
        secure=settings.minio_secure,
    )
    try:
        if not client.bucket_exists(settings.minio_bucket):
            return None, f"MinIO bucket does not exist: {settings.minio_bucket}"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = output_path.with_suffix(output_path.suffix + ".tmp")
        with tarfile.open(temporary, "w:gz") as archive:
            for item in client.list_objects(settings.minio_bucket, recursive=True):
                response = client.get_object(settings.minio_bucket, item.object_name)
                try:
                    info = tarfile.TarInfo(name=item.object_name.lstrip("/"))
                    info.size = int(item.size or 0)
                    info.mtime = int(item.last_modified.timestamp()) if item.last_modified else int(datetime.now().timestamp())
                    archive.addfile(info, response)
                finally:
                    response.close()
                    response.release_conn()
        os.replace(temporary, output_path)
        return output_path, None
    except Exception as exc:
        output_path.with_suffix(output_path.suffix + ".tmp").unlink(missing_ok=True)
        return None, f"MinIO backup failed: {exc}"


def _lock_path() -> Path:
    return backup_root() / ".backup.lock"


def acquire_lock() -> Path:
    lock = _lock_path()
    if lock.exists():
        age = datetime.now().timestamp() - lock.stat().st_mtime
        if age > settings.backup_stale_lock_seconds:
            lock.unlink(missing_ok=True)
        else:
            raise RuntimeError("Another backup is already running.")
    try:
        descriptor = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        os.write(descriptor, f"pid={os.getpid()} started={datetime.now().isoformat()}".encode())
        os.close(descriptor)
    except FileExistsError as exc:
        raise RuntimeError("Another backup is already running.") from exc
    return lock


def _relative(root: Path, path: Path | None) -> str | None:
    if path is None:
        return None
    return str(path.resolve().relative_to(root.resolve())).replace("\\", "/")


def create_backup(
    db: Session,
    backup_type: str,
    scope: str,
    period_value: str | None,
    created_by: str,
    include_database: bool = True,
    include_minio: bool = False,
    include_admin_data: bool = False,
) -> BackupRun:
    period = resolve_period(backup_type, period_value)
    root = backup_root()
    lock = acquire_lock()
    code = f"BKP-{local_now().strftime('%Y%m%d%H%M%S%f')}-{backup_type.upper()}-{scope.upper()}"
    run = BackupRun(
        backup_code=code,
        backup_type=backup_type,
        scope=scope,
        status="running",
        period_start=period.start,
        period_end=period.end,
        created_by=created_by,
    )
    db.add(run)
    db.commit()
    db.refresh(run)

    warnings: list[str] = []
    try:
        stream, resolved_period, row_counts, workbook_warnings = build_backup_workbook(
            db, backup_type, scope, period_value, created_by, include_admin_data=include_admin_data
        )
        warnings.extend(workbook_warnings)
        directory = _period_directory(root, resolved_period)
        prefix = _filename_prefix(resolved_period, scope)
        timestamp = local_now().strftime("%Y%m%d_%H%M%S")
        excel_path = directory / "Excel" / f"{prefix}_{timestamp}.xlsx"
        _atomic_write(excel_path, stream.getvalue())

        database_path = None
        if include_database and scope == "all":
            database_target = directory / "Database" / f"NakshaTech_Database_{timestamp}.dump"
            database_path, warning = _database_dump(database_target)
            if warning:
                warnings.append(warning)

        minio_path = None
        if include_minio and scope == "all":
            minio_target = directory / "MinIO" / f"NakshaTech_MinIO_{timestamp}.tar.gz"
            minio_path, warning = _minio_dump(minio_target)
            if warning:
                warnings.append(warning)

        checksums = {"excel": file_checksum(excel_path)}
        if database_path:
            checksums["database"] = file_checksum(database_path)
        if minio_path:
            checksums["minio"] = file_checksum(minio_path)

        manifest = {
            "backup_code": code,
            "backup_type": backup_type,
            "scope": scope,
            "period": {"start": str(period.start) if period.start else None, "end": str(period.end) if period.end else None, "label": period.label},
            "created_by": created_by,
            "created_at": local_now().isoformat(),
            "files": {
                "excel": _relative(root, excel_path),
                "database": _relative(root, database_path),
                "minio": _relative(root, minio_path),
            },
            "checksums": checksums,
            "row_counts": row_counts,
            "warnings": warnings,
        }
        manifest_path = directory / "Manifests" / f"{prefix}_{timestamp}.json"
        _atomic_write(manifest_path, json.dumps(manifest, indent=2, ensure_ascii=False).encode("utf-8"))

        run.status = "completed_with_warnings" if warnings else "completed"
        run.excel_filename = _relative(root, excel_path)
        run.database_filename = _relative(root, database_path)
        run.minio_filename = _relative(root, minio_path)
        run.manifest_filename = _relative(root, manifest_path)
        run.excel_size_bytes = excel_path.stat().st_size
        run.database_size_bytes = database_path.stat().st_size if database_path else None
        run.minio_size_bytes = minio_path.stat().st_size if minio_path else None
        run.checksums = checksums
        run.row_counts = row_counts
        run.message = "\n".join(warnings) if warnings else "Backup completed successfully."
        run.completed_at = utc_now()
        db.commit()
        db.refresh(run)
        cleanup_retention(root)
        return run
    except Exception as exc:
        db.rollback()
        run = db.get(BackupRun, run.id)
        if run:
            run.status = "failed"
            run.message = str(exc)[:4000]
            run.completed_at = utc_now()
            db.commit()
            db.refresh(run)
        raise
    finally:
        lock.unlink(missing_ok=True)


def cleanup_retention(root: Path | None = None) -> None:
    root = root or backup_root()
    now = local_now()
    daily_root = root / "Daily"
    if daily_root.exists():
        excel_cutoff = now - timedelta(days=settings.backup_daily_retention_days)
        database_cutoff = now - timedelta(days=settings.backup_database_retention_days)
        for path in sorted(daily_root.rglob("*"), reverse=True):
            try:
                if path.is_file():
                    modified = datetime.fromtimestamp(path.stat().st_mtime, now.tzinfo)
                    is_database = "Database" in path.parts or path.suffix.lower() in {".dump", ".sqlite3"}
                    cutoff = database_cutoff if is_database else excel_cutoff
                    if modified < cutoff:
                        path.unlink(missing_ok=True)
                elif path.is_dir() and not any(path.iterdir()):
                    path.rmdir()
            except OSError:
                continue


def resolve_backup_file(run: BackupRun, file_type: str) -> Path:
    field_map = {
        "excel": run.excel_filename,
        "database": run.database_filename,
        "minio": run.minio_filename,
        "manifest": run.manifest_filename,
    }
    relative = field_map.get(file_type)
    if not relative:
        raise FileNotFoundError(f"No {file_type} file is available for this backup.")
    root = backup_root()
    path = (root / relative).resolve()
    if root not in path.parents and path != root:
        raise PermissionError("Invalid backup file path.")
    if not path.exists() or not path.is_file():
        raise FileNotFoundError("Backup file is missing from storage.")
    return path


def pg_dump_available() -> bool:
    return bool(shutil.which(settings.pg_dump_bin) or Path(settings.pg_dump_bin).exists())


def storage_writable() -> bool:
    try:
        root = backup_root()
        test_path = root / ".write-test"
        test_path.write_text("ok", encoding="utf-8")
        test_path.unlink(missing_ok=True)
        return True
    except OSError:
        return False


def run_scheduled_backups(db: Session, created_by: str = "cPanel cron / scheduler") -> list[BackupRun]:
    today = local_now().date()
    runs = [
        create_backup(db, "daily", "all", today.isoformat(), created_by, include_database=True, include_minio=settings.backup_minio_enabled, include_admin_data=True),
        create_backup(db, "current_month", "all", today.strftime("%Y-%m"), created_by, include_database=False, include_minio=False, include_admin_data=True),
    ]
    if today.day == 1:
        previous_day = today - timedelta(days=1)
        runs.append(create_backup(db, "monthly", "all", previous_day.strftime("%Y-%m"), created_by, include_database=False, include_minio=False, include_admin_data=True))
        if today.month == 1:
            runs.append(create_backup(db, "yearly", "all", str(today.year - 1), created_by, include_database=False, include_minio=False, include_admin_data=True))
        if today.month == 4:
            fy_start = today.year - 1
            runs.append(create_backup(db, "financial_year", "all", f"{fy_start}-{str(today.year)[-2:]}", created_by, include_database=False, include_minio=False, include_admin_data=True))
    return runs
