from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.entities import User, utc_now
from app.modules.finance.models import (
    ExpenseClaim,
    FinanceClient,
    FinanceClientMasterProfile,
    FinanceProject,
    FinanceProjectAssignment,
)
from app.modules.travel_km.models import TravelKmClaim

_ALLOWED_CODE = re.compile(r"^[A-Z0-9 ._/\-]+$")
_HEADER_ALIASES = {
    "vendor code": "vendor_code",
    "vendorcode": "vendor_code",
    "client code": "client_code",
    "clientcode": "client_code",
    "uav code": "client_code",
    "uavcode": "client_code",
    "client name": "client_name",
    "clientname": "client_name",
    "task": "task",
    "bd name": "bd_name",
    "bdname": "bd_name",
    "contact person": "contact_person",
    "contactperson": "contact_person",
    "contact person mail id": "contact_email",
    "contact person mailid": "contact_email",
    "contactpersonmailid": "contact_email",
    "country": "country",
}


def _text(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _header_key(value: Any) -> str:
    text = _text(value) or ""
    return re.sub(r"\s+", " ", text.lower()).strip()


def _email(value: Any) -> str | None:
    text = (_text(value) or "").lower()
    if not text:
        return None
    if "@" not in text or text.startswith("@") or text.endswith("@") or "." not in text.rsplit("@", 1)[-1]:
        return None
    return text[:255]


def parse_client_master_workbook(raw: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Parse CLIENT/UAV sheets without trusting row order beyond the header row."""
    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    seen_codes: set[str] = set()

    for ws in wb.worksheets:
        iterator = ws.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            continue
        mapping: dict[int, str] = {}
        for index, value in enumerate(header):
            alias = _HEADER_ALIASES.get(_header_key(value))
            if alias:
                mapping[index] = alias
        if "client_code" not in mapping.values():
            issues.append({"sheet": ws.title, "row": 1, "reason": "Client Code/UAV Code column not found"})
            continue

        client_type = "uav" if ws.title.strip().lower() == "uav" else "client"
        for row_number, values in enumerate(iterator, start=2):
            record: dict[str, Any] = {"source_sheet": ws.title, "client_type": client_type}
            for index, target in mapping.items():
                if index < len(values):
                    record[target] = _text(values[index])
            code = (record.get("client_code") or "").upper()
            if not code and not any(record.get(key) for key in ("client_name", "task", "bd_name", "contact_person", "contact_email", "country", "vendor_code")):
                continue
            if not code:
                issues.append({"sheet": ws.title, "row": row_number, "reason": "Missing Client Code"})
                continue
            if not _ALLOWED_CODE.match(code):
                issues.append({"sheet": ws.title, "row": row_number, "client_code": code, "reason": "Invalid Client Code characters"})
                continue
            if code in seen_codes:
                issues.append({"sheet": ws.title, "row": row_number, "client_code": code, "reason": "Duplicate Client Code inside workbook"})
                continue
            seen_codes.add(code)
            record["client_code"] = code
            record["contact_email"] = _email(record.get("contact_email"))
            record["row_number"] = row_number
            rows.append(record)
    return rows, issues


def import_client_master_workbook(
    db: Session,
    *,
    raw: bytes,
    actor: User | None,
    source_name: str,
    overwrite_existing: bool = False,
) -> dict[str, Any]:
    parsed, issues = parse_client_master_workbook(raw)
    created = 0
    updated = 0
    skipped_existing = 0
    skipped_incomplete = 0
    now = utc_now()

    for record in parsed:
        code = record["client_code"]
        name = record.get("client_name")
        if not name:
            skipped_incomplete += 1
            issues.append({
                "sheet": record.get("source_sheet"),
                "row": record.get("row_number"),
                "client_code": code,
                "reason": "Client Name is blank; row was not imported",
            })
            continue

        client = db.scalar(select(FinanceClient).where(func.lower(FinanceClient.client_code) == code.lower()).limit(1))
        if client is not None and not overwrite_existing:
            skipped_existing += 1
            continue

        contact_name = record.get("contact_person") or "Not provided"
        contact_email = record.get("contact_email")
        bd_name = record.get("bd_name")
        country = record.get("country") or "Not specified"
        source_person = bd_name or "Client Codes Excel Import"

        if client is None:
            client = FinanceClient(
                client_code=code,
                client_name=name[:255],
                primary_phone=None,
                client_email=contact_email,
                contact_person_name=contact_name[:255],
                contact_person_phone="Not provided",
                address=None,
                description=None,
                country=country[:100],
                gst_number=None,
                source_team="bd_team",
                source_person_name=source_person[:255],
                is_active=True,
                created_by_id=actor.id if actor else None,
                updated_by_id=actor.id if actor else None,
            )
            db.add(client)
            db.flush()
            client.master_profile = FinanceClientMasterProfile(
                client_id=client.id,
                vendor_code=(record.get("vendor_code") or None),
                client_type=record.get("client_type") or "client",
                task=record.get("task"),
                bd_name=bd_name,
                contact_person_email=contact_email,
                import_source=source_name[:255],
                imported_at=now,
                created_by_id=actor.id if actor else None,
                updated_by_id=actor.id if actor else None,
            )
            db.add(client.master_profile)
            created += 1
        else:
            client.client_name = name[:255]
            client.client_email = contact_email
            client.contact_person_name = contact_name[:255]
            client.country = country[:100]
            client.source_person_name = source_person[:255]
            client.updated_by_id = actor.id if actor else client.updated_by_id
            client.updated_at = now
            profile = client.master_profile or FinanceClientMasterProfile(client_id=client.id, created_by_id=actor.id if actor else None)
            profile.vendor_code = record.get("vendor_code") or None
            profile.client_type = record.get("client_type") or "client"
            profile.task = record.get("task")
            profile.bd_name = bd_name
            profile.contact_person_email = contact_email
            profile.import_source = source_name[:255]
            profile.imported_at = now
            profile.updated_by_id = actor.id if actor else profile.updated_by_id
            profile.updated_at = now
            if client.master_profile is None:
                client.master_profile = profile
                db.add(profile)
            updated += 1

    db.flush()
    return {
        "source_name": source_name,
        "parsed_rows": len(parsed),
        "created": created,
        "updated": updated,
        "skipped_existing": skipped_existing,
        "skipped_incomplete": skipped_incomplete,
        "issues": issues[:250],
        "issue_count": len(issues),
    }


def client_tracking(db: Session, client: FinanceClient) -> dict[str, Any]:
    project_ids = [project.id for project in client.projects]
    assignments = [] if not project_ids else list(db.scalars(
        select(FinanceProjectAssignment).where(
            FinanceProjectAssignment.project_id.in_(project_ids),
            FinanceProjectAssignment.is_active.is_(True),
        )
    ).all())
    travel = [] if not project_ids else list(db.scalars(
        select(TravelKmClaim).where(TravelKmClaim.project_id.in_(project_ids))
    ).all())
    expenses = [] if not project_ids else list(db.scalars(
        select(ExpenseClaim).where(ExpenseClaim.project_id.in_(project_ids))
    ).all())
    return {
        "project_count": len(project_ids),
        "active_project_count": sum(1 for p in client.projects if p.is_active),
        "assigned_employee_count": len({a.user_id for a in assignments}),
        "travel_claim_count": len(travel),
        "total_travel_km": round(sum(float(c.odometer_km or 0) for c in travel), 2),
        "total_travel_allowance": round(sum(float(c.final_allowance or c.calculated_allowance or 0) for c in travel), 2),
        "expense_claim_count": len(expenses),
        "expense_claim_value": round(sum(float(c.total_amount or 0) for c in expenses), 2),
    }


def project_tracking(db: Session, project: FinanceProject) -> dict[str, Any]:
    assignments = [a for a in project.assignments if a.is_active]
    travel = list(db.scalars(select(TravelKmClaim).where(TravelKmClaim.project_id == project.id)).all())
    expenses = list(db.scalars(select(ExpenseClaim).where(ExpenseClaim.project_id == project.id)).all())
    return {
        "assigned_employee_count": len({a.user_id for a in assignments}),
        "travel_claim_count": len(travel),
        "total_travel_km": round(sum(float(c.odometer_km or 0) for c in travel), 2),
        "total_travel_allowance": round(sum(float(c.final_allowance or c.calculated_allowance or 0) for c in travel), 2),
        "expense_claim_count": len(expenses),
        "expense_claim_value": round(sum(float(c.total_amount or 0) for c in expenses), 2),
    }


def _style_sheet(ws, headers: list[str]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for index, header in enumerate(headers, start=1):
        width = max(12, min(36, len(header) + 4))
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _append_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_sheet(ws, headers)


def build_crm_workbook(db: Session, *, client_id: int | None = None, project_id: int | None = None) -> bytes:
    clients = list(db.scalars(select(FinanceClient).order_by(FinanceClient.client_code.asc())).all())
    projects = list(db.scalars(select(FinanceProject).order_by(FinanceProject.project_code.asc())).all())
    if client_id is not None:
        clients = [c for c in clients if c.id == client_id]
        projects = [p for p in projects if p.client_id == client_id]
    if project_id is not None:
        projects = [p for p in projects if p.id == project_id]
        client_ids = {p.client_id for p in projects if p.client_id is not None}
        clients = [c for c in clients if c.id in client_ids]

    project_ids = {p.id for p in projects}
    assignments = [] if not project_ids else list(db.scalars(
        select(FinanceProjectAssignment).where(FinanceProjectAssignment.project_id.in_(project_ids)).order_by(FinanceProjectAssignment.project_id, FinanceProjectAssignment.id)
    ).all())
    travel = [] if not project_ids else list(db.scalars(
        select(TravelKmClaim).where(TravelKmClaim.project_id.in_(project_ids)).order_by(TravelKmClaim.travel_date, TravelKmClaim.id)
    ).all())
    expenses = [] if not project_ids else list(db.scalars(
        select(ExpenseClaim).where(ExpenseClaim.project_id.in_(project_ids)).order_by(ExpenseClaim.created_at, ExpenseClaim.id)
    ).all())

    wb = Workbook()
    wb.remove(wb.active)

    client_rows = []
    for client in clients:
        profile = client.master_profile
        metrics = client_tracking(db, client)
        client_rows.append([
            client.client_code,
            profile.vendor_code if profile else None,
            client.client_name,
            profile.client_type if profile else "client",
            profile.task if profile else None,
            profile.bd_name if profile else None,
            client.contact_person_name,
            profile.contact_person_email if profile else client.client_email,
            client.country,
            "Active" if client.is_active else "Inactive",
            metrics["project_count"], metrics["active_project_count"], metrics["assigned_employee_count"],
            metrics["travel_claim_count"], metrics["total_travel_km"], metrics["total_travel_allowance"],
            metrics["expense_claim_count"], metrics["expense_claim_value"],
            profile.import_source if profile else None,
            client.created_at, client.updated_at,
        ])
    _append_sheet(wb, "Client Summary", [
        "Client ID / Code", "Vendor Code", "Client Name", "Client Type", "Task / Scope", "BD Name",
        "Contact Person", "Contact Person Mail ID", "Country", "Status", "Total Projects", "Active Projects",
        "Assigned Employees", "Travel Claims", "Total Travel KM", "Travel Allowance (INR)", "Expense Claims",
        "Expense Claim Value (INR)", "Import Source", "Created At", "Updated At",
    ], client_rows)

    project_rows = []
    for project in projects:
        profile = project.master_profile
        metrics = project_tracking(db, project)
        project_rows.append([
            project.project_code, project.project_name,
            project.client.client_code if project.client else None,
            project.client.client_name if project.client else project.client_name,
            profile.task if profile else None,
            profile.project_status if profile else ("active" if project.is_active else "inactive"),
            profile.project_manager.full_name if profile and profile.project_manager else None,
            profile.reporting_manager.full_name if profile and profile.reporting_manager else None,
            project.start_date, project.end_date,
            metrics["assigned_employee_count"], metrics["travel_claim_count"], metrics["total_travel_km"],
            metrics["total_travel_allowance"], metrics["expense_claim_count"], metrics["expense_claim_value"],
            project.created_at, project.updated_at,
        ])
    _append_sheet(wb, "Project Summary", [
        "Project ID / Number", "Project Name", "Client ID / Code", "Client Name", "Task / Scope", "Status",
        "Project Manager", "Reporting Manager", "Start Date", "End Date", "Assigned Employees", "Travel Claims",
        "Total Travel KM", "Travel Allowance (INR)", "Expense Claims", "Expense Claim Value (INR)", "Created At", "Updated At",
    ], project_rows)

    assignment_rows = []
    for assignment in assignments:
        user = assignment.user
        project = assignment.project
        assignment_rows.append([
            project.project_code if project else None,
            project.project_name if project else None,
            project.client.client_name if project and project.client else None,
            user.employee_id if user else None,
            user.full_name if user else None,
            user.email if user else None,
            user.department if user else None,
            "Active" if assignment.is_active else "Inactive",
            assignment.created_at, assignment.updated_at,
        ])
    _append_sheet(wb, "Employee Allocation", [
        "Project ID / Number", "Project Name", "Client Name", "Employee ID", "Employee Name", "Employee Email",
        "Department", "Assignment Status", "Assigned At", "Updated At",
    ], assignment_rows)

    project_by_id = {p.id: p for p in projects}
    travel_rows = []
    for claim in travel:
        project = project_by_id.get(claim.project_id)
        travel_rows.append([
            claim.claim_code, claim.travel_date,
            project.project_code if project else claim.project_code_snapshot,
            project.project_name if project else claim.project_name_snapshot,
            project.client.client_name if project and project.client else claim.client_name_snapshot,
            claim.requester_id, claim.purpose_description,
            float(claim.start_km or 0), float(claim.end_km) if claim.end_km is not None else None,
            float(claim.odometer_km) if claim.odometer_km is not None else None,
            float(claim.final_eligible_km) if claim.final_eligible_km is not None else None,
            float(claim.final_allowance or claim.calculated_allowance or 0), claim.status,
            claim.submitted_at, claim.created_at,
        ])
    _append_sheet(wb, "Travel KM", [
        "Claim Code", "Travel Date", "Project ID / Number", "Project Name", "Client Name", "Employee User ID",
        "Purpose", "Start KM", "End KM", "Odometer KM", "Final Eligible KM", "Allowance (INR)", "Status",
        "Submitted At", "Created At",
    ], travel_rows)

    expense_rows = []
    for claim in expenses:
        project = project_by_id.get(claim.project_id)
        expense_rows.append([
            claim.claim_code,
            project.project_code if project else None,
            project.project_name if project else None,
            project.client.client_name if project and project.client else None,
            claim.requester_id, claim.claim_type, claim.purpose_description,
            float(claim.total_amount or 0), claim.status, claim.submitted_at, claim.created_at,
        ])
    _append_sheet(wb, "Expense Claims", [
        "Claim Code", "Project ID / Number", "Project Name", "Client Name", "Employee User ID", "Claim Type",
        "Purpose", "Amount (INR)", "Status", "Submitted At", "Created At",
    ], expense_rows)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
