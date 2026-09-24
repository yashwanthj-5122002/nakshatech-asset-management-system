from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
import math
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, case, func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.models.entities import User, utc_now
from app.modules.finance.models import (
    ExpenseClaim,
    ExpenseClaimAttachment,
    ExpenseClaimItem,
    ExpenseClaimPayment,
    ExpenseSettlement,
    ExpenseSettlementAttachment,
    FinanceProject,
)
from app.modules.finance.service import (
    CLAIM_TYPE_LABELS,
    STATUS_LABELS,
    approved_amount,
    claim_paid_amount,
    money,
    remaining_amount,
)
from app.modules.finance.visibility import (
    hidden_client_ids,
    hidden_project_ids,
)

MAX_REPORT_ROWS = 50_000
REPORT_PAGE_SIZE_MAX = 250


@dataclass(frozen=True)
class ResolvedPeriod:
    period: str
    label: str
    start_date: date | None
    end_date: date | None
    start_datetime: datetime | None
    end_exclusive_datetime: datetime | None


def resolve_report_period(period: str, *, year: int | None = None, month: int | None = None, quarter: int | None = None) -> ResolvedPeriod:
    normalized = (period or "month").strip().lower()
    if normalized not in {"month", "quarter", "year", "all"}:
        raise ValueError("Report period must be month, quarter, year, or all")
    if normalized == "all":
        return ResolvedPeriod("all", "All Time", None, None, None, None)

    current = utc_now().date()
    resolved_year = year or current.year
    if resolved_year < 2000 or resolved_year > 2100:
        raise ValueError("Report year must be between 2000 and 2100")

    if normalized == "month":
        resolved_month = month or (current.month if resolved_year == current.year else 1)
        if resolved_month < 1 or resolved_month > 12:
            raise ValueError("Report month must be between 1 and 12")
        start = date(resolved_year, resolved_month, 1)
        next_month = date(resolved_year + 1, 1, 1) if resolved_month == 12 else date(resolved_year, resolved_month + 1, 1)
        end = next_month - timedelta(days=1)
        label = start.strftime("%B %Y")
    elif normalized == "quarter":
        resolved_quarter = quarter or (((current.month - 1) // 3 + 1) if resolved_year == current.year else 1)
        if resolved_quarter < 1 or resolved_quarter > 4:
            raise ValueError("Report quarter must be between 1 and 4")
        first_month = (resolved_quarter - 1) * 3 + 1
        start = date(resolved_year, first_month, 1)
        next_start = date(resolved_year + 1, 1, 1) if resolved_quarter == 4 else date(resolved_year, first_month + 3, 1)
        end = next_start - timedelta(days=1)
        label = f"Q{resolved_quarter} {resolved_year}"
    else:
        start = date(resolved_year, 1, 1)
        end = date(resolved_year, 12, 31)
        label = str(resolved_year)

    start_dt = datetime.combine(start, time.min)
    end_exclusive = datetime.combine(end + timedelta(days=1), time.min)
    return ResolvedPeriod(normalized, label, start, end, start_dt, end_exclusive)


def available_finance_years(db: Session) -> list[int]:
    values = db.scalars(
        select(func.coalesce(ExpenseClaim.submitted_at, ExpenseClaim.created_at)).where(ExpenseClaim.status != "draft")
    ).all()
    years = sorted({value.year for value in values if value is not None}, reverse=True)
    if not years:
        years = [utc_now().year]
    return years


def _claim_period_condition(period: ResolvedPeriod):
    if period.start_datetime is None or period.end_exclusive_datetime is None:
        return None
    return or_(
        and_(
            ExpenseClaim.submitted_at.is_not(None),
            ExpenseClaim.submitted_at >= period.start_datetime,
            ExpenseClaim.submitted_at < period.end_exclusive_datetime,
        ),
        and_(
            ExpenseClaim.submitted_at.is_(None),
            ExpenseClaim.created_at >= period.start_datetime,
            ExpenseClaim.created_at < period.end_exclusive_datetime,
        ),
    )


def _claim_filter_conditions(
    *,
    period: ResolvedPeriod,
    include_claim_period: bool = True,
    project_id: int | None = None,
    claim_type: str | None = None,
    status: str | None = None,
    employee_query: str | None = None,
    category: str | None = None,
    search: str | None = None,
):
    conditions = [ExpenseClaim.status != "draft"]
    if include_claim_period:
        period_condition = _claim_period_condition(period)
        if period_condition is not None:
            conditions.append(period_condition)
    if project_id is not None:
        conditions.append(ExpenseClaim.project_id == project_id)
    if claim_type:
        conditions.append(ExpenseClaim.claim_type == claim_type.strip().lower())
    if status:
        conditions.append(ExpenseClaim.status == status.strip().lower())
    if employee_query and employee_query.strip():
        needle = f"%{employee_query.strip().lower()}%"
        conditions.append(or_(func.lower(User.full_name).like(needle), func.lower(User.email).like(needle)))
    if category and category.strip():
        normalized_category = category.strip().lower()
        conditions.append(
            ExpenseClaim.items.any(or_(
                func.lower(ExpenseClaimItem.category) == normalized_category,
                func.lower(ExpenseClaimItem.other_category) == normalized_category,
            ))
        )
    if search and search.strip():
        needle = f"%{search.strip().lower()}%"
        conditions.append(or_(
            func.lower(ExpenseClaim.claim_code).like(needle),
            func.lower(User.full_name).like(needle),
            func.lower(User.email).like(needle),
            func.lower(FinanceProject.project_code).like(needle),
            func.lower(FinanceProject.project_name).like(needle),
            func.lower(ExpenseClaim.purpose_description).like(needle),
        ))
    return conditions


def _hidden_project_conditions(db: Session):
    project_ids = hidden_project_ids(db)
    client_ids = hidden_client_ids(db)
    conditions = []
    if project_ids:
        conditions.append(FinanceProject.id.not_in(project_ids))
    if client_ids:
        conditions.append(or_(
            FinanceProject.client_id.is_(None),
            FinanceProject.client_id.not_in(client_ids),
        ))
    return conditions


def _filtered_claim_query(
    *,
    db: Session,
    period: ResolvedPeriod,
    project_id: int | None = None,
    claim_type: str | None = None,
    status: str | None = None,
    employee_query: str | None = None,
    category: str | None = None,
    search: str | None = None,
):
    report_date = func.coalesce(ExpenseClaim.submitted_at, ExpenseClaim.created_at)
    conditions = _claim_filter_conditions(
        period=period,
        project_id=project_id,
        claim_type=claim_type,
        status=status,
        employee_query=employee_query,
        category=category,
        search=search,
    )
    conditions.extend(_hidden_project_conditions(db))
    return (
        select(ExpenseClaim)
        .join(User, User.id == ExpenseClaim.requester_id)
        .join(FinanceProject, FinanceProject.id == ExpenseClaim.project_id)
        .where(*conditions)
        .options(
            selectinload(ExpenseClaim.project),
            selectinload(ExpenseClaim.items),
            selectinload(ExpenseClaim.attachments),
            selectinload(ExpenseClaim.events),
            selectinload(ExpenseClaim.payments),
        )
        .order_by(report_date.desc(), ExpenseClaim.id.desc())
    )


def filtered_finance_claims(db: Session, **filters) -> list[ExpenseClaim]:
    query = _filtered_claim_query(db=db, **filters)
    return list(db.scalars(query).unique().all())


def _payment_filter_conditions(
    *,
    period: ResolvedPeriod,
    project_id: int | None = None,
    claim_type: str | None = None,
    status: str | None = None,
    employee_query: str | None = None,
    category: str | None = None,
    search: str | None = None,
):
    conditions = _claim_filter_conditions(
        period=period,
        include_claim_period=False,
        project_id=project_id,
        claim_type=claim_type,
        status=status,
        employee_query=employee_query,
        category=category,
        search=search,
    )
    if period.start_date is not None:
        conditions.append(ExpenseClaimPayment.payment_date >= period.start_date)
    if period.end_date is not None:
        conditions.append(ExpenseClaimPayment.payment_date <= period.end_date)
    return conditions


def filtered_finance_payments(
    db: Session,
    *,
    period: ResolvedPeriod,
    project_id: int | None = None,
    claim_type: str | None = None,
    status: str | None = None,
    employee_query: str | None = None,
    category: str | None = None,
    search: str | None = None,
) -> list[ExpenseClaimPayment]:
    conditions = _payment_filter_conditions(
        period=period,
        project_id=project_id,
        claim_type=claim_type,
        status=status,
        employee_query=employee_query,
        category=category,
        search=search,
    )
    conditions.extend(_hidden_project_conditions(db))
    query = (
        select(ExpenseClaimPayment)
        .join(ExpenseClaim, ExpenseClaim.id == ExpenseClaimPayment.claim_id)
        .join(User, User.id == ExpenseClaim.requester_id)
        .join(FinanceProject, FinanceProject.id == ExpenseClaim.project_id)
        .where(*conditions)
        .options(
            selectinload(ExpenseClaimPayment.claim).selectinload(ExpenseClaim.project),
        )
        .order_by(ExpenseClaimPayment.payment_date.asc(), ExpenseClaimPayment.id.asc())
    )
    return list(db.scalars(query).unique().all())


def claim_integrity_issues(claim: ExpenseClaim) -> list[str]:
    issues: list[str] = []
    item_total = money(sum((money(item.amount) for item in claim.items), Decimal("0.00")))
    requested = money(claim.total_amount)
    approved = approved_amount(claim)
    paid = claim_paid_amount(claim)
    remaining = remaining_amount(claim)

    if item_total != requested:
        issues.append(f"Item total INR {float(item_total):,.2f} does not match claim total INR {float(requested):,.2f}")
    if claim.status != "draft" and claim.submitted_at is None:
        issues.append("Submitted workflow record is missing submitted_at")
    if claim.status in {"finance_approved", "partially_paid", "paid"} and claim.finance_decision_at is None:
        issues.append("Finance-approved workflow record is missing finance_decision_at")
    if approved > requested:
        issues.append("Approved amount is greater than requested amount")
    if paid > approved and approved > 0:
        issues.append("Paid amount is greater than approved amount")
    if claim.status == "paid" and remaining > 0:
        issues.append("Claim is marked paid but still has an outstanding balance")
    if claim.status == "partially_paid" and (paid <= 0 or remaining <= 0):
        issues.append("Partial-payment status does not match paid/outstanding balances")
    if claim.claim_type in {"reimbursement", "additional_advance"} and not claim.attachments:
        issues.append("Required reimbursement/additional-advance proof is missing")
    if claim.requested_work_start_date and claim.requested_work_end_date and claim.requested_work_end_date < claim.requested_work_start_date:
        issues.append("Requested work end date is earlier than work start date")
    if claim.approved_work_start_date and claim.approved_work_end_date and claim.approved_work_end_date < claim.approved_work_start_date:
        issues.append("Approved work end date is earlier than work start date")
    if claim.settlement_due_date and claim.approved_work_end_date and claim.settlement_due_date < claim.approved_work_end_date:
        issues.append("Settlement due date is earlier than approved work end date")
    if claim.claim_type == "additional_advance" and claim.parent_advance_claim_id is None:
        issues.append("Additional advance is not linked to an original Advance Request")
    if claim.claim_type == "advance" and claim.settlement and claim.settlement.status == "finance_finalized":
        if money(claim.settlement.total_advance_received) != money(claim.settlement.total_expense_amount):
            issues.append("Finalized settlement still has an advance/expense difference")
    return issues


def _breakdown(rows: dict[str, dict[str, Decimal | int | str]]) -> list[dict]:
    result = []
    for key, values in rows.items():
        result.append({
            "key": key,
            "label": str(values.get("label") or key.replace("_", " ").title()),
            "amount": float(values.get("amount") or Decimal("0.00")),
            "count": int(values.get("count") or 0),
        })
    return sorted(result, key=lambda row: (-row["amount"], row["label"]))


def finance_report_payload(
    db: Session,
    *,
    period: str,
    year: int | None = None,
    month: int | None = None,
    quarter: int | None = None,
    project_id: int | None = None,
    claim_type: str | None = None,
    status: str | None = None,
    employee_query: str | None = None,
    category: str | None = None,
    search: str | None = None,
    page: int = 1,
    page_size: int = 50,
) -> dict:
    resolved = resolve_report_period(period, year=year, month=month, quarter=quarter)
    filter_kwargs = {
        "period": resolved,
        "project_id": project_id,
        "claim_type": claim_type,
        "status": status,
        "employee_query": employee_query,
        "category": category,
        "search": search,
    }
    conditions = _claim_filter_conditions(**filter_kwargs)
    conditions.extend(_hidden_project_conditions(db))

    payment_sum = (
        select(
            ExpenseClaimPayment.claim_id.label("claim_id"),
            func.sum(ExpenseClaimPayment.amount).label("ledger_paid"),
        )
        .group_by(ExpenseClaimPayment.claim_id)
        .subquery()
    )
    approved_expr = case(
        (
            ExpenseClaim.status.in_(("finance_approved", "partially_paid", "paid")),
            func.coalesce(ExpenseClaim.finance_approved_amount, ExpenseClaim.total_amount),
        ),
        else_=Decimal("0.00"),
    )
    paid_expr = func.coalesce(payment_sum.c.ledger_paid, ExpenseClaim.paid_amount, Decimal("0.00"))
    outstanding_expr = case(
        (
            and_(
                ExpenseClaim.status.in_(("finance_approved", "partially_paid", "paid")),
                approved_expr > paid_expr,
            ),
            approved_expr - paid_expr,
        ),
        else_=Decimal("0.00"),
    )
    pending_admin_expr = case(
        (ExpenseClaim.status == "submitted", ExpenseClaim.total_amount),
        else_=Decimal("0.00"),
    )
    pending_finance_expr = case(
        (ExpenseClaim.status == "admin_approved", ExpenseClaim.total_amount),
        else_=Decimal("0.00"),
    )

    totals = db.execute(
        select(
            func.count(ExpenseClaim.id).label("record_count"),
            func.sum(ExpenseClaim.total_amount).label("requested"),
            func.sum(approved_expr).label("approved"),
            func.sum(paid_expr).label("paid"),
            func.sum(outstanding_expr).label("outstanding"),
            func.sum(pending_admin_expr).label("pending_admin"),
            func.sum(pending_finance_expr).label("pending_finance"),
        )
        .select_from(ExpenseClaim)
        .join(User, User.id == ExpenseClaim.requester_id)
        .join(FinanceProject, FinanceProject.id == ExpenseClaim.project_id)
        .outerjoin(payment_sum, payment_sum.c.claim_id == ExpenseClaim.id)
        .where(*conditions)
    ).one()

    total_records = int(totals.record_count or 0)
    requested_total = money(totals.requested)
    approved_total = money(totals.approved)
    paid_total = money(totals.paid)
    outstanding_total = money(totals.outstanding)
    pending_admin_total = money(totals.pending_admin)
    pending_finance_total = money(totals.pending_finance)

    payment_conditions = _payment_filter_conditions(**filter_kwargs)
    payment_conditions.extend(_hidden_project_conditions(db))
    payment_period_row = db.execute(
        select(
            func.count(ExpenseClaimPayment.id).label("payment_count"),
            func.sum(ExpenseClaimPayment.amount).label("payment_amount"),
        )
        .select_from(ExpenseClaimPayment)
        .join(ExpenseClaim, ExpenseClaim.id == ExpenseClaimPayment.claim_id)
        .join(User, User.id == ExpenseClaim.requester_id)
        .join(FinanceProject, FinanceProject.id == ExpenseClaim.project_id)
        .where(*payment_conditions)
    ).one()
    payment_period_count = int(payment_period_row.payment_count or 0)
    payment_period_amount = money(payment_period_row.payment_amount)

    project_rows = db.execute(
        select(
            FinanceProject.project_code,
            FinanceProject.project_name,
            func.sum(ExpenseClaim.total_amount).label("amount"),
            func.count(ExpenseClaim.id).label("count"),
        )
        .select_from(ExpenseClaim)
        .join(User, User.id == ExpenseClaim.requester_id)
        .join(FinanceProject, FinanceProject.id == ExpenseClaim.project_id)
        .where(*conditions)
        .group_by(FinanceProject.project_code, FinanceProject.project_name)
    ).all()
    by_project = sorted([
        {
            "key": row.project_code,
            "label": f"{row.project_code} \u00b7 {row.project_name}",
            "amount": float(money(row.amount)),
            "count": int(row.count or 0),
        }
        for row in project_rows
    ], key=lambda row: (-row["amount"], row["label"]))

    employee_rows = db.execute(
        select(
            User.email,
            User.full_name,
            func.sum(ExpenseClaim.total_amount).label("amount"),
            func.count(ExpenseClaim.id).label("count"),
        )
        .select_from(ExpenseClaim)
        .join(User, User.id == ExpenseClaim.requester_id)
        .join(FinanceProject, FinanceProject.id == ExpenseClaim.project_id)
        .where(*conditions)
        .group_by(User.email, User.full_name)
    ).all()
    by_employee = sorted([
        {
            "key": row.email.lower(),
            "label": row.full_name,
            "amount": float(money(row.amount)),
            "count": int(row.count or 0),
        }
        for row in employee_rows
    ], key=lambda row: (-row["amount"], row["label"]))

    type_rows = db.execute(
        select(
            ExpenseClaim.claim_type,
            func.sum(ExpenseClaim.total_amount).label("amount"),
            func.count(ExpenseClaim.id).label("count"),
        )
        .select_from(ExpenseClaim)
        .join(User, User.id == ExpenseClaim.requester_id)
        .join(FinanceProject, FinanceProject.id == ExpenseClaim.project_id)
        .where(*conditions)
        .group_by(ExpenseClaim.claim_type)
    ).all()
    by_type = sorted([
        {
            "key": row.claim_type,
            "label": CLAIM_TYPE_LABELS.get(row.claim_type, row.claim_type.replace("_", " ").title()),
            "amount": float(money(row.amount)),
            "count": int(row.count or 0),
        }
        for row in type_rows
    ], key=lambda row: (-row["amount"], row["label"]))

    status_rows = db.execute(
        select(
            ExpenseClaim.status,
            func.sum(ExpenseClaim.total_amount).label("amount"),
            func.count(ExpenseClaim.id).label("count"),
        )
        .select_from(ExpenseClaim)
        .join(User, User.id == ExpenseClaim.requester_id)
        .join(FinanceProject, FinanceProject.id == ExpenseClaim.project_id)
        .where(*conditions)
        .group_by(ExpenseClaim.status)
    ).all()
    by_status = sorted([
        {
            "key": row.status,
            "label": STATUS_LABELS.get(row.status, row.status.replace("_", " ").title()),
            "amount": float(money(row.amount)),
            "count": int(row.count or 0),
        }
        for row in status_rows
    ], key=lambda row: (-row["amount"], row["label"]))

    selected_ids = (
        select(ExpenseClaim.id.label("claim_id"))
        .join(User, User.id == ExpenseClaim.requester_id)
        .join(FinanceProject, FinanceProject.id == ExpenseClaim.project_id)
        .where(*conditions)
        .subquery()
    )
    category_key = case(
        (
            and_(
                func.lower(ExpenseClaimItem.category) == "other",
                ExpenseClaimItem.other_category.is_not(None),
                ExpenseClaimItem.other_category != "",
            ),
            ExpenseClaimItem.other_category,
        ),
        else_=ExpenseClaimItem.category,
    )
    category_rows = db.execute(
        select(
            category_key.label("category_key"),
            func.sum(ExpenseClaimItem.amount).label("amount"),
            func.count(ExpenseClaimItem.id).label("count"),
        )
        .select_from(ExpenseClaimItem)
        .join(selected_ids, selected_ids.c.claim_id == ExpenseClaimItem.claim_id)
        .group_by(category_key)
    ).all()
    by_category = sorted([
        {
            "key": (row.category_key or "uncategorized"),
            "label": (row.category_key or "uncategorized").replace("_", " ").title(),
            "amount": float(money(row.amount)),
            "count": int(row.count or 0),
        }
        for row in category_rows
    ], key=lambda row: (-row["amount"], row["label"]))

    item_totals = (
        select(ExpenseClaimItem.claim_id.label("claim_id"), func.sum(ExpenseClaimItem.amount).label("item_total"))
        .group_by(ExpenseClaimItem.claim_id)
        .subquery()
    )
    attachment_counts = (
        select(
            ExpenseClaimAttachment.claim_id.label("claim_id"),
            func.count(ExpenseClaimAttachment.id).label("attachment_count"),
        )
        .group_by(ExpenseClaimAttachment.claim_id)
        .subquery()
    )
    integrity_rows = db.execute(
        select(
            ExpenseClaim.id,
            ExpenseClaim.claim_type,
            ExpenseClaim.status,
            ExpenseClaim.total_amount,
            ExpenseClaim.finance_approved_amount,
            ExpenseClaim.paid_amount,
            ExpenseClaim.submitted_at,
            ExpenseClaim.finance_decision_at,
            func.coalesce(item_totals.c.item_total, Decimal("0.00")).label("item_total"),
            func.coalesce(attachment_counts.c.attachment_count, 0).label("attachment_count"),
            func.coalesce(payment_sum.c.ledger_paid, ExpenseClaim.paid_amount, Decimal("0.00")).label("paid_total"),
        )
        .select_from(ExpenseClaim)
        .join(User, User.id == ExpenseClaim.requester_id)
        .join(FinanceProject, FinanceProject.id == ExpenseClaim.project_id)
        .outerjoin(item_totals, item_totals.c.claim_id == ExpenseClaim.id)
        .outerjoin(attachment_counts, attachment_counts.c.claim_id == ExpenseClaim.id)
        .outerjoin(payment_sum, payment_sum.c.claim_id == ExpenseClaim.id)
        .where(*conditions)
    ).all()
    integrity_count = 0
    approved_statuses = {"finance_approved", "partially_paid", "paid"}
    for row in integrity_rows:
        requested = money(row.total_amount)
        item_total = money(row.item_total)
        approved = money(row.finance_approved_amount) if row.finance_approved_amount is not None else (requested if row.status in approved_statuses else Decimal("0.00"))
        paid = money(row.paid_total)
        remaining = max(Decimal("0.00"), approved - paid)
        has_issue = (
            item_total != requested
            or (row.status != "draft" and row.submitted_at is None)
            or (row.status in approved_statuses and row.finance_decision_at is None)
            or approved > requested
            or (approved > 0 and paid > approved)
            or (row.status == "paid" and remaining > 0)
            or (row.status == "partially_paid" and (paid <= 0 or remaining <= 0))
            or (row.claim_type in {"reimbursement", "additional_advance"} and int(row.attachment_count or 0) == 0)
        )
        if has_issue:
            integrity_count += 1

    page_size = max(1, min(REPORT_PAGE_SIZE_MAX, int(page_size)))
    page = max(1, int(page))
    total_pages = max(1, math.ceil(total_records / page_size)) if total_records else 1
    if page > total_pages:
        page = total_pages
    start_index = (page - 1) * page_size
    visible_claims = list(db.scalars(
        _filtered_claim_query(db=db, **filter_kwargs).offset(start_index).limit(page_size)
    ).unique().all())

    requester_cache: dict[int, User | None] = {}
    claim_rows = []
    for claim in visible_claims:
        requester = requester_cache.setdefault(claim.requester_id, db.get(User, claim.requester_id))
        claim_rows.append({
            "id": claim.id,
            "claim_code": claim.claim_code,
            "submitted_at": claim.submitted_at,
            "requester_name": requester.full_name if requester else "Unknown employee",
            "requester_email": requester.email if requester else "unknown@nakshatech.com",
            "requester_department": requester.department if requester else None,
            "project_code": claim.project.project_code,
            "project_name": claim.project.project_name,
            "claim_type": claim.claim_type,
            "status": claim.status,
            "requested_amount": float(money(claim.total_amount)),
            "approved_amount": float(approved_amount(claim)),
            "paid_amount": float(claim_paid_amount(claim)),
            "outstanding_amount": float(remaining_amount(claim)),
            "attachment_count": len(claim.attachments),
            "payment_count": len(claim.payments) if claim.payments else (1 if claim.paid_amount else 0),
            "updated_at": claim.updated_at,
        })

    return {
        "period": resolved.period,
        "period_label": resolved.label,
        "start_date": resolved.start_date,
        "end_date": resolved.end_date,
        "available_years": available_finance_years(db),
        "total_records": total_records,
        "page": page,
        "page_size": page_size,
        "total_pages": total_pages,
        "requested_amount": float(requested_total),
        "approved_amount": float(approved_total),
        "paid_amount": float(paid_total),
        "outstanding_amount": float(outstanding_total),
        "payment_period_amount": float(payment_period_amount),
        "payment_period_count": payment_period_count,
        "pending_admin_amount": float(pending_admin_total),
        "pending_finance_amount": float(pending_finance_total),
        "integrity_issue_count": integrity_count,
        "by_project": by_project,
        "by_employee": by_employee,
        "by_category": by_category,
        "by_type": by_type,
        "by_status": by_status,
        "claims": claim_rows,
    }


_HEADER_FILL = PatternFill("solid", fgColor="0B2D4F")
_SECTION_FILL = PatternFill("solid", fgColor="102E4B")
_SUBTLE_FILL = PatternFill("solid", fgColor="EAF4FB")
_WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=18, bold=True, color="0B2D4F")
_SUBTITLE_FONT = Font(size=10, color="52677A")
_THIN_GRAY = Side(style="thin", color="CBD5E1")
_MONEY_FORMAT = '₹#,##0.00;[Red](₹#,##0.00);-'
_DATE_FORMAT = "dd-mmm-yyyy"
_DATETIME_FORMAT = "dd-mmm-yyyy hh:mm"


def _safe_sheet_title(value: str) -> str:
    return value[:31]


def _excel_text(value: object | None) -> str:
    """Write untrusted text without allowing Excel formula injection."""
    text_value = "" if value is None else str(value)
    if text_value.startswith(("=", "+", "-", "@")):
        return "'" + text_value
    return text_value


def _write_header(ws, row: int, columns: Iterable[str]) -> None:
    for col, value in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=_THIN_GRAY)
    ws.row_dimensions[row].height = 28


def _auto_width(ws, *, min_width: int = 10, max_width: int = 42) -> None:
    for column_cells in ws.columns:
        width = min_width
        letter = get_column_letter(column_cells[0].column)
        for cell in column_cells[:250]:
            if cell.value is None:
                continue
            width = max(width, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def _format_money_column(ws, column: int, start_row: int, end_row: int) -> None:
    for row in range(start_row, end_row + 1):
        ws.cell(row=row, column=column).number_format = _MONEY_FORMAT


def _claims_ledger_sheet(wb: Workbook, db: Session, claims: list[ExpenseClaim]) -> int:
    ws = wb.create_sheet("Claims Ledger")
    columns = [
        "Claim Code", "Submitted At", "Employee", "Employee Email", "Department", "Project ID", "Project Name", "Client",
        "Request Type", "Purpose", "Currency", "Requested Amount", "Approved Amount", "Paid Amount", "Outstanding Amount", "Status",
        "Admin Verified By", "Admin Decision At", "Admin Comments", "Finance Verified By", "Finance Decision At", "Finance Comments",
        "Attachments", "Payments", "Last Updated", "Integrity Flag",
        "Employee Work Start", "Employee Work End", "Approved Work Start", "Approved Work End",
        "Settlement Due", "Settlement Status", "Parent Advance", "Settlement Tally", "Settlement Expense Total",
    ]
    _write_header(ws, 1, columns)
    for row_idx, claim in enumerate(claims, start=2):
        requester = db.get(User, claim.requester_id)
        issues = claim_integrity_issues(claim)
        values = [
            _excel_text(claim.claim_code),
            claim.submitted_at,
            _excel_text(requester.full_name if requester else "Unknown employee"),
            _excel_text(requester.email if requester else ""),
            _excel_text(requester.department if requester else ""),
            _excel_text(claim.project.project_code),
            _excel_text(claim.project.project_name),
            _excel_text(claim.project.client_name or ""),
            _excel_text(CLAIM_TYPE_LABELS.get(claim.claim_type, claim.claim_type)),
            _excel_text(claim.purpose_description),
            _excel_text(claim.currency),
            float(money(claim.total_amount)),
            float(approved_amount(claim)),
            float(claim_paid_amount(claim)),
            float(remaining_amount(claim)),
            _excel_text(STATUS_LABELS.get(claim.status, claim.status)),
            _excel_text(db.get(User, claim.admin_decision_by_id).full_name if claim.admin_decision_by_id and db.get(User, claim.admin_decision_by_id) else ""),
            claim.admin_decision_at,
            _excel_text(claim.admin_comments or ""),
            _excel_text(db.get(User, claim.finance_decision_by_id).full_name if claim.finance_decision_by_id and db.get(User, claim.finance_decision_by_id) else ""),
            claim.finance_decision_at,
            _excel_text(claim.finance_comments or ""),
            len(claim.attachments),
            len(claim.payments) if claim.payments else (1 if claim.paid_amount else 0),
            claim.updated_at,
            _excel_text("CHECK" if issues else "OK"),
            claim.requested_work_start_date,
            claim.requested_work_end_date,
            claim.approved_work_start_date,
            claim.approved_work_end_date,
            claim.settlement_due_date,
            _excel_text((claim.settlement_status or "not_required").replace("_", " ").title()),
            _excel_text(claim.parent_advance.claim_code if claim.parent_advance else ""),
            _excel_text((
                "Shortage" if claim.settlement and money(claim.settlement.shortage_amount) > 0
                else "Balance To Return" if claim.settlement and money(claim.settlement.balance_to_return) > 0
                else "Tallied" if claim.settlement
                else ""
            )),
            float(claim.settlement.total_expense_amount) if claim.settlement else 0.0,
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col, value=value)
        for col in (2, 18, 21, 25):
            ws.cell(row=row_idx, column=col).number_format = _DATETIME_FORMAT
        for col in (27, 28, 29, 30, 31):
            ws.cell(row=row_idx, column=col).number_format = _DATE_FORMAT
        for col in (12, 13, 14, 15, 35):
            ws.cell(row=row_idx, column=col).number_format = _MONEY_FORMAT
        if issues:
            ws.cell(row=row_idx, column=26).fill = _WARNING_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:AI{max(1, len(claims) + 1)}"
    _auto_width(ws)
    return len(claims) + 1


def _items_sheet(wb: Workbook, claims: list[ExpenseClaim]) -> int:
    ws = wb.create_sheet("Expense Items")
    columns = ["Claim Code", "Submitted At", "Project ID", "Request Type", "Category", "Other Category", "Description", "Amount", "Payment Mode", "Expense Date"]
    _write_header(ws, 1, columns)
    row = 2
    for claim in claims:
        for item in claim.items:
            values = [_excel_text(claim.claim_code), claim.submitted_at, _excel_text(claim.project.project_code), _excel_text(CLAIM_TYPE_LABELS.get(claim.claim_type, claim.claim_type)), _excel_text(item.other_category or item.category), _excel_text(item.other_category or ""), _excel_text(item.description), float(item.amount), _excel_text((item.payment_mode or "").replace("_", " ").title()), item.expense_date]
            for col, value in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=value)
            ws.cell(row=row, column=2).number_format = _DATETIME_FORMAT
            ws.cell(row=row, column=8).number_format = _MONEY_FORMAT
            ws.cell(row=row, column=10).number_format = _DATE_FORMAT
            row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{max(1, row - 1)}"
    _auto_width(ws)
    return row - 1


def _events_sheet(wb: Workbook, claims: list[ExpenseClaim]) -> int:
    ws = wb.create_sheet("Approval History")
    columns = ["Claim Code", "Event Time", "Action", "Actor", "Actor Email", "Actor Role", "From Status", "To Status", "Comments"]
    _write_header(ws, 1, columns)
    row = 2
    for claim in claims:
        for event in claim.events:
            values = [_excel_text(claim.claim_code), event.created_at, _excel_text(event.action), _excel_text(event.actor_name or "System"), _excel_text(event.actor_email or ""), _excel_text(event.actor_role or ""), _excel_text(event.from_status or ""), _excel_text(event.to_status), _excel_text(event.comments or "")]
            for col, value in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=value)
            ws.cell(row=row, column=2).number_format = _DATETIME_FORMAT
            row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(1, row - 1)}"
    _auto_width(ws)
    return row - 1


def _payments_sheet(wb: Workbook, db: Session, payments: list[ExpenseClaimPayment]) -> int:
    ws = wb.create_sheet("Payments in Period")
    columns = [
        "Claim Code", "Claim Submitted At", "Project ID", "Project Name", "Employee", "Employee Email",
        "Payment Date", "Payment Mode", "Payment Reference", "Amount", "Recorded By", "Recorded At", "Comments",
    ]
    _write_header(ws, 1, columns)
    row = 2
    for payment in payments:
        claim = payment.claim
        requester = db.get(User, claim.requester_id)
        recorder = db.get(User, payment.recorded_by_id)
        values = [
            _excel_text(claim.claim_code),
            claim.submitted_at,
            _excel_text(claim.project.project_code),
            _excel_text(claim.project.project_name),
            _excel_text(requester.full_name if requester else ""),
            _excel_text(requester.email if requester else ""),
            payment.payment_date,
            _excel_text(payment.payment_mode.replace("_", " ").title()),
            _excel_text(payment.payment_reference),
            float(payment.amount),
            _excel_text(recorder.full_name if recorder else ""),
            payment.created_at,
            _excel_text(payment.comments or ""),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        ws.cell(row=row, column=2).number_format = _DATETIME_FORMAT
        ws.cell(row=row, column=7).number_format = _DATE_FORMAT
        ws.cell(row=row, column=10).number_format = _MONEY_FORMAT
        ws.cell(row=row, column=12).number_format = _DATETIME_FORMAT
        row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{max(1, row - 1)}"
    _auto_width(ws)
    return row - 1


def _attachments_sheet(wb: Workbook, db: Session, claims: list[ExpenseClaim]) -> int:
    ws = wb.create_sheet("Attachment Register")
    columns = ["Claim Code", "Project ID", "Employee", "Filename", "MIME Type", "Size (KB)", "SHA-256", "Uploaded By", "Uploaded At", "Secure API Path"]
    _write_header(ws, 1, columns)
    row = 2
    for claim in claims:
        requester = db.get(User, claim.requester_id)
        for attachment in claim.attachments:
            uploader = db.get(User, attachment.uploaded_by_id)
            values = [_excel_text(claim.claim_code), _excel_text(claim.project.project_code), _excel_text(requester.full_name if requester else ""), _excel_text(attachment.original_filename), _excel_text(attachment.mime_type), round(attachment.file_size / 1024, 2), _excel_text(attachment.content_sha256 or "Legacy / checksum not captured"), _excel_text(uploader.full_name if uploader else ""), attachment.created_at, _excel_text(f"/api/finance/claims/{claim.id}/attachments/{attachment.id}/content")]
            for col, value in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=value)
            ws.cell(row=row, column=9).number_format = _DATETIME_FORMAT
            row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{max(1, row - 1)}"
    _auto_width(ws)
    return row - 1


def _settlements_sheet(wb: Workbook, db: Session, claims: list[ExpenseClaim]) -> int:
    ws = wb.create_sheet("Settlement Ledger")
    columns = [
        "Settlement Code", "Root Advance", "Project ID", "Employee", "Status", "Submitted At", "Finalized At",
        "Advance Received", "Actual Expense", "Balance To Return", "Shortage", "Tally",
        "Settlement Due", "Admin Verified By", "Admin Decision At", "Finance Verified By", "Finance Decision At",
        "Bill Count",
    ]
    _write_header(ws, 1, columns)
    row = 2
    for claim in claims:
        settlement = claim.settlement if claim.claim_type == "advance" else None
        if settlement is None:
            continue
        requester = db.get(User, claim.requester_id)
        tally = "Shortage" if money(settlement.shortage_amount) > 0 else "Balance To Return" if money(settlement.balance_to_return) > 0 else "Tallied"
        values = [
            _excel_text(settlement.settlement_code), _excel_text(claim.claim_code), _excel_text(claim.project.project_code),
            _excel_text(requester.full_name if requester else ""), _excel_text(settlement.status.replace("_", " ").title()),
            settlement.submitted_at, settlement.finalized_at, float(money(settlement.total_advance_received)),
            float(money(settlement.total_expense_amount)), float(money(settlement.balance_to_return)),
            float(money(settlement.shortage_amount)), _excel_text(tally), claim.settlement_due_date,
            _excel_text(db.get(User, settlement.admin_decision_by_id).full_name if settlement.admin_decision_by_id and db.get(User, settlement.admin_decision_by_id) else ""),
            settlement.admin_decision_at,
            _excel_text(db.get(User, settlement.finance_decision_by_id).full_name if settlement.finance_decision_by_id and db.get(User, settlement.finance_decision_by_id) else ""),
            settlement.finance_decision_at, len(settlement.attachments),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        for col in (6, 7, 15, 17):
            ws.cell(row=row, column=col).number_format = _DATETIME_FORMAT
        ws.cell(row=row, column=13).number_format = _DATE_FORMAT
        for col in (8, 9, 10, 11):
            ws.cell(row=row, column=col).number_format = _MONEY_FORMAT
        row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:R{max(1, row - 1)}"
    _auto_width(ws)
    return row - 1


def _settlement_items_sheet(wb: Workbook, claims: list[ExpenseClaim]) -> int:
    ws = wb.create_sheet("Settlement Actual Items")
    columns = ["Settlement Code", "Root Advance", "Project ID", "Category", "Description", "Payment Mode", "Expense Date", "Amount"]
    _write_header(ws, 1, columns)
    row = 2
    for claim in claims:
        settlement = claim.settlement if claim.claim_type == "advance" else None
        if settlement is None:
            continue
        for item in settlement.items:
            values = [
                _excel_text(settlement.settlement_code), _excel_text(claim.claim_code), _excel_text(claim.project.project_code),
                _excel_text(item.other_category or item.category), _excel_text(item.description),
                _excel_text((item.payment_mode or "").replace("_", " ").title()), item.expense_date, float(money(item.amount)),
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=value)
            ws.cell(row=row, column=7).number_format = _DATE_FORMAT
            ws.cell(row=row, column=8).number_format = _MONEY_FORMAT
            row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(1, row - 1)}"
    _auto_width(ws)
    return row - 1


def _settlement_bills_sheet(wb: Workbook, db: Session, claims: list[ExpenseClaim]) -> int:
    ws = wb.create_sheet("Settlement Bill Register")
    columns = ["Settlement Code", "Root Advance", "Project ID", "Filename", "MIME Type", "Size (KB)", "SHA-256", "Uploaded By", "Uploaded At", "Secure API Path"]
    _write_header(ws, 1, columns)
    row = 2
    for claim in claims:
        settlement = claim.settlement if claim.claim_type == "advance" else None
        if settlement is None:
            continue
        for attachment in settlement.attachments:
            uploader = db.get(User, attachment.uploaded_by_id)
            values = [
                _excel_text(settlement.settlement_code), _excel_text(claim.claim_code), _excel_text(claim.project.project_code),
                _excel_text(attachment.original_filename), _excel_text(attachment.mime_type), round(attachment.file_size / 1024, 2),
                _excel_text(attachment.content_sha256 or "Legacy / checksum not captured"), _excel_text(uploader.full_name if uploader else ""),
                attachment.created_at, _excel_text(f"/api/finance/settlements/{settlement.id}/attachments/{attachment.id}/content"),
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=value)
            ws.cell(row=row, column=9).number_format = _DATETIME_FORMAT
            row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{max(1, row - 1)}"
    _auto_width(ws)
    return row - 1


def _settlement_events_sheet(wb: Workbook, claims: list[ExpenseClaim]) -> int:
    ws = wb.create_sheet("Settlement History")
    columns = ["Settlement Code", "Event Time", "Action", "Actor", "Actor Role", "From Status", "To Status", "Comments"]
    _write_header(ws, 1, columns)
    row = 2
    for claim in claims:
        settlement = claim.settlement if claim.claim_type == "advance" else None
        if settlement is None:
            continue
        for event in settlement.events:
            values = [
                _excel_text(settlement.settlement_code), event.created_at, _excel_text(event.action),
                _excel_text(event.actor_name or "System"), _excel_text(event.actor_role or ""),
                _excel_text(event.from_status or ""), _excel_text(event.to_status), _excel_text(event.comments or ""),
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=value)
            ws.cell(row=row, column=2).number_format = _DATETIME_FORMAT
            row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(1, row - 1)}"
    _auto_width(ws)
    return row - 1


def _integrity_sheet(wb: Workbook, claims: list[ExpenseClaim]) -> int:
    ws = wb.create_sheet("Data Integrity")
    columns = ["Claim Code", "Status", "Issue"]
    _write_header(ws, 1, columns)
    row = 2
    for claim in claims:
        for issue in claim_integrity_issues(claim):
            ws.cell(row=row, column=1, value=_excel_text(claim.claim_code))
            ws.cell(row=row, column=2, value=_excel_text(STATUS_LABELS.get(claim.status, claim.status)))
            ws.cell(row=row, column=3, value=_excel_text(issue))
            ws.cell(row=row, column=3).fill = _WARNING_FILL
            row += 1
    if row == 2:
        ws.cell(row=2, column=1, value="No integrity issues detected in the selected report scope.")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=3)
    _auto_width(ws)
    return row - 1


def _summary_breakdown_sheet(wb: Workbook, name: str, key_header: str, rows: list[tuple[str, str]], *, source_sheet: str, source_key_column: str, source_amount_column: str) -> None:
    ws = wb.create_sheet(_safe_sheet_title(name))
    _write_header(ws, 1, [key_header, "Display Name", "Claim / Item Count", "Amount"])
    source_max = max(2, wb[source_sheet].max_row)
    for row_idx, (key, label) in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=_excel_text(key))
        ws.cell(row=row_idx, column=2, value=_excel_text(label))
        ws.cell(row=row_idx, column=3, value=f'=COUNTIF(\'{source_sheet}\'!${source_key_column}$2:${source_key_column}${source_max},A{row_idx})')
        ws.cell(row=row_idx, column=4, value=f'=SUMIF(\'{source_sheet}\'!${source_key_column}$2:${source_key_column}${source_max},A{row_idx},\'{source_sheet}\'!${source_amount_column}$2:${source_amount_column}${source_max})')
        ws.cell(row=row_idx, column=4).number_format = _MONEY_FORMAT
    ws.freeze_panes = "A2"
    _auto_width(ws)


def build_finance_report_workbook(
    db: Session,
    *,
    period: str,
    year: int | None = None,
    month: int | None = None,
    quarter: int | None = None,
    project_id: int | None = None,
    claim_type: str | None = None,
    status: str | None = None,
    employee_query: str | None = None,
    category: str | None = None,
    search: str | None = None,
) -> tuple[BytesIO, ResolvedPeriod, int]:
    resolved = resolve_report_period(period, year=year, month=month, quarter=quarter)
    claims = filtered_finance_claims(
        db,
        period=resolved,
        project_id=project_id,
        claim_type=claim_type,
        status=status,
        employee_query=employee_query,
        category=category,
        search=search,
    )
    if len(claims) > MAX_REPORT_ROWS:
        raise ValueError(f"The selected report contains {len(claims):,} claims. Narrow the filters below {MAX_REPORT_ROWS:,} claims for one Excel export.")

    payments = filtered_finance_payments(
        db,
        period=resolved,
        project_id=project_id,
        claim_type=claim_type,
        status=status,
        employee_query=employee_query,
        category=category,
        search=search,
    )
    if len(payments) > 200_000:
        raise ValueError("The selected period contains more than 200,000 payment rows. Narrow the filters for one Excel export.")

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.sheet_view.showGridLines = False

    claims_last_row = _claims_ledger_sheet(wb, db, claims)
    items_last_row = _items_sheet(wb, claims)
    _events_sheet(wb, claims)
    payments_last_row = _payments_sheet(wb, db, payments)
    _attachments_sheet(wb, db, claims)
    _settlements_sheet(wb, db, claims)
    _settlement_items_sheet(wb, claims)
    _settlement_bills_sheet(wb, db, claims)
    _settlement_events_sheet(wb, claims)
    _integrity_sheet(wb, claims)

    project_rows = sorted({(_excel_text(claim.project.project_code), _excel_text(f"{claim.project.project_code} \u00b7 {claim.project.project_name}")) for claim in claims})
    _summary_breakdown_sheet(wb, "Project Summary", "Project ID", project_rows, source_sheet="Claims Ledger", source_key_column="F", source_amount_column="L")

    employee_map: dict[str, str] = {}
    for claim in claims:
        requester = db.get(User, claim.requester_id)
        if requester:
            employee_map[_excel_text(requester.email)] = _excel_text(requester.full_name)
    employee_rows = sorted(employee_map.items(), key=lambda item: item[1].lower())
    _summary_breakdown_sheet(wb, "Employee Summary", "Employee Email", employee_rows, source_sheet="Claims Ledger", source_key_column="D", source_amount_column="L")

    category_map: dict[str, str] = {}
    for claim in claims:
        for item in claim.items:
            key = item.other_category.strip() if item.category == "other" and item.other_category else item.category
            category_map[_excel_text(key)] = _excel_text(key.replace("_", " ").title())
    category_rows = sorted(category_map.items(), key=lambda item: item[1].lower())
    _summary_breakdown_sheet(wb, "Category Summary", "Category", category_rows, source_sheet="Expense Items", source_key_column="E", source_amount_column="H")

    summary["A1"] = "NakshaTech Finance & Project Expense Report"
    summary["A1"].font = _TITLE_FONT
    summary.merge_cells("A1:D1")
    summary["A2"] = f"Period: {resolved.label}"
    summary["A2"].font = Font(bold=True, color="0B7285")
    summary["A3"] = "Source of truth: Finance CRM database. Excel is a reporting/export view; approval and payment history remains authoritative in the application."
    summary["A3"].font = _SUBTITLE_FONT
    summary.merge_cells("A3:D3")
    summary["A4"] = f"Generated: {utc_now().strftime('%d-%b-%Y %H:%M UTC')}"
    summary["A4"].font = _SUBTITLE_FONT

    _write_header(summary, 6, ["Metric", "Value", "Control", "Definition"])
    metrics = [
        ("Total Claims", f"=COUNTA('Claims Ledger'!A2:A{max(2, claims_last_row)})", "Black formula", "Non-draft claims submitted in the selected claim period"),
        ("Requested Amount", f"=SUM('Claims Ledger'!L2:L{max(2, claims_last_row)})", "Black formula", "Original employee requested amount for claims submitted in this period"),
        ("Finance Approved Amount", f"=SUM('Claims Ledger'!M2:M{max(2, claims_last_row)})", "Black formula", "Current Finance-approved amount for this claim cohort"),
        ("Paid Against Claim Cohort", f"=SUM('Claims Ledger'!N2:N{max(2, claims_last_row)})", "Black formula", "Cumulative payments against claims submitted in this period, regardless of payment date"),
        ("Outstanding Claim Cohort", f"=SUM('Claims Ledger'!O2:O{max(2, claims_last_row)})", "Black formula", "Approved less cumulative paid for claims submitted in this period"),
        ("Cash Paid During Period", f"=SUM('Payments in Period'!J2:J{max(2, payments_last_row)})", "Black formula", "Payments whose payment date falls inside the selected month / quarter / year"),
        ("Payment Transactions in Period", f"=COUNTA('Payments in Period'!I2:I{max(2, payments_last_row)})", "Black formula", "Individual immutable payment-ledger transactions dated in this period"),
        ("Pending Admin Amount", f"=SUMIF('Claims Ledger'!P2:P{max(2, claims_last_row)},\"Pending Admin Verification\",'Claims Ledger'!L2:L{max(2, claims_last_row)})", "Black formula", "Submitted but not Admin-verified"),
        ("Pending Finance Amount", f"=SUMIF('Claims Ledger'!P2:P{max(2, claims_last_row)},\"Pending Finance Verification\",'Claims Ledger'!L2:L{max(2, claims_last_row)})", "Black formula", "Admin-verified, awaiting Finance"),
        ("Integrity Flags", f"=COUNTIF('Claims Ledger'!Z2:Z{max(2, claims_last_row)},\"CHECK\")", "Review", "Records requiring Finance data-quality review"),
    ]
    for row_idx, (label, formula, control, definition) in enumerate(metrics, start=7):
        summary.cell(row=row_idx, column=1, value=label)
        summary.cell(row=row_idx, column=2, value=formula)
        summary.cell(row=row_idx, column=3, value=control)
        summary.cell(row=row_idx, column=4, value=definition)
        summary.cell(row=row_idx, column=2).font = Font(color="000000")
        if "Amount" in label:
            summary.cell(row=row_idx, column=2).number_format = _MONEY_FORMAT
        if label == "Integrity Flags":
            summary.cell(row=row_idx, column=3).fill = _WARNING_FILL

    summary["A18"] = "Reporting Rules"
    summary["A18"].fill = _SECTION_FILL
    summary["A18"].font = _HEADER_FONT
    summary.merge_cells("A18:D18")
    rules = [
        "Claim-period totals use submitted_at; created_at is used only for legacy records missing submitted_at.",
        "Cash Paid During Period uses each immutable payment transaction's payment_date, so a later payment for an older claim appears in the correct cash period.",
        "Requested Amount is never replaced by Finance Approved Amount. Both values remain visible for audit.",
        "Partial payments remain separate ledger rows and outstanding until the approved amount is fully settled.",
        "Previously submitted proof is retained for audit; corrected evidence is added rather than silently replacing historical evidence.",
        "New Finance proof captures SHA-256 in the Attachment Register; legacy V1 proof remains identified as legacy where no checksum was captured.",
        "Advance settlements are exported separately with actual bill items, payment modes, bill checksums, tally, balance and shortage values.",
        "Data Integrity flags do not modify records; they identify anomalies for controlled correction in the application.",
    ]
    for idx, rule in enumerate(rules, start=19):
        summary.cell(row=idx, column=1, value=f"{idx - 18}.")
        summary.cell(row=idx, column=2, value=rule)
        summary.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=4)
        summary.cell(row=idx, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 24
    summary.column_dimensions["C"].width = 18
    summary.column_dimensions["D"].width = 62
    summary.freeze_panes = "A6"

    # Ask Excel to recalculate formulas on open so totals are always based on the
    # exact downloaded ledger rows rather than cached application numbers.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output, resolved, len(claims)


def finance_report_filename(resolved: ResolvedPeriod) -> str:
    safe_label = resolved.label.replace(" ", "_").replace("/", "-")
    return f"NakshaTech_Finance_{resolved.period.title()}_{safe_label}.xlsx"
