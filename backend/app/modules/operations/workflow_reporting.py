from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models.entities import User
from app.modules.finance.models import FinanceClient, FinanceProject
from app.modules.operations.models import ProjectWorkflow, ProjectWorkflowEvent


REPORT_TYPES = {"clients", "projects", "clients_projects", "finance_approval_history", "project_closure"}
PERIODS = {"daily", "weekly", "monthly", "yearly", "custom", "all"}


def _bounds(period: str, start_date: date | None, end_date: date | None) -> tuple[date | None, date | None]:
    today = date.today()
    if period not in PERIODS:
        raise ValueError("Unsupported report period")
    if period == "all":
        return None, None
    if period == "daily":
        return today, today
    if period == "weekly":
        return today - timedelta(days=today.weekday()), today
    if period == "monthly":
        return today.replace(day=1), today
    if period == "yearly":
        return today.replace(month=1, day=1), today
    if start_date is None or end_date is None:
        raise ValueError("Custom reports require start_date and end_date")
    if end_date < start_date:
        raise ValueError("Custom end_date cannot be earlier than start_date")
    return start_date, end_date


def _in_period(value: date | datetime | None, start: date | None, end: date | None) -> bool:
    if start is None or end is None:
        return True
    if value is None:
        return False
    actual = value.date() if isinstance(value, datetime) else value
    return start <= actual <= end


def _sheet(workbook: Workbook, title: str, headers: list[str]):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).coordinate}"
    return sheet


def _finish(workbook: Workbook) -> bytes:
    if "Sheet" in workbook.sheetnames:
        del workbook["Sheet"]
    for sheet in workbook.worksheets:
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_project_operations_workbook(
    db: Session,
    *,
    report_type: str,
    period: str,
    start_date: date | None = None,
    end_date: date | None = None,
) -> tuple[bytes, str]:
    if report_type not in REPORT_TYPES:
        raise ValueError("Unsupported project operations report type")
    start, end = _bounds(period, start_date, end_date)
    workbook = Workbook()

    clients = list(db.scalars(
        select(FinanceClient)
        .options(selectinload(FinanceClient.master_profile), selectinload(FinanceClient.projects))
        .order_by(FinanceClient.client_code.asc())
    ).unique().all())
    projects = list(db.scalars(
        select(FinanceProject)
        .options(selectinload(FinanceProject.client), selectinload(FinanceProject.master_profile))
        .order_by(FinanceProject.project_code.asc())
    ).unique().all())
    workflows = {row.project_id: row for row in db.scalars(select(ProjectWorkflow)).all()}
    user_ids = {
        value
        for workflow in workflows.values()
        for value in (workflow.bd_owner_user_id, workflow.finance_reviewer_id)
        if value
    }
    user_ids.update(
        project.master_profile.project_manager_id
        for project in projects
        if project.master_profile and project.master_profile.project_manager_id
    )
    users = {row.id: row for row in db.scalars(select(User).where(User.id.in_(user_ids))).all()} if user_ids else {}

    if report_type in {"clients", "clients_projects"}:
        sheet = _sheet(workbook, "Clients", [
            "Client ID", "Organization Name", "Location", "GST Number", "Client Email",
            "Organization Email", "Contact Person", "Contact Email", "Contact Phone", "BD Person",
            "Project Count", "Created Date", "Last Updated Date",
        ])
        for client in clients:
            if not _in_period(client.created_at, start, end):
                continue
            profile = client.master_profile
            sheet.append([
                client.client_code, client.client_name, client.address or client.country, client.gst_number,
                client.client_email, profile.organization_email if profile else None, client.contact_person_name,
                profile.contact_person_email if profile else None,
                None if client.contact_person_phone == "Not provided" else client.contact_person_phone,
                (profile.bd_name if profile and profile.bd_name else client.source_person_name), len(client.projects),
                client.created_at, client.updated_at,
            ])

    if report_type in {"projects", "clients_projects"}:
        sheet = _sheet(workbook, "Projects", [
            "Client ID", "Organization Name", "Project ID", "Project Scope", "Start Date", "End Date",
            "Project Value", "Currency", "PO / WO", "Finance Status", "Submitted Date", "Approved Date",
            "Returned Date", "Finance Feedback", "Finance Reviewer", "Submission Count", "Project Manager",
            "Operational Status", "Closure Status", "Closure Date",
        ])
        for project in projects:
            workflow = workflows.get(project.id)
            reference_date = workflow.submitted_at if workflow and workflow.submitted_at else project.created_at
            if not _in_period(reference_date, start, end):
                continue
            client = project.client
            pm_id = project.master_profile.project_manager_id if project.master_profile else None
            sheet.append([
                client.client_code if client else None, client.client_name if client else project.client_name,
                project.project_code, workflow.scope_text if workflow else (project.master_profile.task if project.master_profile else project.description),
                project.start_date, project.end_date, workflow.commercial_value if workflow else None,
                workflow.currency if workflow else None, workflow.po_wo_number if workflow else None,
                workflow.status if workflow else "legacy", workflow.submitted_at if workflow else None,
                workflow.approved_at if workflow else None, workflow.returned_at if workflow else None,
                workflow.finance_feedback if workflow else None,
                users[workflow.finance_reviewer_id].full_name if workflow and workflow.finance_reviewer_id in users else None,
                workflow.submission_count if workflow else 0,
                users[pm_id].full_name if pm_id in users else None,
                workflow.status if workflow else (project.master_profile.project_status if project.master_profile else None),
                "closed" if workflow and workflow.status == "closed" else ("closure_pending" if workflow and workflow.status == "finance_closure_pending" else "open"),
                workflow.finance_closed_at if workflow else None,
            ])

    if report_type == "finance_approval_history":
        sheet = _sheet(workbook, "Finance Approval History", [
            "Project ID", "Client ID", "Event", "From Status", "To Status", "Feedback / Comment",
            "Reviewer", "Reviewer Role", "Review Date",
        ])
        events = list(db.scalars(select(ProjectWorkflowEvent).where(
            ProjectWorkflowEvent.event_type.in_(["submitted_to_finance", "finance_returned", "finance_approved"])
        ).order_by(ProjectWorkflowEvent.created_at.asc())).all())
        event_users = {row.id: row for row in db.scalars(select(User).where(User.id.in_({event.actor_user_id for event in events}))).all()} if events else {}
        project_map = {project.id: project for project in projects}
        for event in events:
            if not _in_period(event.created_at, start, end):
                continue
            project = project_map.get(event.project_id)
            actor = event_users.get(event.actor_user_id)
            sheet.append([
                project.project_code if project else event.project_id,
                project.client.client_code if project and project.client else None,
                event.event_type, event.from_status, event.to_status, event.comments,
                actor.full_name if actor else None, actor.role if actor else None, event.created_at,
            ])

    if report_type == "project_closure":
        sheet = _sheet(workbook, "Project Closure", [
            "Project ID", "Client ID", "Project Manager", "Operational Completion Date", "Delivery Reference",
            "Finance Status", "Closure Date", "Closure Remarks",
        ])
        project_map = {project.id: project for project in projects}
        for workflow in workflows.values():
            reference_date = workflow.finance_closed_at or workflow.operational_completed_at
            if reference_date is None or not _in_period(reference_date, start, end):
                continue
            project = project_map.get(workflow.project_id)
            pm_id = project.master_profile.project_manager_id if project and project.master_profile else None
            sheet.append([
                project.project_code if project else workflow.project_id,
                project.client.client_code if project and project.client else None,
                users[pm_id].full_name if pm_id in users else None,
                workflow.operational_completed_at, workflow.final_delivery_reference, workflow.status,
                workflow.finance_closed_at, workflow.finance_closure_remarks,
            ])

    filename = f"Naksha_Project_Operations_{report_type}_{period}.xlsx"
    return _finish(workbook), filename
