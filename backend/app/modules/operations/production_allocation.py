"""Team Lead Excel-driven daily activity allocation (V8.3, additive).

The company already maintains a work-file master. Instead of typing Lot / Code / Area /
File / Planned Length / UOM by hand, the Team Lead uploads that workbook for the selected
project and the ERP transforms it into the EXISTING work-package / file architecture:

    Project -> MARCHE LOTTO 3 (lot) -> CODE -> AREA NAME -> FILE NUMBER

Accepted master columns (case-insensitive, flexible spacing):

    MARCHE LOTTO 3   (required)   lot / batch
    CODE             (required)
    AREA NAME        (required)
    FILE NUMBER      (required)
    PLANNED LENGTH   (required)
    UOM              (required)
    REMARKS          (optional)
    FILE NAME        (optional)

Flow: upload -> parse -> validate -> PREVIEW (nothing saved) -> Team Lead confirms ->
transactional import. Repeated uploads are idempotent: rows already present WITHOUT work
history are skipped, rows WITH production/QC/QA history are reported as conflicts and are
never overwritten or reset. Nothing is ever hard-deleted.
"""

from __future__ import annotations

import csv
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, Iterable

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.entities import User
from app.modules.finance.models import FinanceProject
from app.modules.notifications.service import create_global_notification
from app.modules.operations.models import (
    OrthoProjectMember,
    OrthoWorkPackage,
    WorkPackageItem,
    WorkPackageItemActivity,
)
from app.modules.operations import production_items as items_service
from app.modules.operations.uom import convert_quantity, measurement_type, units_compatible

ADMIN_ROLE = "admin"
MANAGEMENT_ROLE = "management"
EMPLOYEE_ROLE = "employee"

# Header aliases -> canonical field.
_HEADER_ALIASES: dict[str, str] = {
    "marche lotto 3": "lot", "marche lotto": "lot", "lot": "lot", "lot no": "lot",
    "lot number": "lot", "batch": "lot", "lotto": "lot", "marche lotto 3 (lot / batch)": "lot",
    "code": "code", "area code": "code",
    "area name": "area_name", "area": "area_name",
    "file number": "item_number", "file no": "item_number", "file_number": "item_number",
    "item number": "item_number", "item_number": "item_number", "file": "item_number",
    "planned length": "planned_quantity", "planned_length": "planned_quantity",
    "length": "planned_quantity", "planned quantity": "planned_quantity",
    "planned_quantity": "planned_quantity", "quantity": "planned_quantity", "qty": "planned_quantity",
    "uom": "unit", "unit": "unit", "planned unit": "unit", "planned_unit": "unit",
    "remarks": "remarks", "remark": "remarks", "note": "remarks", "notes": "remarks",
    "file name": "item_name", "file_name": "item_name", "item name": "item_name", "item_name": "item_name",
    # V8.4 declared physical-file count for a Lot/Code/Area (additive).
    "total files": "total_files", "total_files": "total_files", "totalfiles": "total_files",
    "no of files": "total_files", "no. of files": "total_files", "no of file": "total_files",
    "file count": "total_files", "file_count": "total_files", "files": "total_files",
    "total no of files": "total_files", "total number of files": "total_files",
}

_REQUIRED_FIELDS = ("lot", "code", "area_name", "item_number", "planned_quantity", "unit")
_FIELD_LABELS = {
    "lot": "MARCHE LOTTO 3", "code": "CODE", "area_name": "AREA NAME",
    "item_number": "FILE NUMBER", "planned_quantity": "PLANNED LENGTH", "unit": "UOM",
}
MAX_GENERATED_FILES = 5000

_UNIT_TOKEN = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789./ _-")


# ===========================================================================================
# Parsing / validation
# ===========================================================================================

def _map_headers(header_row: Iterable[Any]) -> tuple[dict[int, str], list[str]]:
    mapping: dict[int, str] = {}
    seen: set[str] = set()
    for index, cell in enumerate(header_row):
        name = str(cell or "").strip().lower()
        if not name:
            continue
        field = _HEADER_ALIASES.get(name)
        if field and field not in seen:
            mapping[index] = field
            seen.add(field)
    missing_fields = [f for f in _REQUIRED_FIELDS if f not in seen]
    # FILE NUMBER is only required when TOTAL FILES is absent (V8.4 count-based areas).
    if "item_number" in missing_fields and "total_files" in seen:
        missing_fields.remove("item_number")
    missing = [_FIELD_LABELS[f] for f in missing_fields]
    return mapping, missing


def _read_rows(filename: str, content: bytes) -> list[list[Any]]:
    name = (filename or "").lower()
    if name.endswith(".csv") or content[:2] != b"PK":
        text = None
        for encoding in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("Unable to decode the uploaded file")
        return [list(row) for row in csv.reader(text.splitlines())]
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # malformed workbook
        raise ValueError(f"The uploaded workbook could not be read: {exc}") from exc
    try:
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def parse_master_workbook(filename: str, content: bytes) -> list[dict]:
    """Parse + validate the master workbook. Raises ValueError with row-level detail."""
    if not content:
        raise ValueError("The uploaded file is empty")
    rows = [row for row in _read_rows(filename, content)
            if any(cell is not None and str(cell).strip() for cell in row)]
    if not rows:
        raise ValueError("The uploaded workbook is empty")
    mapping, missing = _map_headers(rows[0])
    if missing:
        raise ValueError("Missing required column(s): " + ", ".join(missing))

    entries: list[dict] = []
    errors: list[str] = []
    seen_keys: set[tuple[str, str, str, str]] = set()
    for index, row in enumerate(rows[1:], start=2):
        def cell(field: str) -> Any:
            for column, mapped in mapping.items():
                if mapped == field:
                    return row[column] if column < len(row) else None
            return None

        lot = str(cell("lot") or "").strip()
        code = str(cell("code") or "").strip()
        area_name = str(cell("area_name") or "").strip()
        number = str(cell("item_number") or "").strip()
        unit = str(cell("unit") or "").strip()
        raw_qty = cell("planned_quantity")
        raw_total = cell("total_files")
        remarks = str(cell("remarks") or "").strip() or None
        item_name = str(cell("item_name") or "").strip() or None
        row_errors: list[str] = []
        if not lot:
            row_errors.append("blank MARCHE LOTTO 3")
        if not code:
            row_errors.append("blank CODE")
        if not area_name:
            row_errors.append("blank AREA NAME")
        # V8.4: TOTAL FILES (count) OR FILE NUMBER (explicit) — at least one is required.
        total_files = None
        if raw_total is not None and str(raw_total).strip():
            try:
                total_files = int(Decimal(str(raw_total).strip()))
                if total_files < 1:
                    row_errors.append(f"TOTAL FILES must be at least 1 (got '{raw_total}')")
                elif total_files > MAX_GENERATED_FILES:
                    row_errors.append(f"TOTAL FILES {total_files} exceeds the safe limit ({MAX_GENERATED_FILES})")
            except (InvalidOperation, ValueError):
                row_errors.append(f"invalid TOTAL FILES '{raw_total}'")
        if not number and total_files is None:
            row_errors.append("blank FILE NUMBER (or provide TOTAL FILES for the area)")
        if total_files is not None and number:
            row_errors.append("provide EITHER TOTAL FILES for the area OR a FILE NUMBER per row, not both")
        if not unit:
            row_errors.append("blank UOM")
        elif len(unit) > 30 or any(ch not in _UNIT_TOKEN for ch in unit):
            row_errors.append(f"invalid UOM '{unit}'")
        quantity = None
        if raw_qty is None or (isinstance(raw_qty, str) and not raw_qty.strip()):
            row_errors.append("blank PLANNED LENGTH")
        else:
            try:
                quantity = Decimal(str(raw_qty).strip())
                if quantity < 0:
                    row_errors.append(f"negative PLANNED LENGTH '{raw_qty}'")
            except (InvalidOperation, ValueError):
                row_errors.append(f"invalid PLANNED LENGTH '{raw_qty}'")
        # Deterministic physical-file identity: explicit FILE NUMBER, else File 1..N.
        numbers = [number] if number else [str(i) for i in range(1, (total_files or 0) + 1)]
        if not row_errors:
            for candidate in numbers:
                key = (lot, code, area_name, candidate)
                if key in seen_keys:
                    row_errors.append("duplicate row (same Lot + Code + Area + File Number)")
                    break
                seen_keys.add(key)
        if row_errors:
            errors.append(f"row {index}: " + "; ".join(row_errors))
            if len(errors) >= 25:
                errors.append("… too many errors (stopped after 25)")
                break
            continue
        entries.append({
            "row": index, "lot": lot, "code": code, "area_name": area_name,
            "item_number": numbers[0], "item_numbers": numbers,
            # Per-file planned length is only meaningful for explicit FILE NUMBER rows;
            # a count-based area must NOT fabricate per-file KM. The area's declared
            # length is carried separately (applied once, never divided).
            "planned_quantity": quantity if number else None,
            "area_planned_quantity": quantity,
            "total_files": total_files,
            "unit": unit, "remarks": remarks, "item_name": item_name,
        })
    if errors:
        raise ValueError("Import validation failed — " + " | ".join(errors))
    if not entries:
        raise ValueError("No data rows were found in the uploaded workbook")
    return entries


def group_planned_by_unit(entries: Iterable[dict]) -> list[dict]:
    """Planned quantity grouped by COMPATIBLE unit family (never summed across families).

    The total is expressed in the family's dominant original unit (e.g. "3.15 km");
    conversion only happens INSIDE a family, and opaque/unknown legacy units are
    never mixed with anything else.
    """
    families: dict[str, dict] = {}
    for entry in entries:
        unit = entry["unit"]
        family = measurement_type(unit) or f"opaque:{unit}"
        group = families.setdefault(family, {"items": [], "units": {}, "rows": 0})
        group["items"].append((entry["planned_quantity"], unit))
        group["units"][unit] = group["units"].get(unit, 0) + 1
        group["rows"] += 1
    rows: list[dict] = []
    for family, group in sorted(families.items()):
        dominant = max(group["units"].items(), key=lambda kv: kv[1])[0]
        total = Decimal("0")
        for quantity, unit in group["items"]:
            if unit == dominant:
                total += quantity
                continue
            try:
                total += convert_quantity(quantity, unit, dominant)
            except Exception:
                # Incompatible inside the family (defensive): keep it out of the total
                # rather than corrupting it with a wrong conversion.
                continue
        rows.append({
            "measurement_type": family, "unit": dominant,
            "quantity": float(total), "rows": group["rows"],
        })
    return rows


# ===========================================================================================
# Access / lookup helpers
# ===========================================================================================

def _project(db: Session, project_id: int) -> FinanceProject:
    project = db.get(FinanceProject, project_id)
    if project is None:
        raise ValueError("Project not found")
    return project


def _project_packages(db: Session, project_id: int) -> list[OrthoWorkPackage]:
    return list(db.scalars(select(OrthoWorkPackage).where(
        OrthoWorkPackage.project_id == project_id
    ).order_by(OrthoWorkPackage.id.asc())).all())


def _package_actor_flags(db: Session, project_id: int, actor: User, role: str) -> dict:
    project = _project(db, project_id)
    pm_id = None
    if project.master_profile is not None:
        pm_id = project.master_profile.project_manager_id
    packages = _project_packages(db, project_id)
    is_pm = pm_id == actor.id
    is_tl = any(pkg.team_leader_user_id == actor.id for pkg in packages)
    if not is_tl:
        # Before any area exists (or before allocation sets the per-package TL) the
        # PM-approved project Team Lead membership is authoritative.
        is_tl = db.scalar(select(OrthoProjectMember.id).where(
            OrthoProjectMember.project_id == project_id,
            OrthoProjectMember.user_id == actor.id,
            OrthoProjectMember.member_role == "team_leader",
            OrthoProjectMember.is_active.is_(True),
        ).limit(1)) is not None
    read_only_role = role in {ADMIN_ROLE, MANAGEMENT_ROLE}
    return {
        "project": project, "packages": packages, "pm_id": pm_id,
        "is_pm": is_pm, "is_tl": is_tl, "read_only_role": read_only_role,
        "can_manage": (is_pm or is_tl) and not read_only_role,
        "privileged": is_pm or is_tl or read_only_role,
    }


def ensure_master_manage(db: Session, *, actor: User, role: str, project_id: int) -> dict:
    """Master import / allocation requires the project Team Lead or PM (Management read-only)."""
    ctx = _package_actor_flags(db, project_id, actor, role)
    if ctx["read_only_role"]:
        raise PermissionError("Management/Admin access is read-only")
    if not ctx["can_manage"]:
        raise PermissionError("Only the project Team Lead or Project Manager can import the work file master")
    return ctx


def _find_package(packages: list[OrthoWorkPackage], lot: str, code: str, area_name: str) -> OrthoWorkPackage | None:
    for pkg in packages:
        if (pkg.lot_batch or "").strip() == lot and (pkg.package_code or "").strip() == code \
                and (pkg.package_name or "").strip() == area_name:
            return pkg
    return None


def _item_has_history(db: Session, item: WorkPackageItem) -> bool:
    return db.scalar(select(WorkPackageItemActivity.id).where(
        WorkPackageItemActivity.work_package_item_id == item.id
    ).limit(1)) is not None


def _existing_item(db: Session, package_id: int, number: str) -> WorkPackageItem | None:
    return db.scalar(select(WorkPackageItem).where(
        WorkPackageItem.work_package_id == package_id,
        WorkPackageItem.item_number == number,
    ))


# ===========================================================================================
# Preview / commit
# ===========================================================================================

def _analyse(db: Session, *, ctx: dict, entries: list[dict]) -> tuple[list[dict], list[dict], int]:
    """Classify each row's FILE NUMBERS against existing data.

    Per row: numbers that do not exist yet are "new"; numbers already present with
    history are conflicts (never overwritten); numbers already present without history
    are skipped (idempotent re-import). Row-level planned length is applied at most once.
    """
    packages = ctx["packages"]
    fresh: list[dict] = []
    conflicts: list[dict] = []
    skipped_existing = 0
    for entry in entries:
        package = _find_package(packages, entry["lot"], entry["code"], entry["area_name"])
        new_numbers: list[str] = []
        any_existing = False
        for number in entry["item_numbers"]:
            if package is None:
                new_numbers.append(number)
                continue
            existing = _existing_item(db, package.id, number)
            if existing is None:
                new_numbers.append(number)
            elif _item_has_history(db, existing):
                any_existing = True
                conflicts.append({
                    "row": entry["row"], "lot": entry["lot"], "code": entry["code"],
                    "area_name": entry["area_name"], "item_number": number,
                    "reason": "file already has production/QC/QA history — skipped, never overwritten",
                    "work_package_item_id": existing.id,
                })
            else:
                any_existing = True
                skipped_existing += 1
        if new_numbers:
            row = dict(entry)
            row["new_numbers"] = new_numbers
            row["any_number_existing"] = any_existing
            fresh.append(row)
    return fresh, conflicts, skipped_existing


def preview_master_import(
    db: Session, *, actor: User, role: str, project_id: int, filename: str, content: bytes
) -> dict:
    """Parse + validate + show a preview. NOTHING is written to the database."""
    ctx = ensure_master_manage(db, actor=actor, role=role, project_id=project_id)
    entries = parse_master_workbook(filename, content)
    fresh, conflicts, skipped = _analyse(db, ctx=ctx, entries=entries)
    lots = sorted({e["lot"] for e in entries})
    codes = sorted({(e["lot"], e["code"]) for e in entries})
    areas = sorted({(e["lot"], e["code"], e["area_name"]) for e in entries})
    return {
        "filename": filename or "work-file-master",
        "project_id": project_id,
        "project_code": ctx["project"].project_code,
        "rows": len(entries),
        "valid": len(entries),
        "errors": 0,
        "lots": len(lots),
        "codes": len(codes),
        "areas": len(areas),
        "files": sum(len(e["item_numbers"]) for e in entries),
        "new_rows": sum(len(e["new_numbers"]) for e in fresh),
        "skipped_existing": skipped,
        "conflicts": conflicts,
        "planned_by_unit": group_planned_by_unit([
            {"planned_quantity": e["area_planned_quantity"], "unit": e["unit"]}
            for e in entries if e.get("area_planned_quantity") is not None
        ]),
        "lot_list": lots,
    }


def commit_master_import(
    db: Session, *, actor: User, role: str, project_id: int, filename: str, content: bytes
) -> dict:
    """Transactional import of the confirmed master workbook (idempotent, history-safe)."""
    ctx = ensure_master_manage(db, actor=actor, role=role, project_id=project_id)
    entries = parse_master_workbook(filename, content)
    fresh, conflicts, skipped = _analyse(db, ctx=ctx, entries=entries)

    packages = list(ctx["packages"])
    packages_by_area: dict[tuple[str, str, str], OrthoWorkPackage] = {
        ((p.lot_batch or "").strip(), (p.package_code or "").strip(), (p.package_name or "").strip()): p
        for p in packages
    }
    created_packages = 0
    created_items = 0
    source = (filename or "work-file-master").strip()[:255]

    for entry in fresh:
        key = (entry["lot"], entry["code"], entry["area_name"])
        package = packages_by_area.get(key)
        if package is None:
            package = OrthoWorkPackage(
                project_id=project_id,
                package_code=entry["code"],
                package_name=entry["area_name"],
                lot_batch=entry["lot"],
                area=Decimal("0"),
                area_unit=entry["unit"],
                instructions=None,
                current_stage="not_started",
                production_state="not_started",
                qc_state="not_started",
                qa_state="not_started",
                created_by_id=actor.id,
            )
            db.add(package)
            db.flush()
            packages_by_area[key] = package
            packages.append(package)
            created_packages += 1
        # Declared physical-file count is source metadata on the area (additive).
        if entry.get("total_files") is not None:
            package.total_files = max(package.total_files or 0, int(entry["total_files"]))
        for number in entry["new_numbers"]:
            if _existing_item(db, package.id, number) is not None:
                skipped += 1  # re-checked by number within the same area
                continue
            item = WorkPackageItem(
                work_package_id=package.id,
                project_id=project_id,
                item_number=number,
                item_name=entry["item_name"] if number == entry["item_numbers"][0] else None,
                sequence_no=int(number) if str(number).isdigit() else None,
                # Count-based areas carry NO fabricated per-file length (NULL); explicit
                # FILE NUMBER rows keep the imported planned length.
                planned_quantity=entry["planned_quantity"],
                planned_unit=entry["unit"],
                status="not_started",
                is_active=True,
                import_source=source,
                created_by_id=actor.id,
            )
            db.add(item)
            created_items += 1
        # Area planned length: applied ONCE per row (never divided by file count), and only
        # on the import that first creates this row's files (idempotent re-import adds nothing).
        area_planned = entry.get("area_planned_quantity")
        if (area_planned is not None and not entry.get("any_number_existing")
                and units_compatible(package.area_unit, entry["unit"])):
            package.area = (package.area or Decimal("0")) + area_planned
        # Imported REMARKS accumulate on the area (deduped) so the AREA MASTER export
        # and the allocation screen show the master remarks.
        if entry.get("remarks"):
            existing_remarks = (package.remarks or "")
            if entry["remarks"] not in existing_remarks:
                package.remarks = (f"{existing_remarks}\n{entry['remarks']}".strip())[:5000]
    db.flush()
    return {
        "created_packages": created_packages,
        "created_items": created_items,
        "skipped_existing": skipped,
        "conflicts": conflicts,
        "rows": len(entries),
    }


# ===========================================================================================
# Allocation: cascading options + assignment
# ===========================================================================================

def allocation_options(db: Session, *, actor: User, role: str, project_id: int) -> dict:
    """Cascading Lot -> Code -> Area -> Files data plus the PM-approved team lists."""
    ctx = _package_actor_flags(db, project_id, actor, role)
    if not ctx["privileged"]:
        raise PermissionError("You are not assigned to this project")
    packages = ctx["packages"]
    items = items_service.items_for_packages(db, [p.id for p in packages])
    latest = items_service.latest_activities_by_item(db, [p.id for p in packages])
    items_by_package: dict[int, list[WorkPackageItem]] = {}
    for item in items:
        items_by_package.setdefault(item.work_package_id, []).append(item)

    areas: list[dict] = []
    for package in packages:
        package_items = [i for i in items_by_package.get(package.id, []) if i.is_active]
        if not package_items and not (package.lot_batch and package.package_code):
            continue
        area_done = {"production": 0, "qc": 0, "qa": 0, "delivered": 0}
        for item in package_items:
            per = latest.get(item.id, {})
            if per.get("production") is not None and per["production"].status == "completed":
                area_done["production"] += 1
            if per.get("qc") is not None and per["qc"].status == "passed":
                area_done["qc"] += 1
            if per.get("qa") is not None and per["qa"].status == "passed":
                area_done["qa"] += 1
            if per.get("delivery") is not None and per["delivery"].status == "delivered":
                area_done["delivered"] += 1
        total = len(package_items)
        assigned = len([i for i in package_items if i.production_user_id])
        areas.append({
            "work_package_id": package.id,
            "lot": package.lot_batch,
            "code": package.package_code,
            "area_name": package.package_name,
            "planned_quantity": float(package.area) if package.area is not None else None,
            "planned_unit": package.area_unit,
            "supply_date": package.supply_date.isoformat() if package.supply_date else None,
            "declared_total_files": package.total_files,
            "file_count": total,
            "assigned_files": assigned,
            "available_files": total - assigned,
            "files_completed": area_done["production"],
            "production_status": ("completed" if total and area_done["production"] == total
                                  else "in_progress" if area_done["production"] else "not_started"),
            "qc_status": ("passed" if total and area_done["qc"] == total
                          else "in_progress" if area_done["qc"] else "not_started"),
            "qa_status": ("passed" if total and area_done["qa"] == total
                          else "in_progress" if area_done["qa"] else "not_started"),
            "ready_for_delivery": bool(total) and area_done["qa"] == total,
            "delivered_files": area_done["delivered"],
        })

    member_rows = db.scalars(select(OrthoProjectMember).where(
        OrthoProjectMember.project_id == project_id,
        OrthoProjectMember.is_active.is_(True),
    )).all()
    user_ids = {m.user_id for m in member_rows}
    users = {u.id: u for u in db.scalars(select(User).where(User.id.in_(user_ids))).all()} if user_ids else {}

    def members(role_name: str) -> list[dict]:
        return [
            {"user_id": uid, "full_name": users[uid].full_name if uid in users else None}
            for uid in sorted({m.user_id for m in member_rows if m.member_role == role_name})
        ]

    lots: dict[str, list[dict]] = {}
    for area in areas:
        lots.setdefault(area["lot"] or "", []).append(area)
    return {
        "project_id": project_id,
        "project_code": ctx["project"].project_code,
        "team_lead_user_id": next((p.team_leader_user_id for p in packages if p.team_leader_user_id), None),
        "lots": [
            {
                "lot": lot,
                "areas": [
                    {"code": a["code"], "area_name": a["area_name"], "work_package_id": a["work_package_id"]}
                    for a in sorted(group, key=lambda x: (x["code"] or "", x["area_name"] or ""))
                ],
            }
            for lot, group in sorted(lots.items())
        ],
        "areas": areas,
        "team": {"production": members("production"), "qc": members("qc"), "qa": members("qa")},
    }


def area_files(db: Session, *, actor: User, role: str, work_package_id: int) -> dict:
    """Files of one area for the allocation file list (planned quantity + live status)."""
    package = items_service.load_package(db, work_package_id)
    ctx = items_service.actor_context(db, package, actor, role)
    if not ctx["privileged"]:
        raise PermissionError("You are not assigned to this project")
    rows = items_service.build_item_rows(db, [package.id])
    active = [r for r in rows if r["is_active"]]
    assigned = [r for r in active if r.get("production_user_id")]
    groups = group_planned_by_unit([
        {"planned_quantity": Decimal(str(r["planned_quantity"])), "unit": r["planned_unit"] or ""}
        for r in active if r["planned_quantity"] is not None
    ])
    return {
        "work_package_id": package.id,
        "package_code": package.package_code,
        "area_name": package.package_name,
        "lot": package.lot_batch,
        "planned_quantity": float(package.area) if package.area is not None else None,
        "planned_unit": package.area_unit,
        "supply_date": package.supply_date.isoformat() if package.supply_date else None,
        "declared_total_files": package.total_files,
        "total_files": len(active),
        "assigned_files": len(assigned),
        "available_files": len(active) - len(assigned),
        # Imported master remarks are exposed READ-ONLY so the allocation UI can show them
        # beside the files; allocation never edits them.
        "remarks": package.remarks,
        "files": active,
        "planned_groups": groups,
    }


def assign_files(
    db: Session,
    *,
    actor: User,
    role: str,
    work_package_id: int,
    item_ids: list[int],
    production_user_id: int | None = None,
    production_target_date: date | None = None,
    qc_user_id: int | None = None,
    qc_target_date: date | None = None,
    qa_user_id: int | None = None,
    qa_target_date: date | None = None,
    supply_date: date | None = None,
    note: str | None = None,
    reassign: bool = False,
) -> dict:
    """Team Lead allocation: assign selected files + stage targets (server-validated).

    Physical-file EXCLUSIVITY is guaranteed by the backend, not the UI: a file already
    owned by another Production employee is rejected (``reassign`` only permits taking
    over an UNSTARTED file with no history). Selected rows are locked (``FOR UPDATE``)
    so two concurrent allocations of the same file cannot both succeed.
    """
    package = items_service.load_package(db, work_package_id)
    ctx = items_service.actor_context(db, package, actor, role)
    items_service._ensure_mutable(ctx)
    if not item_ids:
        raise ValueError("Select at least one file to assign")

    # All three target dates are required and must be in a possible order.
    if production_target_date is None or qc_target_date is None or qa_target_date is None:
        raise ValueError("Production target, QC target and QA target dates are all required")
    if not (production_target_date <= qc_target_date <= qa_target_date):
        raise ValueError("Target dates must satisfy Production target ≤ QC target ≤ QA target")

    if production_user_id:
        items_service._ensure_assignable(db, package, ctx, int(production_user_id), "production")
    if qc_user_id:
        items_service._ensure_assignable(db, package, ctx, int(qc_user_id), "qc")
    if qa_user_id:
        items_service._ensure_assignable(db, package, ctx, int(qa_user_id), "qa")
    if not (production_user_id and qc_user_id and qa_user_id):
        raise ValueError("Production, QC and QA employees are all required")

    # Lock the rows first so concurrent allocation cannot double-assign a physical file.
    items = db.scalars(select(WorkPackageItem).where(
        WorkPackageItem.id.in_(item_ids),
        WorkPackageItem.work_package_id == package.id,
    ).with_for_update()).all()
    if len(items) != len(set(item_ids)):
        raise ValueError("Some selected files do not belong to this area")

    for item in items:
        if not item.is_active:
            raise ValueError(f"File {item.item_number} is inactive and cannot be allocated")
        owner = item.production_user_id
        if owner and int(owner) != int(production_user_id):
            if items_service.open_session_for(db, item_id=item.id, stage="production") is not None \
                    or _item_has_history(db, item):
                raise ValueError(
                    f"File {item.item_number} is already assigned and has work history — "
                    f"it cannot be reallocated"
                )
            if not reassign:
                raise ValueError(
                    f"File {item.item_number} is already assigned to another Production employee"
                )
        item.production_user_id = int(production_user_id)
        item.qc_user_id = int(qc_user_id)
        item.qa_user_id = int(qa_user_id)
        item.production_target_date = production_target_date
        item.qc_target_date = qc_target_date
        item.qa_target_date = qa_target_date

    if supply_date is not None:
        package.supply_date = supply_date
    # Keep the existing per-package Team Lead semantics: the allocating TL becomes the
    # package Team Lead when none is set yet.
    if package.team_leader_user_id is None and ctx["is_tl"]:
        package.team_leader_user_id = actor.id
    # Area rollup of the stage targets (earliest per stage across active files).
    active_items = [i for i in items_service.items_for_packages(db, [package.id]) if i.is_active]
    package.production_target_date = min(
        (i.production_target_date for i in active_items if i.production_target_date), default=None
    )
    package.qc_target_date = min((i.qc_target_date for i in active_items if i.qc_target_date), default=None)
    package.qa_target_date = min((i.qa_target_date for i in active_items if i.qa_target_date), default=None)
    if note:
        existing = (package.remarks or "").strip()
        package.remarks = (f"{existing}\n{note.strip()}" if existing else note.strip())[:5000]
    db.flush()

    url = f"/production-tracking?project={package.project_id}"
    summary = f"{package.lot_batch or '—'} · {package.package_code} · {package.package_name}"
    planned_groups = group_planned_by_unit([
        {"planned_quantity": i.planned_quantity, "unit": i.planned_unit or package.area_unit or ""}
        for i in items if i.planned_quantity is not None
    ])
    planned_text = " + ".join(f"{g['quantity']} {g['unit']}" for g in planned_groups) or "—"
    create_global_notification(
        db,
        event_type="workflow.allocation.production_assigned",
        title=f"New Production Assignment: {package.package_code}",
        message=(
            f"{summary} — {len(items)} file(s), planned {planned_text}. "
            f"Supply date {supply_date or package.supply_date or '—'}, production target {production_target_date}."
        ),
        category="system", target_url=url,
        recipient_user_ids=[int(production_user_id)],
        dedupe_key=f"alloc-production-{package.id}-{int(production_user_id)}-{min(item_ids)}",
    )
    create_global_notification(
        db,
        event_type="workflow.allocation.qc_assigned",
        title=f"QC Assignment Created: {package.package_code}",
        message=f"{summary} — waiting for Production completion. QC target {qc_target_date}.",
        category="system", target_url=url,
        recipient_user_ids=[int(qc_user_id)],
        dedupe_key=f"alloc-qc-{package.id}-{int(qc_user_id)}-{min(item_ids)}",
    )
    create_global_notification(
        db,
        event_type="workflow.allocation.qa_assigned",
        title=f"QA Assignment Created: {package.package_code}",
        message=f"{summary} — waiting for QC pass. QA target {qa_target_date}.",
        category="system", target_url=url,
        recipient_user_ids=[int(qa_user_id)],
        dedupe_key=f"alloc-qa-{package.id}-{int(qa_user_id)}-{min(item_ids)}",
    )
    return {
        "assigned_files": len(items),
        "planned_by_unit": planned_groups,
        "supply_date": package.supply_date.isoformat() if package.supply_date else None,
        "production_target_date": production_target_date.isoformat(),
        "qc_target_date": qc_target_date.isoformat(),
        "qa_target_date": qa_target_date.isoformat(),
    }
