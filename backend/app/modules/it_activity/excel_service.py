from __future__ import annotations

from collections import Counter
from datetime import datetime
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session

from app.modules.it_activity.models import (
    ITHandoverRecord,
    ITPurchaseRecord,
    ITPurchaseRequest,
    ITPurchaseRequestHistory,
)
from app.modules.it_activity.service import (
    IST,
    local_datetime,
    month_bounds,
    monthly_activity_data,
    purchase_request_query,
)

EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NAVY = "0F2744"
GREEN = "059669"
LIGHT_GREEN = "DDF5EB"
LIGHT_BLUE = "EAF2F8"
BORDER = "D5E1EB"
TEXT = "172B3A"
MUTED = "617487"


def _safe(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (int, float, bool, datetime)):
        return value
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text[:32000]


def _header(sheet, columns: list[str]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for col, title in enumerate(columns, 1):
        cell = sheet.cell(1, col, title)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color=BORDER))
    sheet.row_dimensions[1].height = 32


def _write_rows(sheet, columns: list[str], rows: Iterable[dict[str, Any]]) -> int:
    _header(sheet, columns)
    count = 0
    for row_index, row in enumerate(rows, 2):
        count += 1
        for col_index, column in enumerate(columns, 1):
            cell = sheet.cell(row_index, col_index, _safe(row.get(column)))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FBFD")
        sheet.row_dimensions[row_index].height = 30
    for col_index, column in enumerate(columns, 1):
        samples = [len(str(sheet.cell(row, col_index).value or "")) for row in range(1, min(sheet.max_row, 200) + 1)]
        width = min(max([len(column), *samples]) + 2, 45)
        sheet.column_dimensions[get_column_letter(col_index)].width = max(width, 12)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(sheet.max_row, 1)}"
    return count


def _display_item(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "Reporting Month": item.get("reporting_month"),
        "Activity / Business Date": item.get("activity_date"),
        "Activity Time": item.get("activity_time"),
        "System Recorded At": item.get("system_recorded_at"),
        "User": item.get("performed_by"),
        "User Email": item.get("performed_by_email"),
        "Role": item.get("performed_by_role"),
        "Asset Code": item.get("asset_code"),
        "CPU / Asset Tag": item.get("cpu_asset_tag"),
        "Workstation / DC No.": item.get("workstation_no"),
        "Device Category": item.get("device_category"),
        "Department": item.get("department"),
        "Action": item.get("action_label"),
        "Action Type": item.get("action_type"),
        "Field / Component": item.get("field_or_component"),
        "Old Value": item.get("old_value"),
        "New Value": item.get("new_value"),
        "Reason": item.get("reason"),
        "Remarks": item.get("remarks"),
        "Batch / Activity ID": item.get("batch_code"),
        "Source": item.get("source_type"),
        "Time Recorded": "Yes" if item.get("time_recorded") else "No — source Excel had date only",
    }


def build_monthly_it_activity_workbook(db: Session, month_key: str) -> tuple[BytesIO, dict[str, int]]:
    data = monthly_activity_data(db, month_key, limit=100000)
    start_date, end_date, _utc_start, _utc_end = month_bounds(month_key)
    items = data["items"]

    wb = Workbook()
    wb.remove(wb.active)
    counts: dict[str, int] = {}

    summary = wb.create_sheet("Monthly Summary")
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:F1")
    summary["A1"] = f"NakshaTech IT Monthly Activity — {data['month']['label']}"
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    summary["A1"].alignment = Alignment(horizontal="left", vertical="center")
    summary.row_dimensions[1].height = 38
    summary["A3"] = "Reporting Period"
    summary["B3"] = f"{start_date.strftime('%d-%m-%Y')} to {end_date.strftime('%d-%m-%Y')}"
    summary["A4"] = "Timezone"
    summary["B4"] = "Asia/Kolkata"
    summary["A5"] = "Generated At"
    summary["B5"] = datetime.now(IST).strftime("%d-%m-%Y %I:%M:%S %p")
    summary["A7"] = "Metric"
    summary["B7"] = "Count / Value"
    for cell in summary[7]:
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(color="FFFFFF", bold=True)
    metrics = [
        ("Assets Edited", data["summary"]["assets_edited"]),
        ("Asset Edit Operations", data["summary"]["asset_edit_operations"]),
        ("Component Changes", data["summary"]["component_changes"]),
        ("Upgrades", data["summary"]["upgrades"]),
        ("Replacements", data["summary"]["replacements"]),
        ("Downgrades", data["summary"]["downgrades"]),
        ("Upgrade + Replacement", data["summary"]["combined_changes"]),
        ("Laptop Handovers", data["summary"]["laptop_handovers"]),
        ("Desktop Handovers", data["summary"]["desktop_handovers"]),
        ("Returns", data["summary"]["return_operations"]),
        ("Purchases Recorded", data["summary"]["purchases_recorded"]),
        ("Purchase Value (INR)", data["summary"]["purchase_value"]),
        ("Total Detailed Activity Rows", data["summary"]["total_activities"]),
    ]
    for index, (label, value) in enumerate(metrics, 8):
        summary.cell(index, 1, label)
        summary.cell(index, 2, value)
        if index % 2 == 0:
            summary.cell(index, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            summary.cell(index, 2).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 24
    counts["Monthly Summary"] = len(metrics)

    all_columns = [
        "Reporting Month", "Activity / Business Date", "Activity Time", "System Recorded At", "User", "User Email", "Role", "Asset Code", "CPU / Asset Tag",
        "Workstation / DC No.", "Device Category", "Department", "Action", "Action Type",
        "Field / Component", "Old Value", "New Value", "Reason", "Remarks",
        "Batch / Activity ID", "Source", "Time Recorded",
    ]
    all_rows = [_display_item(item) for item in items]
    sheet = wb.create_sheet("Detailed Monthly Activity")
    counts["Detailed Monthly Activity"] = _write_rows(sheet, all_columns, all_rows)

    asset_rows = [_display_item(item) for item in items if item["source_type"] == "asset_edit"]
    sheet = wb.create_sheet("Asset Edit History")
    counts["Asset Edit History"] = _write_rows(sheet, all_columns, asset_rows)

    component_rows = [_display_item(item) for item in items if item["source_type"] == "component_change"]
    sheet = wb.create_sheet("Component Changes")
    counts["Component Changes"] = _write_rows(sheet, all_columns, component_rows)

    handovers = list(db.scalars(
        select(ITHandoverRecord)
        .where(or_(
            ITHandoverRecord.reporting_month == month_key,
            and_(ITHandoverRecord.reporting_month.is_(None), ITHandoverRecord.activity_date >= start_date, ITHandoverRecord.activity_date <= end_date),
        ))
        .order_by(ITHandoverRecord.activity_date, ITHandoverRecord.id)
    ).all())
    handover_columns = [
        "Activity Code", "Device Category", "Asset Code", "Internal Asset No.", "Workstation / DC No.",
        "Employee Name", "Department", "Work Mode", "Specification", "Serial Number",
        "Accessories Provided", "Condition", "Action", "Original Action", "Reporting Month", "Activity Date", "Activity Time",
        "System Recorded At", "Issued / Performed By", "Role", "Remarks", "Asset Updated", "Source File", "Source Sheet", "Source Row",
    ]
    for category, sheet_name in (("laptop", "Laptop Handover Return"), ("desktop", "Desktop Handover Return")):
        rows = []
        for record in handovers:
            if record.device_category != category:
                continue
            rows.append({
                "Activity Code": record.activity_code,
                "Device Category": record.device_category.title(),
                "Asset Code": record.asset_code_snapshot,
                "Internal Asset No.": record.internal_asset_no,
                "Workstation / DC No.": record.dc_number,
                "Employee Name": record.employee_name,
                "Department": record.department,
                "Work Mode": record.work_mode,
                "Specification": record.specification,
                "Serial Number": record.serial_number,
                "Accessories Provided": record.accessories_provided,
                "Condition": record.condition,
                "Action": record.action_type.replace("_", " ").title(),
                "Original Action": record.action_raw,
                "Reporting Month": record.reporting_month or record.activity_date.strftime("%Y-%m"),
                "Activity Date": record.activity_date.strftime("%d-%m-%Y"),
                "Activity Time": record.activity_time.strftime("%I:%M:%S %p") if record.activity_time else "Not recorded in source Excel",
                "System Recorded At": local_datetime(record.created_at).strftime("%d-%m-%Y %I:%M:%S %p") if local_datetime(record.created_at) else None,
                "Issued / Performed By": record.performed_by or record.issued_by,
                "Role": record.performed_by_role,
                "Remarks": record.remarks,
                "Asset Updated": record.asset_updated_status,
                "Source File": record.source_file,
                "Source Sheet": record.source_sheet,
                "Source Row": record.source_row,
            })
        sheet = wb.create_sheet(sheet_name)
        counts[sheet_name] = _write_rows(sheet, handover_columns, rows)

    purchases = list(db.scalars(
        select(ITPurchaseRecord)
        .where(or_(
            ITPurchaseRecord.reporting_month == month_key,
            and_(ITPurchaseRecord.reporting_month.is_(None), ITPurchaseRecord.purchase_date >= start_date, ITPurchaseRecord.purchase_date <= end_date),
        ))
        .order_by(ITPurchaseRecord.purchase_date, ITPurchaseRecord.id)
    ).all())
    purchase_columns = [
        "Purchase Code", "Approval Request Code", "Reporting Month", "Purchase Date", "System Recorded At", "PO Number", "Asset Number", "Linked Asset Code", "Supplier Name",
        "Supplier Contact", "Item Description", "Warranty / Serial No.", "Quantity", "Unit Price (INR)",
        "Total Price (INR)", "Received Date", "Inspection Status", "Approved By", "Department",
        "Remarks", "Recorded By", "Role", "Time", "Source File", "Source Sheet", "Source Row",
    ]
    purchase_rows = []
    for record in purchases:
        purchase_rows.append({
            "Purchase Code": record.purchase_code,
            "Approval Request Code": record.purchase_request_code,
            "Reporting Month": record.reporting_month or record.purchase_date.strftime("%Y-%m"),
            "Purchase Date": record.purchase_date.strftime("%d-%m-%Y"),
            "System Recorded At": local_datetime(record.created_at).strftime("%d-%m-%Y %I:%M:%S %p") if local_datetime(record.created_at) else None,
            "PO Number": record.po_number,
            "Asset Number": record.asset_number,
            "Linked Asset Code": record.linked_asset_code_snapshot,
            "Supplier Name": record.supplier_name,
            "Supplier Contact": record.supplier_contact,
            "Item Description": record.item_description,
            "Warranty / Serial No.": record.warranty_number,
            "Quantity": record.quantity,
            "Unit Price (INR)": record.unit_price,
            "Total Price (INR)": record.total_price,
            "Received Date": record.received_date.strftime("%d-%m-%Y") if record.received_date else None,
            "Inspection Status": record.inspection_status,
            "Approved By": record.approved_by,
            "Department": record.department,
            "Remarks": record.remarks,
            "Recorded By": record.created_by,
            "Role": record.created_by_role,
            "Time": (local_datetime(record.created_at).strftime("%I:%M:%S %p") if local_datetime(record.created_at) else None) if not record.imported else "Not recorded in source Excel",
            "Source File": record.source_file,
            "Source Sheet": record.source_sheet,
            "Source Row": record.source_row,
        })
    sheet = wb.create_sheet("Purchase Details")
    counts["Purchase Details"] = _write_rows(sheet, purchase_columns, purchase_rows)
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, purchase_columns.index("Unit Price (INR)") + 1).number_format = '₹#,##0.00'
        sheet.cell(row, purchase_columns.index("Total Price (INR)") + 1).number_format = '₹#,##0.00'

    user_columns = ["User", "Role", "Asset Edits", "Component Changes", "Handover / Return", "Purchases", "Total Activities"]
    user_data: dict[tuple[str, str], Counter] = {}
    for item in items:
        key = (str(item.get("performed_by") or "Unknown"), str(item.get("performed_by_role") or "Unknown"))
        counter = user_data.setdefault(key, Counter())
        source = item.get("source_type")
        counter[source] += 1
        counter["total"] += 1
    user_rows = [{
        "User": user,
        "Role": role,
        "Asset Edits": counter["asset_edit"],
        "Component Changes": counter["component_change"],
        "Handover / Return": counter["handover_return"],
        "Purchases": counter["purchase"],
        "Total Activities": counter["total"],
    } for (user, role), counter in sorted(user_data.items())]
    sheet = wb.create_sheet("User Activity Summary")
    counts["User Activity Summary"] = _write_rows(sheet, user_columns, user_rows)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream, counts



def build_purchase_request_workbook(
    db: Session,
    *,
    month: str | None = None,
    status: str | None = None,
    department: str | None = None,
    priority: str | None = None,
    search: str | None = None,
) -> tuple[BytesIO, int]:
    requests = list(db.scalars(
        purchase_request_query(
            month,
            status=status,
            department=department,
            priority=priority,
            search=search,
        )
        .order_by(ITPurchaseRequest.requested_at.desc(), ITPurchaseRequest.id.desc())
    ).unique().all())

    wb = Workbook()
    wb.remove(wb.active)

    summary = wb.create_sheet("Approval Summary")
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:F1")
    summary["A1"] = "NakshaTech Purchase Permission & Approval Register"
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    summary["A1"].alignment = Alignment(horizontal="left", vertical="center")
    summary.row_dimensions[1].height = 38
    summary["A3"] = "Reporting Month"
    summary["B3"] = month or "All months"
    summary["A4"] = "Status Filter"
    summary["B4"] = status or "All statuses"
    summary["A5"] = "Department Filter"
    summary["B5"] = department or "All departments"
    summary["A6"] = "Priority Filter"
    summary["B6"] = priority or "All priorities"
    summary["A7"] = "Generated At"
    summary["B7"] = datetime.now(IST).strftime("%d-%m-%Y %I:%M:%S %p")
    summary["A9"] = "Metric"
    summary["B9"] = "Count / Value"
    for cell in summary[9]:
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(color="FFFFFF", bold=True)
    status_counts = Counter(record.status for record in requests)
    metrics = [
        ("Total Requests", len(requests)),
        ("Pending Approval", status_counts.get("pending_approval", 0)),
        ("Approved", status_counts.get("approved", 0)),
        ("Rejected", status_counts.get("rejected", 0)),
        ("Sent Back", status_counts.get("sent_back", 0)),
        ("Purchase Completed", status_counts.get("purchase_completed", 0)),
        ("Estimated Value (INR)", round(sum(float(record.estimated_total_amount or 0) for record in requests), 2)),
        ("Approved Value (INR)", round(sum(float(record.approved_amount or 0) for record in requests), 2)),
        ("Actual Purchase Value (INR)", round(sum(float(record.purchase_record.total_price or 0) for record in requests if record.purchase_record), 2)),
    ]
    for row_index, (label, value) in enumerate(metrics, 10):
        summary.cell(row_index, 1, label)
        summary.cell(row_index, 2, value)
        if row_index % 2 == 0:
            summary.cell(row_index, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            summary.cell(row_index, 2).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 28

    columns = [
        "Request Code",
        "Reporting Month",
        "Requesting Department",
        "Requested Employee",
        "Item Type",
        "Item Name",
        "Item Description",
        "Quantity",
        "Estimated Unit Price (INR)",
        "Estimated Total (INR)",
        "Business Requirement",
        "Required By",
        "Priority",
        "IT Remarks",
        "Status",
        "Branch",
        "Requested By",
        "Requester Email",
        "Requested At",
        "Approved Amount (INR)",
        "Management Remarks",
        "Decision By",
        "Decision Email",
        "Decision At",
        "Purchase Code",
        "Purchase Date",
        "Actual Purchase Amount (INR)",
        "Purchase Completed At",
    ]
    rows: list[dict[str, Any]] = []
    for record in requests:
        purchase = record.purchase_record
        requested_at = local_datetime(record.requested_at)
        decided_at = local_datetime(record.decided_at)
        completed_at = local_datetime(record.purchase_completed_at)
        rows.append({
            "Request Code": record.request_code,
            "Reporting Month": record.reporting_month,
            "Requesting Department": record.requesting_department,
            "Requested Employee": record.requested_employee,
            "Item Type": record.item_type.title(),
            "Item Name": record.item_name,
            "Item Description": record.item_description,
            "Quantity": record.quantity,
            "Estimated Unit Price (INR)": record.estimated_unit_price,
            "Estimated Total (INR)": record.estimated_total_amount,
            "Business Requirement": record.business_reason,
            "Required By": record.required_by_date.strftime("%d-%m-%Y") if record.required_by_date else None,
            "Priority": record.priority.title(),
            "IT Remarks": record.it_remarks,
            "Status": record.status.replace("_", " ").title(),
            "Branch": record.branch,
            "Requested By": record.requested_by_name,
            "Requester Email": record.requested_by_email,
            "Requested At": requested_at.strftime("%d-%m-%Y %I:%M:%S %p") if requested_at else None,
            "Approved Amount (INR)": record.approved_amount,
            "Management Remarks": record.management_remarks,
            "Decision By": record.decided_by_name,
            "Decision Email": record.decided_by_email,
            "Decision At": decided_at.strftime("%d-%m-%Y %I:%M:%S %p") if decided_at else None,
            "Purchase Code": purchase.purchase_code if purchase else None,
            "Purchase Date": purchase.purchase_date.strftime("%d-%m-%Y") if purchase else None,
            "Actual Purchase Amount (INR)": purchase.total_price if purchase else None,
            "Purchase Completed At": completed_at.strftime("%d-%m-%Y %I:%M:%S %p") if completed_at else None,
        })
    sheet = wb.create_sheet("Purchase Requests")
    _write_rows(sheet, columns, rows)
    for column_name in (
        "Estimated Unit Price (INR)",
        "Estimated Total (INR)",
        "Approved Amount (INR)",
        "Actual Purchase Amount (INR)",
    ):
        column_index = columns.index(column_name) + 1
        for row_index in range(2, sheet.max_row + 1):
            sheet.cell(row_index, column_index).number_format = '₹#,##0.00'

    request_ids = [record.id for record in requests]
    histories = list(db.scalars(
        select(ITPurchaseRequestHistory)
        .where(ITPurchaseRequestHistory.request_id.in_(request_ids))
        .order_by(ITPurchaseRequestHistory.created_at, ITPurchaseRequestHistory.id)
    ).all()) if request_ids else []
    request_codes = {record.id: record.request_code for record in requests}
    history_columns = [
        "Request Code",
        "Action",
        "From Status",
        "To Status",
        "Remarks",
        "Performed By",
        "Email",
        "Role",
        "Date & Time",
    ]
    history_rows = []
    for history in histories:
        created_at = local_datetime(history.created_at)
        history_rows.append({
            "Request Code": request_codes.get(history.request_id),
            "Action": history.action.replace("_", " ").title(),
            "From Status": (history.from_status or "").replace("_", " ").title(),
            "To Status": history.to_status.replace("_", " ").title(),
            "Remarks": history.remarks,
            "Performed By": history.performed_by_name,
            "Email": history.performed_by_email,
            "Role": history.performed_by_role.replace("_", " ").title(),
            "Date & Time": created_at.strftime("%d-%m-%Y %I:%M:%S %p") if created_at else None,
        })
    history_sheet = wb.create_sheet("Approval History")
    _write_rows(history_sheet, history_columns, history_rows)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream, len(requests)
