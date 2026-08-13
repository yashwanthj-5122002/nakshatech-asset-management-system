from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from hashlib import sha256
from io import BytesIO
import re
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.models.entities import Asset, User
from app.modules.it_activity.models import ITHandoverRecord, ITPurchaseRecord
from app.modules.it_activity.service import make_code, normalize_action, normalize_text


def _clean_scalar(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.startswith("="):
        text = text[1:]
    return text or None


def _normalise_header(value: Any) -> str:
    text = _clean_scalar(value) or ""
    text = text.lower().replace("&", " and ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _parse_date(value: Any) -> date | None:
    if value in (None, "", "-"):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _parse_number(value: Any) -> float | None:
    if value in (None, "", "-"):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    text = text.replace("₹", "").replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group()) if match else None


def _source_key(*parts: Any) -> str:
    material = "|".join(str(part or "").strip().lower() for part in parts)
    return sha256(material.encode("utf-8")).hexdigest()


def _find_header_row(sheet, required_tokens: tuple[str, ...]) -> tuple[int, dict[str, int]] | None:
    for row_idx in range(1, min(sheet.max_row, 25) + 1):
        mapping: dict[str, int] = {}
        values = []
        for col_idx in range(1, min(sheet.max_column, 40) + 1):
            header = _normalise_header(sheet.cell(row_idx, col_idx).value)
            if header:
                mapping[header] = col_idx
                values.append(header)
        joined = " | ".join(values)
        if all(token in joined for token in required_tokens):
            return row_idx, mapping
    return None


def _column(mapping: dict[str, int], *candidates: str) -> int | None:
    for candidate in candidates:
        normalized = _normalise_header(candidate)
        for header, index in mapping.items():
            if header == normalized or normalized in header or header in normalized:
                return index
    return None


def _cell(sheet, row_idx: int, mapping: dict[str, int], *candidates: str) -> Any:
    col_idx = _column(mapping, *candidates)
    return sheet.cell(row_idx, col_idx).value if col_idx else None


def _match_asset(db: Session, *values: Any) -> Asset | None:
    candidates = [_clean_scalar(value) for value in values]
    candidates = [value for value in candidates if value and value.lower() not in {"own", "n/a", "na", "-"}]
    for value in candidates:
        asset = db.scalar(
            select(Asset).where(or_(
                func.lower(Asset.asset_code) == value.lower(),
                func.lower(Asset.cpu_asset_tag) == value.lower(),
                func.lower(Asset.system_name) == value.lower(),
                func.lower(Asset.workstation_no) == value.lower(),
            )).limit(1)
        )
        if asset:
            return asset
    return None


def import_handover_workbook(
    db: Session,
    content: bytes,
    filename: str,
    device_category: str,
    user: User,
) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
    created = skipped = invalid = 0
    warnings: list[str] = []
    sheets: dict[str, int] = {}

    for sheet in workbook.worksheets:
        header_info = _find_header_row(sheet, ("employee", "handover"))
        if not header_info:
            continue
        header_row, mapping = header_info
        sheet_created = 0
        for row_idx in range(header_row + 1, sheet.max_row + 1):
            employee = _clean_scalar(_cell(sheet, row_idx, mapping, "Employee Name", "Employee / Floor"))
            action_raw = _clean_scalar(_cell(sheet, row_idx, mapping, "Handover/Return", "Handover Return"))
            activity_date = _parse_date(_cell(sheet, row_idx, mapping, "Date"))
            internal_no = _clean_scalar(_cell(sheet, row_idx, mapping, "Meher or Own No.", "Meher or own No.", "Desktop Make & Model"))
            specification = _clean_scalar(_cell(sheet, row_idx, mapping, "Spec", "Specification"))
            accessories = _clean_scalar(_cell(sheet, row_idx, mapping, "Accessories Provided"))
            serial_number = _clean_scalar(_cell(sheet, row_idx, mapping, "Serial Number"))
            dc_number = _clean_scalar(_cell(sheet, row_idx, mapping, "DC NO", "DC NO(SL NO)", "Employee ID"))
            department = _clean_scalar(_cell(sheet, row_idx, mapping, "Department"))
            remarks = _clean_scalar(_cell(sheet, row_idx, mapping, "Remarks"))

            if not any((employee, action_raw, activity_date, internal_no, specification, accessories, serial_number, dc_number, department, remarks)):
                continue
            if not action_raw or not activity_date:
                invalid += 1
                if len(warnings) < 25:
                    warnings.append(f"{sheet.title} row {row_idx}: missing action or valid date")
                continue

            key = _source_key("handover", sheet.title, row_idx, device_category, activity_date, internal_no, employee, action_raw, specification, dc_number)
            if db.scalar(select(ITHandoverRecord.id).where(ITHandoverRecord.source_key == key)):
                skipped += 1
                continue

            asset = _match_asset(db, internal_no, serial_number, dc_number)
            record = ITHandoverRecord(
                activity_code=make_code("ITHR"),
                asset_id=asset.id if asset else None,
                asset_code_snapshot=asset.asset_code if asset else None,
                device_category=device_category,
                employee_name=employee,
                dc_number=dc_number,
                department=department,
                work_mode=_clean_scalar(_cell(sheet, row_idx, mapping, "Work Mode")),
                internal_asset_no=internal_no,
                specification=specification,
                serial_number=serial_number,
                accessories_provided=accessories,
                condition=_clean_scalar(_cell(sheet, row_idx, mapping, "Condition", "Condition at Handover")),
                action_type=normalize_action(action_raw),
                action_raw=action_raw,
                activity_date=activity_date,
                activity_time=None,
                issued_by=_clean_scalar(_cell(sheet, row_idx, mapping, "Issued By")),
                remarks=remarks,
                asset_updated_status=_clean_scalar(_cell(sheet, row_idx, mapping, "asset updated")),
                source_file=filename,
                source_sheet=sheet.title,
                source_row=row_idx,
                source_key=key,
                imported=True,
                performed_by=_clean_scalar(_cell(sheet, row_idx, mapping, "Issued By")) or user.full_name,
                performed_by_email=None,
                performed_by_role="historical_import",
                reporting_month=activity_date.strftime("%Y-%m"),
            )
            db.add(record)
            created += 1
            sheet_created += 1
        if sheet_created:
            sheets[sheet.title] = sheet_created

    db.commit()
    return {"created": created, "skipped": skipped, "invalid": invalid, "sheets": sheets, "warnings": warnings}


def import_purchase_workbook(db: Session, content: bytes, filename: str, user: User) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
    created = skipped = invalid = 0
    warnings: list[str] = []
    sheets: dict[str, int] = {}

    for sheet in workbook.worksheets:
        header_info = _find_header_row(sheet, ("date", "supplier", "item"))
        if not header_info:
            continue
        header_row, mapping = header_info
        sheet_created = 0
        for row_idx in range(header_row + 1, sheet.max_row + 1):
            purchase_date = _parse_date(_cell(sheet, row_idx, mapping, "Date"))
            supplier_name = _clean_scalar(_cell(sheet, row_idx, mapping, "Supplier Name"))
            item_description = _clean_scalar(_cell(sheet, row_idx, mapping, "Item Description"))
            if not any((purchase_date, supplier_name, item_description)):
                continue
            if not purchase_date or not supplier_name or not item_description:
                invalid += 1
                if len(warnings) < 25:
                    warnings.append(f"{sheet.title} row {row_idx}: missing date, supplier or item description")
                continue

            po_number = _clean_scalar(_cell(sheet, row_idx, mapping, "PO Number"))
            asset_number = _clean_scalar(_cell(sheet, row_idx, mapping, "Asset Number"))
            warranty = _clean_scalar(_cell(sheet, row_idx, mapping, "Warranty No.", "Warrenty No.", "S/N", "Additional Details"))
            key = _source_key("purchase", sheet.title, row_idx, purchase_date, po_number, asset_number, supplier_name, item_description, warranty)
            if db.scalar(select(ITPurchaseRecord.id).where(ITPurchaseRecord.source_key == key)):
                skipped += 1
                continue

            quantity = _parse_number(_cell(sheet, row_idx, mapping, "Quantity")) or 1
            unit_price = _parse_number(_cell(sheet, row_idx, mapping, "Unit Price"))
            total_price = _parse_number(_cell(sheet, row_idx, mapping, "Total Price"))
            asset = _match_asset(db, asset_number, warranty)
            record = ITPurchaseRecord(
                purchase_code=make_code("ITPO"),
                linked_asset_id=asset.id if asset else None,
                linked_asset_code_snapshot=asset.asset_code if asset else None,
                purchase_date=purchase_date,
                po_number=po_number,
                asset_number=asset_number,
                supplier_name=supplier_name,
                supplier_contact=_clean_scalar(_cell(sheet, row_idx, mapping, "Supplier Contact")),
                item_description=item_description,
                warranty_number=warranty,
                quantity=quantity,
                unit_price=unit_price,
                total_price=total_price if total_price is not None else (unit_price * quantity if unit_price is not None else None),
                received_date=_parse_date(_cell(sheet, row_idx, mapping, "Received Date")),
                inspection_status=_clean_scalar(_cell(sheet, row_idx, mapping, "Inspection Status")),
                approved_by=_clean_scalar(_cell(sheet, row_idx, mapping, "Approved By")),
                department=_clean_scalar(_cell(sheet, row_idx, mapping, "Department")),
                remarks=_clean_scalar(_cell(sheet, row_idx, mapping, "Remarks")),
                source_file=filename,
                source_sheet=sheet.title,
                source_row=row_idx,
                source_key=key,
                imported=True,
                created_by=_clean_scalar(_cell(sheet, row_idx, mapping, "Approved By")) or user.full_name,
                created_by_email=None,
                created_by_role="historical_import",
                reporting_month=purchase_date.strftime("%Y-%m"),
            )
            db.add(record)
            created += 1
            sheet_created += 1
        if sheet_created:
            sheets[sheet.title] = sheet_created

    db.commit()
    return {"created": created, "skipped": skipped, "invalid": invalid, "sheets": sheets, "warnings": warnings}


def backfill_it_activity_reporting_months(db: Session) -> dict[str, int]:
    """Populate missing effective months from each record's business date.

    Older imports predate the reporting_month column. Their source dates are the
    authoritative business dates, so this backfill is deterministic and safe to
    run repeatedly.
    """
    handovers_updated = 0
    purchases_updated = 0

    handovers = db.scalars(
        select(ITHandoverRecord).where(
            or_(
                ITHandoverRecord.reporting_month.is_(None),
                func.trim(ITHandoverRecord.reporting_month) == "",
            )
        )
    ).all()
    for record in handovers:
        record.reporting_month = record.activity_date.strftime("%Y-%m")
        handovers_updated += 1

    purchases = db.scalars(
        select(ITPurchaseRecord).where(
            or_(
                ITPurchaseRecord.reporting_month.is_(None),
                func.trim(ITPurchaseRecord.reporting_month) == "",
            )
        )
    ).all()
    for record in purchases:
        record.reporting_month = record.purchase_date.strftime("%Y-%m")
        purchases_updated += 1

    if handovers_updated or purchases_updated:
        db.commit()

    return {
        "handovers_updated": handovers_updated,
        "purchases_updated": purchases_updated,
    }


def ensure_bundled_it_activity_reference_data(
    db: Session,
    data_dir: Path,
    user: User,
) -> dict[str, Any]:
    """Load the bundled historical IT activity workbooks exactly once.

    The importers use stable source keys, and this guard also checks the source
    filename before opening a workbook. This keeps application startup fast after
    the first successful import while preserving idempotency.
    """
    result: dict[str, Any] = {
        "backfill": backfill_it_activity_reporting_months(db),
        "files": {},
    }

    handover_jobs = (
        ("laptop", data_dir / "Laptop_Handover_and_Returned.xlsx"),
        ("desktop", data_dir / "Desktop_Handover_and_Returned.xlsx"),
    )
    for category, path in handover_jobs:
        if not path.exists():
            result["files"][path.name] = {"status": "missing"}
            continue
        already_loaded = db.scalar(
            select(ITHandoverRecord.id)
            .where(ITHandoverRecord.source_file == path.name)
            .limit(1)
        )
        if already_loaded is not None:
            result["files"][path.name] = {"status": "already_loaded"}
            continue
        imported = import_handover_workbook(db, path.read_bytes(), path.name, category, user)
        result["files"][path.name] = {"status": "imported", **imported}

    purchase_path = data_dir / "Purchase_Details.xlsx"
    if not purchase_path.exists():
        result["files"][purchase_path.name] = {"status": "missing"}
    else:
        already_loaded = db.scalar(
            select(ITPurchaseRecord.id)
            .where(ITPurchaseRecord.source_file == purchase_path.name)
            .limit(1)
        )
        if already_loaded is not None:
            result["files"][purchase_path.name] = {"status": "already_loaded"}
        else:
            imported = import_purchase_workbook(
                db,
                purchase_path.read_bytes(),
                purchase_path.name,
                user,
            )
            result["files"][purchase_path.name] = {"status": "imported", **imported}

    # A second backfill also covers any rows created by older importer versions
    # or records inserted concurrently before this guarded import completed.
    result["post_import_backfill"] = backfill_it_activity_reporting_months(db)
    return result
