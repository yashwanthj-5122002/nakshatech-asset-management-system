from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

MAIN_SHEET = "May-2026"
TRINITY_SHEET = "Trinity Equipment details"
AMRUT_SHEET = "Amrut 2.0_Inventroy Sheet"
HDD_SHEET = "Hard disk delivered to SoI"
REQUIRED_SHEETS = [MAIN_SHEET, TRINITY_SHEET, AMRUT_SHEET, HDD_SHEET]

MAIN_HEADERS = [
    "S. NO", "DGPS Details", "QTY", "Equipment Manufacturer", "Model Number", "Serial No",
    "Date of Initial Verification", "Instrument | Equipment (ID Number)", "Calibration Required",
    "Maintenance Required", "Frequency", "Responsible Function", "Calibrated by",
    "Last Calibration Date", "Next Calibration Date", "Location of Location",
    "Location of Equipment", "Equipment Tolerance", "Remarks",
]
UIN_HEADERS = ["Application Number", "Serial Number", "Category", "UIN Status", "UIN", "UAV", "Issue Date", "Class"]
TRINITY_HEADERS = ["S. No", "Model", "Description", "Unit No", "Serial No", "UIN No", "Qty ", "Remarks"]
AMRUT_HEADERS = ["S. No", "Assisgn Name", "Equipment Name", "Serial No", "Quantity", "Storage ", "Working", "Remarks"]
HDD_HEADERS = ["S.No", "Date", "Ulbs", "Sq.Km", "Hdd Serial Number", "Storage", "Courier Name", "Courier Details", "Remarks"]

ALL_ATTRIBUTE_GROUPS = {
    MAIN_SHEET: MAIN_HEADERS,
    f"{MAIN_SHEET}::UAV/UIN": UIN_HEADERS,
    TRINITY_SHEET: TRINITY_HEADERS,
    AMRUT_SHEET: AMRUT_HEADERS,
    HDD_SHEET: HDD_HEADERS,
}

DISPLAY_LABELS = {
    "Assisgn Name": "Assigned Name",
    "Storage ": "Storage",
    "Qty ": "Quantity",
    "Ulbs": "ULBs",
    "Hdd Serial Number": "HDD Serial Number",
}

NORMALIZED_FIELDS = {
    "DGPS Details": "asset_name",
    "QTY": "quantity",
    "Equipment Manufacturer": "manufacturer",
    "Model Number": "model_number",
    "Serial No": "serial_number",
    "Date of Initial Verification": "date_of_initial_verification",
    "Instrument | Equipment (ID Number)": "imported_equipment_id",
    "Calibration Required": "calibration_required",
    "Maintenance Required": "maintenance_required",
    "Frequency": "technical_frequency",
    "Responsible Function": "responsible_function",
    "Calibrated by": "calibrated_by",
    "Last Calibration Date": "last_calibration_date",
    "Next Calibration Date": "next_calibration_date",
    "Location of Location": "location_of_location",
    "Location of Equipment": "location_of_equipment",
    "Equipment Tolerance": "equipment_tolerance",
    "Remarks": "remarks",
    "Model": "model_number",
    "Description": "asset_name",
    "Unit No": "unit_number",
    "UIN No": "uin",
    "Equipment Name": "asset_name",
    "Assisgn Name": "current_custodian",
    "Working": "working_condition",
    "Storage ": "storage",
    "Date": "delivery_date",
    "Ulbs": "ulbs",
    "Sq.Km": "square_km",
    "Hdd Serial Number": "hdd_serial_number",
    "Courier Name": "courier_name",
    "Courier Details": "courier_details",
    "Application Number": "application_number",
    "Serial Number": "serial_number",
    "UIN Status": "uin_status",
    "UIN": "uin",
    "UAV": "uav",
    "Issue Date": "issue_date",
    "Class": "uav_class",
}


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return None
    if isinstance(value, (int, float, bool, str)):
        return value
    return str(value)


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _clean(value)
    if not text or text.upper() in {"NA", "N/A", "-"}:
        return None
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d-%B-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def _quantity(value: Any) -> tuple[float, str | None, str | None]:
    raw = _clean(value)
    if value is None:
        return 1.0, raw, None
    if isinstance(value, (int, float)):
        return float(value), raw, None
    match = re.search(r"-?\d+(?:\.\d+)?", str(value))
    if not match:
        return 1.0, raw, None
    qty = float(match.group())
    unit = str(value)[match.end():].strip() or None
    return qty, raw, unit


def _has_multiple_serials(value: Any) -> bool:
    text = _clean(value) or ""
    return any(separator in text for separator in ("/", "&", ",", "\n"))


def _valid_serial(value: Any) -> bool:
    text = (_clean(value) or "").upper()
    return bool(text and text not in {"NA", "N/A", "-", "NONE"})


def _category(asset_name: str | None, model: str | None, section: str | None, sheet: str) -> str:
    text = " ".join(filter(None, [asset_name, model, section, sheet])).lower()
    if "hdd" in text or "hard disk" in text:
        return "Data Storage"
    if "drone" in text or "uav" in text or "vtol" in text:
        return "Drone"
    if "dgps" in text or "rtk station" in text or "trimble da2" in text or "rover" in text:
        return "Survey Equipment"
    if "battery" in text:
        return "Battery"
    if "camera" in text or "insta 360" in text:
        return "Camera / Sensor"
    if "controller" in text or "remote control" in text:
        return "Controller"
    if "laptop" in text or "mobile" in text:
        return "Computing Device"
    if "sd card" in text or "memory card" in text:
        return "Memory Card"
    if "modem" in text or "sim" in text or "air fiber" in text:
        return "Connectivity"
    if "charger" in text or "cable" in text or "adapter" in text:
        return "Power / Cable"
    return "Accessory"


def _status(quantity: float, remarks: str | None, working: str | None = None) -> str:
    text = " ".join(filter(None, [remarks, working])).lower()
    if "sent for service" in text or "sent for services" in text:
        return "sent_for_service"
    if "out" in text or "project" in text:
        return "deployed_to_project"
    if working and working.strip().lower() in {"no", "not working", "non working"}:
        return "damaged"
    if quantity == 0:
        return "pending_verification"
    return "available"


def _record(source_sheet: str, source_row: int, source_section: str | None, record_type: str,
            raw_payload: dict[str, Any], normalized: dict[str, Any], issues: list[str]) -> dict[str, Any]:
    return {
        "source_sheet": source_sheet,
        "source_row": source_row,
        "source_section": source_section,
        "record_type": record_type,
        "original_raw_payload": {key: _json_value(value) for key, value in raw_payload.items()},
        "original_header_map": {key: DISPLAY_LABELS.get(key, key.strip()) for key in raw_payload},
        "normalized_payload": normalized,
        "issues": issues,
        "classification": "new_record",
        "reconciliation_status": "needs_review" if issues else "ready",
    }


def _parse_main(ws) -> list[dict[str, Any]]:
    actual_headers = [ws.cell(1, col).value for col in range(1, 20)]
    if actual_headers != MAIN_HEADERS:
        raise ValueError("May-2026 main headers do not match the required 19-column structure")
    rows: list[dict[str, Any]] = []
    section: str | None = None
    for row_number in range(2, 105):
        values = [ws.cell(row_number, col).value for col in range(1, 20)]
        if not any(value is not None for value in values):
            continue
        first = values[0]
        if not isinstance(first, (int, float)):
            label = _clean(values[1]) or _clean(values[0])
            if label:
                section = label
            continue
        raw = dict(zip(MAIN_HEADERS, values))
        qty, raw_qty, unit = _quantity(raw["QTY"])
        serial = _clean(raw["Serial No"])
        equipment_id = _clean(raw["Instrument | Equipment (ID Number)"])
        name = _clean(raw["DGPS Details"]) or "Unnamed Equipment"
        model = _clean(raw["Model Number"])
        remarks = _clean(raw["Remarks"])
        issues: list[str] = []
        if qty == 0:
            issues.append("quantity_zero")
        if not _valid_serial(serial):
            issues.append("missing_serial_number")
        if _has_multiple_serials(serial):
            issues.append("multiple_serial_numbers_require_split")
        if not equipment_id:
            issues.append("missing_imported_equipment_id")
        normalized = {
            "asset_name": name,
            "category": _category(name, model, section, MAIN_SHEET),
            "subcategory": section,
            "manufacturer": _clean(raw["Equipment Manufacturer"]),
            "model_number": model,
            "serial_number": serial if _valid_serial(serial) and not _has_multiple_serials(serial) else None,
            "raw_serial_number": serial,
            "imported_equipment_id": equipment_id,
            "quantity": qty,
            "raw_quantity": raw_qty,
            "unit_of_measure": unit,
            "tracking_type": "serialized_asset" if _valid_serial(serial) and not _has_multiple_serials(serial) and qty <= 1 else "quantity_asset",
            "date_of_initial_verification": _parse_date(raw["Date of Initial Verification"]),
            "calibration_required": _clean(raw["Calibration Required"]),
            "maintenance_required": _clean(raw["Maintenance Required"]),
            "technical_frequency": _clean(raw["Frequency"]),
            "responsible_function": _clean(raw["Responsible Function"]),
            "calibrated_by": _clean(raw["Calibrated by"]),
            "last_calibration_date": _parse_date(raw["Last Calibration Date"]),
            "next_calibration_date": _parse_date(raw["Next Calibration Date"]),
            "location_of_location": _clean(raw["Location of Location"]),
            "location_of_equipment": _clean(raw["Location of Equipment"]),
            "equipment_tolerance": _clean(raw["Equipment Tolerance"]),
            "remarks": remarks,
            "current_status": _status(qty, remarks),
            "working_condition": None,
            "current_location": _clean(raw["Location of Equipment"]) or _clean(raw["Location of Location"]),
            "current_custodian": None,
            "associated_people": [],
            "is_serialized": bool(_valid_serial(serial) and not _has_multiple_serials(serial)),
            "is_telemetry_capable": name.lower() == "drone",
        }
        rows.append(_record(MAIN_SHEET, row_number, section, "asset", raw, normalized, issues))

    equipment_counts = Counter(
        row["normalized_payload"].get("imported_equipment_id")
        for row in rows if row["normalized_payload"].get("imported_equipment_id")
    )
    for row in rows:
        equipment_id = row["normalized_payload"].get("imported_equipment_id")
        if equipment_id and equipment_counts[equipment_id] > 1:
            row["issues"].append("duplicate_imported_equipment_id")
            row["reconciliation_status"] = "needs_review"
    return rows


def _parse_telecom(ws) -> list[dict[str, Any]]:
    records = []
    for row_number in range(2, ws.max_row + 1):
        connection = ws.cell(row_number, 2).value
        device = ws.cell(row_number, 3).value
        if _clean(device) != "Airtel Modem" or not connection:
            continue
        raw = {"Connection Number": connection, "Device Type": device}
        normalized = {
            "connection_number": str(connection).strip(),
            "device_type": _clean(device),
            "provider": "Airtel",
            "current_status": "available",
            "remarks": None,
        }
        records.append(_record(MAIN_SHEET, row_number, "Airtel modem records", "telecom", raw, normalized, []))
    return records


def _parse_uin(ws) -> list[dict[str, Any]]:
    header_row = None
    for row_number in range(1, ws.max_row + 1):
        values = [ws.cell(row_number, col).value for col in range(2, 10)]
        if values == UIN_HEADERS:
            header_row = row_number
            break
    if header_row is None:
        raise ValueError("Embedded UAV/UIN table was not found in May-2026")
    records = []
    for row_number in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(row_number, col).value for col in range(2, 10)]
        if not any(value is not None for value in values):
            continue
        raw = dict(zip(UIN_HEADERS, values))
        normalized = {
            "application_number": _clean(raw["Application Number"]),
            "serial_number": _clean(raw["Serial Number"]),
            "category": _clean(raw["Category"]),
            "uin_status": _clean(raw["UIN Status"]),
            "uin": _clean(raw["UIN"]),
            "uav": _clean(raw["UAV"]),
            "issue_date": _parse_date(raw["Issue Date"]),
            "uav_class": _clean(raw["Class"]),
        }
        issues = []
        if not normalized["uin"]:
            issues.append("uin_missing")
        if not normalized["serial_number"]:
            issues.append("uin_serial_missing")
        records.append(_record(MAIN_SHEET, row_number, "Embedded UAV/UIN table", "uin", raw, normalized, issues))
    return records


def _parse_trinity(ws) -> list[dict[str, Any]]:
    if [ws.cell(1, col).value for col in range(1, 9)] != TRINITY_HEADERS:
        raise ValueError("Trinity Equipment details headers do not match the expected structure")
    rows = []
    current_model = None
    current_unit = None
    current_uin = None
    current_remarks = None
    current_description = None
    for row_number in range(2, ws.max_row + 1):
        values = [ws.cell(row_number, col).value for col in range(1, 9)]
        if not any(value is not None for value in values):
            continue
        if values == TRINITY_HEADERS:
            current_model = current_unit = current_uin = current_remarks = current_description = None
            continue
        raw = dict(zip(TRINITY_HEADERS, values))
        if _clean(raw["Model"]):
            current_model = _clean(raw["Model"])
        if _clean(raw["Description"]):
            current_description = _clean(raw["Description"])
        if raw["Unit No"] is not None:
            current_unit = _clean(raw["Unit No"])
        if raw["UIN No"] is not None:
            current_uin = _clean(raw["UIN No"])
        if raw["Remarks"] is not None:
            current_remarks = _clean(raw["Remarks"])
        if not current_description and raw["Serial No"] is None:
            continue
        qty, raw_qty, unit = _quantity(raw["Qty "])
        serial = _clean(raw["Serial No"])
        # Serial-only continuation rows represent one physical battery each.
        continuation = raw["S. No"] is None and serial is not None
        normalized_qty = 1.0 if continuation else qty
        issues = []
        if normalized_qty == 0:
            issues.append("quantity_zero")
        if not _valid_serial(serial):
            issues.append("missing_serial_number")
        normalized = {
            "asset_name": current_description or "Trinity Component",
            "category": _category(current_description, current_model, "Trinity Kit", TRINITY_SHEET),
            "subcategory": "Trinity Kit Component",
            "manufacturer": "Quantum-Systems" if current_model and "Trinity" in current_model else None,
            "model_number": current_model,
            "serial_number": serial if _valid_serial(serial) else None,
            "raw_serial_number": serial,
            "imported_equipment_id": None,
            "quantity": normalized_qty,
            "raw_quantity": raw_qty,
            "unit_of_measure": unit,
            "tracking_type": "serialized_asset" if _valid_serial(serial) else "quantity_asset",
            "current_status": _status(normalized_qty, current_remarks),
            "working_condition": None,
            "current_location": "NakshaTech" if current_remarks == "IN" else None,
            "current_custodian": None,
            "associated_people": [],
            "remarks": current_remarks,
            "unit_number": current_unit,
            "uin": current_uin,
            "kit_name": f"Trinity Unit {current_unit}" if current_unit else None,
            "kit_status": "deployed_to_project" if current_remarks and "OUT" in current_remarks.upper() else "available",
            "kit_project": "ICON" if current_remarks and "ICON" in current_remarks.upper() else None,
            "is_serialized": _valid_serial(serial),
            "is_telemetry_capable": bool(current_description and "drone" in current_description.lower()),
        }
        rows.append(_record(TRINITY_SHEET, row_number, f"Trinity Unit {current_unit}" if current_unit else None,
                            "trinity_component", raw, normalized, issues))
    return rows


def _parse_amrut(ws) -> list[dict[str, Any]]:
    if [ws.cell(1, col).value for col in range(1, 9)] != AMRUT_HEADERS:
        raise ValueError("Amrut 2.0_Inventroy Sheet headers do not match the expected structure")
    rows = []
    for row_number in range(2, ws.max_row + 1):
        values = [ws.cell(row_number, col).value for col in range(1, 9)]
        if not any(value is not None for value in values):
            continue
        raw = dict(zip(AMRUT_HEADERS, values))
        name = _clean(raw["Equipment Name"])
        if not name:
            continue
        qty, raw_qty, unit = _quantity(raw["Quantity"])
        serial = _clean(raw["Serial No"])
        custodian_raw = _clean(raw["Assisgn Name"])
        people = [part.strip() for part in re.split(r",|&", custodian_raw or "") if part.strip()]
        working = _clean(raw["Working"])
        remarks = _clean(raw["Remarks"])
        issues = []
        if raw["Quantity"] is None:
            issues.append("missing_quantity")
        if not _valid_serial(serial):
            issues.append("missing_serial_number")
        normalized = {
            "asset_name": name,
            "category": _category(name, None, "Amrut 2.0", AMRUT_SHEET),
            "subcategory": "Amrut 2.0 Equipment",
            "manufacturer": None,
            "model_number": None,
            "serial_number": serial if _valid_serial(serial) else None,
            "raw_serial_number": serial,
            "imported_equipment_id": None,
            "quantity": qty,
            "raw_quantity": raw_qty,
            "unit_of_measure": unit,
            "tracking_type": "serialized_asset" if _valid_serial(serial) else "quantity_asset",
            "current_status": _status(qty, remarks, working),
            "working_condition": working,
            "current_custodian": people[0] if people else custodian_raw,
            "associated_people": people,
            "current_location": None,
            "remarks": remarks,
            "storage": _clean(raw["Storage "]),
            "is_serialized": _valid_serial(serial),
            "is_telemetry_capable": False,
        }
        rows.append(_record(AMRUT_SHEET, row_number, "Amrut 2.0 custody", "amrut_asset", raw, normalized, issues))
    return rows


def _parse_hdd(ws) -> list[dict[str, Any]]:
    header_row = None
    for row_number in range(1, ws.max_row + 1):
        values = [ws.cell(row_number, col).value for col in range(2, 11)]
        if values == HDD_HEADERS:
            header_row = row_number
            break
    if header_row is None:
        raise ValueError("Hard disk delivered to SoI header row was not found")
    rows = []
    for row_number in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(row_number, col).value for col in range(2, 11)]
        if not any(value is not None for value in values):
            continue
        raw = dict(zip(HDD_HEADERS, values))
        normalized = {
            "delivery_date": _parse_date(raw["Date"]),
            "ulbs": _clean(raw["Ulbs"]),
            "square_km": float(raw["Sq.Km"]) if isinstance(raw["Sq.Km"], (int, float)) else None,
            "hdd_serial_number": _clean(raw["Hdd Serial Number"]),
            "storage": _clean(raw["Storage"]),
            "courier_name": _clean(raw["Courier Name"]),
            "courier_details": _clean(raw["Courier Details"]),
            "remarks": _clean(raw["Remarks"]),
            "delivery_status": "delivered",
        }
        issues = []
        if not normalized["hdd_serial_number"]:
            issues.append("missing_hdd_serial_number")
        if not normalized["courier_details"]:
            issues.append("missing_courier_tracking")
        if normalized["square_km"] is None:
            issues.append("missing_square_km")
        rows.append(_record(HDD_SHEET, row_number, "Gujarat Amrut 2.0 Delivery status", "hdd_delivery", raw, normalized, issues))
    return rows


def parse_workbook(content: bytes, filename: str) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), data_only=True)
    missing = [sheet for sheet in REQUIRED_SHEETS if sheet not in workbook.sheetnames]
    if missing:
        raise ValueError(f"Missing required workbook sheets: {', '.join(missing)}")

    records = []
    records.extend(_parse_main(workbook[MAIN_SHEET]))
    records.extend(_parse_telecom(workbook[MAIN_SHEET]))
    records.extend(_parse_uin(workbook[MAIN_SHEET]))
    records.extend(_parse_trinity(workbook[TRINITY_SHEET]))
    records.extend(_parse_amrut(workbook[AMRUT_SHEET]))
    records.extend(_parse_hdd(workbook[HDD_SHEET]))

    type_counts = Counter(record["record_type"] for record in records)
    issue_counts = Counter(issue for record in records for issue in record["issues"])
    summary = {
        "source_workbook": filename,
        "reporting_month": "2026-05",
        "total_records": len(records),
        "main_inventory_records": type_counts["asset"],
        "telecom_connections": type_counts["telecom"],
        "uin_registrations": type_counts["uin"],
        "trinity_components": type_counts["trinity_component"],
        "amrut_assignment_records": type_counts["amrut_asset"],
        "hdd_delivery_transactions": type_counts["hdd_delivery"],
        "trinity_kits": len({r["normalized_payload"].get("unit_number") for r in records if r["record_type"] == "trinity_component" and r["normalized_payload"].get("unit_number")}),
        "records_needing_review": sum(bool(record["issues"]) for record in records),
        "issue_counts": dict(issue_counts),
        "named_source_attributes": sum(len(headers) for headers in ALL_ATTRIBUTE_GROUPS.values()),
        "sheet_names": workbook.sheetnames,
    }
    return {"summary": summary, "records": records, "attribute_groups": ALL_ATTRIBUTE_GROUPS}


def load_template_bytes() -> bytes:
    template = Path(__file__).resolve().parents[2] / "data" / "drone_hardware_inventory_template.xlsx"
    return template.read_bytes()
