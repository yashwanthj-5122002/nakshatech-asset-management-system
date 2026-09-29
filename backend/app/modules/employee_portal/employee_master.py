from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.models.entities import User, utc_now
from app.modules.employee_portal.models import EmployeeMaster, EmployeeMasterPublication


EXPECTED_COLUMNS = (
    "SL.No",
    "Accesscardno",
    "Employee Number",
    "Employee Name",
    "Phone",
    "Curr.Department",
    "Curr.Designation",
    "Email",
)

NOT_REGISTERED = "not_registered"
EMAIL_OTP_PENDING = "email_otp_pending"
EMAIL_VERIFIED = "email_verified"
ACCOUNT_SETUP_PENDING = "account_setup_pending"
AUTHENTICATOR_PENDING = "authenticator_pending"
ACTIVE = "active"
DISABLED = "disabled"
NEEDS_REVIEW = "needs_review"

ACTIVE_EMPLOYMENT = "active"


def _department_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()


# Profile scope only. These values are never copied to User.role.
DEPARTMENT_NORMALIZATION: dict[str, str] = {
    "business development": "bd",
    "civil": "civil",
    "digital marketing": "digital_marketing",
    "finance admin": "finance_admin",
    "housekeeping": "housekeeping",
    "hr": "hr",
    "it": "it",
    "laser scanning": "laser_scanning",
    "lidar": "lidar",
    "mobile mapping": "mobile_mapping",
    "mobile mapping gis": "mobile_mapping",
    "orthophoto": "ortho",
    "photogrammetry": "photogrammetry",
    "project coordination": "project_coordination",
    "reality capture and bim solutions": "reality_capture_bim",
    "software development": "software_team",
    "survey": "survey",
    "tender": "tender",
}

DEPARTMENT_LABELS: dict[str, str] = {
    "bd": "Business Development",
    "civil": "Civil",
    "digital_marketing": "Digital Marketing",
    "finance_admin": "Finance & Admin",
    "housekeeping": "Housekeeping",
    "hr": "HR",
    "it": "IT",
    "laser_scanning": "Laser Scanning",
    "lidar": "LiDAR",
    "mobile_mapping": "Mobile Mapping",
    "ortho": "Orthophoto",
    "photogrammetry": "Photogrammetry",
    "project_coordination": "Project Coordination",
    "reality_capture_bim": "Reality Capture and BIM Solutions",
    "software_team": "Software Development",
    "survey": "Survey",
    "tender": "Tender",
}

TECHNICAL_DEPARTMENT_CODES = {"laser_scanning", "lidar", "civil", "ortho", "mobile_mapping"}


def text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_email(value: str) -> str:
    return value.strip().casefold()


def normalize_identifier(value: str) -> str:
    return re.sub(r"\s+", "", value).casefold()


def normalize_phone(value: str) -> str:
    digits = re.sub(r"\D", "", value)
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    elif len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    return digits


def normalize_department(value: str) -> str | None:
    return DEPARTMENT_NORMALIZATION.get(_department_key(value))


def department_label(code: str | None) -> str:
    return DEPARTMENT_LABELS.get((code or "").strip().casefold(), (code or "").replace("_", " ").title())


def employee_profile(master: EmployeeMaster) -> dict[str, Any]:
    return {
        "id": master.id,
        "employee_name": master.employee_name,
        "employee_number": master.employee_number,
        "access_card_no": master.access_card_no,
        "phone": master.phone,
        "department": master.department_raw,
        "department_code": master.department_code,
        "designation": master.designation_raw,
        "email": master.email,
        "employment_status": master.employment_status,
        "crm_account_status": master.crm_account_status,
    }


def employee_master_for_user(db: Session, user: User) -> EmployeeMaster | None:
    return db.scalar(select(EmployeeMaster).where(EmployeeMaster.linked_user_id == user.id))


def authenticated_department_code(db: Session, user: User) -> str | None:
    master = employee_master_for_user(db, user)
    if master is not None:
        return master.department_code
    return normalize_department(user.department or "")


def is_imported_master(master: EmployeeMaster) -> bool:
    return not master.source_batch_id.startswith("new-joiner-")


def imported_master_is_published(db: Session) -> bool:
    control = db.get(EmployeeMasterPublication, 1)
    return bool(control and control.is_published)


def set_imported_master_publication(db: Session, published: bool, *, actor_id: int | None = None) -> None:
    control = db.get(EmployeeMasterPublication, 1)
    if control is None:
        control = EmployeeMasterPublication(id=1)
        db.add(control)
    control.is_published = published
    control.updated_by_user_id = actor_id
    control.updated_at = utc_now()
    db.flush()


def has_blocking_errors(report: dict[str, Any]) -> bool:
    return any(report[key] for key in (
        "duplicate_emails", "duplicate_employee_numbers", "duplicate_access_cards",
        "missing_mandatory", "unknown_departments", "invalid_email_rows", "invalid_phones",
    ))


def inspect_workbook(path: str | Path | BytesIO, *, source_name: str | None = None) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    source = Path(path) if not isinstance(path, BytesIO) else path
    workbook = load_workbook(source, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        workbook.close()
        raise ValueError(f"Expected one worksheet, found {len(workbook.sheetnames)}")
    sheet = workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    try:
        first_row = next(iterator)
    except StopIteration as exc:
        workbook.close()
        raise ValueError("Employee Master workbook is empty") from exc
    headers = tuple(text_value(value) for value in first_row)
    if headers != EXPECTED_COLUMNS:
        workbook.close()
        raise ValueError(
            "Employee Master columns do not match the required schema. "
            f"Expected {list(EXPECTED_COLUMNS)}, received {list(headers)}"
        )

    records: list[dict[str, Any]] = []
    for excel_row, values in enumerate(iterator, start=2):
        if all(value is None or text_value(value) == "" for value in values):
            continue
        raw = dict(zip(headers, values))
        record = {column: text_value(raw.get(column)) for column in EXPECTED_COLUMNS}
        record.update(
            excel_row=excel_row,
            email_normalized=normalize_email(record["Email"]),
            employee_number_normalized=normalize_identifier(record["Employee Number"]),
            access_card_normalized=normalize_identifier(record["Accesscardno"]),
            phone_normalized=normalize_phone(record["Phone"]),
            department_code=normalize_department(record["Curr.Department"]),
        )
        records.append(record)
        if len(records) > 20000:
            workbook.close()
            raise ValueError("Employee Master workbook exceeds the 20,000 row limit")
    workbook.close()

    duplicate_fields = {
        "duplicate_emails": "email_normalized",
        "duplicate_employee_numbers": "employee_number_normalized",
        "duplicate_access_cards": "access_card_normalized",
    }
    duplicates: dict[str, list[dict[str, Any]]] = {}
    duplicate_rows: set[int] = set()
    for report_name, key in duplicate_fields.items():
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for record in records:
            if record[key]:
                grouped[record[key]].append(record)
        duplicate_groups = [
            {"normalized": value, "rows": [item["excel_row"] for item in items], "count": len(items)}
            for value, items in sorted(grouped.items())
            if len(items) > 1
        ]
        duplicates[report_name] = duplicate_groups
        duplicate_rows.update(row for group in duplicate_groups for row in group["rows"])

    mandatory = ("Accesscardno", "Employee Number", "Employee Name", "Phone", "Curr.Department", "Curr.Designation", "Email")
    missing = [
        {"row": record["excel_row"], "fields": [field for field in mandatory if not record[field]]}
        for record in records
        if any(not record[field] for field in mandatory)
    ]
    invalid_email_rows = [
        record["excel_row"]
        for record in records
        if record["Email"] and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", record["email_normalized"])
    ]
    external = [
        record["excel_row"]
        for record in records
        if record["email_normalized"] and not record["email_normalized"].endswith("@nakshatech.com")
    ]
    unknown_departments = [record["excel_row"] for record in records if record["department_code"] is None]
    invalid_phones = [
        record["excel_row"]
        for record in records
        if len(record["phone_normalized"]) != 10 or record["phone_normalized"][:1] not in {"6", "7", "8", "9"}
    ]
    review_rows = set(external) | set(unknown_departments) | set(invalid_email_rows) | set(invalid_phones) | duplicate_rows
    review_rows.update(item["row"] for item in missing)
    sl_numbers = [int(record["SL.No"]) for record in records if record["SL.No"].isdigit()]
    missing_source_sl_numbers = (
        sorted(set(range(min(sl_numbers), max(sl_numbers) + 1)).difference(sl_numbers)) if sl_numbers else []
    )
    report = {
        "file": source_name or (source.name if isinstance(source, Path) else "uploaded.xlsx"),
        "sheet": sheet.title,
        "headers": list(headers),
        "total_rows": len(records),
        "valid_rows": len(records) - len(review_rows),
        **duplicates,
        "blank_emails": [record["excel_row"] for record in records if not record["Email"]],
        "missing_mandatory": missing,
        "external_emails": external,
        "unknown_departments": unknown_departments,
        "invalid_email_rows": invalid_email_rows,
        "invalid_phones": invalid_phones,
        "rows_requiring_review": sorted(review_rows),
        "missing_source_sl_numbers": missing_source_sl_numbers,
        "department_counts": dict(sorted(Counter(record["Curr.Department"] for record in records).items())),
        "database_changes": 0,
    }
    return records, report


def _status_for_linked_user(user: User) -> str:
    status = (user.account_status or "").strip().casefold()
    if user.is_active and status in {"", "active"}:
        return ACTIVE
    if status in {"pending_mfa", "pending_password_change"}:
        return AUTHENTICATOR_PENDING
    return DISABLED


def _reconcile_user(db: Session, master: EmployeeMaster, warnings: list[dict[str, Any]]) -> None:
    master.linked_user_id = None
    if not master.email_normalized.endswith("@nakshatech.com"):
        master.crm_account_status = NEEDS_REVIEW
        master.review_reason = "External email is not eligible for automatic self-registration"
        return

    users = list(db.scalars(select(User).where(func.lower(User.email) == master.email_normalized)).all())
    if not users:
        master.crm_account_status = NOT_REGISTERED
        master.review_reason = None
        return
    if len(users) != 1:
        master.crm_account_status = NEEDS_REVIEW
        master.review_reason = "Normalized email matches more than one CRM account"
        return

    user = users[0]
    existing_link = db.scalar(
        select(EmployeeMaster).where(EmployeeMaster.linked_user_id == user.id, EmployeeMaster.id != master.id)
    )
    if existing_link is not None:
        master.crm_account_status = NEEDS_REVIEW
        master.review_reason = "CRM account is already linked to a different Employee Master record"
        return

    legacy_id = normalize_identifier(user.employee_id or "")
    allowed_ids = {master.employee_number_normalized, master.access_card_no_normalized}
    if legacy_id and legacy_id not in allowed_ids:
        master.crm_account_status = NEEDS_REVIEW
        master.review_reason = "Existing CRM Employee ID matches neither Employee Number nor Access Card"
        warnings.append(
            {
                "type": "existing_identity_mismatch",
                "email": master.email_normalized,
                "crm_employee_id": user.employee_id,
                "employee_number": master.employee_number,
                "access_card": master.access_card_no,
            }
        )
        return

    if legacy_id:
        collisions = list(
            db.scalars(
                select(User).where(func.lower(User.employee_id) == legacy_id, User.id != user.id)
            ).all()
        )
        if collisions:
            warnings.append(
                {
                    "type": "preserved_legacy_employee_id_collision",
                    "email": master.email_normalized,
                    "crm_employee_id": user.employee_id,
                    "other_user_ids": [item.id for item in collisions],
                }
            )

    master.linked_user_id = user.id
    master.crm_account_status = _status_for_linked_user(user)
    master.review_reason = None


def import_workbook(db: Session, path: str | Path | BytesIO, *, source_batch_id: str | None = None, source_name: str | None = None, actor_id: int | None = None) -> dict[str, Any]:
    if source_batch_id and source_batch_id.startswith("new-joiner-"):
        raise ValueError("The new-joiner batch prefix is reserved for approved onboarding")
    records, report = inspect_workbook(path, source_name=source_name)
    if has_blocking_errors(report):
        raise ValueError("Employee Master validation contains blocking identity or format errors; run --dry-run for details")

    # Every import returns the Excel directory to private review, even when it was published before.
    set_imported_master_publication(db, False, actor_id=actor_id)

    batch_id = source_batch_id or f"employee-master-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{uuid4().hex[:8]}"
    now = utc_now()
    created = 0
    updated = 0
    warnings: list[dict[str, Any]] = []
    for record in records:
        candidates = list(
            db.scalars(
                select(EmployeeMaster).where(
                    or_(
                        EmployeeMaster.employee_number_normalized == record["employee_number_normalized"],
                        EmployeeMaster.access_card_no_normalized == record["access_card_normalized"],
                        EmployeeMaster.email_normalized == record["email_normalized"],
                    )
                )
            ).all()
        )
        unique_candidates = {candidate.id: candidate for candidate in candidates}
        if len(unique_candidates) > 1:
            raise ValueError(
                f"Row {record['excel_row']} matches multiple existing Employee Master records; manual reconciliation is required"
            )
        master = next(iter(unique_candidates.values()), None)
        if master is None:
            master = EmployeeMaster(
                source_sl_no=record["SL.No"],
                access_card_no=record["Accesscardno"],
                access_card_no_normalized=record["access_card_normalized"],
                employee_number=record["Employee Number"],
                employee_number_normalized=record["employee_number_normalized"],
                employee_name=record["Employee Name"],
                phone=record["Phone"],
                phone_normalized=record["phone_normalized"],
                department_raw=record["Curr.Department"],
                department_code=record["department_code"],
                designation_raw=record["Curr.Designation"],
                email=record["Email"],
                email_normalized=record["email_normalized"],
                employment_status=ACTIVE_EMPLOYMENT,
                crm_account_status=NOT_REGISTERED,
                source_batch_id=batch_id,
                imported_at=now,
            )
            db.add(master)
            db.flush()
            created += 1
        else:
            master.source_sl_no = record["SL.No"]
            master.access_card_no = record["Accesscardno"]
            master.access_card_no_normalized = record["access_card_normalized"]
            master.employee_number = record["Employee Number"]
            master.employee_number_normalized = record["employee_number_normalized"]
            master.employee_name = record["Employee Name"]
            master.phone = record["Phone"]
            master.phone_normalized = record["phone_normalized"]
            master.department_raw = record["Curr.Department"]
            master.department_code = record["department_code"]
            master.designation_raw = record["Curr.Designation"]
            master.email = record["Email"]
            master.email_normalized = record["email_normalized"]
            master.source_batch_id = batch_id
            master.imported_at = now
            updated += 1

        _reconcile_user(db, master, warnings)
        if master.email_normalized in {
            records[index]["email_normalized"] for index in range(len(records)) if records[index]["excel_row"] in report["external_emails"]
        }:
            master.crm_account_status = NEEDS_REVIEW
            master.review_reason = "External email is not eligible for automatic self-registration"
        db.flush()

    status_counts = dict(
        Counter(master.crm_account_status for master in db.scalars(select(EmployeeMaster)).all())
    )
    return {
        **report,
        "source_batch_id": batch_id,
        "created": created,
        "updated": updated,
        "preserved_missing_from_source": int(db.scalar(select(func.count(EmployeeMaster.id))) or 0) - len(records),
        "status_counts": status_counts,
        "warnings": warnings,
    }
