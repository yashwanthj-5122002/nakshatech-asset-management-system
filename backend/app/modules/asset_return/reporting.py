from __future__ import annotations

from collections import Counter
from datetime import date, datetime, timedelta
from io import BytesIO
from math import ceil
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.entities import Asset, AssetHistory, ComponentReplacement, ReplacementRecord, WorkRecord
from app.modules.asset_return.models import AssetVendorReturn, SpareMonitor
from app.modules.asset_return.service import active_inventory_assets
from app.services.asset_drilldown_service import (
    SORT_FIELDS,
    VALID_SCOPES,
    _distribution,
    _options,
    _scope_label,
    _scope_matches,
    _search_matches,
    _same,
    _sort_key,
    _status_distribution,
    _status_matches,
    _summary,
    build_asset_drilldown_workbook,
)
from app.services.asset_lifecycle_service import (
    inventory_summary,
    is_primary_device_type,
    lifecycle_status_distribution,
)
from app.services.excel_service import (
    HEADER_FILL,
    HEADER_FONT,
    LIVE_ASSET_HEADERS,
    LIVE_ASSET_WIDTHS,
    NAVY,
    _asset_row,
    _history_fields,
    _latest_history_map,
    _month_bounds,
    _style_header,
    _write_asset_register_sheet,
    build_monthly_summary_report,
)
from app.services.monthly_snapshot_service import (
    assets_for_month,
    get_snapshot_run,
    month_end,
    month_start,
    parse_month_key,
    parse_template_sheet_month,
)
from app.services.periodic_reporting_service import (
    ReportPeriod,
    _approval_rows,
    _asset_rows,
    _asset_source_label,
    _asset_summary_row,
    _build_reconciliation_sheet,
    _build_summary_sheet,
    _component_rows,
    _handover_rows,
    _history_rows,
    _purchase_request_rows,
    _purchase_rows,
    _query_period_records,
    _replacement_rows,
    _work_rows,
    _write_rows,
)
from app.modules.it_activity.models import ITPurchaseRequest


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def query_active_asset_drilldown(
    db: Session,
    *,
    month_key: str,
    scope: str = "all",
    scope_value: str | None = None,
    search: str | None = None,
    department: str | None = None,
    device_type: str | None = None,
    status: str | None = None,
    location: str | None = None,
    work_mode: str | None = None,
    ownership: str | None = None,
    client_name: str | None = None,
    project_id: str | None = None,
    current_holder: str | None = None,
    brand: str | None = None,
    capacity: str | None = None,
    sort_by: str = "asset_code",
    sort_dir: str = "asc",
    page: int = 1,
    page_size: int = 20,
) -> dict[str, Any]:
    if scope not in VALID_SCOPES:
        raise ValueError("Scope must be all, primary, device, status or department")
    if scope not in {"all", "primary"} and not str(scope_value or "").strip():
        raise ValueError("A scope value is required for this drill-down")
    if sort_by not in SORT_FIELDS:
        raise ValueError("Unsupported asset sort field")
    if sort_dir not in {"asc", "desc"}:
        raise ValueError("Sort direction must be asc or desc")
    if page < 1:
        raise ValueError("Page must be at least 1")
    if page_size < 1 or page_size > 100:
        raise ValueError("Page size must be between 1 and 100")

    selected_start = parse_month_key(month_key)
    assets, source = assets_for_month(db, selected_start)
    if source == "missing":
        raise ValueError("No asset register is available for the selected month")
    assets = active_inventory_assets(list(assets))

    scoped_assets = [asset for asset in assets if _scope_matches(asset, scope, scope_value)]
    filtered_assets = []
    for asset in scoped_assets:
        asset_department = str(getattr(asset, "department", None) or "Unassigned").strip()
        asset_location = str(getattr(asset, "location", None) or "Unknown").strip()
        if not _search_matches(asset, search):
            continue
        if department and not _same(asset_department, department):
            continue
        if device_type and not _same(getattr(asset, "device_type", None), device_type):
            # Preserve canonical aliases handled by scope matching through the
            # service-level device filter only when exact values are not used.
            from app.services.asset_lifecycle_service import canonical_device_type
            if canonical_device_type(getattr(asset, "device_type", None)) != canonical_device_type(device_type):
                continue
        if not _status_matches(getattr(asset, "status", None), status):
            continue
        if location and not _same(asset_location, location):
            continue
        if work_mode and not _same(getattr(asset, "work_mode", None), work_mode):
            continue
        if ownership and not _same(getattr(asset, "ownership", None), ownership):
            continue
        if client_name and not _same(getattr(asset, "client_name", None), client_name):
            continue
        if project_id and not _same(getattr(asset, "project_id", None), project_id):
            continue
        if current_holder and not _same(getattr(asset, "current_holder", None), current_holder):
            continue
        if brand and not _same(getattr(asset, "brand", None), brand):
            continue
        if capacity and not _same(getattr(asset, "capacity", None), capacity):
            continue
        filtered_assets.append(asset)

    filtered_assets.sort(key=lambda asset: _sort_key(asset, sort_by), reverse=sort_dir == "desc")
    filtered_total = len(filtered_assets)
    pages = max(1, ceil(filtered_total / page_size))
    safe_page = min(page, pages)
    start_index = (safe_page - 1) * page_size
    page_assets = filtered_assets[start_index:start_index + page_size]
    current = month_start()
    return {
        "month": {
            "key": selected_start.strftime("%Y-%m"),
            "label": selected_start.strftime("%B %Y"),
            "source": source,
            "is_live": selected_start == current,
        },
        "scope": {"type": scope, "value": scope_value, "label": _scope_label(scope, scope_value)},
        "scope_total": len(scoped_assets),
        "filtered_total": filtered_total,
        "page": safe_page,
        "page_size": page_size,
        "pages": pages,
        "summary": _summary(scoped_assets),
        "filtered_summary": _summary(filtered_assets),
        "visuals": {
            "device_distribution": _distribution(filtered_assets, "device_type", "Other"),
            "status_distribution": _status_distribution(filtered_assets),
            "department_distribution": _distribution(filtered_assets, "department", "Unassigned"),
        },
        "filter_options": _options(scoped_assets),
        "assets": page_assets,
        "all_filtered_assets": filtered_assets,
    }


def build_vendor_return_workbook(db: Session) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    returns = list(db.scalars(select(AssetVendorReturn).order_by(AssetVendorReturn.created_at.desc())).all())
    spares = list(db.scalars(select(SpareMonitor).order_by(SpareMonitor.monitor_tag)).all())

    summary.append(["Metric", "Count"])
    _style_header(summary)
    summary.append(["Returned Assets", len(returns)])
    summary.append(["Complete Returns", sum(row.return_mode == "complete_return" for row in returns)])
    summary.append(["Desktop Returns Without Monitor", sum(row.return_mode == "return_without_monitor" for row in returns)])
    summary.append(["Retained Monitors", len(spares)])
    summary.append(["Available Spare Monitors", sum(row.status == "available" for row in spares)])
    summary.append(["Spare Monitors In Use", sum(row.status == "in_use" for row in spares)])
    summary.append(["Damaged Spare Monitors", sum(row.status == "damaged" for row in spares)])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 16

    returned = workbook.create_sheet("Returned Assets")
    returned_headers = [
        "Return ID", "Internal Asset ID", "CPU / Asset Tag", "Device Type", "Return Mode",
        "Return Date", "Vendor", "Return Reference", "Condition", "Previous Status", "Previous User",
        "Previous Department", "Previous Workstation", "Monitor Tags Before Return", "Retained Monitor Tags",
        "Reason", "Remarks", "Reporting Month", "Returned By", "User Email", "User Role", "System Recorded At",
    ]
    returned.append(returned_headers)
    _style_header(returned)
    for row in returns:
        returned.append([
            row.return_code,
            row.asset_code_snapshot,
            row.cpu_asset_tag_snapshot,
            row.device_type_snapshot,
            row.return_mode.replace("_", " ").title(),
            row.return_date,
            row.vendor_name,
            row.return_reference,
            row.condition,
            (row.previous_status or "").replace("_", " ").title(),
            row.previous_used_by,
            row.previous_department,
            row.previous_workstation_no,
            row.monitor_tags_snapshot,
            row.retained_monitor_tags,
            row.reason,
            row.remarks,
            row.reporting_month,
            row.performed_by_name,
            row.performed_by_email,
            row.performed_by_role.replace("_", " ").title(),
            row.created_at,
        ])

    spare_sheet = workbook.create_sheet("Retained Monitors")
    spare_headers = [
        "Monitor Asset Tag", "Status", "Location", "Source Return ID", "Source Desktop Asset ID",
        "Retained Date", "Current Asset ID", "Current CPU / Asset Tag", "Assigned At", "Assigned By",
        "Assigned By Email", "Remarks", "Created At", "Updated At",
    ]
    spare_sheet.append(spare_headers)
    _style_header(spare_sheet)
    asset_map = {asset.id: asset for asset in db.scalars(select(Asset)).all()}
    for row in spares:
        current = asset_map.get(row.current_asset_id) if row.current_asset_id else None
        spare_sheet.append([
            row.monitor_tag,
            row.status.replace("_", " ").title(),
            row.location,
            row.source_return_id,
            row.source_asset_code,
            row.retained_date,
            current.asset_code if current else None,
            current.cpu_asset_tag if current else None,
            row.assigned_at,
            row.assigned_by_name,
            row.assigned_by_email,
            row.remarks,
            row.created_at,
            row.updated_at,
        ])

    audit = workbook.create_sheet("Return Audit History")
    audit.append(["Return ID", "Event", "Asset", "Return Mode", "Performed By", "User Email", "Role", "At", "Reason"])
    _style_header(audit)
    for row in reversed(returns):
        audit.append([
            row.return_code,
            "Vendor Return Recorded",
            row.cpu_asset_tag_snapshot or row.asset_code_snapshot,
            row.return_mode.replace("_", " ").title(),
            row.performed_by_name,
            row.performed_by_email,
            row.performed_by_role.replace("_", " ").title(),
            row.created_at,
            row.reason,
        ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        if sheet.max_column:
            sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{max(sheet.max_row, 1)}"
        for column in range(1, sheet.max_column + 1):
            values = [len(str(sheet.cell(row, column).value or "")) for row in range(1, min(sheet.max_row, 300) + 1)]
            sheet.column_dimensions[get_column_letter(column)].width = min(max(max(values, default=10) + 2, 12), 42)
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, datetime):
                    cell.number_format = "DD-MM-YYYY HH:MM:SS"
                elif isinstance(cell.value, date):
                    cell.number_format = "DD-MM-YYYY"

    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream


def build_active_monthly_asset_report(db: Session, month_key: str) -> BytesIO:
    start = parse_month_key(month_key)
    template_path = Path(settings.seed_excel_path)
    assets, source = assets_for_month(db, start)
    assets = active_inventory_assets(list(assets))
    desired_sheet = start.strftime("%B %Y")

    if source == "template":
        # Original historical workbooks predate this controlled return workflow.
        from app.services.excel_service import build_monthly_asset_report
        return build_monthly_asset_report(db, month_key)

    workbook = load_workbook(template_path) if template_path.exists() else Workbook()
    source_sheet = workbook[workbook.sheetnames[0]]
    ws = workbook[desired_sheet] if desired_sheet in workbook.sheetnames else workbook.copy_worksheet(source_sheet)
    audit_history = _latest_history_map(db) if source == "live" else {}
    _write_asset_register_sheet(workbook, ws, assets, desired_sheet, audit_history)
    for sheet in list(workbook.worksheets):
        if sheet.title != desired_sheet:
            workbook.remove(sheet)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream


def build_active_asset_report(db: Session) -> BytesIO:
    workbook = Workbook()
    ws = workbook.active
    current_sheet_name = datetime.now().astimezone().strftime("%B %Y")
    ws.title = current_sheet_name
    ws.append(LIVE_ASSET_HEADERS)
    yellow_fill = PatternFill("solid", fgColor="FFF200")
    for cell in ws[1]:
        cell.fill = yellow_fill
        cell.font = Font(color="000000", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    assets = active_inventory_assets(list(db.scalars(select(Asset).order_by(Asset.id)).all()))
    latest_history = _latest_history_map(db)
    office_assets = [asset for asset in assets if asset.work_mode == "office"]
    remote_assets = [asset for asset in assets if asset.work_mode in {"wfh", "field"}]
    serial = 1
    for asset in office_assets:
        ws.append(_asset_row(asset, serial, latest_history.get(asset.id)))
        serial += 1
    ws.append([])
    ws.append(["Work From Home Systems and laptops"])
    heading_row = ws.max_row
    ws.merge_cells(start_row=heading_row, start_column=1, end_row=heading_row, end_column=len(LIVE_ASSET_HEADERS))
    ws.cell(heading_row, 1).fill = HEADER_FILL
    ws.cell(heading_row, 1).font = HEADER_FONT
    serial = 1
    for asset in remote_assets:
        ws.append(_asset_row(asset, serial, latest_history.get(asset.id)))
        serial += 1
    ws.append([])
    counts = inventory_summary(assets)
    ws.append([None, "Current Inventory Summary", "Computer", "Laptop", "Mobile", "Printer"])
    ws.append([None, "NakshaTech", counts["computers"], counts["laptops"], counts["smartphones"], counts["printers"]])
    ws.append([None, "Total", counts["computers"], counts["laptops"], counts["smartphones"], counts["printers"]])
    ws.freeze_panes = "A2"
    last_column = get_column_letter(len(LIVE_ASSET_HEADERS))
    ws.auto_filter.ref = f"A1:{last_column}{max(len(office_assets) + 1, 2)}"
    ws.sheet_view.showGridLines = False
    for index, width in enumerate(LIVE_ASSET_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    template_path = Path(settings.seed_excel_path)
    if template_path.exists():
        source_book = load_workbook(template_path, read_only=True, data_only=False, keep_links=False)
        for source_sheet in source_book.worksheets:
            if source_sheet.title == current_sheet_name:
                continue
            target = workbook.create_sheet(source_sheet.title)
            last_nonempty = 0
            cached_rows: list[tuple] = []
            for row_index, values in enumerate(source_sheet.iter_rows(values_only=True), start=1):
                values = tuple(values[:25])
                cached_rows.append(values)
                if any(value not in (None, "") for value in values):
                    last_nonempty = row_index
            for values in cached_rows[:last_nonempty]:
                target.append(values)
            if target.max_row:
                for cell in target[1]:
                    if cell.value is not None:
                        cell.fill = yellow_fill
                        cell.font = Font(color="000000", bold=True)
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            target.freeze_panes = "A2"
            target.sheet_view.showGridLines = False
            for index, width in enumerate(LIVE_ASSET_WIDTHS[:25], start=1):
                target.column_dimensions[get_column_letter(index)].width = width
        source_book.close()

    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream


def build_active_dashboard_report(db: Session) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Dashboard Summary"
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:H2")
    summary["A1"] = "NAKSHATECH IT ASSET MANAGEMENT DASHBOARD"
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")

    all_assets = list(db.scalars(select(Asset)).all())
    assets = active_inventory_assets(all_assets)
    works = list(db.scalars(select(WorkRecord).where(WorkRecord.module == "it")).all())
    primary_assets = [asset for asset in assets if is_primary_device_type(asset.device_type)]
    primary_summary = inventory_summary(primary_assets)
    all_summary = inventory_summary(assets)
    department_counts = Counter(asset.department or "Unassigned" for asset in primary_assets)
    kpis = [
        ("Total IT Assets", primary_summary["total"]), ("Computers", primary_summary["computers"]),
        ("Laptops", primary_summary["laptops"]), ("Smartphones", primary_summary["smartphones"]),
        ("Printers", primary_summary["printers"]), ("External HDDs", all_summary["external_hdds"]),
        ("Servers", primary_summary["servers"]), ("Network Devices", primary_summary["network_devices"]),
        ("Other Device Types", primary_summary["other"]), ("Assigned / In Use", primary_summary["assigned"]),
        ("Available", primary_summary["available"]), ("Under Repair", primary_summary["repair"]),
        ("Replacement Pending", primary_summary["replacement_pending"]), ("Damaged / At Risk", primary_summary["damaged"]),
        ("Retired / Finalized", primary_summary["terminal"]),
    ]
    summary.append([])
    summary.append(["KPI", "Value"])
    _style_header(summary, 4)
    for row in kpis:
        summary.append(row)

    asset_sheet = workbook.create_sheet("Asset Register")
    asset_sheet.append([
        "Asset ID", "Used By", "Workstation", "Department", "Device Type", "CPU Asset Tag",
        "Monitor Asset Tags", "Mouse Asset Tag", "Keyboard Asset Tag", "System Name", "Processor",
        "Memory", "SSD", "HDD", "IP Address", "MAC Address", "Graphics Card", "Operating System",
        "Antivirus", "Network Type", "Work Mode", "Location", "Status", "Remarks", "Asset Date",
        "Brand", "Model", "Serial Number", "Connection Type",
    ])
    _style_header(asset_sheet)
    for asset in sorted(assets, key=lambda item: item.asset_code):
        asset_sheet.append([
            asset.asset_code, asset.used_by, asset.workstation_no, asset.department, asset.device_type,
            asset.cpu_asset_tag, asset.monitor_asset_tags, asset.mouse_asset_tag, asset.keyboard_asset_tag,
            asset.system_name, asset.processor, asset.memory_gb, asset.ssd, asset.hdd, asset.ip_address,
            asset.mac_address, asset.graphics_card, asset.operating_system, asset.antivirus,
            asset.network_type, asset.work_mode, asset.location, asset.status.replace("_", " ").title(),
            asset.remarks, asset.asset_date, asset.brand, asset.model, asset.serial_number, asset.connection_type,
        ])

    status_sheet = workbook.create_sheet("Asset Status")
    status_sheet.append(["Status", "Count"])
    _style_header(status_sheet)
    for item in lifecycle_status_distribution(primary_assets):
        status_sheet.append([item["name"], item["value"]])
    pie = PieChart()
    pie.title = "Asset Status Distribution"
    pie.add_data(Reference(status_sheet, min_col=2, min_row=1, max_row=status_sheet.max_row), titles_from_data=True)
    pie.set_categories(Reference(status_sheet, min_col=1, min_row=2, max_row=status_sheet.max_row))
    status_sheet.add_chart(pie, "D2")

    department_sheet = workbook.create_sheet("Department Summary")
    department_sheet.append(["Department", "Asset Count"])
    _style_header(department_sheet)
    for department, count in department_counts.most_common():
        department_sheet.append([department, count])
    bar = BarChart()
    bar.type = "bar"
    bar.title = "Assets by Department"
    bar.add_data(Reference(department_sheet, min_col=2, min_row=1, max_row=department_sheet.max_row), titles_from_data=True)
    bar.set_categories(Reference(department_sheet, min_col=1, min_row=2, max_row=department_sheet.max_row))
    department_sheet.add_chart(bar, "D2")

    work_sheet = workbook.create_sheet("Work Records")
    work_sheet.append(["Work ID", "Asset", "Title", "Type", "Technician", "Priority", "Status", "Approval", "Created"])
    _style_header(work_sheet)
    asset_by_id = {asset.id: asset.asset_code for asset in all_assets}
    for work in works:
        work_sheet.append([work.work_code, asset_by_id.get(work.asset_id), work.title, work.work_type, work.technician, work.priority, work.status, work.approval_status, work.created_at])

    data_quality_sheet = workbook.create_sheet("Data Quality")
    data_quality_sheet.append(["Check", "Count", "Details"])
    _style_header(data_quality_sheet)
    ip_counts = Counter(asset.ip_address for asset in assets if asset.ip_address and asset.ip_address not in {"-", "Dynamic"})
    duplicate_ips = sorted(ip for ip, count in ip_counts.items() if count > 1)
    for check in [
        ("Missing employee assignment", sum(not asset.used_by for asset in assets), "Complete employee allocation"),
        ("Missing MAC address", sum(not asset.mac_address for asset in assets), "Capture device MAC address"),
        ("Missing IP address", sum(not asset.ip_address for asset in assets), "Capture IP or mark DHCP"),
        ("Missing asset date", sum(not asset.asset_date for asset in assets), "Verify purchase or assignment date"),
        ("Duplicate IP addresses", len(duplicate_ips), ", ".join(duplicate_ips) or "None"),
    ]:
        data_quality_sheet.append(check)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2" if sheet.title != "Dashboard Summary" else "A4"
        for column_index, column in enumerate(sheet.columns, start=1):
            cells = list(column)
            if not cells:
                continue
            max_length = max(len(str(getattr(cell, "value", None) or "")) for cell in cells)
            sheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 42)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream


def build_active_monthly_summary_report(db: Session, month_key: str) -> BytesIO:
    start, _end, start_dt, end_dt = _month_bounds(month_key)
    assets, source = assets_for_month(db, start)
    if source == "template":
        return build_monthly_summary_report(db, month_key)
    closing_assets = active_inventory_assets(list(assets))
    previous_start = (start - timedelta(days=1)).replace(day=1)
    previous_assets, previous_source = assets_for_month(db, previous_start)
    opening_count = len(active_inventory_assets(list(previous_assets))) if previous_source != "missing" else len(closing_assets)

    histories = list(db.scalars(select(AssetHistory).where(or_(
        AssetHistory.reporting_month == month_key,
        and_(AssetHistory.reporting_month.is_(None), AssetHistory.created_at >= start_dt, AssetHistory.created_at < end_dt),
    ))).all())
    manual_added_ids = {item.asset_id for item in histories if item.change_type == "asset_created"}
    manual_added = [asset for asset in closing_assets if getattr(asset, "id", None) in manual_added_ids]
    component_records = list(db.scalars(select(ComponentReplacement).where(or_(
        ComponentReplacement.reporting_month == month_key,
        and_(ComponentReplacement.reporting_month.is_(None), ComponentReplacement.created_at >= start_dt, ComponentReplacement.created_at < end_dt),
    ))).all())
    complete_replacements = db.scalar(select(func.count(ReplacementRecord.id)).where(or_(
        ReplacementRecord.reporting_month == month_key,
        and_(ReplacementRecord.reporting_month.is_(None), ReplacementRecord.created_at >= start_dt, ReplacementRecord.created_at < end_dt),
    ))) or 0
    change_counts = Counter((record.change_type or "replacement") for record in component_records)
    full_edits = [history for history in histories if history.change_type == "full_edit"]
    vendor_returns = db.scalar(select(func.count(AssetVendorReturn.id)).where(or_(
        AssetVendorReturn.reporting_month == month_key,
        and_(AssetVendorReturn.reporting_month.is_(None), AssetVendorReturn.return_date >= start, AssetVendorReturn.return_date <= month_end(start)),
    ))) or 0

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Monthly Asset Summary"
    ws.merge_cells("A1:F2")
    ws["A1"] = f"NAKSHATECH MONTHLY ASSET SUMMARY - {start.strftime('%B %Y').upper()}"
    ws["A1"].fill = HEADER_FILL
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.append([])
    ws.append(["Metric", "Count", "Explanation"])
    _style_header(ws, 4)
    rows = [
        ("Opening Active Assets", opening_count, "Active assets carried into the selected month"),
        ("New Assets Added", len(manual_added), "Manually registered active assets during the selected month"),
        ("Assets Edited", len({history.asset_id for history in full_edits}), "Unique assets updated through Full Edit"),
        ("Full Asset Edit Operations", len(full_edits), "Successful Full Edit save operations"),
        ("Asset Fields Edited", sum(history.field_count or len(_history_fields(history)) for history in full_edits), "Old-to-new fields captured by audit history"),
        ("Component Upgrades", change_counts.get("upgrade", 0), "Individual upgraded fields/components"),
        ("Component Replacements", change_counts.get("replacement", 0), "Individual replacement fields/components"),
        ("Component Downgrades", change_counts.get("downgrade", 0), "Individual downgraded fields/components"),
        ("Upgrade + Replacement Items", change_counts.get("upgrade_replacement", 0), "Combined upgrade/replacement items"),
        ("Complete Asset Replacements", complete_replacements, "Whole computer/laptop replacement workflows"),
        ("Employee / Custody Returns", sum(history.change_type == "asset_returned" for history in histories), "Assets returned from employee custody to IT"),
        ("Rental / Vendor Returns", vendor_returns, "Rental desktops removed from active inventory and returned to vendor"),
        ("Retired / Archived", sum(history.change_type == "asset_retired" for history in histories), "Assets retired from active use"),
        ("Closing Active Assets", len(closing_assets), "Active inventory count after vendor returns"),
    ]
    for row in rows:
        ws.append(row)

    added = workbook.create_sheet("New Assets Added")
    added.append(["Asset ID", "CPU / Asset Tag", "Workstation", "Device", "Used By", "Department", "Created At"])
    _style_header(added)
    for asset in manual_added:
        added.append([asset.asset_code, asset.cpu_asset_tag, asset.workstation_no, asset.device_type, asset.used_by, asset.department, asset.created_at])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2" if sheet.title != "Monthly Asset Summary" else "A4"
        for column in range(1, sheet.max_column + 1):
            letter = get_column_letter(column)
            max_length = max((len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)), default=10)
            sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 45)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream


def _active_assets_for_closing_period(db: Session, period: ReportPeriod) -> tuple[list[Any], str, date | None]:
    current = month_start()
    for month in reversed([item for item in period.months if item <= current]):
        assets, source = assets_for_month(db, month)
        if source != "missing":
            return active_inventory_assets(list(assets)), source, month
    return [], "missing", None


def build_active_complete_period_report(db: Session, period: ReportPeriod) -> BytesIO:
    monthly_rows: list[dict[str, Any]] = []
    for month in period.months:
        assets, source = assets_for_month(db, month)
        monthly_rows.append(_asset_summary_row(month, active_inventory_assets(list(assets)), source))

    closing_assets, closing_source, closing_month = _active_assets_for_closing_period(db, period)
    records = _query_period_records(db, period)
    live_assets = list(db.scalars(select(Asset).order_by(Asset.id)).all())
    assets_by_id = {asset.id: asset for asset in live_assets}
    requests_by_id = {item.id: item for item in db.scalars(select(ITPurchaseRequest)).unique().all()}

    workbook = Workbook()
    workbook.remove(workbook.active)
    _build_summary_sheet(
        workbook,
        period=period,
        closing_month=closing_month,
        closing_source=closing_source,
        closing_assets=closing_assets,
        records=records,
        monthly_rows=monthly_rows,
    )
    row_counts: dict[str, int] = {}
    monthly_columns = [
        "Month", "Month Key", "Source", "Source Available", "Primary IT Assets", "Computers", "Laptops",
        "Smartphones", "Printers", "External HDDs", "Assigned / In Use", "Available", "Under Repair",
        "Replacement Pending", "All Asset Rows",
    ]
    row_counts["Monthly Asset Totals"] = _write_rows(workbook.create_sheet("Monthly Asset Totals"), monthly_columns, monthly_rows)
    asset_columns = [
        "Internal Asset ID", "CPU / Asset Tag", "Device Type", "Assigned User", "Current Holder", "Workstation",
        "Department", "Location / Floor", "Work Mode", "Status", "Brand", "Model", "Serial Number", "Connection",
        "Capacity", "Ownership", "Client", "Project ID", "Processor", "Memory", "SSD", "HDD", "Monitor Asset Tags",
        "Mouse Asset Tag", "Keyboard Asset Tag", "System Name", "IP Address", "MAC Address", "Graphics Card",
        "Operating System", "Antivirus", "Network Type", "Price", "Original Asset Date", "Last Asset Date", "Remarks",
    ]
    row_counts["Primary IT Assets"] = _write_rows(workbook.create_sheet("Primary IT Assets"), asset_columns, _asset_rows(closing_assets, include="primary"))
    row_counts["Printers"] = _write_rows(workbook.create_sheet("Printers"), asset_columns, _asset_rows(closing_assets, include="printers"))
    row_counts["External HDDs"] = _write_rows(workbook.create_sheet("External HDDs"), asset_columns, _asset_rows(closing_assets, include="external_hdds"))

    history_columns = ["Reporting Month", "System Recorded At", "Internal Asset ID", "CPU / Asset Tag", "Device Type", "Assigned User", "Department", "Action", "Change Type", "Batch ID", "Fields Changed", "Old Value", "New Value", "Reason", "Remarks", "Changed By", "Changed By Email", "Changed By Role"]
    row_counts["Asset Change History"] = _write_rows(workbook.create_sheet("Asset Change History"), history_columns, _history_rows(records["histories"], assets_by_id))
    work_columns = ["Reporting Month", "Work Code", "Internal Asset ID", "CPU / Asset Tag", "Title", "Work Type", "Project", "Assigned To", "Technician", "Priority", "Issue Description", "Details", "Status", "Root Cause", "Resolution", "Replaced Component", "Replacement Asset Tag", "Cost", "Approval Status", "Start Date", "Expected Completion", "Completed At", "System Recorded At", "Last Updated At"]
    row_counts["Work Records"] = _write_rows(workbook.create_sheet("Work Records"), work_columns, _work_rows(records["work_records"], assets_by_id))
    component_columns = ["Reporting Month", "Replacement Code", "Batch ID", "Sequence", "Internal Asset ID", "CPU / Asset Tag", "Workstation", "Component Type", "Change Type", "Field", "Old Value", "New Value", "Old Condition", "Reason", "Technician", "Replacement Date", "Performed By", "Performed By Email", "Performed By Role", "Approved By", "Remarks", "System Recorded At"]
    row_counts["Component Changes"] = _write_rows(workbook.create_sheet("Component Changes"), component_columns, _component_rows(records["components"], assets_by_id))
    replacement_columns = ["Reporting Month", "Replacement Code", "Old Asset ID", "Old CPU / Asset Tag", "New Asset ID", "New CPU / Asset Tag", "Reason", "Damage Category", "Inspection Finding", "Approval Status", "Final Action", "Requested By", "Requested By Email", "Requested By Role", "Approved By", "Approved By Email", "Approved By Role", "Requested At", "Approved At"]
    row_counts["Full Replacements"] = _write_rows(workbook.create_sheet("Full Replacements"), replacement_columns, _replacement_rows(records["replacements"], assets_by_id))
    handover_columns = ["Reporting Month", "Activity Code", "Device Category", "Action", "Original Action", "Activity Date", "Activity Time", "Internal Asset ID", "Internal Asset No.", "DC / Workstation No.", "Employee Name", "Department", "Work Mode", "Specification", "Serial Number", "Accessories Provided", "Condition", "Issued / Performed By", "Performed By Email", "Performed By Role", "Remarks", "Asset Updated Status", "Source File", "Source Sheet", "Source Row", "Imported", "System Recorded At"]
    row_counts["Handovers Returns"] = _write_rows(workbook.create_sheet("Handovers Returns"), handover_columns, _handover_rows(records["handovers"]))
    purchase_columns = ["Reporting Month", "Purchase Code", "Purchase Request ID", "Linked Asset ID", "Purchase Date", "PO Number", "Asset Number", "Supplier", "Supplier Contact", "Item Description", "Warranty Number", "Quantity", "Unit Price", "Total Price", "Received Date", "Inspection Status", "Approved By", "Department", "Remarks", "Created By", "Created By Email", "Created By Role", "Source File", "Source Sheet", "Source Row", "Imported", "System Recorded At"]
    row_counts["Purchases"] = _write_rows(workbook.create_sheet("Purchases"), purchase_columns, _purchase_rows(records["purchases"]))
    request_columns = ["Reporting Month", "Request Code", "Requesting Department", "Requested Employee", "Item Type", "Item Name", "Item Description", "Quantity", "Estimated Unit Price", "Estimated Total Amount", "Business Reason", "Required By Date", "Priority", "IT Remarks", "Status", "Branch", "Requested By", "Requested By Email", "Requested By Role", "Requested At", "Approved Amount", "Management Remarks", "Decided By", "Decided By Email", "Decided By Role", "Decided At", "Purchase Completed At", "Last Updated At"]
    row_counts["Purchase Requests"] = _write_rows(workbook.create_sheet("Purchase Requests"), request_columns, _purchase_request_rows(records["requests"]))
    approval_columns = ["Action Month", "Request Code", "Request Reporting Month", "Action", "From Status", "To Status", "Remarks", "Performed By", "Performed By Email", "Performed By Role", "System Recorded At"]
    row_counts["Approval History"] = _write_rows(workbook.create_sheet("Approval History"), approval_columns, _approval_rows(records["request_histories"], requests_by_id))

    vendor_sheet = workbook.create_sheet("Rental Vendor Returns")
    vendor_columns = ["Return ID", "Asset ID", "CPU / Asset Tag", "Return Mode", "Return Date", "Vendor", "Return Reference", "Retained Monitor Tags", "Reason", "Returned By", "Recorded At"]
    vendor_rows = []
    for row in db.scalars(select(AssetVendorReturn).where(AssetVendorReturn.return_date >= period.start_date, AssetVendorReturn.return_date <= period.end_date).order_by(AssetVendorReturn.return_date)).all():
        vendor_rows.append({
            "Return ID": row.return_code, "Asset ID": row.asset_code_snapshot, "CPU / Asset Tag": row.cpu_asset_tag_snapshot,
            "Return Mode": row.return_mode.replace("_", " ").title(), "Return Date": row.return_date, "Vendor": row.vendor_name,
            "Return Reference": row.return_reference, "Retained Monitor Tags": row.retained_monitor_tags, "Reason": row.reason,
            "Returned By": row.performed_by_name, "Recorded At": row.created_at,
        })
    _write_rows(vendor_sheet, vendor_columns, vendor_rows)

    _build_reconciliation_sheet(
        workbook,
        monthly_rows=monthly_rows,
        closing_assets=closing_assets,
        closing_month=closing_month,
        closing_source=closing_source,
        records=records,
        row_counts=row_counts,
    )
    for sheet in workbook.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_view.zoomScale = 85
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream


__all__ = [
    "EXCEL_MIME",
    "build_active_asset_report",
    "build_active_complete_period_report",
    "build_active_dashboard_report",
    "build_active_monthly_asset_report",
    "build_active_monthly_summary_report",
    "build_asset_drilldown_workbook",
    "build_vendor_return_workbook",
    "query_active_asset_drilldown",
]
