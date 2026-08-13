from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from app.models.entities import ApprovalDecisionHistory, Asset, ReplacementRecord, WorkRecord
from app.modules.employee_portal.models import SupportTicket
from app.modules.it_activity.models import ITPurchaseRequest
from app.modules.it_activity.service import month_bounds
from app.services.approval_workflow_service import WORKFLOW_PURCHASE_REQUEST
from app.services.asset_lifecycle_service import inventory_summary, is_primary_device_type
from app.services.ticket_sla_service import SLA_TERMINAL_STATUSES, build_ticket_sla_snapshot_map


def _month_filter(reporting_column, timestamp_column, month: str | None):
    if not month:
        return None
    _start, _end, utc_start, utc_end = month_bounds(month)
    return or_(
        reporting_column == month,
        and_(
            reporting_column.is_(None),
            timestamp_column >= utc_start,
            timestamp_column < utc_end,
        ),
    )


def _history_month_filter(month: str | None):
    conditions = [ApprovalDecisionHistory.workflow_type == WORKFLOW_PURCHASE_REQUEST]
    if month:
        _start, _end, utc_start, utc_end = month_bounds(month)
        conditions.extend([
            ApprovalDecisionHistory.created_at >= utc_start,
            ApprovalDecisionHistory.created_at < utc_end,
        ])
    return and_(*conditions)


def management_control_center(db: Session, month: str | None = None) -> dict[str, Any]:
    assets = list(db.scalars(select(Asset)).all())
    primary_assets = [asset for asset in assets if is_primary_device_type(asset.device_type)]
    inventory = inventory_summary(primary_assets)

    work_filter = _month_filter(WorkRecord.reporting_month, WorkRecord.created_at, month)
    replacement_filter = _month_filter(ReplacementRecord.reporting_month, ReplacementRecord.created_at, month)
    purchase_filter = _month_filter(ITPurchaseRequest.reporting_month, ITPurchaseRequest.requested_at, month)

    active_work_query = select(func.count(WorkRecord.id)).where(
        WorkRecord.module == "it",
        WorkRecord.status.not_in(["completed", "closed"]),
    )
    if work_filter is not None:
        active_work_query = active_work_query.where(work_filter)
    active_it_work = int(db.scalar(active_work_query) or 0)

    replacement_query = select(func.count(ReplacementRecord.id)).where(
        ReplacementRecord.final_action.in_(["replacement_pending", "procurement_required"])
    )
    if replacement_filter is not None:
        replacement_query = replacement_query.where(replacement_filter)
    active_replacements = int(db.scalar(replacement_query) or 0)

    purchase_query = select(ITPurchaseRequest).where(ITPurchaseRequest.status == "pending_approval")
    if purchase_filter is not None:
        purchase_query = purchase_query.where(purchase_filter)
    pending_purchases = list(db.scalars(
        purchase_query.order_by(ITPurchaseRequest.requested_at.asc(), ITPurchaseRequest.id.asc())
    ).all())

    approval_items: list[dict[str, Any]] = []
    for request in pending_purchases:
        approval_items.append({
            "workflow": "purchase_request",
            "id": request.id,
            "code": request.request_code,
            "title": request.item_name,
            "submitted_by": request.requested_by_name,
            "submitted_at": request.requested_at,
            "department": request.requesting_department,
            "priority": request.priority,
            "amount": request.estimated_total_amount,
            "reason": request.business_reason,
            "status": request.status,
            "target_url": "/it/purchase-requests",
            "reporting_month": request.reporting_month,
            "metadata": {
                "requested_employee": request.requested_employee,
                "quantity": request.quantity,
                "item_type": request.item_type,
                "required_by_date": request.required_by_date,
                "it_remarks": request.it_remarks,
            },
        })

    open_tickets = list(db.scalars(
        select(SupportTicket)
        .where(SupportTicket.status.not_in(SLA_TERMINAL_STATUSES))
        .order_by(SupportTicket.created_at.asc(), SupportTicket.id.asc())
        .limit(1000)
    ).all())
    sla_map = build_ticket_sla_snapshot_map(db, open_tickets)
    critical_tickets = [ticket for ticket in open_tickets if ticket.priority == "critical"]
    breached_tickets = [ticket for ticket in open_tickets if sla_map.get(ticket.id, {}).get("sla_breached")]
    warning_tickets = [ticket for ticket in open_tickets if sla_map.get(ticket.id, {}).get("sla_warning")]

    risk_tickets = sorted(
        open_tickets,
        key=lambda ticket: (
            0 if sla_map.get(ticket.id, {}).get("sla_breached") else 1,
            0 if ticket.priority == "critical" else 1,
            ticket.created_at,
        ),
    )[:12]
    ticket_items = [
        {
            "id": ticket.id,
            "ticket_code": ticket.ticket_code,
            "title": ticket.title,
            "department": ticket.department,
            "priority": ticket.priority,
            "status": ticket.status,
            "created_at": ticket.created_at,
            "sla_status": sla_map.get(ticket.id, {}).get("sla_status"),
            "sla_due_at": sla_map.get(ticket.id, {}).get("sla_due_at"),
            "sla_breached": bool(sla_map.get(ticket.id, {}).get("sla_breached")),
            "sla_warning": bool(sla_map.get(ticket.id, {}).get("sla_warning")),
            "target_url": f"/tickets/{ticket.id}",
        }
        for ticket in risk_tickets
    ]

    history_query = (
        select(ApprovalDecisionHistory)
        .where(_history_month_filter(month))
        .order_by(ApprovalDecisionHistory.created_at.desc(), ApprovalDecisionHistory.id.desc())
        .limit(50)
    )
    recent_history = list(db.scalars(history_query).all())

    purchase_value_query = select(func.coalesce(func.sum(ITPurchaseRequest.approved_amount), 0)).where(
        ITPurchaseRequest.status.in_(["approved", "purchase_completed"])
    )
    if purchase_filter is not None:
        purchase_value_query = purchase_value_query.where(purchase_filter)
    approved_purchase_value = float(db.scalar(purchase_value_query) or 0)

    return {
        "generated_at": datetime.now(timezone.utc).replace(tzinfo=None),
        "month": month,
        "authority_model": "purchase_approval_only",
        "executive": {
            "primary_assets": inventory.get("total", 0),
            "assigned_assets": inventory.get("assigned", 0),
            "available_assets": inventory.get("available", 0),
            "repair_assets": inventory.get("repair", 0),
            "replacement_pending_assets": inventory.get("replacement_pending", 0),
            "active_it_work": active_it_work,
            "active_replacements": active_replacements,
            "pending_approvals": len(pending_purchases),
            "pending_purchase_requests": len(pending_purchases),
            "approved_purchase_value": round(approved_purchase_value, 2),
            "open_critical_tickets": len(critical_tickets),
            "sla_warnings": len(warning_tickets),
            "sla_breaches": len(breached_tickets),
        },
        "approvals": approval_items,
        "risk_tickets": ticket_items,
        "recent_decisions": [
            {
                "id": history.id,
                "workflow": history.workflow_type,
                "record_id": history.record_id,
                "code": history.record_code,
                "action": history.action,
                "from_status": history.from_status,
                "to_status": history.to_status,
                "remarks": history.remarks,
                "performed_by": history.performed_by_name,
                "performed_by_role": history.performed_by_role,
                "created_at": history.created_at,
            }
            for history in recent_history
        ],
    }


def build_management_control_workbook(data: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Executive Summary"
    summary.append(["Metric", "Value"])
    for key, value in data["executive"].items():
        summary.append([key.replace("_", " ").title(), value])

    approvals = workbook.create_sheet("Purchase Approvals")
    approvals.append([
        "Request", "Item", "Submitted By", "Submitted At", "Department",
        "Priority", "Estimated Amount", "Business Reason", "Reporting Month",
    ])
    for item in data["approvals"]:
        approvals.append([
            item["code"], item["title"], item["submitted_by"], item["submitted_at"],
            item["department"], item["priority"], item["amount"], item["reason"],
            item["reporting_month"],
        ])

    decisions = workbook.create_sheet("Purchase Decisions")
    decisions.append([
        "Request", "Action", "From", "To", "Performed By", "Role", "At", "Remarks",
    ])
    for item in data["recent_decisions"]:
        decisions.append([
            item["code"], item["action"], item["from_status"], item["to_status"],
            item["performed_by"], item["performed_by_role"], item["created_at"], item["remarks"],
        ])

    risks = workbook.create_sheet("Ticket Risk")
    risks.append(["Ticket", "Title", "Department", "Priority", "Status", "SLA Status", "Due At"])
    for item in data["risk_tickets"]:
        risks.append([
            item["ticket_code"], item["title"], item["department"], item["priority"],
            item["status"], item["sla_status"], item["sla_due_at"],
        ])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
