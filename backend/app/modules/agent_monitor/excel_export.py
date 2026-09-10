from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.api.dependencies import require_roles
from app.models.entities import User
from app.modules.agent_monitor.service import AgentMonitorService

EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NAVY = "0F2744"
BLUE = "0B6F8F"
LIGHT_BLUE = "EAF4F8"
BORDER = "D5E1EB"

router = APIRouter()

AGENT_COLUMNS = [
    "Agent ID",
    "Status",
    "Workstation Number",
    "CPU Asset Tag",
    "CPU Asset Number / Legacy Alias",
    "Hostname",
    "Current NT ID",
    "Current Username",
    "Current Display Name",
    "Department",
    "Operating System",
    "Total RAM (GB)",
    "Primary Local IP",
    "Last Public IP",
    "Building",
    "Floor",
    "Room",
    "Cabin",
    "Current Application",
    "Current Process",
    "Activity State",
    "Idle Seconds",
    "Listening Port Count",
    "First Seen",
    "Last Seen",
]


def _safe_excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (int, float, bool, datetime)):
        return value
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text[:32000]


def _ram_gb(total_ram_bytes: Any) -> float | None:
    try:
        if total_ram_bytes in (None, ""):
            return None
        return round(float(total_ram_bytes) / 1024 / 1024 / 1024, 2)
    except (TypeError, ValueError):
        return None


def _display_row(agent: dict[str, Any]) -> dict[str, Any]:
    return {
        "Agent ID": agent.get("agent_id"),
        "Status": agent.get("status"),
        "Workstation Number": agent.get("workstation_number"),
        "CPU Asset Tag": agent.get("cpu_asset_tag"),
        "CPU Asset Number / Legacy Alias": agent.get("cpu_asset_number"),
        "Hostname": agent.get("hostname"),
        "Current NT ID": agent.get("current_nt_id"),
        "Current Username": agent.get("current_username"),
        "Current Display Name": agent.get("current_display_name"),
        "Department": agent.get("department"),
        "Operating System": agent.get("operating_system"),
        "Total RAM (GB)": _ram_gb(agent.get("total_ram_bytes")),
        "Primary Local IP": agent.get("primary_local_ip"),
        "Last Public IP": agent.get("last_public_ip"),
        "Building": agent.get("building_name"),
        "Floor": agent.get("floor_number"),
        "Room": agent.get("room_number"),
        "Cabin": agent.get("cabin_name"),
        "Current Application": agent.get("current_application"),
        "Current Process": agent.get("current_process"),
        "Activity State": agent.get("activity_state"),
        "Idle Seconds": agent.get("idle_seconds"),
        "Listening Port Count": agent.get("listening_port_count"),
        "First Seen": agent.get("first_seen_at"),
        "Last Seen": agent.get("last_seen_at"),
    }


async def fetch_all_agents() -> list[dict[str, Any]]:
    """Read every System Manager agent through the existing backend-only gateway.

    System Manager caps the admin list endpoint at 100 records per page. Export
    must therefore walk all pages rather than exporting only the first screen.
    Duplicate agent IDs are de-duplicated defensively in case the remote list
    changes while the workbook is being generated.
    """

    page = 1
    page_size = 100
    collected: list[dict[str, Any]] = []
    seen_ids: set[str] = set()

    while True:
        payload = await AgentMonitorService.get_json(
            "/api/v1/admin/agents",
            {"page": page, "page_size": page_size},
        )
        if not isinstance(payload, dict):
            break

        items = payload.get("items") or []
        if not isinstance(items, list):
            break

        for item in items:
            if not isinstance(item, dict):
                continue
            agent_id = str(item.get("agent_id") or "").strip()
            if agent_id and agent_id in seen_ids:
                continue
            if agent_id:
                seen_ids.add(agent_id)
            collected.append(item)

        try:
            total = int(payload.get("total") or 0)
        except (TypeError, ValueError):
            total = 0

        if not items:
            break
        if total:
            if len(collected) >= total:
                break
        elif len(items) < page_size:
            break

        page += 1
        if page > 10000:
            break

    return collected


def build_agent_inventory_workbook(agents: list[dict[str, Any]]) -> BytesIO:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:D1")
    summary["A1"] = "NakshaTech Agent Inventory Export"
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 38

    status_counts: dict[str, int] = {}
    for agent in agents:
        status = str(agent.get("status") or "unknown").strip().lower() or "unknown"
        status_counts[status] = status_counts.get(status, 0) + 1

    summary_rows = [
        ("Generated At (UTC)", datetime.now(timezone.utc).strftime("%d-%m-%Y %H:%M:%S")),
        ("Source", "System Manager Agent Monitoring gateway"),
        ("Total Agents", len(agents)),
        ("Online", status_counts.get("online", 0)),
        ("Delayed", status_counts.get("delayed", 0)),
        ("Offline", status_counts.get("offline", 0)),
        ("Revoked", status_counts.get("revoked", 0)),
        ("Comparison Key", "CPU Asset Tag — compare this column with the IT Asset Register CPU / Asset Tag"),
    ]
    for row_index, (label, value) in enumerate(summary_rows, 3):
        summary.cell(row_index, 1, label).font = Font(bold=True)
        summary.cell(row_index, 2, _safe_excel_value(value))
        if row_index % 2 == 1:
            summary.cell(row_index, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            summary.cell(row_index, 2).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 78

    sheet = wb.create_sheet("Agent Inventory")
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False

    for col_index, title in enumerate(AGENT_COLUMNS, 1):
        cell = sheet.cell(1, col_index, title)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color=BORDER))
    sheet.row_dimensions[1].height = 34

    for row_index, agent in enumerate(agents, 2):
        row = _display_row(agent)
        for col_index, column in enumerate(AGENT_COLUMNS, 1):
            cell = sheet.cell(row_index, col_index, _safe_excel_value(row.get(column)))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FBFD")
        sheet.row_dimensions[row_index].height = 28

    for col_index, column in enumerate(AGENT_COLUMNS, 1):
        samples = [
            len(str(sheet.cell(row, col_index).value or ""))
            for row in range(1, min(sheet.max_row, 200) + 1)
        ]
        width = min(max([len(column), *samples]) + 2, 42)
        sheet.column_dimensions[get_column_letter(col_index)].width = max(width, 12)

    last_column = get_column_letter(len(AGENT_COLUMNS))
    sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@router.get("/export.xlsx")
async def download_agent_inventory_excel(
    user: User = Depends(require_roles("software_team")),
) -> StreamingResponse:
    agents = await fetch_all_agents()
    stream = build_agent_inventory_workbook(agents)
    filename = "NakshaTech Agent Inventory.xlsx"
    return StreamingResponse(
        stream,
        media_type=EXCEL_MIME,
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
            "X-Agent-Count": str(len(agents)),
            "Cache-Control": "no-store",
        },
    )
