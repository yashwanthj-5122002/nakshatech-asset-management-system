from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from hashlib import sha256
from io import BytesIO
import json
from typing import Any, Iterable
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import Date, DateTime, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.entities import (
    Asset,
    AssetHistory,
    ComponentReplacement,
    Drone,
    DroneLocation,
    ReplacementRecord,
    User,
    WorkRecord,
)
from app.modules.drone.models import (
    DroneAssetKit,
    DroneAssetMovement,
    DroneAuditLog,
    DroneHDDDelivery,
    DroneImportBatch,
    DroneImportException,
    DroneKitComponent,
    DroneOperation,
    DroneOperationItem,
    DroneProject,
    DroneSurveyAsset,
    DroneTelecomConnection,
    DroneUINRegistration,
    DroneWorkRecord,
)


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
VALID_BACKUP_ROLES = {"it", "drone", "management", "admin"}
MAX_EXCEL_TEXT = 32000


@dataclass(frozen=True)
class CurrentMonthPeriod:
    start: date
    end: date
    key: str
    label: str


@dataclass(frozen=True)
class ExportSpec:
    sheet_name: str
    model: type
    area: str
    date_field: str | None = None
    master: bool = False
    exclude: tuple[str, ...] = ()
    admin_only: bool = False
    special: str | None = None


EXPORT_SPECS: tuple[ExportSpec, ...] = (
    ExportSpec("IT Asset Register", Asset, "it", master=True),
    ExportSpec("IT Work Records", WorkRecord, "it", date_field="created_at"),
    ExportSpec("IT Component Changes", ComponentReplacement, "it", date_field="created_at"),
    ExportSpec("IT Replacements", ReplacementRecord, "it", date_field="created_at"),
    ExportSpec("IT Asset History", AssetHistory, "it", date_field="created_at"),
    ExportSpec(
        "Drone Asset Register",
        DroneSurveyAsset,
        "drone",
        master=True,
        exclude=("original_raw_payload", "original_header_map"),
    ),
    ExportSpec("Drone Kits", DroneAssetKit, "drone", master=True, exclude=("original_raw_payload",)),
    ExportSpec("Drone Kit Components", DroneKitComponent, "drone", master=True),
    ExportSpec("Drone Projects", DroneProject, "drone", master=True),
    ExportSpec("Drone Operations", DroneOperation, "drone", date_field="operation_date"),
    ExportSpec(
        "Drone Operation Items",
        DroneOperationItem,
        "drone",
        date_field="created_at",
        special="operation_items",
    ),
    ExportSpec("Drone Movements", DroneAssetMovement, "drone", date_field="occurred_at"),
    ExportSpec("Drone Work Records", DroneWorkRecord, "drone", date_field="created_at"),
    ExportSpec(
        "Drone UIN Register",
        DroneUINRegistration,
        "drone",
        master=True,
        exclude=("original_raw_payload",),
    ),
    ExportSpec(
        "Drone HDD Deliveries",
        DroneHDDDelivery,
        "drone",
        date_field="delivery_date",
        exclude=("original_raw_payload",),
    ),
    ExportSpec(
        "Drone Telecom",
        DroneTelecomConnection,
        "drone",
        master=True,
        exclude=("original_raw_payload",),
    ),
    ExportSpec("Drone Audit", DroneAuditLog, "drone", date_field="created_at"),
    ExportSpec("Drone Import Batches", DroneImportBatch, "drone", date_field="created_at"),
    ExportSpec("Drone Import Exceptions", DroneImportException, "drone", date_field="created_at"),
    ExportSpec("Telemetry Drones", Drone, "drone", master=True),
    ExportSpec("Drone Locations", DroneLocation, "drone", date_field="recorded_at"),
    ExportSpec(
        "System Users",
        User,
        "admin",
        master=True,
        exclude=("password_hash",),
        admin_only=True,
    ),
)


def local_now() -> datetime:
    try:
        zone = ZoneInfo(settings.backup_timezone)
    except Exception:
        zone = timezone.utc
    return datetime.now(zone)


def current_month_period(now: datetime | None = None) -> CurrentMonthPeriod:
    selected = (now or local_now()).date().replace(day=1)
    next_month = (selected.replace(day=28) + timedelta(days=4)).replace(day=1)
    end = next_month - timedelta(days=1)
    return CurrentMonthPeriod(
        start=selected,
        end=end,
        key=selected.strftime("%Y-%m"),
        label=selected.strftime("%B %Y"),
    )


def _model_columns(model: type, excluded: Iterable[str] = ()) -> list[str]:
    excluded_set = set(excluded)
    return [column.name for column in model.__table__.columns if column.name not in excluded_set]


def _serialize(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return value
    if isinstance(value, (dict, list, tuple, set)):
        text = json.dumps(value, ensure_ascii=False, default=str)
        return text[:MAX_EXCEL_TEXT]
    if isinstance(value, bytes):
        return f"<binary {len(value)} bytes>"
    text = str(value) if not isinstance(value, (int, float, bool)) else value
    if isinstance(text, str):
        if len(text) > MAX_EXCEL_TEXT:
            text = text[:MAX_EXCEL_TEXT] + "..."
        # Prevent user-entered values from becoming executable spreadsheet formulae.
        if text.startswith(("=", "+", "-", "@")):
            text = "'" + text
    return text


def _filter_query(model: type, field_name: str | None, period: CurrentMonthPeriod):
    query = select(model)
    if not field_name:
        return query
    column = getattr(model, field_name)
    column_type = model.__table__.columns[field_name].type
    if isinstance(column_type, DateTime):
        start_dt = datetime.combine(period.start, time.min)
        end_dt = datetime.combine(period.end + timedelta(days=1), time.min)
        return query.where(column >= start_dt, column < end_dt)
    if isinstance(column_type, Date):
        return query.where(column >= period.start, column <= period.end)
    return query


def _rows_for_spec(db: Session, spec: ExportSpec, period: CurrentMonthPeriod) -> list[Any]:
    if spec.special == "operation_items":
        operation_ids = list(
            db.scalars(
                _filter_query(DroneOperation, "operation_date", period)
                .with_only_columns(DroneOperation.id)
                .order_by(DroneOperation.id)
            ).all()
        )
        if not operation_ids:
            return []
        return list(
            db.scalars(
                select(DroneOperationItem)
                .where(DroneOperationItem.operation_id.in_(operation_ids))
                .order_by(DroneOperationItem.id)
            ).all()
        )

    query = _filter_query(spec.model, None if spec.master else spec.date_field, period)
    primary_key = list(spec.model.__table__.primary_key.columns)[0]
    return list(db.scalars(query.order_by(primary_key)).all())


def _sheet_title(value: str) -> str:
    invalid = set('[]:*?/\\')
    title = "".join("_" if char in invalid else char for char in value)
    return title[:31]


def _friendly_header(name: str) -> str:
    return (
        name.replace("_", " ")
        .strip()
        .title()
        .replace("Uin", "UIN")
        .replace("Hdd", "HDD")
        .replace("Id", "ID")
    )


def _write_table_sheet(workbook: Workbook, sheet_name: str, columns: list[str], rows: list[Any]) -> int:
    sheet = workbook.create_sheet(_sheet_title(sheet_name))
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="082A52")
    header_font = Font(color="FFFFFF", bold=True)
    thin_blue = Side(style="thin", color="D8E4EF")

    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(1, index, _friendly_header(column))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_blue)
    sheet.row_dimensions[1].height = 30

    for row_index, row in enumerate(rows, start=2):
        for column_index, column in enumerate(columns, start=1):
            value = _serialize(getattr(row, column, None))
            cell = sheet.cell(row_index, column_index, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
        if row_index % 2 == 0:
            for column_index in range(1, len(columns) + 1):
                sheet.cell(row_index, column_index).fill = PatternFill("solid", fgColor="F8FBFE")

    if columns:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    for index, column in enumerate(columns, start=1):
        max_length = len(_friendly_header(column))
        for row_index in range(2, min(len(rows) + 2, 202)):
            value = sheet.cell(row_index, index).value
            if value is not None:
                max_length = max(max_length, min(len(str(value)), 60))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 12), 45)
    return len(rows)


def _write_summary_sheet(
    workbook: Workbook,
    role: str,
    period: CurrentMonthPeriod,
    generated_for: str,
    row_counts: dict[str, int],
) -> None:
    sheet = workbook.active
    sheet.title = "Backup Summary"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 31
    sheet.column_dimensions["B"].width = 72

    sheet.merge_cells("A1:B1")
    sheet["A1"] = "NakshaTech Asset Management - Local Monthly Backup"
    sheet["A1"].fill = PatternFill("solid", fgColor="082A52")
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 36

    values = [
        ("Backup Role", role.upper()),
        ("Reporting Month", period.label),
        ("Period Start", period.start),
        ("Period End", period.end),
        ("Generated At", local_now().replace(tzinfo=None)),
        ("Generated For", generated_for),
        ("Timezone", settings.backup_timezone),
        ("Storage Rule", "The local Windows agent replaces the current-month file hourly."),
        ("Month Close Rule", "At the month boundary, the last valid current file is moved to Monthly and never overwritten."),
        ("Data Rule", "Master sheets show current state. Activity and history sheets contain records from this reporting month."),
        ("Security Rule", "Password hashes and secrets are never included in this workbook."),
    ]
    for row_number, (label, value) in enumerate(values, start=3):
        sheet.cell(row_number, 1, label).font = Font(bold=True, color="082A52")
        sheet.cell(row_number, 1).fill = PatternFill("solid", fgColor="EDF4FA")
        sheet.cell(row_number, 2, value)
        sheet.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(value, datetime):
            sheet.cell(row_number, 2).number_format = "yyyy-mm-dd hh:mm:ss"
        elif isinstance(value, date):
            sheet.cell(row_number, 2).number_format = "yyyy-mm-dd"

    start_row = len(values) + 5
    sheet.cell(start_row, 1, "Sheet").font = Font(bold=True, color="FFFFFF")
    sheet.cell(start_row, 2, "Rows Exported").font = Font(bold=True, color="FFFFFF")
    sheet.cell(start_row, 1).fill = PatternFill("solid", fgColor="126AE8")
    sheet.cell(start_row, 2).fill = PatternFill("solid", fgColor="126AE8")
    for offset, (name, count) in enumerate(row_counts.items(), start=1):
        sheet.cell(start_row + offset, 1, name)
        sheet.cell(start_row + offset, 2, count)


def _include_spec(role: str, spec: ExportSpec) -> bool:
    if spec.admin_only:
        return role == "admin"
    if role == "it":
        return spec.area == "it"
    if role == "drone":
        return spec.area == "drone"
    if role in {"management", "admin"}:
        return spec.area in {"it", "drone"}
    return False


def build_current_month_workbook(
    db: Session,
    role: str,
    generated_for: str = "NakshaTech Local Backup Agent",
) -> tuple[bytes, CurrentMonthPeriod, dict[str, int]]:
    role = role.strip().lower()
    if role not in VALID_BACKUP_ROLES:
        raise ValueError("Role must be admin, management, it or drone")

    period = current_month_period()
    workbook = Workbook()
    row_counts: dict[str, int] = {}

    for spec in EXPORT_SPECS:
        if not _include_spec(role, spec):
            continue
        rows = _rows_for_spec(db, spec, period)
        columns = _model_columns(spec.model, spec.exclude)
        row_counts[spec.sheet_name] = _write_table_sheet(workbook, spec.sheet_name, columns, rows)

    _write_summary_sheet(workbook, role, period, generated_for, row_counts)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue(), period, row_counts


def content_sha256(data: bytes) -> str:
    return sha256(data).hexdigest()
