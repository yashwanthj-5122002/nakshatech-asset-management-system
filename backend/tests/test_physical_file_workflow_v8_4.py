"""Focused Phase-1 implementation tests: V8.4 physical-file workflow.

Isolated TEST DATA only. Proves the NEW physical-file behaviour can safely enter
manual UAT: TOTAL FILES expansion + idempotency, backend-exclusive allocation,
Start -> Complete (server time, INITIAL/REWORK attempts), immediate per-file QC
handoff, rework preserving history, ready/delivery/PM-completion gates, distinct
physical-file counters and one-source-of-truth across /ortho + tracking + workbook.

This is NOT the release regression (that runs after human UAT).
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.database import SessionLocal
from app.modules.commercial import models as _commercial_models  # noqa: F401
from app.modules.operations import production_allocation as alloc
from app.modules.operations import production_items as items_service
from app.modules.operations.models import (
    OrthoProjectMember,
    WorkPackageItem,
    WorkPackageItemActivity,
)
from app.modules.operations.production_items import (
    activity_rows_for_packages,
    build_item_rows,
    employee_productivity_rows,
    file_aggregates_for_packages,
)
from app.modules.operations.production_reporting import build_production_workbook
from app.modules.operations.production_service import tracking_summary
from app.modules.operations.workflow_service import ortho_dashboard
from test_production_file_tracking_v8_2 import make_user
from test_production_master_excel_allocation_v8_3 import (
    items_of,
    packages_of,
    run_review,
    setup_project,
    xlsx_bytes,
)

TL = "employee"
PM = "lidar"
COUNT_HEADERS = ["MARCHE LOTTO 3", "CODE", "AREA NAME", "TOTAL FILES", "PLANNED LENGTH", "UOM", "REMARKS"]


def _rows(n, lot="LOS-031", code="LOS1045287", area="FOGNATURA 11", length=6.0):
    return [[lot, code, area, n, length, "km", ""]]


def _import(db, ctx, n, filename="COUNT.xlsx", **kw):
    return alloc.commit_master_import(
        db, actor=ctx["team"]["tl"], role=TL, project_id=ctx["project"].id,
        filename=filename, content=xlsx_bytes(COUNT_HEADERS, _rows(n, **kw)))


def _setup(db, code, n=3, **kw):
    ctx = setup_project(db, code)
    _import(db, ctx, n, **kw)
    pkg = packages_of(db, ctx)[0]
    items = sorted([i for i in items_of(db, ctx) if i.work_package_id == pkg.id],
                   key=lambda i: int(i.item_number))
    return ctx, pkg, items


def _prod2(db, ctx):
    """A second PM-approved Production member for exclusivity tests."""
    user = make_user(db, f"pf2.{ctx['project'].id}@nakshatech.com", "employee", "LiDAR")
    db.flush()
    db.add(OrthoProjectMember(project_id=ctx["project"].id, user_id=user.id,
                              member_role="production", is_active=True,
                              assigned_by_id=ctx["pm"].id))
    db.flush()
    return user


def _assign(db, ctx, pkg, items, *, prod=None, qc=None, qa=None, reassign=False, targets=None):
    t = ctx["team"]
    prod_t, qc_t, qa_t = targets or (date(2026, 10, 20), date(2026, 10, 22), date(2026, 10, 24))
    return alloc.assign_files(
        db, actor=t["tl"], role=TL, work_package_id=pkg.id, item_ids=[i.id for i in items],
        production_user_id=(prod or t["prod"]).id, production_target_date=prod_t,
        qc_user_id=(qc or t["qc"]).id, qc_target_date=qc_t,
        qa_user_id=(qa or t["qa"]).id, qa_target_date=qa_t,
        supply_date=date(2026, 10, 15), reassign=reassign)


def _session(db, ctx, item, *, start, end, finalize=True, qty=None, stage="production",
             actor_key="prod", **kw):
    activity = items_service.start_activity(
        db, actor=ctx["team"][actor_key], role=TL, work_package_item_id=item.id,
        stage=stage, started_at=start)
    return items_service.stop_activity(
        db, actor=ctx["team"][actor_key], activity_id=activity.id, status="completed",
        finalize=finalize, quantity_completed=(Decimal(qty) if qty else None),
        unit=item.planned_unit, ended_at=end, **kw)


def _row(db, pkg, item):
    return next(r for r in build_item_rows(db, [pkg.id])
                if r["work_package_item_id"] == item.id)


def _agg(db, pkg):
    return file_aggregates_for_packages(db, [pkg.id])[pkg.id]


# =============================== A/B/C/D physical pool + import ==========================

def test_a_total_files_generates_exact_physical_pool():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-A", n=30)
        assert [i.item_number for i in items] == [str(i) for i in range(1, 31)]
        assert len(items) == 30
        assert pkg.total_files == 30
        assert float(pkg.area) == 6.0                 # declared length applied ONCE
        assert all(i.planned_quantity is None for i in items)   # no fabricated per-file KM
        assert _agg(db, pkg)["files_expected"] == 30


def test_b_reimport_is_idempotent():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-B", n=30)
        _import(db, ctx, 30)                          # same workbook again
        again = [i for i in items_of(db, ctx) if i.work_package_id == pkg.id]
        assert len(again) == 30
        assert float(pkg.area) == 6.0                 # length NOT re-added


def test_c_increase_adds_only_missing_files():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-C", n=30)
        result = _import(db, ctx, 35)
        now = sorted([i.item_number for i in items_of(db, ctx) if i.work_package_id == pkg.id],
                     key=int)
        assert len(now) == 35
        assert result["created_items"] == 5
        assert now[-5:] == [str(i) for i in range(31, 36)]
        assert pkg.total_files == 35


def test_d_reduction_never_deletes_history():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-D", n=30)
        _assign(db, ctx, pkg, items[:1])
        _session(db, ctx, items[0], start=datetime(2026, 10, 13, 9, 0),
                 end=datetime(2026, 10, 13, 11, 0), qty="1.0")
        _import(db, ctx, 5)                           # unsafe reduction
        still = [i for i in items_of(db, ctx) if i.work_package_id == pkg.id]
        assert len(still) == 30                       # nothing deleted
        assert any(i.item_number == "30" for i in still)


# =============================== E/F/G exclusive allocation ==============================

def test_e_exclusive_allocation_blocks_second_employee():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-E", n=10)
        other = _prod2(db, ctx)
        _assign(db, ctx, pkg, items[:3])              # A owns 1..3
        with pytest.raises(ValueError, match="already assigned"):
            _assign(db, ctx, pkg, items[:3], prod=other)
        _assign(db, ctx, pkg, items[3:5], prod=other)  # B gets only available files


def test_f_direct_duplicate_allocation_rejected():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-F", n=5)
        other = _prod2(db, ctx)
        _assign(db, ctx, pkg, items[:2])
        with pytest.raises(ValueError):
            alloc.assign_files(db, actor=ctx["team"]["tl"], role=TL, work_package_id=pkg.id,
                               item_ids=[items[0].id], production_user_id=other.id,
                               production_target_date=date(2026, 10, 20),
                               qc_user_id=ctx["team"]["qc"].id, qc_target_date=date(2026, 10, 22),
                               qa_user_id=ctx["team"]["qa"].id, qa_target_date=date(2026, 10, 24))


def test_g_concurrent_duplicate_assignment_only_one_wins():
    """Two independent transactions racing for the same file: the second is rejected."""
    with SessionLocal() as db1:
        ctx, pkg, items = _setup(db1, "PF-G", n=5)
        other = _prod2(db1, ctx)
        pkg_id, item_id = pkg.id, items[0].id
        ids = {"tl": ctx["team"]["tl"].id, "prod": ctx["team"]["prod"].id,
               "qc": ctx["team"]["qc"].id, "qa": ctx["team"]["qa"].id, "other": other.id}
        project_id = ctx["project"].id
        # Session 1 wins (locks + commits).
        alloc.assign_files(db1, actor=ctx["team"]["tl"], role=TL, work_package_id=pkg_id,
                           item_ids=[item_id], production_user_id=ids["prod"],
                           production_target_date=date(2026, 10, 20),
                           qc_user_id=ids["qc"], qc_target_date=date(2026, 10, 22),
                           qa_user_id=ids["qa"], qa_target_date=date(2026, 10, 24))
        db1.commit()
        db1.expire_all()
        owner = db1.get(WorkPackageItem, item_id).production_user_id

    with SessionLocal() as db2:
        from app.models.entities import User
        tl = db2.get(User, ids["tl"])
        with pytest.raises(ValueError):
            alloc.assign_files(db2, actor=tl, role=TL, work_package_id=pkg_id,
                               item_ids=[item_id], production_user_id=ids["other"],
                               production_target_date=date(2026, 10, 20),
                               qc_user_id=ids["qc"], qc_target_date=date(2026, 10, 22),
                               qa_user_id=ids["qa"], qa_target_date=date(2026, 10, 24))
    assert owner == ids["prod"]


# =============================== H employee visibility ===================================

def test_h_employee_sees_only_own_files():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-H", n=6)
        other = _prod2(db, ctx)
        _assign(db, ctx, pkg, items[:2])
        _assign(db, ctx, pkg, items[2:4], prod=other)
        payload = items_service.item_list(db, actor=ctx["team"]["prod"], role=TL,
                                          work_package_id=pkg.id)
        assert {i["item_number"] for i in payload["items"]} == {"1", "2"}
        with pytest.raises(PermissionError):
            items_service.start_activity(db, actor=other, role=TL,
                                         work_package_item_id=items[0].id, stage="production")


# =============================== I/J/K production attempt ===============================

def test_i_start_records_server_attempt_and_blocks_duplicate():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-I", n=2)
        _assign(db, ctx, pkg, items)
        activity = items_service.start_activity(db, actor=ctx["team"]["prod"], role=TL,
                                                work_package_item_id=items[0].id,
                                                stage="production")
        assert activity.started_at is not None
        assert activity.attempt_type == "INITIAL" and activity.attempt_number == 1
        with pytest.raises(ValueError, match="open"):
            items_service.start_activity(db, actor=ctx["team"]["prod"], role=TL,
                                         work_package_item_id=items[0].id, stage="production")


def test_j_complete_records_duration_and_counts_one_file():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-J", n=2)
        _assign(db, ctx, pkg, items)
        done = _session(db, ctx, items[0], qty="1.0",
                        start=datetime(2026, 10, 13, 9, 10), end=datetime(2026, 10, 13, 10, 5))
        assert done.duration_seconds == 55 * 60
        row = _row(db, pkg, items[0])
        assert row["production_status"] == "completed"
        assert row["production_files"] == 1
        assert _agg(db, pkg)["files_production_done"] == 1


def test_k_cross_midnight_duration_is_positive():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-K", n=1)
        _assign(db, ctx, pkg, items)
        done = _session(db, ctx, items[0], qty="1.0",
                        start=datetime(2026, 9, 24, 18, 27, 25),
                        end=datetime(2026, 9, 25, 10, 9, 25))
        assert done.duration_seconds > 0
        assert done.duration_seconds == 15 * 3600 + 42 * 60


# =============================== L immediate QC handoff ==================================

def test_l_qc_ready_per_file_before_siblings_complete():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-L", n=3)
        _assign(db, ctx, pkg, items)
        _session(db, ctx, items[0], qty="1.0",
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 10, 0))
        agg = _agg(db, pkg)
        assert agg["files_production_done"] == 1
        assert agg["qc_pending"] == 1          # File 1 immediately QC-ready
        # QC may start File 1 while Files 2-3 are still incomplete
        assert items_service.start_activity(db, actor=ctx["team"]["qc"], role=TL,
                                            work_package_item_id=items[0].id, stage="qc").stage == "qc"
        with pytest.raises(ValueError, match="Production must be completed"):
            items_service.start_activity(db, actor=ctx["team"]["qc"], role=TL,
                                         work_package_item_id=items[1].id, stage="qc")


# =============================== M/N/O/P/S/T rework ======================================

def test_m_n_o_p_s_t_qc_reject_rework_preserves_history_and_one_file():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-M", n=1)
        item = items[0]
        _assign(db, ctx, pkg, items)
        _session(db, ctx, item, qty="1.0",
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 10, 0))
        run_review(db, ctx, item, "qc", "rejected", remarks="Misaligned")   # N
        assert _row(db, pkg, item)["qc_status"] == "rejected"
        assert _agg(db, pkg)["production_rework"] == 1                      # T: rollup rework
        assert _agg(db, pkg)["files_production_done"] == 0                  # not falsely complete
        # O: production rework is a NEW attempt; the INITIAL attempt is preserved
        rework = _session(db, ctx, item, qty="0.2",
                          start=datetime(2026, 10, 14, 9, 0), end=datetime(2026, 10, 14, 10, 0))
        assert rework.attempt_type == "REWORK" and rework.attempt_number == 2
        history = items_service.item_history(db, actor=ctx["team"]["tl"], role=TL,
                                             work_package_item_id=item.id)
        prod = [h for h in history if h["stage"] == "production"]
        assert len(prod) == 2 and prod[0]["attempt_type"] == "INITIAL"
        # S: still exactly ONE physical file
        assert _agg(db, pkg)["files_expected"] == 1
        assert _agg(db, pkg)["files_production_done"] == 1
        # P: QC retry passes
        run_review(db, ctx, item, "qc", "passed", remarks="Fixed")
        assert _row(db, pkg, item)["qc_status"] == "passed"
        assert len([h for h in items_service.item_history(
            db, actor=ctx["team"]["tl"], role=TL, work_package_item_id=item.id)
            if h["stage"] == "qc"]) == 2


# =============================== Q/R QA ==================================================

def test_q_r_qa_pass_and_reject():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-Q", n=2)
        item = items[0]
        _assign(db, ctx, pkg, items)
        _session(db, ctx, item, qty="1.0",
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 10, 0))
        run_review(db, ctx, item, "qc", "passed")
        run_review(db, ctx, item, "qa", "rejected", remarks="Attributes")   # R
        assert _row(db, pkg, item)["qa_status"] == "rejected"
        assert _agg(db, pkg)["qa_rework"] == 1
        run_review(db, ctx, item, "qa", "passed")                            # Q
        assert _row(db, pkg, item)["qa_status"] == "passed"
        row = _row(db, pkg, item)
        assert row["production_files"] == 1                                  # one file
        assert len([h for h in items_service.item_history(
            db, actor=ctx["team"]["tl"], role=TL, work_package_item_id=item.id)
            if h["stage"] == "qa"]) == 2                                     # attempts preserved


# =============================== U/V ready gate ==========================================

def _qa_all(db, ctx, items):
    for item in items:
        _session(db, ctx, item, qty="1.0", start=datetime(2026, 10, 13, 9, 0),
                 end=datetime(2026, 10, 13, 10, 0))
        run_review(db, ctx, item, "qc", "passed")
        run_review(db, ctx, item, "qa", "passed")


def test_u_ready_blocked_until_all_qa_pass():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-U", n=3)
        _assign(db, ctx, pkg, items)
        _qa_all(db, ctx, items[:2])                 # 2 of 3
        with pytest.raises(ValueError, match="Not all files passed QA"):
            items_service.mark_ready_for_delivery(db, actor=ctx["team"]["tl"], role=TL,
                                                  work_package_id=pkg.id)
        assert _agg(db, pkg)["ready_for_delivery"] is False
        _qa_all(db, ctx, items[2:])                 # 3 of 3
        assert _agg(db, pkg)["ready_for_delivery"] is True
        assert _agg(db, pkg)["ready"] == 3          # 3 distinct ready files
        readiness = items_service.mark_ready_for_delivery(
            db, actor=ctx["team"]["tl"], role=TL, work_package_id=pkg.id)   # V
        assert readiness["ready"] is True
        assert pkg.ready_for_delivery_at is not None
        assert pkg.ready_for_delivery_by_id == ctx["team"]["tl"].id


# =============================== W/X delivery ============================================

def test_w_x_delivery_requires_qa_pass_then_is_explicit():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-W", n=1)
        item = items[0]
        _assign(db, ctx, pkg, items)
        _session(db, ctx, item, qty="1.0", start=datetime(2026, 10, 13, 9, 0),
                 end=datetime(2026, 10, 13, 10, 0))
        run_review(db, ctx, item, "qc", "passed")
        with pytest.raises(ValueError, match="QA must pass"):
            items_service.deliver_item(db, actor=ctx["team"]["tl"], role=TL,
                                       work_package_item_id=item.id)         # W
        run_review(db, ctx, item, "qa", "passed")
        delivered = items_service.deliver_item(db, actor=ctx["team"]["tl"], role=TL,
                                               work_package_item_id=item.id, remarks="Sent")
        assert delivered.status == "delivered"                               # X
        assert _agg(db, pkg)["files_delivered"] == 1
        assert ctx["workflow"].finance_closed_at is None                     # Z


# =============================== Y/Z PM completion + finance =============================

def test_y_z_pm_completion_requires_delivery_and_finance_stays_open():
    from app.modules.operations.schemas import WorkflowOperationalCompletion
    from app.modules.operations.workflow_service import operational_complete

    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-Y", n=2)
        _assign(db, ctx, pkg, items)
        _qa_all(db, ctx, items)
        items_service.mark_ready_for_delivery(db, actor=ctx["team"]["tl"], role=TL,
                                              work_package_id=pkg.id)
        db.flush()
        with pytest.raises(ValueError, match="must be delivered"):
            operational_complete(db, actor=ctx["pm"], project_id=ctx["project"].id,
                                 payload=WorkflowOperationalCompletion(
                                     completion_date=date(2026, 10, 28),
                                     final_delivery_reference="REF-PF"))
        for item in items:
            items_service.deliver_item(db, actor=ctx["team"]["tl"], role=TL,
                                       work_package_item_id=item.id, remarks="Sent")
        db.flush()
        db.expire_all()
        workflow = operational_complete(
            db, actor=ctx["pm"], project_id=ctx["project"].id,
            payload=WorkflowOperationalCompletion(completion_date=date(2026, 10, 28),
                                                  final_delivery_reference="REF-PF"))
        assert workflow.operational_completed_at is not None                 # Y
        assert workflow.finance_closed_at is None and workflow.status != "closed"  # Z


# =============================== AA daily activity =======================================

def test_aa_daily_activity_is_file_based_with_attempts():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-AA", n=1)
        item = items[0]
        _assign(db, ctx, pkg, items)
        _session(db, ctx, item, qty="1.0", start=datetime(2026, 10, 13, 9, 10),
                 end=datetime(2026, 10, 13, 10, 5))
        run_review(db, ctx, item, "qc", "rejected", remarks="redo")
        _session(db, ctx, item, qty="0.3", start=datetime(2026, 10, 14, 9, 0),
                 end=datetime(2026, 10, 14, 9, 30))
        acts = activity_rows_for_packages(db, [pkg.id])
        prod = [a for a in acts if a["stage"] == "production"]
        assert [a["attempt_type"] for a in prod] == ["INITIAL", "REWORK"]
        assert prod[0]["item_number"] == item.item_number
        assert prod[0]["started_at"] and prod[0]["ended_at"]


# =============================== AB employee productivity ================================

def test_ab_employee_productivity_counts_distinct_files():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-AB", n=2)
        _assign(db, ctx, pkg, items)
        _session(db, ctx, items[0], qty="1.0", start=datetime(2026, 10, 13, 9, 0),
                 end=datetime(2026, 10, 13, 10, 0))
        run_review(db, ctx, items[0], "qc", "rejected", remarks="redo")
        _session(db, ctx, items[0], qty="0.5", start=datetime(2026, 10, 13, 11, 0),
                 end=datetime(2026, 10, 13, 11, 30))     # rework same file
        rows = employee_productivity_rows(db, [pkg.id])
        prod_rows = [r for r in rows if r["employee_user_id"] == ctx["team"]["prod"].id]
        # rework must NOT create a second physical-file row/count for the same day+area
        assert sum(r["files_worked"] for r in prod_rows) == 1
        assert sum(r["files_completed"] for r in prod_rows) == 1


# =============================== AC reports ==============================================

def test_ac_report_no_of_files_matches_physical_count():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-AC", n=3)
        _assign(db, ctx, pkg, items)
        _qa_all(db, ctx, items)
        content, _ = build_production_workbook(
            db, report_type="project_workbook", actor=ctx["pm"], role=PM,
            project_id=ctx["project"].id)
        wb = load_workbook(BytesIO(content), data_only=True)
        area = next(r for r in wb["AREA MASTER"].iter_rows(min_row=2, values_only=True)
                    if r[1] == pkg.package_code)
        assert area[20] == 3                     # NO OF FILES = 3 distinct physical files
        history = wb["ACTIVITY HISTORY"]
        h = [c.value for c in history[1]]
        prod_rows = [r for r in history.iter_rows(min_row=2, values_only=True)
                     if r[h.index("STAGE")] == "PRODUCTION"]
        assert len(prod_rows) == 3               # one attempt each


# =============================== AD cross-read consistency ===============================

def test_ad_ortho_tracking_workbook_agree():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-AD", n=3)
        _assign(db, ctx, pkg, items)
        _qa_all(db, ctx, items)
        db.flush()
        db.expire_all()
        ortho = next(p for p in ortho_dashboard(db, actor=ctx["pm"], role=PM)["projects"]
                     if p["project_id"] == ctx["project"].id)["packages"][0]
        track = next(r for r in tracking_summary(db, actor=ctx["pm"], role=PM,
                                                project_id=ctx["project"].id)["projects"][0]["package_rows"]
                     if r["work_package_id"] == pkg.id)
        agg = _agg(db, pkg)
        for key in ("files_expected", "files_production_done", "files_qc_passed",
                    "files_qa_passed", "files_delivered"):
            assert ortho[key] == track[key] == agg[key], key
        assert ortho["files_ready"] == track["files_ready"] == agg["ready"] == 3


# =============================== AE/AF authorization =====================================

def test_ae_af_unassigned_and_management_rejected():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-AE", n=2)
        _assign(db, ctx, pkg, items[:1])
        other = _prod2(db, ctx)
        with pytest.raises(PermissionError):
            items_service.start_activity(db, actor=other, role=TL,
                                         work_package_item_id=items[0].id, stage="production")
        mgmt = make_user(db, "pf.mgmt@nakshatech.com", "management")
        db.flush()
        with pytest.raises(PermissionError, match="read-only"):
            items_service.start_activity(db, actor=mgmt, role="management",
                                         work_package_item_id=items[0].id, stage="production")


# =============================== AG legacy multi-session readable ========================

def test_ag_legacy_multi_session_activity_remains_readable():
    with SessionLocal() as db:
        ctx, pkg, items = _setup(db, "PF-AG", n=1)
        item = items[0]
        _assign(db, ctx, pkg, items)
        # legacy partial "save daily progress" still accepted + readable
        _session(db, ctx, item, qty="0.4", finalize=False,
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 10, 0))
        _session(db, ctx, item, qty="0.6", finalize=True,
                 start=datetime(2026, 10, 13, 11, 0), end=datetime(2026, 10, 13, 12, 0))
        acts = activity_rows_for_packages(db, [pkg.id])
        prod = [a for a in acts if a["stage"] == "production"]
        assert len(prod) == 2
        assert prod[0]["finalized"] is False and prod[1]["finalized"] is True
        assert _agg(db, pkg)["files_production_done"] == 1        # still one file
