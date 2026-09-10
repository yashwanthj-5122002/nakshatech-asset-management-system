from __future__ import annotations

from io import BytesIO
import asyncio

import pytest
from openpyxl import load_workbook

from app.modules.agent_monitor import excel_export


@pytest.fixture(autouse=True)
def isolated_application_database():
    """This export unit test does not use the application database."""
    yield


def test_fetch_all_agents_walks_every_page_and_deduplicates(monkeypatch):
    calls: list[int] = []

    async def fake_get_json(path: str, query: dict | None = None):
        assert path == "/api/v1/admin/agents"
        page = int((query or {}).get("page", 1))
        calls.append(page)
        if page == 1:
            return {
                "items": [
                    {"agent_id": "a-1", "cpu_asset_tag": "NT1001", "status": "online"},
                    {"agent_id": "a-2", "cpu_asset_tag": "NT1002", "status": "offline"},
                ],
                "page": 1,
                "page_size": 100,
                "total": 3,
            }
        return {
            "items": [
                {"agent_id": "a-2", "cpu_asset_tag": "NT1002", "status": "offline"},
                {"agent_id": "a-3", "cpu_asset_tag": "NT1003", "status": "delayed"},
            ],
            "page": 2,
            "page_size": 100,
            "total": 3,
        }

    monkeypatch.setattr(excel_export.AgentMonitorService, "get_json", fake_get_json)
    rows = asyncio.run(excel_export.fetch_all_agents())

    assert calls == [1, 2]
    assert [row["agent_id"] for row in rows] == ["a-1", "a-2", "a-3"]


def test_workbook_contains_cpu_asset_tag_and_all_agent_rows():
    stream = excel_export.build_agent_inventory_workbook([
        {
            "agent_id": "a-1",
            "status": "online",
            "workstation_number": "WS-01",
            "cpu_asset_tag": "NT1001",
            "hostname": "DESKTOP-01",
            "current_nt_id": "NT\\user1",
            "department": "Software",
            "operating_system": "Windows 11",
            "total_ram_bytes": 16 * 1024 * 1024 * 1024,
            "last_seen_at": "2026-08-29T11:00:00Z",
        }
    ])

    workbook = load_workbook(BytesIO(stream.getvalue()), read_only=True)
    assert workbook.sheetnames == ["Summary", "Agent Inventory"]
    sheet = workbook["Agent Inventory"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert "CPU Asset Tag" in headers
    assert "Workstation Number" in headers
    assert sheet.max_row == 2
    cpu_col = headers.index("CPU Asset Tag") + 1
    assert sheet.cell(2, cpu_col).value == "NT1001"
