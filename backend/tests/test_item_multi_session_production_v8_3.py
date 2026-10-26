"""Focused tests: file-level Production with MULTIPLE sessions (partial progress vs final
completion) on the shared Daily Activity workflow.

Business rule pinned here: ending a work session ("Save Daily Progress", finalize=False) and
FINISHING a physical file ("Complete Production", finalize=True) are DISTINCT operations.
  - quantity/time SUM across closed sessions (never overwrite a previous session)
  - ONE physical file is counted exactly ONCE (partial sessions contribute files_completed=0)
  - QC only unlocks on FINAL production completion
  - actual may differ from planned (below/equal/above) without blocking completion
  - cross-midnight sessions compute a correct duration from full datetimes
  - ACTIVITY HISTORY / LEGACY TILE ACTIVITY / EMPLOYEE PRODUCTIVITY reflect reality
"""

from __future__ import annotations

from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.database import SessionLocal
from app.modules.commercial import models as _commercial_models  # noqa: F401
from app.modules.operations import production_items as items_service
from app.modules.operations.models import WorkPackageItemActivity
from app.modules.operations.production_reporting import build_production_workbook
from test_production_master_excel_allocation_v8_3 import (
    MASTER_HEADERS,
    MASTER_ROWS,
    assign_all,
    import_master,
    items_of,
    packages_of,
    run_review,
    setup_project,
    xlsx_bytes,
)

TL_ROLE = "employee"


def _prepare(db, code):
    ctx = setup_project(db, code)
    import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
    package = packages_of(db, ctx)[0]
    items = [i for i in items_of(db, ctx) if i.work_package_id == package.id]
    assign_all(db, ctx, package, [i.id for i in items])
    return ctx, package, items


def _session(db, ctx, item, *, quantity, finalize, start, end, stage="production", **kw):
    """Run one production session with explicit server timestamps."""
    activity = items_service.start_activity(
        db, actor=ctx["team"]["prod"], role=TL_ROLE, work_package_item_id=item.id,
        stage=stage, started_at=start)
    return items_service.stop_activity(
        db, actor=ctx["team"]["prod"], activity_id=activity.id, status="completed",
        finalize=finalize, quantity_completed=Decimal(quantity), unit=item.planned_unit,
        ended_at=end, **kw)


def _row(db, package, item):
    return next(r for r in items_service.build_item_rows(db, [package.id])
                if r["work_package_item_id"] == item.id)


# ===========================================================================================
# partial save does NOT finish the file / unlock QC
# ===========================================================================================

def test_partial_save_does_not_complete_file_or_unlock_qc():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS1")
        item = items[0]
        base = datetime(2026, 10, 13, 9, 0, 0)
        activity = _session(db, ctx, item, quantity="0.40", finalize=False,
                            start=base, end=base + timedelta(hours=2))
        assert activity.files_completed == 0
        assert activity.status == "in_progress"
        row = _row(db, package, item)
        assert row["production_quantity"] == 0.4      # cumulative (partial) actual
        assert row["production_sessions"] == 1
        assert row["production_files"] == 0           # file NOT completed
        assert row["production_status"] == "in_progress"
        assert row["production_completed_at"] is None
        assert row["status"] == "in_progress"
        assert row["qc_status"] == "not_started"
        # QC gate is enforced server-side, not just hidden in the UI
        with pytest.raises(ValueError, match="Production must be completed"):
            items_service.start_activity(db, actor=ctx["team"]["qc"], role=TL_ROLE,
                                         work_package_item_id=item.id, stage="qc")


def test_restart_after_partial_save_creates_a_new_session_never_overwrites():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS2")
        item = items[0]
        s1 = _session(db, ctx, item, quantity="0.40", finalize=False,
                      start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 11, 0))
        s2 = _session(db, ctx, item, quantity="0.30", finalize=False,
                      start=datetime(2026, 10, 14, 9, 0), end=datetime(2026, 10, 14, 11, 0))
        assert s1.id != s2.id
        assert s1.quantity_completed == Decimal("0.40")   # first session preserved
        assert s1.ended_at is not None
        history = items_service.item_history(db, actor=ctx["team"]["tl"], role=TL_ROLE,
                                             work_package_item_id=item.id)
        assert len([h for h in history if h["stage"] == "production"]) == 2


def test_multiple_sessions_accumulate_then_final_completion_counts_one_file():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS3")
        item = items[0]
        _session(db, ctx, item, quantity="0.40", finalize=False,
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 12, 0))
        _session(db, ctx, item, quantity="0.30", finalize=False,
                 start=datetime(2026, 10, 14, 9, 30), end=datetime(2026, 10, 14, 13, 30))
        final = _session(db, ctx, item, quantity="0.50", finalize=True,
                         start=datetime(2026, 10, 15, 10, 0), end=datetime(2026, 10, 15, 12, 30))
        assert final.files_completed == 1
        row = _row(db, package, item)
        assert row["production_quantity"] == 1.2      # 0.4 + 0.3 + 0.5
        assert row["production_sessions"] == 3
        assert row["production_files"] == 1           # ONE physical file, not 3
        assert row["production_status"] == "completed"
        assert row["production_completed_at"] is not None
        assert str(row["production_completed_at"]).startswith("2026-10-15T12:30")
        assert row["production_duration_seconds"] == 3 * 3600 + 4 * 3600 + 2 * 3600 + 1800
        assert row["status"] == "production_completed"
        # QC unlocked only now
        activity = items_service.start_activity(db, actor=ctx["team"]["qc"], role=TL_ROLE,
                                                work_package_item_id=item.id, stage="qc")
        assert activity.stage == "qc"


def test_duplicate_finalization_is_blocked():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS4")
        item = items[0]
        _session(db, ctx, item, quantity="1.20", finalize=True,
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 11, 0))
        with pytest.raises(ValueError, match="already completed"):
            items_service.start_activity(db, actor=ctx["team"]["prod"], role=TL_ROLE,
                                         work_package_item_id=item.id, stage="production")
        row = _row(db, package, item)
        assert row["production_files"] == 1
        assert db.scalar(select(func.count(WorkPackageItemActivity.id)).where(
            WorkPackageItemActivity.work_package_item_id == item.id,
            WorkPackageItemActivity.stage == "production")) == 1


# ===========================================================================================
# rework loop: QC reject leaves production re-openable, file still counted once
# ===========================================================================================

def test_qc_reject_allows_production_rework_without_double_counting_file():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS12")
        item = items[0]
        _session(db, ctx, item, quantity="1.20", finalize=True,
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 11, 0))
        run_review(db, ctx, item, "qc", "rejected", remarks="Needs rework")
        # production re-open is allowed because rework is outstanding
        rework = items_service.start_activity(
            db, actor=ctx["team"]["prod"], role=TL_ROLE, work_package_item_id=item.id,
            stage="production", started_at=datetime(2026, 10, 14, 9, 0))
        items_service.stop_activity(
            db, actor=ctx["team"]["prod"], activity_id=rework.id, status="completed",
            finalize=True, quantity_completed=Decimal("0.20"), unit=item.planned_unit,
            ended_at=datetime(2026, 10, 14, 10, 0))
        row = _row(db, package, item)
        assert row["production_quantity"] == 1.4        # 1.2 + 0.2, sessions accumulate
        assert row["production_sessions"] == 2
        assert row["production_files"] == 1             # re-finalization does NOT count twice
        run_review(db, ctx, item, "qc", "passed", remarks="Fixed")
        assert _row(db, package, item)["qc_status"] == "passed"


# ===========================================================================================
# actual vs planned
# ===========================================================================================

def test_actual_below_planned_still_finalizes():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS5")
        item = items[0]  # planned 1.20 km
        _session(db, ctx, item, quantity="0.90", finalize=True,
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 11, 0))
        row = _row(db, package, item)
        assert row["production_quantity"] == 0.9
        assert row["production_files"] == 1
        assert row["planned_quantity"] == 1.2          # planned quantity unchanged


def test_actual_above_planned_is_preserved():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS6")
        item = items[0]  # planned 1.20 km
        _session(db, ctx, item, quantity="1.60", finalize=True,
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 11, 0))
        row = _row(db, package, item)
        assert row["production_quantity"] == 1.6       # actual kept, not clamped
        assert row["planned_quantity"] == 1.2
        assert row["production_files"] == 1


# ===========================================================================================
# cross-midnight
# ===========================================================================================

def test_cross_midnight_partial_session_duration():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS7")
        item = items[0]
        activity = _session(db, ctx, item, quantity="0.40", finalize=False,
                            start=datetime(2026, 9, 24, 18, 7, 2),
                            end=datetime(2026, 9, 25, 10, 19, 18))
        assert activity.duration_seconds == (datetime(2026, 9, 25, 10, 19, 18)
                                             - datetime(2026, 9, 24, 18, 7, 2)).total_seconds()
        row = _row(db, package, item)
        assert row["production_duration_seconds"] == activity.duration_seconds
        assert row["production_files"] == 0            # still only a saved session
        assert row["production_quantity"] == 0.4


# ===========================================================================================
# multi-employee attribution
# ===========================================================================================

def test_multi_employee_activity_attributed_to_exact_employee():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS8")
        item = items[0]
        _session(db, ctx, item, quantity="0.50", finalize=False,
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 11, 0))
        second = items_service.start_activity(
            db, actor=ctx["team"]["tl"], role=TL_ROLE, work_package_item_id=item.id,
            stage="production", started_at=datetime(2026, 10, 14, 9, 0))
        items_service.stop_activity(
            db, actor=ctx["team"]["tl"], activity_id=second.id, status="completed",
            finalize=True, quantity_completed=Decimal("0.70"), unit=item.planned_unit,
            ended_at=datetime(2026, 10, 14, 11, 0))
        history = items_service.item_history(db, actor=ctx["team"]["tl"], role=TL_ROLE,
                                             work_package_item_id=item.id)
        prod = [h for h in history if h["stage"] == "production"]
        assert all(h["employee_user_id"] is not None for h in prod)
        assert {h["employee_user_id"] for h in prod} == {
            ctx["team"]["prod"].id, ctx["team"]["tl"].id}
        row = _row(db, package, item)
        assert row["production_quantity"] == 1.2
        assert row["production_sessions"] == 2
        assert row["production_files"] == 1


# ===========================================================================================
# QC / QA gates + history preserved
# ===========================================================================================

def test_qc_and_qa_gates_and_reject_history():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS9")
        item = items[0]
        with pytest.raises(ValueError, match="Production must be completed"):
            items_service.start_activity(db, actor=ctx["team"]["qc"], role=TL_ROLE,
                                         work_package_item_id=item.id, stage="qc")
        _session(db, ctx, item, quantity="1.20", finalize=True,
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 11, 0))
        with pytest.raises(ValueError, match="QC must pass"):
            items_service.start_activity(db, actor=ctx["team"]["qa"], role=TL_ROLE,
                                         work_package_item_id=item.id, stage="qa")
        run_review(db, ctx, item, "qc", "rejected", remarks="Misaligned")
        run_review(db, ctx, item, "qc", "passed", remarks="Fixed")
        run_review(db, ctx, item, "qa", "rework", remarks="Attributes")
        history = items_service.item_history(db, actor=ctx["team"]["tl"], role=TL_ROLE,
                                             work_package_item_id=item.id)
        assert [h["status"] for h in history if h["stage"] == "qc"] == ["rejected", "passed"]
        assert [h["status"] for h in history if h["stage"] == "qa"] == ["rework"]
        row = _row(db, package, item)
        assert row["qc_attempts"] == 2 and row["qa_attempts"] == 1


# ===========================================================================================
# reports: area rollup, activity history, legacy tile activity, employee productivity
# ===========================================================================================

def _workbook(db, ctx):
    content, _ = build_production_workbook(
        db, report_type="project_workbook", actor=ctx["pm"], role="lidar",
        project_id=ctx["project"].id)
    return load_workbook(BytesIO(content), data_only=True)


def test_workbook_area_rollup_and_activity_history_and_legacy_sheets():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS10")
        item = items[0]
        _session(db, ctx, item, quantity="0.40", finalize=False,
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 12, 0))
        _session(db, ctx, item, quantity="0.30", finalize=False,
                 start=datetime(2026, 10, 14, 9, 30), end=datetime(2026, 10, 14, 13, 30))
        _session(db, ctx, item, quantity="0.50", finalize=True,
                 start=datetime(2026, 10, 15, 10, 0), end=datetime(2026, 10, 15, 12, 30))
        workbook = _workbook(db, ctx)

        assert workbook.sheetnames == [
            "AREA MASTER", "FILE ACTIVITY", "EMPLOYEE PRODUCTIVITY",
            "ACTIVITY HISTORY", "LEGACY TILE ACTIVITY",
        ]

        # AREA MASTER: ACTUAL KM = cumulative across sessions; NO OF FILES still 1
        area = next(r for r in workbook["AREA MASTER"].iter_rows(min_row=2, values_only=True)
                    if r[1] == package.package_code)
        assert area[3] == 1.2          # planned LENGTH
        assert area[6] == 1.2          # ACTUAL KM (0.4 + 0.3 + 0.5)
        assert area[11] == "completed"
        assert area[20] == 1           # NO OF FILES = distinct physical file

        # ACTIVITY HISTORY: three production session rows, never merged
        history = workbook["ACTIVITY HISTORY"]
        headers = [c.value for c in history[1]]
        session_rows = [r for r in history.iter_rows(min_row=2, values_only=True)
                        if r[headers.index("STAGE")] == "PRODUCTION"
                        and r[headers.index("FILE NUMBER")] == item.item_number]
        assert len(session_rows) == 3
        qty_idx = headers.index("QUANTITY THIS SESSION")
        assert [r[qty_idx] for r in session_rows] == [0.4, 0.3, 0.5]
        fin_idx = headers.index("FINALIZED")
        assert [r[fin_idx] for r in session_rows] == ["NO", "NO", "YES"]
        assert session_rows[2][headers.index("TIME TAKEN")] == "02:30:00"

        # LEGACY TILE ACTIVITY: production rows with Length, accurate data only
        legacy = workbook["LEGACY TILE ACTIVITY"]
        legacy_rows = [r for r in legacy.iter_rows(min_row=2, values_only=True)
                       if r[0] == "PRODUCTION"]
        assert len(legacy_rows) == 3
        assert sum(r[8] for r in legacy_rows) == pytest.approx(1.2)   # LENGTH column

        # EMPLOYEE PRODUCTIVITY: day-grouped rows for the same employee; total actual = 1.2,
        # and only the FINAL session's row contributes the completed file.
        prod = workbook["EMPLOYEE PRODUCTIVITY"]
        pheaders = [c.value for c in prod[1]]
        prows = [r for r in prod.iter_rows(min_row=2, values_only=True)
                 if r[0] and "prod.ft.MS10" in str(r[0])]
        assert len(prows) == 3
        assert sum(r[pheaders.index("ACTUAL QUANTITY")] for r in prows) == pytest.approx(1.2)
        final_row = next(r for r in prows if r[pheaders.index("DATE")] == "2026-10-15")
        assert final_row[pheaders.index("FILES WORKED")] == 1
        assert final_row[pheaders.index("COMPLETED FILES")] == 1
        assert sum(r[pheaders.index("COMPLETED FILES")] for r in prows) == 1   # file counted once


def test_ready_for_delivery_and_delivery_after_full_chain():
    with SessionLocal() as db:
        ctx, package, items = _prepare(db, "MS11")
        item = items[0]
        with pytest.raises(ValueError, match="Not all files passed QA"):
            items_service.mark_ready_for_delivery(
                db, actor=ctx["team"]["tl"], role=TL_ROLE, work_package_id=package.id)
        _session(db, ctx, item, quantity="1.20", finalize=True,
                 start=datetime(2026, 10, 13, 9, 0), end=datetime(2026, 10, 13, 11, 0))
        run_review(db, ctx, item, "qc", "passed")
        run_review(db, ctx, item, "qa", "passed")
        readiness = items_service.mark_ready_for_delivery(
            db, actor=ctx["team"]["tl"], role=TL_ROLE, work_package_id=package.id)
        assert readiness["ready"] is True
        delivered = items_service.deliver_item(
            db, actor=ctx["team"]["tl"], role=TL_ROLE, work_package_item_id=item.id,
            remarks="Sent", reference="DLV-MS11")
        assert delivered.status == "delivered"
        # QA pass / delivery must NOT close Finance
        assert ctx["workflow"].finance_closed_at is None
        assert ctx["workflow"].status != "closed"
