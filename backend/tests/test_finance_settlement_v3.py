from __future__ import annotations

from datetime import date
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models.entities import User
from app.modules.finance.models import ExpenseClaimAttachment, FinanceProject
from app.modules.finance.reporting import build_finance_report_workbook
from app.modules.finance import documents
from app.modules.finance.documents import build_claim_a4_pdf
from app.modules.finance.schemas import ExpenseClaimCreateRequest, SettlementUpsertRequest
from app.modules.finance.service import admin_decision, create_claim, ensure_finance_seed_data, finance_decision, mark_paid, submit_claim
from app.modules.finance.settlement_service import admin_settlement_decision, finance_settlement_decision, submit_settlement, upsert_settlement


def user(db, email: str, role: str) -> User:
    row = User(email=email, full_name=email.split('@')[0], password_hash='x', role=role, branch='Head Office', email_verified=True, account_status='active', is_active=True)
    db.add(row); db.flush(); return row


def advance_request(project_id: int, amount: float = 10000) -> ExpenseClaimCreateRequest:
    return ExpenseClaimCreateRequest(project_id=project_id, claim_type='advance', purpose_description='Field work advance for production Finance V3 settlement test.', requested_work_start_date=date(2026,8,28), requested_work_end_date=date(2026,8,30), items=[{'category':'travel','description':'Planned field travel','amount':amount}])


def release_advance(db, employee, admin, finance, project, amount=10000):
    claim=create_claim(db, requester=employee, payload=advance_request(project.id, amount))
    claim=submit_claim(db, claim=claim, requester=employee)
    claim=admin_decision(db, claim=claim, actor=admin, action='approve', comments='Verified', approved_work_start_date=date(2026,8,28), approved_work_end_date=date(2026,8,30), settlement_due_date=date(2026,9,2))
    claim=finance_decision(db, claim=claim, actor=finance, action='approve', comments='Approved', approved_value=amount, approved_work_start_date=date(2026,8,28), approved_work_end_date=date(2026,8,29), settlement_due_date=date(2026,9,1))
    claim=mark_paid(db, claim=claim, actor=finance, payment_reference=f'UTR-{claim.id}', paid_amount=amount, payment_mode='bank_transfer', payment_date=date(2026,8,28), comments='Released')
    return claim


def test_advance_settlement_tallies_and_finance_finalizes(monkeypatch):
    with SessionLocal() as db:
        ensure_finance_seed_data(db)
        project=db.scalar(select(FinanceProject).order_by(FinanceProject.id)); assert project
        employee=user(db,'v3employee@nakshatech.com','employee'); admin=user(db,'v3admin@nakshatech.com','admin'); finance=user(db,'v3finance@nakshatech.com','finance'); db.commit()
        claim=release_advance(db, employee, admin, finance, project, 10000)
        settlement=upsert_settlement(db, root_claim=claim, requester=employee, payload=SettlementUpsertRequest(items=[{'category':'hotel','description':'Site hotel','amount':6000,'payment_mode':'upi','expense_date':date(2026,8,28)},{'category':'travel','description':'Site travel','amount':4000,'payment_mode':'cash','expense_date':date(2026,8,29)}]))
        db.add(ExpenseClaimAttachment(claim_id=claim.id, uploaded_by_id=employee.id, original_filename='advance-proof.pdf', storage_key=f'test/{claim.id}', mime_type='application/pdf', file_size=10))
        # Settlement submission requires settlement attachment specifically; append an in-memory metadata row without reading storage in this test.
        from app.modules.finance.models import ExpenseSettlementAttachment
        db.add(ExpenseSettlementAttachment(settlement_id=settlement.id, uploaded_by_id=employee.id, original_filename='bill.pdf', storage_key=f'test-set/{settlement.id}', mime_type='application/pdf', file_size=10, content_sha256='0'*64)); db.commit(); db.refresh(settlement)
        settlement=submit_settlement(db, settlement=settlement, requester=employee)
        assert settlement.status=='submitted'; assert float(settlement.total_expense_amount)==10000; assert float(settlement.balance_to_return)==0; assert float(settlement.shortage_amount)==0
        settlement=admin_settlement_decision(db, settlement=settlement, actor=admin, action='approve', comments='Bills verified')
        settlement=finance_settlement_decision(db, settlement=settlement, actor=finance, action='approve', comments='Tally verified')
        assert settlement.status=='finance_finalized'; assert settlement.root_claim.settlement_status=='settled'
        report=build_claim_a4_pdf(db, claim)
        assert report.getvalue().startswith(b'%PDF')
        # Verify the complete A4 pack and all-bills ZIP using deterministic in-memory proof bytes.
        monkeypatch.setattr(documents, 'read_finance_attachment_bytes', lambda _key: report.getvalue())
        pack=documents.build_complete_a4_pack(db, claim)
        assert len(PdfReader(pack).pages) >= 3
        bills=documents.build_all_bills_zip(claim)
        with ZipFile(bills) as archive:
            assert len(archive.namelist()) == 2


def test_additional_advance_is_linked_and_included_in_final_advance_total():
    with SessionLocal() as db:
        ensure_finance_seed_data(db)
        project=db.scalar(select(FinanceProject).order_by(FinanceProject.id)); assert project
        employee=user(db,'v3add@nakshatech.com','employee'); admin=user(db,'v3addadmin@nakshatech.com','admin'); finance=user(db,'v3addfin@nakshatech.com','finance'); db.commit()
        root=release_advance(db, employee, admin, finance, project, 10000)
        payload=ExpenseClaimCreateRequest(project_id=project.id, claim_type='additional_advance', purpose_description='Additional field amount required after the original advance was insufficient.', requested_work_start_date=date(2026,8,29), requested_work_end_date=date(2026,8,30), parent_advance_claim_id=root.id, previous_advance_amount=10000, amount_already_used=9500, items=[{'category':'fuel','description':'Additional site fuel','amount':2000}])
        child=create_claim(db, requester=employee, payload=payload)
        db.add(ExpenseClaimAttachment(claim_id=child.id, uploaded_by_id=employee.id, original_filename='additional-proof.pdf', storage_key=f'additional/{child.id}', mime_type='application/pdf', file_size=10, content_sha256='1'*64)); db.commit(); db.refresh(child)
        child=submit_claim(db, claim=child, requester=employee); child=admin_decision(db, claim=child, actor=admin, action='approve', comments='Verified'); child=finance_decision(db, claim=child, actor=finance, action='approve', comments='Approved', approved_value=2000, settlement_due_date=date(2026,9,2)); child=mark_paid(db, claim=child, actor=finance, payment_reference=f'ADD-{child.id}', paid_amount=2000, payment_date=date(2026,8,29), comments='Additional released')
        db.refresh(root)
        settlement=upsert_settlement(db, root_claim=root, requester=employee, payload=SettlementUpsertRequest(items=[{'category':'travel','description':'Final supported project cost','amount':12000,'payment_mode':'upi','expense_date':date(2026,8,30)}]))
        assert float(settlement.total_advance_received)==12000
        assert child.parent_advance_claim_id==root.id


def test_v3_excel_contains_settlement_ledgers():
    with SessionLocal() as db:
        ensure_finance_seed_data(db)
        project=db.scalar(select(FinanceProject).order_by(FinanceProject.id)); assert project
        employee=user(db,'v3excel@nakshatech.com','employee'); admin=user(db,'v3exceladmin@nakshatech.com','admin'); finance=user(db,'v3excelfin@nakshatech.com','finance'); db.commit()
        claim=release_advance(db, employee, admin, finance, project, 1000)
        upsert_settlement(db, root_claim=claim, requester=employee, payload=SettlementUpsertRequest(items=[{'category':'travel','description':'Actual','amount':900,'payment_mode':'cash','expense_date':date(2026,8,29)}]))
        output, _, _=build_finance_report_workbook(db, period='year', year=2026)
        wb=load_workbook(BytesIO(output.getvalue()), data_only=False)
        for name in ['Settlement Ledger','Settlement Actual Items','Settlement Bill Register','Settlement History']:
            assert name in wb.sheetnames
