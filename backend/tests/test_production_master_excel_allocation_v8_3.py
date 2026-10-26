"""V8.3 Team Lead Excel-driven daily activity allocation — focused tests.

Covers the master-workbook import (exact sample, required/optional columns, validation,
duplicate rows, idempotency, transactional rollback, history never overwritten), the
cascading Lot -> Code -> Area -> Files allocation with automatic planned-quantity rollup,
the three stage targets + Supply Date, assignment notifications and permissions, the
Production -> QC -> QA gated file workflow with automatic completed-file counting and
actual quantities, Ready-for-Delivery calculation/confirmation, and the live workbook
exports (AREA MASTER exact 26 columns, FILE ACTIVITY targets/timings, EMPLOYEE PRODUCTIVITY).

Reuses the V8.2 fixtures (BD -> Finance -> PM -> team) so the whole flow runs on the real
existing workflow engine.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from sqlalchemy import func, select

from app.core.database import SessionLocal
from app.modules.commercial import models as _commercial_models  # noqa: F401
from app.modules.notifications.models import GlobalNotification
from app.modules.operations import production_allocation as alloc
from app.modules.operations import production_items as items_service
from app.modules.operations.models import OrthoWorkPackage, WorkPackageItem, WorkPackageItemActivity
from app.modules.operations.production_reporting import (
    AREA_MASTER_HEADERS,
    EMPLOYEE_PRODUCTIVITY_HEADERS,
    FILE_ACTIVITY_HEADERS,
    build_production_workbook,
)
from test_production_file_tracking_v8_2 import approved_project, full_team, make_user

TL_ROLE = "employee"

MASTER_HEADERS = ["MARCHE LOTTO 3", "CODE", "AREA NAME", "FILE NUMBER", "PLANNED LENGTH", "UOM", "REMARKS"]
MASTER_ROWS = [
    ["LOT-031", "DW1045287", "FOGNATURA 1", 1, 1.20, "km", ""],
    ["LOT-032", "DW1045288", "FOGNATURA 2", 2, 0.85, "km", ""],
    ["LOT-033", "DW1045289", "FOGNATURA 3", 3, 1.10, "km", ""],
]


def xlsx_bytes(headers: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def csv_bytes(headers: list[str], rows: list[list]) -> bytes:
    lines = [",".join(str(c) for c in headers)]
    lines += [",".join("" if c is None else str(c) for c in row) for row in rows]
    return ("\n".join(lines) + "\n").encode("utf-8")


def setup_project(db, code: str):
    bd, finance, pm, project, workflow = approved_project(db, code=code)
    team = full_team(db, pm, project, code)
    return {"bd": bd, "finance": finance, "pm": pm, "project": project, "workflow": workflow, "team": team}


def import_master(db, ctx, content: bytes, filename: str = "TESTFILESDAILYACTIVITY.xlsx"):
    return alloc.commit_master_import(
        db, actor=ctx["team"]["tl"], role=TL_ROLE, project_id=ctx["project"].id,
        filename=filename, content=content,
    )


def packages_of(db, ctx) -> list[OrthoWorkPackage]:
    return list(db.scalars(select(OrthoWorkPackage).where(
        OrthoWorkPackage.project_id == ctx["project"].id).order_by(OrthoWorkPackage.id)).all())


def items_of(db, ctx) -> list[WorkPackageItem]:
    return list(db.scalars(select(WorkPackageItem).where(
        WorkPackageItem.project_id == ctx["project"].id).order_by(WorkPackageItem.id)).all())


def assign_all(db, ctx, package, item_ids, *, prod=None, qc=None, qa=None,
               prod_t=date(2026, 10, 20), qc_t=date(2026, 10, 22), qa_t=date(2026, 10, 24),
               supply=date(2026, 10, 15)):
    team = ctx["team"]
    return alloc.assign_files(
        db, actor=team["tl"], role=TL_ROLE, work_package_id=package.id, item_ids=item_ids,
        production_user_id=(prod or team["prod"]).id, production_target_date=prod_t,
        qc_user_id=(qc or team["qc"]).id, qc_target_date=qc_t,
        qa_user_id=(qa or team["qa"]).id, qa_target_date=qa_t, supply_date=supply,
        note="Imported master allocation",
    )


def run_production(db, ctx, item, *, quantity="1.20"):
    activity = items_service.start_activity(
        db, actor=ctx["team"]["prod"], role=TL_ROLE, work_package_item_id=item.id, stage="production")
    return items_service.stop_activity(
        db, actor=ctx["team"]["prod"], activity_id=activity.id, status="completed",
        quantity_completed=Decimal(quantity), unit=item.planned_unit, remarks="done")


def run_review(db, ctx, item, stage: str, status: str, *, remarks=None):
    actor = ctx["team"][stage]
    activity = items_service.start_activity(
        db, actor=actor, role=TL_ROLE, work_package_item_id=item.id, stage=stage)
    return items_service.stop_activity(
        db, actor=actor, activity_id=activity.id, status=status, remarks=remarks)


# ===========================================================================================
# master workbook: parse / preview / import
# ===========================================================================================

def test_exact_sample_xlsx_preview():
    with SessionLocal() as db:
        ctx = setup_project(db, "MX1")
        preview = alloc.preview_master_import(
            db, actor=ctx["team"]["tl"], role=TL_ROLE, project_id=ctx["project"].id,
            filename="TESTFILESDAILYACTIVITY.xlsx",
            content=xlsx_bytes(MASTER_HEADERS, MASTER_ROWS),
        )
        assert preview["rows"] == 3 and preview["valid"] == 3 and preview["errors"] == 0
        assert preview["lots"] == 3 and preview["codes"] == 3 and preview["areas"] == 3
        assert preview["files"] == 3
        assert preview["planned_by_unit"] == [
            {"measurement_type": "distance", "unit": "km", "quantity": 3.15, "rows": 3}
        ]
        # preview writes nothing
        assert packages_of(db, ctx) == [] and items_of(db, ctx) == []


def test_exact_sample_commit_creates_areas_and_files():
    with SessionLocal() as db:
        ctx = setup_project(db, "MX2")
        summary = import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        assert summary["created_packages"] == 3 and summary["created_items"] == 3
        assert summary["conflicts"] == []
        packages = packages_of(db, ctx)
        assert [(p.lot_batch, p.package_code, p.package_name) for p in packages] == [
            ("LOT-031", "DW1045287", "FOGNATURA 1"),
            ("LOT-032", "DW1045288", "FOGNATURA 2"),
            ("LOT-033", "DW1045289", "FOGNATURA 3"),
        ]
        items = items_of(db, ctx)
        assert [i.item_number for i in items] == ["1", "2", "3"]
        assert [float(i.planned_quantity) for i in items] == [1.2, 0.85, 1.1]
        assert all(i.planned_unit == "km" for i in items)
        assert all(i.import_source == "TESTFILESDAILYACTIVITY.xlsx" for i in items)
        assert [float(p.area) for p in packages] == [1.2, 0.85, 1.1]


def test_sample_csv_same_structure():
    with SessionLocal() as db:
        ctx = setup_project(db, "MX3")
        summary = import_master(db, ctx, csv_bytes(MASTER_HEADERS, MASTER_ROWS), "master.csv")
        assert summary["created_items"] == 3


def test_missing_required_header_rejected():
    with SessionLocal() as db:
        ctx = setup_project(db, "MX4")
        headers = ["MARCHE LOTTO 3", "CODE", "AREA NAME", "FILE NUMBER", "PLANNED LENGTH"]
        with pytest.raises(ValueError, match="UOM"):
            import_master(db, ctx, xlsx_bytes(headers, [[r[i] for i in range(5)] for r in MASTER_ROWS]))
        assert items_of(db, ctx) == []


def test_optional_remarks_and_file_name_accepted():
    with SessionLocal() as db:
        ctx = setup_project(db, "MX5")
        headers = MASTER_HEADERS + ["FILE NAME"]
        rows = [
            ["LOT-031", "DW1045287", "FOGNATURA 1", 1, 1.20, "km", "Check boundaries", "FOGNATURA1.dxf"],
        ]
        import_master(db, ctx, xlsx_bytes(headers, rows))
        item = items_of(db, ctx)[0]
        assert item.item_name == "FOGNATURA1.dxf"
        assert "Check boundaries" in (packages_of(db, ctx)[0].remarks or "")


def test_field_validation_errors():
    with SessionLocal() as db:
        ctx = setup_project(db, "MX6")
        rows = [
            ["", "DW1", "AREA A", 1, 1.0, "km", ""],          # blank lot
            ["LOT-X", "", "AREA A", 2, 1.0, "km", ""],        # blank code
            ["LOT-X", "DW2", "", 3, 1.0, "km", ""],           # blank area
            ["LOT-X", "DW3", "AREA A", "", 1.0, "km", ""],    # blank file number
            ["LOT-X", "DW4", "AREA A", 4, "abc", "km", ""],   # invalid length
            ["LOT-X", "DW5", "AREA A", 5, -2, "km", ""],      # negative length
            ["LOT-X", "DW6", "AREA A", 6, 1.0, "", ""],       # blank uom
        ]
        with pytest.raises(ValueError, match="row 2"):
            import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, rows))
        assert items_of(db, ctx) == []


def test_duplicate_rows_rejected():
    with SessionLocal() as db:
        ctx = setup_project(db, "MX7")
        rows = [MASTER_ROWS[0], MASTER_ROWS[0]]
        with pytest.raises(ValueError, match="duplicate row"):
            import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, rows))


def test_empty_and_malformed_workbook():
    with SessionLocal() as db:
        ctx = setup_project(db, "MX8")
        with pytest.raises(ValueError, match="empty"):
            import_master(db, ctx, b"")
        with pytest.raises(ValueError):
            import_master(db, ctx, b"this is not a workbook", "broken.xlsx")


def test_repeated_import_is_idempotent():
    with SessionLocal() as db:
        ctx = setup_project(db, "MX9")
        content = xlsx_bytes(MASTER_HEADERS, MASTER_ROWS)
        first = import_master(db, ctx, content)
        second = import_master(db, ctx, content)
        assert first["created_items"] == 3
        assert second["created_items"] == 0 and second["created_packages"] == 0
        assert second["skipped_existing"] == 3
        assert len(items_of(db, ctx)) == 3 and len(packages_of(db, ctx)) == 3


def test_reimport_preserves_work_history():
    with SessionLocal() as db:
        ctx = setup_project(db, "MX10")
        content = xlsx_bytes(MASTER_HEADERS, MASTER_ROWS)
        import_master(db, ctx, content)
        package = packages_of(db, ctx)[0]
        item = items_of(db, ctx)[0]
        assign_all(db, ctx, package, [item.id])
        run_production(db, ctx, item, quantity="1.16")
        before_planned = float(item.planned_quantity)
        summary = import_master(db, ctx, content)
        assert len(summary["conflicts"]) == 1  # the file with history is reported, not overwritten
        assert float(item.planned_quantity) == before_planned
        db.refresh(item)
        assert item.status == "production_completed"
        assert len(items_of(db, ctx)) == 3
        assert db.scalar(select(func.count(WorkPackageItemActivity.id)).where(
            WorkPackageItemActivity.work_package_item_id == item.id)) >= 1


def test_validation_failure_is_transactional():
    with SessionLocal() as db:
        ctx = setup_project(db, "MX11")
        bad = xlsx_bytes(MASTER_HEADERS, [["LOT-A", "DWA", "AREA A", 1, 1.0, "km", ""],
                                          ["LOT-A", "DWA", "AREA A", "", 1.0, "km", ""]])
        with pytest.raises(ValueError):
            import_master(db, ctx, bad)
        assert packages_of(db, ctx) == [] and items_of(db, ctx) == []


# ===========================================================================================
# allocation: cascading options, auto quantity, targets, notifications
# ===========================================================================================

def test_allocation_options_cascading():
    with SessionLocal() as db:
        ctx = setup_project(db, "AL1")
        import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        options = alloc.allocation_options(db, actor=ctx["team"]["tl"], role=TL_ROLE,
                                           project_id=ctx["project"].id)
        assert [l["lot"] for l in options["lots"]] == ["LOT-031", "LOT-032", "LOT-033"]
        lot31 = options["lots"][0]
        assert [a["code"] for a in lot31["areas"]] == ["DW1045287"]
        assert [a["area_name"] for a in lot31["areas"]] == ["FOGNATURA 1"]
        assert {list(u.keys())[0] for u in [options["team"]["production"][0]]} == {"user_id"}
        assert options["team"]["production"] and options["team"]["qc"] and options["team"]["qa"]
        # every area appears once (no duplicates from repeated imports)
        assert len(options["areas"]) == 3


def test_area_files_filtering_and_planned_groups():
    with SessionLocal() as db:
        ctx = setup_project(db, "AL2")
        import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        packages = packages_of(db, ctx)
        data = alloc.area_files(db, actor=ctx["team"]["tl"], role=TL_ROLE,
                                work_package_id=packages[0].id)
        assert [f["item_number"] for f in data["files"]] == ["1"]
        assert data["planned_groups"] == [
            {"measurement_type": "distance", "unit": "km", "quantity": 1.2, "rows": 1}
        ]
        assert data["lot"] == "LOT-031" and data["area_name"] == "FOGNATURA 1"


def test_assign_sets_targets_supply_and_assignees():
    with SessionLocal() as db:
        ctx = setup_project(db, "AL3")
        import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        team = ctx["team"]
        packages = packages_of(db, ctx)
        package = packages[0]
        item = items_of(db, ctx)[0]
        summary = alloc.assign_files(
            db, actor=team["tl"], role=TL_ROLE, work_package_id=package.id, item_ids=[item.id],
            production_user_id=team["prod"].id, production_target_date=date(2026, 10, 20),
            qc_user_id=team["qc"].id, qc_target_date=date(2026, 10, 22),
            qa_user_id=team["qa"].id, qa_target_date=date(2026, 10, 24),
            supply_date=date(2026, 10, 15), note="From master",
        )
        assert summary["planned_by_unit"] == [
            {"measurement_type": "distance", "unit": "km", "quantity": 1.2, "rows": 1}
        ]
        assert summary["supply_date"] == "2026-10-15"
        assert item.production_user_id == team["prod"].id
        assert item.qc_user_id == team["qc"].id and item.qa_user_id == team["qa"].id
        assert item.production_target_date == date(2026, 10, 20)
        assert item.qc_target_date == date(2026, 10, 22) and item.qa_target_date == date(2026, 10, 24)
        assert package.supply_date == date(2026, 10, 15)
        assert package.production_target_date == date(2026, 10, 20)
        assert package.qc_target_date == date(2026, 10, 22) and package.qa_target_date == date(2026, 10, 24)
        assert package.team_leader_user_id == team["tl"].id
        # area payload exposes the targets + supply date
        payload = items_service.item_list(db, actor=team["tl"], role=TL_ROLE, work_package_id=package.id)
        assert payload["supply_date"] == "2026-10-15"
        assert payload["stage_targets"]["qa"] == "2026-10-24"


def test_assign_invalid_target_order_rejected():
    with SessionLocal() as db:
        ctx = setup_project(db, "AL4")
        import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        package = packages_of(db, ctx)[0]
        item = items_of(db, ctx)[0]
        team = ctx["team"]
        with pytest.raises(ValueError, match="Target dates must satisfy"):
            alloc.assign_files(
                db, actor=team["tl"], role=TL_ROLE, work_package_id=package.id, item_ids=[item.id],
                production_user_id=team["prod"].id, production_target_date=date(2026, 10, 25),
                qc_user_id=team["qc"].id, qc_target_date=date(2026, 10, 22),
                qa_user_id=team["qa"].id, qa_target_date=date(2026, 10, 24),
                supply_date=date(2026, 10, 15), note=None,
            )


def test_assign_requires_team_members():
    with SessionLocal() as db:
        ctx = setup_project(db, "AL5")
        import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        package = packages_of(db, ctx)[0]
        item = items_of(db, ctx)[0]
        outsider = make_user(db, "alloc.outsider@nakshatech.com", "employee", "LiDAR")
        with pytest.raises(ValueError, match="team member"):
            assign_all(db, ctx, package, [item.id], prod=outsider)


def test_assignment_notifications():
    with SessionLocal() as db:
        ctx = setup_project(db, "AL6")
        import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        package = packages_of(db, ctx)[0]
        item = items_of(db, ctx)[0]
        assign_all(db, ctx, package, [item.id])
        for user_key, event in (
            ("prod", "workflow.allocation.production_assigned"),
            ("qc", "workflow.allocation.qc_assigned"),
            ("qa", "workflow.allocation.qa_assigned"),
        ):
            count = db.scalar(select(func.count(GlobalNotification.id)).where(
                GlobalNotification.event_type == event,
                GlobalNotification.recipient_user_id == ctx["team"][user_key].id,
            ))
            assert count == 1, f"{event} should notify exactly once"


# ===========================================================================================
# file workflow: production -> QC -> QA with automatic counting and gates
# ===========================================================================================

def test_production_complete_auto_files_and_actual_km():
    with SessionLocal() as db:
        ctx = setup_project(db, "WF1")
        import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        package = packages_of(db, ctx)[0]
        item = items_of(db, ctx)[0]
        assign_all(db, ctx, package, [item.id])
        activity = run_production(db, ctx, item, quantity="1.16")
        assert activity.files_completed == 1  # automatic
        assert item.status == "production_completed"
        payload = items_service.item_list(db, actor=ctx["team"]["tl"], role=TL_ROLE,
                                          work_package_id=package.id)
        row = payload["items"][0]
        assert row["production_quantity"] == 1.16
        assert row["production_files"] == 1
        assert row["production_deadline"]["status"] == "completed_on_time"  # 20-Oct target
        totals = items_service.employee_file_totals(db, [package.id])
        assert totals[0]["production_completed"] == 1 and totals[0]["production_pending"] == 0


def test_qc_gate_and_ready_after_production():
    with SessionLocal() as db:
        ctx = setup_project(db, "WF2")
        import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        package = packages_of(db, ctx)[0]
        item = items_of(db, ctx)[0]
        assign_all(db, ctx, package, [item.id])
        with pytest.raises(ValueError, match="Production must be completed"):
            items_service.start_activity(db, actor=ctx["team"]["qc"], role=TL_ROLE,
                                         work_package_item_id=item.id, stage="qc")
        run_production(db, ctx, item)
        qc_ready = db.scalar(select(func.count(GlobalNotification.id)).where(
            GlobalNotification.event_type == "workflow.item.qc_ready",
            GlobalNotification.recipient_user_id == ctx["team"]["qc"].id))
        assert qc_ready == 1
        activity = items_service.start_activity(db, actor=ctx["team"]["qc"], role=TL_ROLE,
                                                work_package_item_id=item.id, stage="qc")
        assert activity.stage == "qc"


def test_qa_gate_after_qc_pass():
    with SessionLocal() as db:
        ctx = setup_project(db, "WF3")
        import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        package = packages_of(db, ctx)[0]
        item = items_of(db, ctx)[0]
        assign_all(db, ctx, package, [item.id])
        run_production(db, ctx, item)
        run_review(db, ctx, item, "qc", "rejected", remarks="Not aligned")
        with pytest.raises(ValueError, match="QC must pass"):
            items_service.start_activity(db, actor=ctx["team"]["qa"], role=TL_ROLE,
                                         work_package_item_id=item.id, stage="qa")
        run_review(db, ctx, item, "qc", "passed")
        qa_ready = db.scalar(select(func.count(GlobalNotification.id)).where(
            GlobalNotification.event_type == "workflow.item.qa_ready",
            GlobalNotification.recipient_user_id == ctx["team"]["qa"].id))
        assert qa_ready == 1
        activity = items_service.start_activity(db, actor=ctx["team"]["qa"], role=TL_ROLE,
                                                work_package_item_id=item.id, stage="qa")
        assert activity.stage == "qa"


def test_qc_reject_history_and_notification():
    with SessionLocal() as db:
        ctx = setup_project(db, "WF4")
        import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        package = packages_of(db, ctx)[0]
        item = items_of(db, ctx)[0]
        assign_all(db, ctx, package, [item.id])
        run_production(db, ctx, item)
        run_review(db, ctx, item, "qc", "rejected", remarks="Rework needed")
        history = items_service.item_history(db, actor=ctx["team"]["tl"], role=TL_ROLE,
                                            work_package_item_id=item.id)
        assert [h["status"] for h in history if h["stage"] == "qc"] == ["rejected"]
        assert item.status == "qc_rejected"
        rework_notice = db.scalar(select(func.count(GlobalNotification.id)).where(
            GlobalNotification.event_type == "workflow.item.qc_rework",
            GlobalNotification.recipient_user_id == ctx["team"]["tl"].id))
        assert rework_notice == 1


def test_qa_reject_history():
    with SessionLocal() as db:
        ctx = setup_project(db, "WF5")
        import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        package = packages_of(db, ctx)[0]
        item = items_of(db, ctx)[0]
        assign_all(db, ctx, package, [item.id])
        run_production(db, ctx, item)
        run_review(db, ctx, item, "qc", "passed")
        run_review(db, ctx, item, "qa", "rework", remarks="Attributes need correction")
        history = items_service.item_history(db, actor=ctx["team"]["tl"], role=TL_ROLE,
                                            work_package_item_id=item.id)
        assert [h["status"] for h in history if h["stage"] == "qa"] == ["rework"]
        assert item.status == "qa_rejected"


# ===========================================================================================
# delivery readiness
# ===========================================================================================

def _prepare_area(db, code):
    ctx = setup_project(db, code)
    import_master(db, ctx, xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
    package = packages_of(db, ctx)[0]
    items = [i for i in items_of(db, ctx) if i.work_package_id == package.id]
    assign_all(db, ctx, package, [i.id for i in items])
    return ctx, package, items


def test_readiness_counts_while_pending():
    with SessionLocal() as db:
        ctx, package, items = _prepare_area(db, "RD1")
        readiness = items_service.delivery_readiness(db, package)
        assert readiness["applicable"] is True and readiness["ready"] is False
        assert readiness["total_files"] == 1
        assert readiness["production_pending"] == 1 and readiness["qa_passed"] == 0


def test_ready_when_all_qa_pass_and_tl_confirms():
    with SessionLocal() as db:
        ctx, package, items = _prepare_area(db, "RD2")
        for item in items:
            run_production(db, ctx, item)
            run_review(db, ctx, item, "qc", "passed")
            run_review(db, ctx, item, "qa", "passed")
        readiness = items_service.delivery_readiness(db, package)
        assert readiness["ready"] is True and readiness["qa_passed"] == readiness["total_files"]
        ready_notice = db.scalar(select(func.count(GlobalNotification.id)).where(
            GlobalNotification.event_type == "workflow.package.ready_for_delivery",
            GlobalNotification.recipient_user_id == ctx["team"]["tl"].id))
        assert ready_notice == 1
        confirmed = items_service.mark_ready_for_delivery(
            db, actor=ctx["team"]["tl"], role=TL_ROLE, work_package_id=package.id)
        assert confirmed["confirmed"] is True and confirmed["ready_by_id"] == ctx["team"]["tl"].id
        assert package.ready_for_delivery_at is not None
        pm_notice = db.scalar(select(func.count(GlobalNotification.id)).where(
            GlobalNotification.event_type == "workflow.package.ready_confirmed",
            GlobalNotification.recipient_user_id == ctx["pm"].id))
        assert pm_notice == 1
        # idempotent: confirming again keeps the original timestamp
        first_at = package.ready_for_delivery_at
        again = items_service.mark_ready_for_delivery(
            db, actor=ctx["team"]["tl"], role=TL_ROLE, work_package_id=package.id)
        assert again["ready_at"] == first_at.isoformat()


def test_mark_ready_requires_all_passed():
    with SessionLocal() as db:
        ctx, package, items = _prepare_area(db, "RD3")
        run_production(db, ctx, items[0])
        run_review(db, ctx, items[0], "qc", "passed")  # QA still pending
        with pytest.raises(ValueError, match="Not all files passed QA"):
            items_service.mark_ready_for_delivery(
                db, actor=ctx["team"]["tl"], role=TL_ROLE, work_package_id=package.id)


def test_no_file_package_legacy_readiness():
    with SessionLocal() as db:
        ctx = setup_project(db, "RD4")
        # a legacy package created by the existing area-level allocation flow (no files)
        from app.modules.operations.schemas import WorkflowWorkAllocation
        from app.modules.operations.workflow_service import allocate_work
        team = ctx["team"]
        package = allocate_work(db, actor=team["tl"], project_id=ctx["project"].id,
                                payload=WorkflowWorkAllocation(
                                    package_code="LEG-01", area_name="Legacy area",
                                    quantity=Decimal("5"), quantity_unit="km",
                                    target_date=date(2026, 10, 20), instructions="Legacy",
                                    production_user_id=team["prod"].id, qc_user_id=team["qc"].id,
                                    qa_user_id=team["qa"].id))
        db.flush()
        readiness = items_service.delivery_readiness(db, package)
        assert readiness["applicable"] is False and readiness["ready"] is False
        with pytest.raises(ValueError, match="tracks no files"):
            items_service.mark_ready_for_delivery(
                db, actor=team["tl"], role=TL_ROLE, work_package_id=package.id)


# ===========================================================================================
# permissions
# ===========================================================================================

def test_management_read_only():
    with SessionLocal() as db:
        ctx = setup_project(db, "PM1")
        content = xlsx_bytes(MASTER_HEADERS, MASTER_ROWS)
        management = make_user(db, "v83.mgmt@nakshatech.com", "management")
        with pytest.raises(PermissionError, match="read-only"):
            alloc.commit_master_import(db, actor=management, role="management",
                                       project_id=ctx["project"].id,
                                       filename="master.xlsx", content=content)
        import_master(db, ctx, content)
        package = packages_of(db, ctx)[0]
        item = items_of(db, ctx)[0]
        with pytest.raises(PermissionError, match="read-only"):
            alloc.assign_files(
                db, actor=management, role="management", work_package_id=package.id,
                item_ids=[item.id], production_user_id=ctx["team"]["prod"].id,
                production_target_date=date(2026, 10, 20), qc_user_id=ctx["team"]["qc"].id,
                qc_target_date=date(2026, 10, 22), qa_user_id=ctx["team"]["qa"].id,
                qa_target_date=date(2026, 10, 24), supply_date=None, note=None)
        with pytest.raises(PermissionError, match="read-only"):
            items_service.mark_ready_for_delivery(
                db, actor=management, role="management", work_package_id=package.id)
        # read-only visibility still works
        options = alloc.allocation_options(db, actor=management, role="management",
                                           project_id=ctx["project"].id)
        assert len(options["areas"]) == 3


def test_pm_can_import_and_monitor():
    with SessionLocal() as db:
        ctx = setup_project(db, "PM2")
        summary = alloc.commit_master_import(
            db, actor=ctx["pm"], role="lidar", project_id=ctx["project"].id,
            filename="master.xlsx", content=xlsx_bytes(MASTER_HEADERS, MASTER_ROWS))
        assert summary["created_items"] == 3
        package = packages_of(db, ctx)[0]
        item = items_of(db, ctx)[0]
        assign_all(db, ctx, package, [item.id])
        run_production(db, ctx, item)
        payload = items_service.item_list(db, actor=ctx["pm"], role="lidar",
                                          work_package_id=package.id)
        assert payload["can_manage"] is True and len(payload["items"]) == 1


# ===========================================================================================
# workbook exports
# ===========================================================================================

def _workbook(db, actor, role, project_id):
    content, filename = build_production_workbook(
        db, report_type="project_workbook", actor=actor, role=role, project_id=project_id)
    from openpyxl import load_workbook
    return load_workbook(BytesIO(content), data_only=True), filename


def test_area_master_exact_26_and_file_level_values():
    with SessionLocal() as db:
        ctx, package, items = _prepare_area(db, "RP1")
        run_production(db, ctx, items[0], quantity="1.16")
        run_review(db, ctx, items[0], "qc", "passed")
        run_review(db, ctx, items[0], "qa", "passed")
        workbook, _ = _workbook(db, ctx["pm"], "lidar", ctx["project"].id)
        sheet = workbook["AREA MASTER"]
        assert [c.value for c in sheet[1]] == AREA_MASTER_HEADERS
        assert len(AREA_MASTER_HEADERS) == 26
        row = next(r for r in sheet.iter_rows(min_row=2, values_only=True) if r[1] == package.package_code)
        assert row[0] == "LOT-031" and row[2] == "FOGNATURA 1"
        assert row[3] == 1.2            # LENGTH (planned)
        assert str(row[5]).startswith("2026-10-15")   # SUPPLY DATE from allocation
        assert row[6] == 1.16           # ACTUAL KM (actual production quantity)
        assert "prod.ft.RP1" in str(row[10])   # PROD_ALLOTMENT = assignee
        assert row[11] == "completed"   # production STATUS
        assert "qc.ft.RP1" in str(row[13]) and row[14] == "approved"
        assert "qa.ft.RP1" in str(row[16]) and row[17] == "approved"
        assert row[20] == 1             # NO OF FILES = active files


def test_file_activity_targets_and_deadlines():
    with SessionLocal() as db:
        ctx, package, items = _prepare_area(db, "RP2")
        run_production(db, ctx, items[0], quantity="1.16")
        workbook, _ = _workbook(db, ctx["pm"], "lidar", ctx["project"].id)
        sheet = workbook["FILE ACTIVITY"]
        headers = [c.value for c in sheet[1]]
        assert headers == FILE_ACTIVITY_HEADERS
        row = next(r for r in sheet.iter_rows(min_row=2, values_only=True) if r[4] == items[0].item_number)
        assert row[6] == 1.2 and row[7] == "km"          # PLANNED LENGTH / UOM
        assert row[8] == "2026-10-15"                    # SUPPLY DATE
        assert row[10] == "2026-10-20"                   # PRODUCTION TARGET DATE
        assert row[16] == 1.16                           # ACTUAL QUANTITY
        assert row[18] == "completed_on_time"            # PRODUCTION DEADLINE STATUS
        assert row[26] == "2026-10-22"                   # QC TARGET DATE
        assert row[36] == "2026-10-24"                   # QA TARGET DATE
        assert row[45] == "NO"                           # READY FOR DELIVERY (QA pending)


def test_employee_productivity_counts():
    with SessionLocal() as db:
        ctx, package, items = _prepare_area(db, "RP3")
        run_production(db, ctx, items[0])
        run_review(db, ctx, items[0], "qc", "passed")
        workbook, _ = _workbook(db, ctx["pm"], "lidar", ctx["project"].id)
        sheet = workbook["EMPLOYEE PRODUCTIVITY"]
        headers = [c.value for c in sheet[1]]
        assert headers == EMPLOYEE_PRODUCTIVITY_HEADERS
        prod_row = next(r for r in sheet.iter_rows(min_row=2, values_only=True)
                        if r[0] and "prod.ft.RP3" in str(r[0]))
        assert prod_row[headers.index("MARCHE LOTTO 3")] == "LOT-031"
        assert prod_row[headers.index("CODE")] == "DW1045287"
        assert prod_row[headers.index("ASSIGNED FILES")] == 1
        assert prod_row[headers.index("COMPLETED FILES")] == 1
        assert prod_row[headers.index("PENDING FILES")] == 0
        assert prod_row[headers.index("FILES WORKED")] == 1
        assert prod_row[headers.index("ON-TIME COUNT")] == 1
        assert prod_row[headers.index("SYSTEM MALFUNCTION COUNT")] == 0


def test_workbook_project_isolation():
    with SessionLocal() as db:
        ctx_a, package_a, items_a = _prepare_area(db, "RP4")
        ctx_b = setup_project(db, "RP5")
        b_rows = [
            ["LOT-900", "DWB-001", "OTHER AREA 1", 1, 2.00, "km", ""],
            ["LOT-901", "DWB-002", "OTHER AREA 2", 2, 3.00, "km", ""],
        ]
        import_master(db, ctx_b, xlsx_bytes(MASTER_HEADERS, b_rows))
        workbook, _ = _workbook(db, ctx_a["pm"], "lidar", ctx_a["project"].id)
        codes = [r[1] for r in workbook["AREA MASTER"].iter_rows(min_row=2, values_only=True)]
        assert package_a.package_code in codes
        assert all(code not in {"DWB-001", "DWB-002"} for code in codes)  # project B absent


def test_workbook_is_a_live_snapshot():
    with SessionLocal() as db:
        ctx, package, items = _prepare_area(db, "RP6")
        first, _ = _workbook(db, ctx["pm"], "lidar", ctx["project"].id)
        first_status = next(r[17] for r in first["FILE ACTIVITY"].iter_rows(min_row=2, values_only=True))
        assert first_status == "not_started"
        run_production(db, ctx, items[0])
        second, _ = _workbook(db, ctx["pm"], "lidar", ctx["project"].id)
        second_row = next(r for r in second["FILE ACTIVITY"].iter_rows(min_row=2, values_only=True))
        assert second_row[17] == "completed"          # regenerated from the database
        assert second_row[13] is not None             # production end date present
