from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.core.database import SessionLocal
from app.models.entities import User
from app.modules.finance.client_master_io import (
    build_crm_workbook,
    import_client_master_workbook,
    parse_client_master_workbook,
)
from app.modules.finance.service import client_payload


def _actor(db) -> User:
    user = User(
        email="crm.v61.finance@nakshatech.com",
        full_name="CRM V6.1 Finance",
        password_hash="x",
        role="finance",
        branch="Head Office",
        department="Finance",
        email_verified=True,
        account_status="active",
        is_active=True,
    )
    db.add(user)
    db.flush()
    return user


def _sample_excel() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "CLIENT"
    ws.append(["VENDOR CODE", "Client Code", "Client Name", "Task", "BD Name", "Contact person", "Contact person Mail Id", "Country"])
    ws.append(["NV500", "NT100", "Sample Survey Client", "LiDAR Mapping", "BD Owner", "Ruben", "ruben@example.com", "Spain"])
    ws.append(["NV501", "NT101", "Second Client", "GIS", "BD Two", "Jane", "jane@example.com", "India"])
    uav = wb.create_sheet("UAV")
    uav.append(["UAV Code", "Client Name", "Task", "BD Name", "Contact person", "Contact person Mail Id", "Country"])
    uav.append(["UAV-01", "UAV Sample Client", "Drone Survey", "BD UAV", "Pilot", "pilot@example.com", "India"])
    out = BytesIO(); wb.save(out); return out.getvalue()


def test_v6_1_excel_import_manual_codes_and_exports_are_deterministic():
    with SessionLocal() as db:
        actor = _actor(db)
        result = import_client_master_workbook(
            db, raw=_sample_excel(), actor=actor, source_name="unit-client-codes.xlsx", overwrite_existing=False
        )
        db.commit()
        assert result["created"] == 3
        assert result["updated"] == 0
        assert result["skipped_existing"] == 0

        duplicate = import_client_master_workbook(
            db, raw=_sample_excel(), actor=actor, source_name="unit-client-codes.xlsx", overwrite_existing=False
        )
        assert duplicate["created"] == 0
        assert duplicate["skipped_existing"] == 3

        from app.modules.finance.models import FinanceClient
        rows = db.query(FinanceClient).order_by(FinanceClient.client_code).all()
        assert [row.client_code for row in rows] == ["NT100", "NT101", "UAV-01"]
        payload = client_payload(rows[0])
        assert payload["vendor_code"] == "NV500"
        assert payload["client_type"] == "client"
        assert payload["task"] == "LiDAR Mapping"
        assert payload["bd_name"] == "BD Owner"

        exported = build_crm_workbook(db)
        workbook = load_workbook(BytesIO(exported), read_only=True, data_only=True)
        assert workbook.sheetnames == ["Client Summary", "Project Summary", "Employee Allocation", "Travel KM", "Expense Claims"]
        summary = workbook["Client Summary"]
        assert summary.cell(1, 1).value == "Client ID / Code"
        assert summary.cell(2, 1).value == "NT100"
        assert summary.cell(2, 2).value == "NV500"


def test_v6_1_bundled_client_codes_workbook_is_parseable_and_keeps_official_ids():
    source = Path(__file__).resolve().parents[1] / "app" / "data" / "CLIENT_CODES_MASTER.xlsx"
    assert source.exists()
    rows, issues = parse_client_master_workbook(source.read_bytes())
    assert len(rows) >= 1000
    codes = {row["client_code"] for row in rows}
    assert "NT100" in codes
    assert len(codes) == len(rows)
    # The supplied workbook intentionally contains some incomplete client-name rows; importer reports rather than inventing names.
    assert isinstance(issues, list)
