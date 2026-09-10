from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
import hashlib

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models.entities import User
from app.modules.finance.attachments import validate_finance_attachment_bytes
from app.modules.finance.models import ExpenseClaim, FinanceProject
from app.modules.finance.reporting import build_finance_report_workbook, finance_report_payload
from app.modules.finance.schemas import ExpenseClaimCreateRequest
from app.modules.finance.service import (
    admin_decision,
    claim_payload,
    create_claim,
    ensure_finance_seed_data,
    finance_decision,
    mark_paid,
    submit_claim,
)


def make_user(db, *, email: str, name: str, role: str, department: str | None = None) -> User:
    user = User(
        email=email,
        full_name=name,
        password_hash="test-hash",
        role=role,
        branch="Head Office",
        department=department,
        email_verified=True,
        account_status="active",
        is_active=True,
    )
    db.add(user)
    db.flush()
    return user


def make_claim(db, *, employee: User, project: FinanceProject, amount: float, claim_type: str = "advance", work_start: date = date(2026, 8, 28), work_end: date = date(2026, 8, 30)) -> ExpenseClaim:
    claim = create_claim(
        db,
        requester=employee,
        payload=ExpenseClaimCreateRequest(
            project_id=project.id,
            claim_type=claim_type,
            purpose_description="Production Finance V2 reporting regression test expense.",
            requested_work_start_date=work_start,
            requested_work_end_date=work_end,
            items=[{"category": "travel", "description": "Project travel expense", "amount": amount}],
        ),
    )
    return claim



def test_finance_attachment_checksum_is_captured_for_audit_integrity():
    data = b"%PDF-1.4\n% Finance V2 test document\n"
    validated = validate_finance_attachment_bytes(
        filename="receipt.pdf",
        declared_mime_type="application/pdf",
        data=data,
    )
    assert validated.content_sha256 == hashlib.sha256(data).hexdigest()


def test_partial_payments_preserve_approved_requested_and_payment_ledger():
    with SessionLocal() as db:
        ensure_finance_seed_data(db)
        project = db.scalar(select(FinanceProject).order_by(FinanceProject.id.asc()))
        assert project is not None
        employee = make_user(db, email="employee.v2@nakshatech.com", name="Employee V2", role="employee", department="Survey")
        admin = make_user(db, email="admin.v2@nakshatech.com", name="Admin V2", role="admin")
        finance = make_user(db, email="finance.v2@nakshatech.com", name="Finance V2", role="finance")
        db.commit()

        claim = make_claim(db, employee=employee, project=project, amount=7500)
        claim = submit_claim(db, claim=claim, requester=employee)
        claim = admin_decision(db, claim=claim, actor=admin, action="approve", comments="Verified")
        claim = finance_decision(db, claim=claim, actor=finance, action="approve", comments="Approved after review", approved_value=7000, settlement_due_date=date(2026, 9, 5))
        assert float(claim.finance_approved_amount or 0) == 7000
        assert float(claim.total_amount) == 7500

        claim = mark_paid(
            db,
            claim=claim,
            actor=finance,
            payment_reference="UTR-V2-001",
            paid_amount=3000,
            payment_mode="bank_transfer",
            payment_date=date(2026, 8, 27),
            comments="First payment",
        )
        assert claim.status == "partially_paid"
        view = claim_payload(db, claim, finance, "finance")
        assert view["paid_amount"] == 3000
        assert view["remaining_amount"] == 4000
        assert len(view["payments"]) == 1
        assert view["can_mark_paid"] is True

        with pytest.raises(ValueError, match="already recorded"):
            mark_paid(
                db,
                claim=claim,
                actor=finance,
                payment_reference="UTR-V2-001",
                paid_amount=1000,
                payment_mode="bank_transfer",
                payment_date=date(2026, 8, 27),
                comments="Duplicate should fail",
            )

        claim = mark_paid(
            db,
            claim=claim,
            actor=finance,
            payment_reference="UTR-V2-002",
            paid_amount=4000,
            payment_mode="upi",
            payment_date=date(2026, 8, 28),
            comments="Final payment",
        )
        assert claim.status == "paid"
        view = claim_payload(db, claim, finance, "finance")
        assert view["paid_amount"] == 7000
        assert view["remaining_amount"] == 0
        assert len(view["payments"]) == 2
        assert view["can_mark_paid"] is False


def test_month_quarter_year_reporting_and_excel_are_reconciled():
    with SessionLocal() as db:
        ensure_finance_seed_data(db)
        projects = list(db.scalars(select(FinanceProject).order_by(FinanceProject.id.asc()).limit(2)).all())
        assert len(projects) == 2
        employee = make_user(db, email="report.employee@nakshatech.com", name="Report Employee", role="employee", department="GIS")
        admin = make_user(db, email="report.admin@nakshatech.com", name="Report Admin", role="admin")
        finance = make_user(db, email="report.finance@nakshatech.com", name="Report Finance", role="finance")
        db.commit()

        august = make_claim(db, employee=employee, project=projects[0], amount=5000)
        august = submit_claim(db, claim=august, requester=employee)
        august.submitted_at = datetime(2026, 8, 10, 9, 0)
        august = admin_decision(db, claim=august, actor=admin, action="approve", comments="Verified")
        august = finance_decision(db, claim=august, actor=finance, action="approve", comments="Approved", approved_value=4500, settlement_due_date=date(2026, 9, 5))
        august = mark_paid(db, claim=august, actor=finance, payment_reference="RPT-AUG-1", paid_amount=4500, payment_mode="bank_transfer", payment_date=date(2026, 8, 12), comments="Paid")

        september = make_claim(db, employee=employee, project=projects[1], amount=2500, work_start=date(2026, 9, 5), work_end=date(2026, 9, 7))
        september = submit_claim(db, claim=september, requester=employee)
        september.submitted_at = datetime(2026, 9, 5, 11, 30)
        september.purpose_description = '=HYPERLINK("https://invalid.example","should-not-run")'
        db.commit()

        january = make_claim(db, employee=employee, project=projects[0], amount=1000, work_start=date(2026, 1, 15), work_end=date(2026, 1, 17))
        january = submit_claim(db, claim=january, requester=employee)
        january.submitted_at = datetime(2026, 1, 15, 8, 0)
        db.commit()

        june = make_claim(db, employee=employee, project=projects[0], amount=1200, work_start=date(2026, 6, 20), work_end=date(2026, 6, 22))
        june = submit_claim(db, claim=june, requester=employee)
        june.submitted_at = datetime(2026, 6, 20, 8, 0)
        june = admin_decision(db, claim=june, actor=admin, action="approve", comments="Verified")
        june = finance_decision(db, claim=june, actor=finance, action="approve", comments="Approved", approved_value=1200, settlement_due_date=date(2026, 6, 30))
        june = mark_paid(db, claim=june, actor=finance, payment_reference="RPT-JUNE-PAID-SEP", paid_amount=1200, payment_mode="bank_transfer", payment_date=date(2026, 9, 10), comments="Paid in September")

        month_report = finance_report_payload(db, period="month", year=2026, month=8, page=1, page_size=50)
        assert month_report["period_label"] == "August 2026"
        assert month_report["total_records"] == 1
        assert month_report["requested_amount"] == 5000
        assert month_report["approved_amount"] == 4500
        assert month_report["paid_amount"] == 4500
        assert month_report["outstanding_amount"] == 0
        assert month_report["payment_period_amount"] == 4500
        assert month_report["payment_period_count"] == 1

        september_report = finance_report_payload(db, period="month", year=2026, month=9, page=1, page_size=50)
        assert september_report["total_records"] == 1
        assert september_report["paid_amount"] == 0
        assert september_report["payment_period_amount"] == 1200
        assert september_report["payment_period_count"] == 1

        q3_report = finance_report_payload(db, period="quarter", year=2026, quarter=3, page=1, page_size=50)
        assert q3_report["period_label"] == "Q3 2026"
        assert q3_report["total_records"] == 2
        assert q3_report["requested_amount"] == 7500
        assert q3_report["payment_period_amount"] == 5700
        assert q3_report["payment_period_count"] == 2

        year_report = finance_report_payload(db, period="year", year=2026, page=1, page_size=50)
        assert year_report["total_records"] == 4
        assert 2026 in year_report["available_years"]

        output, resolved, count = build_finance_report_workbook(db, period="quarter", year=2026, quarter=3)
        assert resolved.label == "Q3 2026"
        assert count == 2
        workbook = load_workbook(BytesIO(output.getvalue()), data_only=False)
        assert workbook.sheetnames[:6] == ["Summary", "Claims Ledger", "Expense Items", "Approval History", "Payments in Period", "Attachment Register"]
        assert "Data Integrity" in workbook.sheetnames
        assert "Project Summary" in workbook.sheetnames
        assert "Employee Summary" in workbook.sheetnames
        assert "Category Summary" in workbook.sheetnames
        assert workbook["Claims Ledger"].max_row == 3
        assert workbook["Claims Ledger"]["J2"].value.startswith("'=HYPERLINK")
        assert workbook["Summary"]["B8"].value.startswith("=SUM")
        assert workbook["Payments in Period"].max_row == 3
        assert workbook["Summary"]["B12"].value.startswith("=SUM")
