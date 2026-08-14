from __future__ import annotations

from datetime import date

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

# Import the cumulative runtime during test collection so the additive return
# models are registered in Base.metadata before the autouse schema fixture runs.
from app.batch4_main import app  # noqa: F401
from app.core.database import SessionLocal
from app.models.entities import Asset, AssetHistory, User
from app.modules.asset_return.models import AssetVendorReturn, SpareMonitor
from app.modules.asset_return.reporting import build_vendor_return_workbook
from app.modules.asset_return.router import create_component_change_batch_with_spares, list_active_assets
from app.modules.asset_return.schemas import AssetVendorReturnCreate
from app.modules.asset_return.service import perform_vendor_return
from app.schemas.component_replacement import ComponentChangeBatchCreate, ComponentChangeItem


def _user(db, email: str = "rental-return-it@nakshatech.com", role: str = "it") -> User:
    user = User(
        email=email,
        full_name="Rental Return IT",
        password_hash="pytest-only-password-hash",
        role=role,
        branch="Head Office",
        is_active=True,
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def _desktop(
    db,
    *,
    code: str,
    cpu: str,
    monitor: str | None,
    used_by: str | None = None,
    workstation: str | None = None,
) -> Asset:
    asset = Asset(
        asset_code=code,
        cpu_asset_tag=cpu,
        monitor_asset_tags=monitor,
        device_type="Computer",
        status="assigned" if used_by else "available",
        work_mode="office",
        used_by=used_by,
        department="GIS" if used_by else "IT",
        workstation_no=workstation,
        location="Head Office",
        asset_date=date(2026, 8, 1),
    )
    db.add(asset)
    db.commit()
    db.refresh(asset)
    return asset


def _return_payload(mode: str, *, vendor: str = "UAT Rental Vendor") -> AssetVendorReturnCreate:
    return AssetVendorReturnCreate(
        return_mode=mode,
        return_date=date(2026, 8, 14),
        vendor_name=vendor,
        return_reference="UAT-DC-001",
        condition="Good / working",
        reason="Rental period completed",
        remarks="Pytest vendor return",
        reporting_month="2026-08",
        spare_location="IT Store",
        confirm_vendor_return=True,
    )


def test_complete_vendor_return_preserves_history_and_removes_asset_from_active_register():
    with SessionLocal() as db:
        it_user = _user(db)
        asset = _desktop(db, code="NT-PC-RET-001", cpu="RENT-CPU-001", monitor="RENT-MON-001")
        asset_id = asset.id

        result = perform_vendor_return(db, asset.id, _return_payload("complete_return"), it_user)
        assert result["return_mode"] == "complete_return"
        assert result["retained_monitor_tags"] is None

        persisted = db.get(Asset, asset_id)
        assert persisted is not None
        assert persisted.status == "returned_to_vendor"
        # Complete return keeps the old monitor tag on the inactive source row as
        # trace evidence; it is not converted into a reusable spare.
        assert persisted.monitor_asset_tags == "RENT-MON-001"
        assert db.scalar(select(SpareMonitor).where(SpareMonitor.source_asset_id == asset_id)) is None
        return_record = db.scalar(select(AssetVendorReturn).where(AssetVendorReturn.asset_id == asset_id))
        assert return_record is not None
        audit = db.scalar(
            select(AssetHistory)
            .where(AssetHistory.asset_id == asset_id, AssetHistory.change_type == "asset_vendor_return")
        )
        assert audit is not None
        assert audit.batch_code == return_record.return_code

        active_rows = list_active_assets(
            search=None,
            department=None,
            device_type=None,
            status=None,
            work_mode=None,
            month=None,
            limit=500,
            db=db,
            user=it_user,
        )
        assert all(row["id"] != asset_id for row in active_rows)
        # Direct DB identity survives, proving this is an audited removal rather
        # than destructive SQL deletion.
        assert db.get(Asset, asset_id) is not None


def test_vendor_return_is_blocked_until_employee_custody_is_closed():
    with SessionLocal() as db:
        it_user = _user(db)
        asset = _desktop(
            db,
            code="NT-PC-RET-002",
            cpu="RENT-CPU-002",
            monitor="RENT-MON-002",
            used_by="Assigned Employee",
            workstation="WS-RET-002",
        )
        with pytest.raises(HTTPException) as exc:
            perform_vendor_return(db, asset.id, _return_payload("complete_return"), it_user)
        assert exc.value.status_code == 409
        assert "Handover & Return" in str(exc.value.detail)
        db.refresh(asset)
        assert asset.status == "assigned"
        assert db.scalar(select(AssetVendorReturn).where(AssetVendorReturn.asset_id == asset.id)) is None


def test_return_without_monitor_creates_available_spare_and_excel_register():
    with SessionLocal() as db:
        it_user = _user(db)
        asset = _desktop(db, code="NT-PC-RET-003", cpu="RENT-CPU-003", monitor="KEEP-MON-003")
        result = perform_vendor_return(db, asset.id, _return_payload("return_without_monitor"), it_user)
        assert result["retained_monitor_tags"] == "KEEP-MON-003"

        db.refresh(asset)
        assert asset.status == "returned_to_vendor"
        assert asset.monitor_asset_tags is None
        spare = db.scalar(select(SpareMonitor).where(SpareMonitor.monitor_tag == "KEEP-MON-003"))
        assert spare is not None
        assert spare.status == "available"
        assert spare.current_asset_id is None
        assert spare.location == "IT Store"

        workbook = load_workbook(build_vendor_return_workbook(db), read_only=True, data_only=True)
        assert workbook.sheetnames == ["Summary", "Returned Assets", "Retained Monitors", "Return Audit History"]
        returned_values = list(workbook["Returned Assets"].values)
        spare_values = list(workbook["Retained Monitors"].values)
        assert any("RENT-CPU-003" in tuple(str(value or "") for value in row) for row in returned_values)
        assert any("KEEP-MON-003" in tuple(str(value or "") for value in row) for row in spare_values)
        workbook.close()


def test_retained_spare_monitor_is_consumed_by_existing_component_change_atomically():
    with SessionLocal() as db:
        it_user = _user(db)
        source = _desktop(db, code="NT-PC-RET-004", cpu="RENT-CPU-004", monitor="SPARE-MON-004")
        perform_vendor_return(db, source.id, _return_payload("return_without_monitor"), it_user)
        target = _desktop(db, code="NT-PC-TARGET-004", cpu="TARGET-CPU-004", monitor="BAD-MON-004")

        response = create_component_change_batch_with_spares(
            payload=ComponentChangeBatchCreate(
                reporting_month="2026-08",
                asset_id=target.id,
                change_type="replacement",
                technician=it_user.full_name,
                replacement_date=date(2026, 8, 14),
                remarks="Use retained rental monitor",
                items=[ComponentChangeItem(
                    component_type="Monitor",
                    change_type="replacement",
                    old_value="BAD-MON-004",
                    new_value="SPARE-MON-004",
                    old_condition="Damaged / not working",
                    reason="Replace failed monitor with retained spare",
                )],
            ),
            db=db,
            user=it_user,
        )
        assert response["records"][0]["new_value"] == "SPARE-MON-004"
        db.refresh(target)
        assert target.monitor_asset_tags == "SPARE-MON-004"
        spare = db.scalar(select(SpareMonitor).where(SpareMonitor.monitor_tag == "SPARE-MON-004"))
        assert spare is not None
        assert spare.status == "in_use"
        assert spare.current_asset_id == target.id
        assert spare.assigned_by_email == it_user.email

        second_target = _desktop(db, code="NT-PC-TARGET-005", cpu="TARGET-CPU-005", monitor="BAD-MON-005")
        with pytest.raises(HTTPException) as exc:
            create_component_change_batch_with_spares(
                payload=ComponentChangeBatchCreate(
                    reporting_month="2026-08",
                    asset_id=second_target.id,
                    change_type="replacement",
                    technician=it_user.full_name,
                    replacement_date=date(2026, 8, 14),
                    items=[ComponentChangeItem(
                        component_type="Monitor",
                        change_type="replacement",
                        old_value="BAD-MON-005",
                        new_value="SPARE-MON-004",
                        old_condition="Damaged / not working",
                        reason="Must not reuse an in-use spare",
                    )],
                ),
                db=db,
                user=it_user,
            )
        assert exc.value.status_code == 409
        db.refresh(second_target)
        assert second_target.monitor_asset_tags == "BAD-MON-005"
