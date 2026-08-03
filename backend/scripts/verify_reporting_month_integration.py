"""Self-contained verification for IT effective reporting-month behavior.

Runs against an in-memory SQLite database. It does not connect to or modify the
real PostgreSQL database. Execute inside the backend container or local venv:

    python scripts/verify_reporting_month_integration.py
"""
from __future__ import annotations

from pathlib import Path
import sys

# Allow this file to be run both as `python scripts/<name>.py` and
# as `python -m scripts.<name>` from the backend container.
BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from datetime import datetime
import os

# Force application imports to use SQLite; never touch the configured PostgreSQL.
os.environ["DATABASE_URL"] = "sqlite+pysqlite:///:memory:"

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

# Register all application tables before create_all.
from app.core.database import Base
from app.models.entities import Asset, AssetHistory, User
from app.modules.backup import models as _backup_models  # noqa: F401
from app.modules.drone import models as _drone_models  # noqa: F401
from app.modules.it_activity import models as _it_activity_models  # noqa: F401
from app.modules.it_activity.excel_service import build_monthly_it_activity_workbook
from app.modules.it_activity.service import monthly_activity_data
from app.api.router import update_asset_status
from app.schemas.asset import AssetStatusUpdate


def main() -> None:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)

    with Session(engine) as db:
        user = User(
            email="reporting-month-test@nakshatech.local",
            full_name="Reporting Month Test",
            password_hash="not-used",
            role="it",
            branch="Head Office",
        )
        asset = Asset(
            asset_code="QA-REPORT-MONTH-001",
            cpu_asset_tag="QA-CPU-001",
            device_type="Computer",
            status="available",
            work_mode="office",
            remarks="Permanent master note",
        )
        db.add_all([user, asset])
        db.flush()

        # This change was entered in August but intentionally assigned to June.
        june_recorded_in_august = AssetHistory(
            asset_id=asset.id,
            action="Asset details updated",
            change_type="full_edit",
            batch_code="EDIT-JUNE-TEST",
            old_value='{"memory_gb": "8 GB"}',
            new_value='{"memory_gb": "16 GB"}',
            reason="June inventory correction",
            remarks="June activity remark only",
            changed_by=user.email,
            changed_by_name=user.full_name,
            changed_by_role=user.role,
            field_count=1,
            reporting_month="2026-06",
            created_at=datetime(2026, 8, 3, 9, 30, 0),
        )
        august_record = AssetHistory(
            asset_id=asset.id,
            action="Asset details updated",
            change_type="full_edit",
            batch_code="EDIT-AUGUST-TEST",
            old_value='{"processor": "i5"}',
            new_value='{"processor": "i7"}',
            reason="August upgrade",
            remarks="August activity remark only",
            changed_by=user.email,
            changed_by_name=user.full_name,
            changed_by_role=user.role,
            field_count=1,
            reporting_month="2026-08",
            created_at=datetime(2026, 8, 3, 10, 30, 0),
        )
        db.add_all([june_recorded_in_august, august_record])
        db.commit()

        june = monthly_activity_data(db, "2026-06")
        august = monthly_activity_data(db, "2026-08")
        assert june["summary"]["asset_edit_operations"] == 1
        assert august["summary"]["asset_edit_operations"] == 1
        assert {item["batch_code"] for item in june["items"]} == {"EDIT-JUNE-TEST"}
        assert {item["batch_code"] for item in august["items"]} == {"EDIT-AUGUST-TEST"}
        june_item = june["items"][0]
        assert june_item["reporting_month"] == "2026-06"
        assert june_item["activity_date"] == "2026-08-03"
        assert june_item["remarks"] == "June activity remark only"

        # Status activity remarks must stay in history and never alter the
        # permanent Asset Master Remarks field.
        update_asset_status(
            asset.id,
            AssetStatusUpdate(
                status="repair",
                remarks="June repair activity remark",
                reporting_month="2026-06",
            ),
            db,
            user,
        )
        db.refresh(asset)
        assert asset.remarks == "Permanent master note"
        latest = db.query(AssetHistory).order_by(AssetHistory.id.desc()).first()
        assert latest is not None
        assert latest.reporting_month == "2026-06"
        assert latest.remarks == "June repair activity remark"

        # The monthly workbook must contain both the effective reporting month
        # and the actual system-recorded timestamp.
        workbook_bytes, counts = build_monthly_it_activity_workbook(db, "2026-06")
        assert counts["Asset Edit History"] >= 1
        workbook_bytes.seek(0)
        workbook = load_workbook(workbook_bytes, read_only=True, data_only=True)
        sheet = workbook["Asset Edit History"]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        assert "Reporting Month" in headers
        assert "System Recorded At" in headers
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        month_col = headers.index("Reporting Month")
        assert rows and all(row[month_col] == "2026-06" for row in rows)

    print("REPORTING MONTH VERIFICATION PASSED")
    print("- June-assigned activity is listed only in June")
    print("- August activity remains separate")
    print("- Actual recorded timestamp is preserved")
    print("- Activity remarks do not modify Asset Master Remarks")
    print("- Monthly Excel contains both reporting month and recorded timestamp")


if __name__ == "__main__":
    main()
