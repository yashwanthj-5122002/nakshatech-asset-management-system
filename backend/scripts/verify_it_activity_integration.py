from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.database import Base, SessionLocal, engine
from app.modules.drone import models as _drone_models  # noqa: F401
from app.modules.backup import models as _backup_models  # noqa: F401
from app.modules.it_activity import models as _models  # noqa: F401
from app.modules.it_activity.excel_service import build_monthly_it_activity_workbook
from app.modules.it_activity.models import ITHandoverRecord, ITPurchaseRecord
from app.modules.it_activity.service import monthly_activity_data


def main() -> None:
    Base.metadata.create_all(bind=engine)
    current_month = datetime.now().strftime("%Y-%m")
    with SessionLocal() as db:
        handovers = db.scalar(select(func.count(ITHandoverRecord.id))) or 0
        purchases = db.scalar(select(func.count(ITPurchaseRecord.id))) or 0
        data = monthly_activity_data(db, current_month, limit=10000)
        stream, counts = build_monthly_it_activity_workbook(db, current_month)
        workbook = load_workbook(stream, read_only=True, data_only=False)
        required = {
            "Monthly Summary", "Detailed Monthly Activity", "Asset Edit History", "Component Changes",
            "Laptop Handover Return", "Desktop Handover Return", "Purchase Details", "User Activity Summary",
        }
        missing = required.difference(workbook.sheetnames)
        if missing:
            raise RuntimeError(f"Monthly workbook is missing sheets: {sorted(missing)}")
        print("IT ACTIVITY INTEGRATION VERIFIED")
        print(f"Historical handover/return records: {handovers}")
        print(f"Historical purchase records: {purchases}")
        print(f"Current month detailed activities: {data['total']}")
        print(f"Workbook sheets: {', '.join(workbook.sheetnames)}")
        print(f"Workbook row counts: {counts}")


if __name__ == "__main__":
    main()
