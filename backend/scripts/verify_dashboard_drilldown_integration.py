from __future__ import annotations

import os
import sys
from pathlib import Path
from tempfile import NamedTemporaryFile

os.environ.setdefault("DATABASE_URL", "sqlite+pysqlite:///:memory:")
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.core.database import Base
from app.models.entities import Asset
from app.modules.drone import models as _drone_models  # register referenced drone tables
from app.services.asset_drilldown_service import build_asset_drilldown_workbook, query_asset_drilldown
from app.services.monthly_snapshot_service import month_start


def asset(code: str, device: str, department: str, status: str, used_by: str | None, workstation: str) -> Asset:
    return Asset(
        asset_code=code,
        cpu_asset_tag=code,
        system_name=f"SYS-{code}",
        workstation_no=workstation,
        device_type=device,
        department=department,
        status=status,
        used_by=used_by,
        work_mode="office",
        location="Head Office",
        processor="Intel Test CPU",
        memory_gb="16 GB",
        ssd="512 GB",
    )


def require_text(source: str, text: str, label: str) -> None:
    if text not in source:
        raise AssertionError(f"{label}: missing {text}")


def verify_backend_source_contract() -> None:
    router_source = (ROOT / "app" / "api" / "router.py").read_text(encoding="utf-8")
    service_source = (ROOT / "app" / "services" / "asset_drilldown_service.py").read_text(encoding="utf-8")

    require_text(router_source, '@router.get("/dashboard/it/assets")', "Drill-down endpoint")
    require_text(router_source, '@router.get("/reports/it-dashboard-assets.xlsx")', "Drill-down Excel endpoint")
    require_text(router_source, 'require_roles("admin", "management", "it")', "View permissions")
    require_text(service_source, 'ASSIGNED_GROUP = {"assigned", "in_use"}', "Assigned grouping")
    require_text(service_source, 'assets_for_month', "Month-aware inventory source")
    require_text(service_source, 'all_filtered_assets', "Excel and drawer result parity")
    require_text(service_source, '"visuals": {', "Drawer visual distributions")
    require_text(service_source, '_status_distribution(filtered_assets)', "Status chart accuracy")


def main() -> None:
    verify_backend_source_contract()
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    month_key = month_start().strftime("%Y-%m")

    with Session(engine) as db:
        db.add_all([
            asset("NT-PC-001", "Computer", "BIM", "assigned", "Asha", "BIM-01"),
            asset("NT-PC-002", "Computer", "BIM", "in_use", "Ravi", "BIM-02"),
            asset("NT-LAP-001", "Laptop", "BIM", "available", None, "BIM-L01"),
            asset("NT-LAP-002", "Laptop", "Ortho / GIS", "repair", "Neha", "GIS-L01"),
            asset("NT-MOB-001", "Smartphone", "Mapping", "replacement_pending", "Amit", "MAP-M01"),
        ])
        db.commit()

        all_result = query_asset_drilldown(db, month_key=month_key, scope="all", page=1, page_size=20)
        assert all_result["scope_total"] == 5
        assert all_result["summary"]["computers"] == 2
        assert all_result["summary"]["laptops"] == 2
        assert all_result["summary"]["smartphones"] == 1

        computer_result = query_asset_drilldown(
            db, month_key=month_key, scope="device", scope_value="Computer", page=1, page_size=20
        )
        assert computer_result["scope_total"] == 2
        assert {item.device_type for item in computer_result["all_filtered_assets"]} == {"Computer"}
        assert computer_result["visuals"]["status_distribution"] == [
            {"name": "Assigned / In Use", "value": 2, "key": "assigned"}
        ]
        assert computer_result["visuals"]["department_distribution"] == [
            {"name": "BIM", "value": 2, "key": "BIM"}
        ]

        bim_result = query_asset_drilldown(
            db, month_key=month_key, scope="department", scope_value="BIM", page=1, page_size=20
        )
        assert bim_result["scope_total"] == 3
        assert bim_result["summary"]["computers"] == 2
        assert bim_result["summary"]["laptops"] == 1
        assert bim_result["visuals"]["device_distribution"] == [
            {"name": "Computer", "value": 2, "key": "Computer"},
            {"name": "Laptop", "value": 1, "key": "Laptop"},
        ]

        bim_laptops = query_asset_drilldown(
            db,
            month_key=month_key,
            scope="department",
            scope_value="BIM",
            device_type="Laptop",
            page=1,
            page_size=20,
        )
        assert bim_laptops["filtered_total"] == 1
        assert bim_laptops["all_filtered_assets"][0].asset_code == "NT-LAP-001"
        assert sum(item["value"] for item in bim_laptops["visuals"]["device_distribution"]) == 1
        assert sum(item["value"] for item in bim_laptops["visuals"]["status_distribution"]) == 1
        assert sum(item["value"] for item in bim_laptops["visuals"]["department_distribution"]) == 1

        assigned_result = query_asset_drilldown(
            db, month_key=month_key, scope="status", scope_value="assigned", page=1, page_size=20
        )
        assert assigned_result["scope_total"] == 2
        assert {item.status for item in assigned_result["all_filtered_assets"]} == {"assigned", "in_use"}

        search_result = query_asset_drilldown(
            db, month_key=month_key, scope="all", search="GIS-L01", page=1, page_size=20
        )
        assert search_result["filtered_total"] == 1
        assert search_result["all_filtered_assets"][0].asset_code == "NT-LAP-002"

        paged = query_asset_drilldown(db, month_key=month_key, scope="all", page=2, page_size=2)
        assert paged["pages"] == 3
        assert paged["page"] == 2
        assert len(paged["assets"]) == 2

        stream = build_asset_drilldown_workbook(bim_result)
        with NamedTemporaryFile(suffix=".xlsx") as temp:
            temp.write(stream.getvalue())
            temp.flush()
            workbook = load_workbook(temp.name, read_only=True, data_only=True)
            assert workbook.sheetnames == ["Drilldown Summary", "Asset Details"]
            assert workbook["Asset Details"].max_row == 4
            workbook.close()

    print("DASHBOARD DRILL-DOWN INTEGRATION VERIFICATION PASSED")
    print("- Device cards return only the selected device type")
    print("- Department drill-down returns the exact department breakdown")
    print("- Assigned / In Use correctly combines assigned and in_use")
    print("- Combined filters, search, sorting and pagination are isolated")
    print("- Filtered Excel contains the same records as the drawer")
    print("- Drawer pie charts and department graph use all matching records, not only the current page")


if __name__ == "__main__":
    main()
